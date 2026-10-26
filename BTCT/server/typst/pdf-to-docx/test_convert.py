#!/usr/bin/env python3
"""Checks for convert.py's readers, against a PDF this file builds itself.

    <venv>/python -I -B test_convert.py          the built-in document
    <venv>/python -I -B test_convert.py a.pdf    plus a real report

Run by hand. It is deliberately not part of `bun run test`: the server's test
suite runs under Bun with no Python, and this needs pymupdf. It needs neither
Word nor the converter's output, because everything it checks is a pure
function of the PDF: which bands repeat, where they sit, what the page-number
offset is, and which lines read as a table of contents.

With a PDF argument it prints what it found there instead of asserting, which
is the quickest way to see why a real report behaves the way it does.
"""

import collections
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from xml.etree import ElementTree

import pymupdf
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Pt

# Isolated mode drops the script's own directory from the import path, and
# this file is run with it so it sees what the server's child process sees.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import convert  # noqa: E402

PAGE_WIDTH, PAGE_HEIGHT = 612.0, 792.0
LEFT, RIGHT = 54.0, 558.0
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
# The run children that put something on the page. A run of none of these is
# a bookmark or a field boundary, and what font it names does not show.
VISIBLE_RUN_CHILDREN = tuple(W + tag for tag in ("t", "tab", "br", "sym", "noBreakHyphen",
                                                 "softHyphen", "instrText"))


def build_pdf():
    """Four pages: a bare cover, then three with a header, footer and a list.

    Page 1 (index 1) carries the table of contents. Printed page numbers start
    at 1 on index 1, which is the offset the converter has to work out so Word
    prints the same numbers.
    """
    doc = pymupdf.open()
    for index in range(4):
        page = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        if index == 0:
            page.insert_text((200, 400), "COVER", fontsize=30)
            continue
        page.insert_text((LEFT, 46), "ACME Security Assessment", fontsize=10)
        page.insert_text((RIGHT - 60, 46), "Contoso Ltd", fontsize=10)
        page.draw_line((LEFT, 53), (RIGHT, 53), width=1)
        page.draw_rect(pymupdf.Rect(0, 770, PAGE_WIDTH, PAGE_HEIGHT), fill=(0.95, 0.95, 0.95))
        page.insert_text((250, 785), "CONFIDENTIAL", fontsize=9)
        page.insert_text((RIGHT - 10, 785), str(index), fontsize=9)
        page.insert_text((LEFT, 90), "Body text on page %d." % index, fontsize=11)
    toc = doc[1]
    leader = " . " * 40
    toc.insert_text((LEFT, 150), "1. Introduction" + leader + "1", fontsize=11)
    toc.insert_text((LEFT + 16, 170), "1.1", fontsize=11)
    toc.insert_text((LEFT + 38, 170), "Scope" + leader + "2", fontsize=11)
    toc.insert_text((LEFT, 190), "2. Findings" + leader + "3", fontsize=11)
    return pymupdf.open("pdf", doc.tobytes())


def solid_leader_pdf():
    """One page whose dot leader is set with no space between the dots."""
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((LEFT, 150), "A title" + "." * 40 + "7", fontsize=11)
    return pymupdf.open("pdf", doc.tobytes())


def check(condition, label, detail=""):
    line = "%-5s %s%s" % ("ok" if condition else "FAIL", label, (" -- " + str(detail)) if detail else "")
    # A Windows console is cp1252 and some of these labels carry Cyrillic or
    # CJK, which would end the run with an encoding error rather than a result.
    print(line.encode("ascii", "backslashreplace").decode("ascii"))
    return bool(condition)


def run_built_in():
    doc = build_pdf()
    passed = True

    left, right = convert.content_bounds(doc)
    passed &= check(abs(left - LEFT) < 2, "content left edge", left)
    passed &= check(abs(right - RIGHT) < 6, "content right edge", right)

    header, header_reason = convert.detect_band(doc, "top")
    footer, footer_reason = convert.detect_band(doc, "bottom")
    passed &= check(header_reason is None and footer_reason is None, "neither edge was declined",
                    (header_reason, footer_reason))
    passed &= check(header is not None, "a running header was found")
    passed &= check(footer is not None, "a running footer was found")
    if not header or not footer:
        return False

    passed &= check(header["pages"] == [1, 2, 3], "the cover carries no header", header["pages"])
    passed &= check(header["rule_under"], "the rule under the header joined the band")
    if header["rule_under"]:
        rule = header["rule_under"]
        passed &= check(abs(rule["x0"] - LEFT) < 2 and abs(rule["x1"] - RIGHT) < 2,
                        "the rule's own ends were measured", (rule["x0"], rule["x1"]))
        passed &= check(convert.border_eighths(rule["width"]) == str(int(round(rule["width"] * 8))),
                        "the rule's weight comes from the stroke", convert.border_eighths(rule["width"]))
        passed &= check(convert.border_eighths(0.0) == "6" and convert.border_eighths(40.0) == "96",
                        "a border weight stays inside what Word accepts")
        passed &= check(convert.band_paint_span(header, 0.0, RIGHT) == (rule["x0"], rule["x1"]),
                        "the header paragraph is indented to the rule")
    passed &= check(header["bottom"] >= 53.0, "the header band reaches the rule", header["bottom"])
    passed &= check(len(header["columns"]) == 2, "the header has a left and a right piece", len(header["columns"]))
    passed &= check(
        all(column["page_offset"] is None for column in header["columns"]),
        "no page number in the header",
    )

    passed &= check(footer["strip_behind"] is not None, "the filled strip joined the footer band")
    passed &= check(footer["top"] <= 771.0, "the footer band reaches the strip", footer["top"])
    numbered = [c for c in footer["columns"] if c["page_offset"] is not None]
    passed &= check(len(numbered) == 1, "exactly one footer piece is a page number", len(numbered))
    if numbered:
        passed &= check(numbered[0]["page_offset"] == 0, "the page-number offset", numbered[0]["page_offset"])
        passed &= check(
            [part["kind"] for part in numbered[0]["parts"]] == ["page"],
            "the page number is the whole piece",
            numbered[0]["parts"],
        )

    entries = convert.detect_toc(doc, [header, footer])
    passed &= check(len(entries) == 3, "three table-of-contents entries", len(entries))
    if len(entries) == 3:
        passed &= check([e["title"] for e in entries] == ["1. Introduction", "1.1 Scope", "2. Findings"],
                        "the titles", [e["title"] for e in entries])
        passed &= check([e["number"] for e in entries] == ["1", "2", "3"], "the page numbers")
        passed &= check(len(entries[1]["segments"]) == 2, "the numbered entry keeps its two columns")
        passed &= check(abs(entries[1]["x0"] - (LEFT + 16)) < 2, "the sub-entry is indented", entries[1]["x0"])
        pitch = convert.toc_line_pitch(entries)
        passed &= check(pitch is not None and abs(pitch - 20) < 2, "the line pitch", pitch)
        # The fixture writes " . " at 11 pt, so a dot advances about 3.1 pt
        # and the next one starts about two spaces later: what Word has to add
        # to its own leader is the pair of spaces between them.
        extras = [e.get("leader_extra", 0.0) for e in entries]
        passed &= check(all(5.0 < extra < 7.5 for extra in extras),
                        "the leader's extra spacing was measured", [round(e, 2) for e in extras])
        solid = convert.leader_spacing(pymupdf.open("pdf", solid_leader_pdf().tobytes())[0])
        passed &= check(solid and all(value < 0.2 for value in solid.values()),
                        "a solid leader asks for no extra spacing", solid)

    # Nothing outside a band may be taken for one.
    body_before = doc[2].get_text().count("Body text")
    convert.strip_bands(doc, [header, footer])
    after = doc[2].get_text()
    passed &= check("ACME Security Assessment" not in after, "the header text is gone")
    passed &= check("CONFIDENTIAL" not in after, "the footer text is gone")
    passed &= check(after.count("Body text") == body_before, "the body text survived")

    # The same pure helpers the rest of the script leans on.
    passed &= check(convert.normalize("Page  12 of 34") == "Page # of #", "digit runs normalize away")
    passed &= check(convert.font_family("ABCDEF+Poppins-Bold") == "Poppins", "font family")
    passed &= check(convert.font_family("ArialMT") == "Arial", "font family, no separator")
    passed &= check(convert.TOC_LINE.match("A title . . . . . . . 12") is not None, "a spaced dot leader")
    passed &= check(convert.TOC_LINE.match("Not a leader ... 12") is None, "an ellipsis is not a leader")
    passed &= check(convert.squash("Глава 1") == "глава1",
                    "squash keeps Cyrillic", convert.squash("Глава 1"))
    passed &= check(convert.squash("第一章") == "第一章", "squash keeps CJK")
    doc.close()
    passed &= run_comparison_checks()
    passed &= run_overprint_checks()
    passed &= run_orphan_char_checks()
    passed &= run_split_cell_checks()
    passed &= run_rotated_checks()
    passed &= run_cell_edge_checks()
    passed &= run_toc_checks()
    passed &= run_thread_checks()
    passed &= run_tail_checks()
    passed &= run_line_shape_checks()
    passed &= run_run_checks()
    passed &= run_token_space_checks()
    passed &= run_column_checks()
    passed &= run_column_page_end_checks()
    passed &= run_join_checks()
    passed &= run_align_checks()
    passed &= run_rhythm_checks()
    passed &= run_lost_walk_checks()
    passed &= run_tagged_member_checks()
    passed &= run_page_start_checks()
    passed &= run_cover_checks()
    passed &= run_side_picture_checks()
    passed &= run_cell_condense_checks()
    passed &= run_split_dash_checks()
    passed &= run_header_bar_checks()
    passed &= run_list_hanging_checks()
    passed &= run_tab_stop_checks()
    passed &= run_list_break_checks()
    passed &= run_margin_note_checks()
    passed &= run_page_fill_band_checks()
    return passed


def run_page_start_checks():
    """A Word page break goes before a PDF page's first line only where Word would not start one."""
    from docx.oxml.ns import qn

    def paragraphs_of(document):
        return list(document.element.body.iter(qn("w:p")))

    def starts_with_break(p):
        run = p.find(qn("w:r"))
        br = None if run is None else run.find(qn("w:br"))
        return br is not None and br.get(qn("w:type")) == "page"

    document = Document()
    for text in ("alpha one", "bravo two", "charlie three", "delta four", "echo five"):
        document.add_paragraph(text)
    wanted = [convert.compare_squash("charlie three"), convert.compare_squash("echo five")]
    declined = collections.Counter()
    inserted = convert.ensure_page_starts(document, wanted, declined)
    ps = paragraphs_of(document)
    passed = check(inserted == 2 and starts_with_break(ps[2]) and starts_with_break(ps[4])
                   and not starts_with_break(ps[1]), "page starts: a break goes before each page's first line", inserted)
    passed &= check(convert.ensure_page_starts(document, wanted, declined) == 0
                    and len(list(document.element.body.iter(qn("w:br")))) == 2,
                    "page starts: a second call adds nothing")
    passed &= check(convert.ensure_page_starts(document, [convert.compare_squash("nowhere to be seen")], declined) == 0
                    and declined["page-start-unfound"] == 1,
                    "page starts: a line that is not there is counted, not guessed", declined)

    # A section break is carried in the previous paragraph's pPr, as pdf2docx writes it.
    document = Document()
    for text in ("alpha one", "bravo two", "charlie three"):
        document.add_paragraph(text)
    ps = paragraphs_of(document)
    ppr = ps[1].get_or_add_pPr() if hasattr(ps[1], "get_or_add_pPr") else ps[1].makeelement(qn("w:pPr"), {})
    ppr.append(ppr.makeelement(qn("w:sectPr"), {}))
    passed &= check(convert.ensure_page_starts(document, [convert.compare_squash("charlie three")]) == 0,
                    "page starts: a section break in the previous paragraph already starts a page")
    continuous = ppr.find(qn("w:sectPr"))
    continuous.append(continuous.makeelement(qn("w:type"), {qn("w:val"): "continuous"}))
    passed &= check(convert.ensure_page_starts(document, [convert.compare_squash("charlie three")]) == 1,
                    "page starts: a continuous section does not")

    # Never a break before the first body paragraph: that is a blank first page.
    document = Document()
    for text in ("alpha one", "bravo two"):
        document.add_paragraph(text)
    passed &= check(convert.ensure_page_starts(document, [convert.compare_squash("alpha one")]) == 0
                    and not list(document.element.body.iter(qn("w:br"))),
                    "page starts: no break before the first paragraph")

    # Section count already equals the PDF page count: nothing is matched, so a
    # wrong text match cannot cut a page. Page 2's first line "Scope" is merged
    # into a longer paragraph here, and a paragraph "Scope" sits on page 1.
    document = Document()
    for text in ("Intro", "Scope", "filler filler", "Scope of the engagement continues over a long line"):
        document.add_paragraph(text)
    ps = paragraphs_of(document)
    ppr = ps[2].get_or_add_pPr()
    ppr.append(ppr.makeelement(qn("w:sectPr"), {}))
    declined = collections.Counter()
    passed &= check(convert.word_page_sections(document) == 2
                    and convert.ensure_page_starts(document, [convert.compare_squash("Scope")], declined, pages=2) == 0
                    and not list(document.element.body.iter(qn("w:br"))) and not declined,
                    "page starts: as many sections as pages means nothing is inserted", declined)
    passed &= check(convert.ensure_page_starts(document, [convert.compare_squash("Scope")], declined, pages=3) == 1,
                    "page starts: fewer sections than pages goes on to match")
    return passed


def run_tagged_member_checks():
    """The tagged path's row matching and its refusal of a header that changes."""
    def row(text, y0, x0=72.0, x1=300.0):
        return {"text": text, "y0": y0, "y1": y0 + 12.0, "x0": x0, "x1": x1}

    passed = True
    header, heading, body, footer = (row("Report", 36.0), row("Report", 90.0),
                                     row("Lorem ipsum", 110.0), row("Report", 760.0))
    # The member is the copy no stripped row still sits on, wherever it is:
    # here the band is the copy nearer the middle of the page.
    members = convert.tagged_members([header, heading, body, footer], [dict(header), dict(body), dict(footer)])
    passed &= check(members is not None and [id(r) for r in members] == [id(heading)],
                    "tagged members: the removed copy is chosen by position, not by nearness to an edge",
                    members and [r["y0"] for r in members])
    members = convert.tagged_members([header, heading, body, footer], [dict(heading), dict(body)])
    passed &= check(members is not None and {id(r) for r in members} == {id(header), id(footer)},
                    "tagged members: a body heading equal to the band keeps its copy",
                    members and [r["y0"] for r in members])
    passed &= check(convert.tagged_members([header, body], [row("Report Lorem", 36.0)]) is None,
                    "tagged members: a stripped row the original lacks refuses the comparison")
    passed &= check(convert.tagged_members([header, body], [row("Report", 60.0), dict(body)]) is None,
                    "tagged members: the same text somewhere else is not the same row")
    # A page of many rows is matched in time proportional to its rows, not
    # their square: 6000 distinct rows took seconds when every stripped row
    # scanned the whole original.
    many = [row("Row %d" % i, 20.0 + (i % 60) * 12.0, 72.0 + (i // 60) * 5.0) for i in range(6000)]
    started = time.perf_counter()
    members = convert.tagged_members(many, [dict(r) for r in reversed(many[1:])])
    took = time.perf_counter() - started
    passed &= check(members is not None and [id(r) for r in members] == [id(many[0])] and took < 0.5,
                    "tagged members: a page of 6000 rows is matched quickly", "%.2f s" % took)
    passed &= run_band_variant_checks(row)
    return passed


def run_band_variant_checks(row):
    """Which groups of a band are one slot with variants, and which are strays."""
    portrait, sideways = (595.3, 841.9), (841.9, 595.3)

    def group(size, text, y0, pages, x0=72.0, x1=300.0):
        return (size, text, round(y0 / convert.BAND_Y_TOLERANCE_PT)), [(i, row(text, y0, x0, x1)) for i in pages]

    def groups(*items):
        return dict(items)

    passed = True
    # Odd and even headers in one place, on pages that never share one: one slot
    # with two variants, each page of the paper size has one of them.
    passed &= check(not convert.has_stray_group(groups(group(portrait, "Acme", 36.0, [0, 2]),
                                                       group(portrait, "Contoso", 36.0, [1, 3]))),
                    "variants: odd and even headers fill their place between them")
    # A chapter's title that changes every few pages, beside a footer on all of them.
    passed &= check(not convert.has_stray_group(groups(
        group(portrait, "One", 36.0, [0, 1, 2]), group(portrait, "Two", 36.0, [3]),
        group(portrait, "Three", 36.0, [4, 5]), group(portrait, "Page #", 800.0, [0, 1, 2, 3, 4, 5]))),
        "variants: a chapter title in one place, a page number in another")
    # A left piece and a right piece on the same pages are two groups that share pages.
    passed &= check(not convert.has_stray_group(groups(
        group(portrait, "Acme", 36.0, [0, 1]), group(portrait, "Contoso", 36.0, [0, 1], 400.0, 540.0))),
        "variants: a left and a right piece on the same pages are not strays")
    # A note set at the top of one page's body, with the header on all four.
    passed &= check(convert.has_stray_group(groups(
        group(portrait, "Header", 36.0, [0, 1, 2, 3]), group(portrait, "Placeholder note", 90.0, [2]))),
        "variants: a row in a place the other pages leave empty is a stray")
    # The sideways page keeps its band somewhere else; that is its own paper.
    passed &= check(not convert.has_stray_group(groups(
        group(portrait, "Footer", 800.0, [0, 1, 2]), group(sideways, "Footer", 560.0, [3]))),
        "variants: a sideways page's band at another height is not a stray")
    # But a stray on the sideways paper is still one.
    passed &= check(convert.has_stray_group(groups(
        group(portrait, "Footer", 800.0, [0, 1, 2]), group(sideways, "Footer", 560.0, [3, 4]),
        group(sideways, "Note", 300.0, [4]))),
        "variants: a stray on the sideways paper is found against that paper's pages")
    return passed


def rhythm_document(before, after, line, text="Second block here"):
    """A two-paragraph document whose spacing the vertical pass can rewrite."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Pt(72)
    section.bottom_margin = Pt(72)
    first = document.add_paragraph()
    convert.set_spacing(first._p, before=0.0, after=0.0, line=12.0)
    first.add_run("First block here")
    second = document.add_paragraph()
    convert.set_spacing(second._p, before=before, after=after, line=line)
    second.add_run(text)
    return document


def exact_table(document, heights, widths, indent=0.0):
    """A table whose rows and cells are the sizes a PDF grid would give them."""
    table = document.add_table(rows=len(heights), cols=len(widths))
    properties = table._tbl.find(convert.qn("w:tblPr"))
    node = convert.OxmlElement("w:tblInd")
    node.set(convert.qn("w:w"), str(indent * 20.0))     # the float the converter writes
    node.set(convert.qn("w:type"), "dxa")
    properties.append(node)
    for span, row in zip(heights, table.rows):
        node = convert.OxmlElement("w:trHeight")
        node.set(convert.qn("w:val"), str(int(round(span * 20))))
        node.set(convert.qn("w:hRule"), "exact")
        row._tr.get_or_add_trPr().append(node)
        for width, cell in zip(widths, row.cells):
            node = convert.OxmlElement("w:tcW")
            node.set(convert.qn("w:w"), str(int(round(width * 20))))
            node.set(convert.qn("w:type"), "dxa")
            cell._tc.get_or_add_tcPr().insert(0, node)
    return table


def fill_cell(cell, text, before, line, after=0.0):
    paragraph = cell.paragraphs[0]
    convert.set_spacing(paragraph._p, before=before, after=after, line=line)
    paragraph.add_run(text)
    return paragraph


def pdf_line(text, base, x0, x1, ascent=9.0):
    return {"key": convert.compare_squash(text), "base": base, "x0": x0, "x1": x1,
            "y0": base - ascent, "y1": base + ascent / 4.0}


def cell_top_border(cell, width):
    borders = convert.OxmlElement("w:tcBorders")
    top = convert.OxmlElement("w:top")
    top.set(convert.qn("w:val"), "single")
    top.set(convert.qn("w:sz"), str(width * 8.0))
    borders.append(top)
    cell._tc.get_or_add_tcPr().append(borders)


WRAPPED_STRIPES = """
#set page(paper: "a4", margin: (x: 2.2cm, top: 2.5cm, bottom: 2.5cm))
#set text(size: 9pt)
= Remediation plan
#table(
  columns: (40pt, 120pt, 50pt, 80pt, 52pt, 35pt),
  inset: (x: 6pt, y: 6pt),
  align: horizon,
  stroke: none,
  fill: (x, y) => if y == 0 { rgb("#12303a") } else if calc.odd(y) { rgb("#f4efe6") },
  table.header(text(fill: white)[ID], text(fill: white)[Action], text(fill: white)[Level],
    text(fill: white)[Owner], text(fill: white)[Due], text(fill: white)[Done]),
  [MCD-01], [Rotate the build server credentials now], [Critical], [Platform team], [2026-05-04], [Yes],
  [MCD-02], [Close the Redis port], [Critical], [Infrastructure], [2026-05-11], [No],
  [MCD-03], [Parameterise the shipment search query], [High], [Application team], [2026-05-18], [No],
  [MCD-04], [Enforce unique admin passwords], [High], [Endpoint team], [2026-05-25], [Yes],
  [MCD-05], [Check object ownership], [High], [Application team], [2026-06-01], [No],
  [MCD-06], [Restrict upload types and the web root], [Medium], [Application team], [2026-06-08], [No],
  [MCD-07], [Update the bastion host packages], [Medium], [Platform team], [2026-06-15], [Yes],
)
"""


ZEBRA_CELLS = """
#set page(paper: "a4", margin: 2cm)
#set text(size: 10pt)
= Hosts
#table(columns: (3cm, 6cm, 3cm), stroke: none, inset: 5pt,
  fill: (x, y) => if y == 0 { rgb("#16263d") } else if calc.odd(y) { luma(240) } else { white },
  text(fill: white)[*Host*], text(fill: white)[*Services*], text(fill: white)[*Severity*],
  [alpha-01], [ssh 22], [High],
  [bravo-02], [smb 445 \\ rdp 3389 \\ winrm 5985], [Medium],
  [charlie-03], [ldap 389], [Low],
  [delta-04], [mssql 1433 \\ http 8080], [Low],
  [echo-05], [ftp 21], [Info],
)
#v(1.5cm)
#table(columns: (3cm, 6cm, 3cm), stroke: none, inset: 5pt,
  fill: (x, y) => if calc.odd(x + y) { luma(235) } else { rgb("#dde8f5") },
  [*Host*], [*Services*], [*Severity*],
  [alpha-01], [ssh 22], [High],
  [bravo-02], [smb 445 \\ rdp 3389], [Medium],
  [charlie-03], [ldap 389], [Low],
)
#v(1.5cm)
#table(columns: (3cm, 6cm, 3cm), stroke: none, inset: 5pt,
  fill: (x, y) => if calc.odd(y) { luma(235) },
  [*Host*], [*Services*], [*Severity*],
  [foxtrot-06], [ssh 22], [High],
  [golf-07], [imap 143 \\ vnc 5900], [Medium],
  [hotel-08], [ldap 389], [Low],
  [india-09], [mssql 1433 \\ pop3 110], [Low],
  [juliet-10], [ftp 21], [Info],
)
"""


NOTES_BETWEEN_TABLES = """
#set page(width: 21cm, height: 29.7cm, margin: 2cm)
#set text(size: 10pt)
#let hosts(..rows) = table(columns: (3cm, 6cm, 1fr), stroke: none,
  fill: (_, y) => if y == 0 { rgb("#334455") } else if calc.even(y) { luma(225) },
  table.header([#text(fill: white)[Host]], [#text(fill: white)[Service]], [#text(fill: white)[Notes]]),
  ..rows.pos().flatten())
= Network
#hosts(([10.0.0.5], [smb 445], [signing off]), ([10.0.0.6], [rdp 3389], [nla on]), ([10.0.0.7], [http 80], [default page]))
#align(center)[Table 1: Hosts]
=== Next
A paragraph of body prose follows the table and runs over two lines so that it is wide enough to cross every column of the tables above and below it.
#hosts(([10.0.0.8], [ssh 22], [keys only]), ([10.0.0.9], [ftp 21], [anon]), ([10.0.0.10], [dns 53], [axfr]))
Short note.
#hosts(([10.0.0.11], [smtp 25], [relay]), ([10.0.0.12], [smb 445 \\ rdp 3389], [two lines]), ([10.0.0.13], [ldap 389], [anon bind]), ([10.0.0.14], [http 8080], [tomcat]))
After the last table.
"""

NOTES_BETWEEN_TABLES_OF_COLUMNS = """
#set page(width: 21cm, height: 29.7cm, margin: 2cm)
#set text(size: 10pt)
#let hosts(cols: (3cm, 6cm, 1fr), ..rows) = table(columns: cols, stroke: none,
  fill: (_, y) => if y == 0 { rgb("#334455") } else if calc.even(y) { luma(225) },
  table.header(..range(cols.len()).map(i => text(fill: white)[Col #i])),
  ..rows.pos().flatten())
== 10.0.0.5 (DC01)
#hosts(([kerberos], [88/tcp], [open]), ([ldap], [389/tcp], [anon bind]), ([smb], [445/tcp], [signing off]))
== 10.0.0.6 (WEB01)
#hosts(([http], [80/tcp], [IIS 10]), ([https], [443/tcp], [weak TLS]), ([rdp], [3389/tcp], [nla on]))
- Short bullet.
#hosts(cols: (3cm, 6cm, 2cm, 1fr), ([ssh], [22/tcp], [x], [keys]), ([ftp], [21/tcp], [y], [anon]), ([dns], [53/tcp], [z], [axfr]))
*Remediation:*
#hosts(([a], [b], [c]))
"""


def stripe_table_pdf(path, widths, rows=8, left=64.8, top=80.0, pitch=16.0):
    """A host index the way typst draws one: a dark header row of filled cells, then beige stripes.

    Every cell is its own filled rectangle touching its neighbour, the
    stripes alternate with unfilled rows, and the numbers sit against each
    column's right edge, so the gap between the words is not the gap between
    the cells (the fourth column here is 34 pt wide but its words begin 106 pt
    after the third column's).
    """
    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    page.insert_text((left, 60.0), "Index of hosts", fontsize=12)
    edges = [left]
    for width in widths:
        edges.append(edges[-1] + width)
    labels = ["Host", "Address", "Open ports", "Open", "Findings"]
    for row in range(rows + 1):
        y = top + pitch * row
        height = pitch
        header = row == 0
        shaded = header or row % 2 == 0
        for column in range(len(widths)):
            if shaded:
                page.draw_rect(pymupdf.Rect(edges[column], y, edges[column + 1], y + height), color=None,
                               fill=(0.07, 0.19, 0.23) if header else (0.96, 0.94, 0.9))
            word = labels[column] if header else ("mcd-fs%d" % row, "10.10.5.%d" % (140 + row), "1433, 3389",
                                                  str(row), str(row % 4))[column]
            x = edges[column] + 6.0
            if column >= 3:
                x = edges[column + 1] - 6.0 - pymupdf.get_text_length(word, fontsize=9)
            page.insert_text((x, y + 11.0), word, fontsize=9, color=(1, 1, 1) if header else (0, 0, 0))
    page.insert_text((left, top + pitch * (rows + 1) + 24.0), "Hosts with open findings are listed first.",
                     fontsize=10)
    doc.save(path)
    doc.close()


def table_column_widths(docx_path, needle):
    """Each cell's width in points along the first row of the table that holds `needle`."""
    for table in Document(docx_path).tables:
        if any(needle in cell.text for row in table.rows for cell in row.cells):
            row = table.rows[0]
            return [int(cell._tc.tcPr.tcW.w) / 20.0 for cell in row.cells]
    return None


def table_rows_and_cell(docx_path, needle, inside):
    """(rows in the table that holds `needle`, the text of the cell in it that holds `inside`)."""
    for table in Document(docx_path).tables:
        if any(needle in cell.text for row in table.rows for cell in row.cells):
            cells = [cell.text for row in table.rows for cell in row.cells if inside in cell.text]
            return len(table.rows), cells[0] if cells else None
    return None, None


def run_cell_edge_checks():
    """A borderless table's columns follow the PDF's cell fills, not the gaps between its words."""
    from types import SimpleNamespace
    passed = True

    # The edges a row of touching cells proves, and what does not count.
    row = [(60.0, 100.0, 110.0, 116.0), (110.0, 100.0, 180.0, 116.0), (180.0, 100.0, 300.0, 116.0)]
    stripe = [(60.0, 132.0, 110.0, 148.0), (110.0, 132.0, 180.0, 148.0), (180.0, 132.0, 300.0, 148.0)]
    edges = convert.abutting_fill_edges(row + stripe)
    passed &= check([round(edge["x"], 1) for edge in edges] == [110.0, 180.0] and [edge["rows"] for edge in edges] == [2, 2],
                    "cell edges: two touching rows vote twice for each edge they share, and no other",
                    [(edge["x"], edge["rows"]) for edge in edges])
    passed &= check(convert.abutting_fill_edges([(60.0, 100.0, 110.0, 116.0), (114.0, 100.0, 180.0, 116.0)]) == [],
                    "cell edges: two boxes with a gutter between them prove no edge")
    passed &= check(convert.abutting_fill_edges([(60.0, 100.0, 110.0, 116.0), (60.0, 116.0, 110.0, 132.0)]) == [],
                    "cell edges: boxes stacked in one column prove no edge")
    passed &= check(convert.abutting_fill_edges([(60.0, 100.0, 110.0, 116.0), (110.0, 116.0, 180.0, 132.0)]) == [],
                    "cell edges: boxes that meet only at a corner prove no edge")
    near = convert.abutting_fill_edges([(60.0, 100.0, 110.0, 116.0), (110.1, 100.0, 180.0, 116.0),
                                        (60.0, 132.0, 110.2, 148.0), (110.0, 132.0, 180.0, 148.0)])
    passed &= check(len(near) == 1 and near[0]["rows"] == 2,
                    "cell edges: edges a tenth of a point apart are one edge", near)

    # The files the search looks in must not hide a pair a bucket boundary falls between.
    passed &= check(len(convert.touching_cells([(60.0, 100.0, 110.25, 116.0), (110.75, 100.0, 180.0, 116.0)])) == 1
                    and not convert.touching_cells([(60.0, 100.0, 110.25, 116.0), (110.8, 100.0, 180.0, 116.0)]),
                    "cell edges: boxes half a point apart touch, a hair more does not")

    # Which free border takes which edge.
    class Border:
        def __init__(self, kind, low, high, finalized=False):
            self.border_type, self.LRange, self.URange, self.finalized = kind, low, high, finalized
            self.pinned = None
            self._boundary = (SimpleNamespace(LRange=60.0, URange=98.0), SimpleNamespace(LRange=700.0, URange=720.0))

        @property
        def is_vertical(self):
            return "V" in self.border_type

        def is_valid(self, value):
            return self.LRange - 1.0 <= value <= self.URange + 1.0

        def get_boundary_borders(self):
            return self._boundary

        def finalize_by_value(self, value):
            self.finalized, self.pinned = True, value

    class Borders:
        def __init__(self, instances):
            self._instances = instances

    def fill(box, parts=None):
        shape = SimpleNamespace(bbox=box)
        if parts:
            shape._btct_parts = parts
        return shape

    fused = [fill((60.0, 100.0, 300.0, 116.0), row), fill((60.0, 132.0, 300.0, 148.0), stripe)]
    first, second = Border("VI", 105.0, 120.0), Border("VI", 170.0, 200.0)
    none_in_gap, done, outer = Border("VI", 320.0, 340.0), Border("VI", 105.0, 120.0, True), Border("VL", 105.0, 120.0)
    wide = Border("VI", 100.0, 200.0)
    convert.place_borders_on_cell_edges(Borders([first, second, none_in_gap, done, outer]), fused)
    passed &= check(first.pinned == 110.0 and second.pinned == 180.0,
                    "cell edges: a free inner border takes the edge inside its gap", (first.pinned, second.pinned))
    passed &= check(none_in_gap.pinned is None and done.pinned is None and outer.pinned is None,
                    "cell edges: a gap with no edge, a border already placed and an outer border are left alone")
    # Two edges in one gap: the one more rows vote for.
    three = fused + [fill((60.0, 164.0, 300.0, 180.0), [(60.0, 164.0, 110.0, 180.0), (110.0, 164.0, 180.0, 180.0),
                                                       (180.0, 164.0, 300.0, 180.0)])]
    convert.place_borders_on_cell_edges(Borders([wide]), three)
    passed &= check(wide.pinned in (110.0, 180.0), "cell edges: a gap holding two edges takes one of them", wide.pinned)
    busy = fused + [fill((60.0, 164.0, 180.0, 180.0), [(60.0, 164.0, 110.0, 180.0), (110.0, 164.0, 180.0, 180.0)])]
    uneven = Border("VI", 100.0, 200.0)
    convert.place_borders_on_cell_edges(Borders([uneven]), busy)
    passed &= check(uneven.pinned == 110.0, "cell edges: the edge more rows vote for wins the gap", uneven.pinned)

    # A row border between two lines of one cell is a line break, not a row.
    class RowBorder:
        border_type, is_vertical, finalized = "HI", False, False

        def __init__(self, low, high, left, right, reference=False):
            self.LRange, self.URange, self.is_reference = low, high, reference
            self._boundary = (SimpleNamespace(URange=left), SimpleNamespace(LRange=right))

        def get_boundary_borders(self):
            return self._boundary

    cells = [fill((60.0, 100.0, 300.0, 148.0), [(60.0, 100.0, 110.0, 148.0), (110.0, 100.0, 180.0, 148.0),
                                                (180.0, 100.0, 300.0, 148.0)])]
    inside = RowBorder(122.0, 124.0, 112.0, 178.0)
    outside = RowBorder(160.0, 162.0, 112.0, 178.0)
    wider = RowBorder(122.0, 124.0, 112.0, 250.0)
    straddles = RowBorder(96.0, 104.0, 112.0, 178.0)
    already = RowBorder(122.0, 124.0, 112.0, 178.0, reference=True)
    convert.place_borders_on_cell_edges(Borders([inside, outside, wider, straddles, already]), cells)
    passed &= check(inside.is_reference and already.is_reference,
                    "cell edges: a gap wholly inside one cell of a row of cells stops being a row boundary")
    passed &= check(not outside.is_reference and not wider.is_reference and not straddles.is_reference,
                    "cell edges: a gap below the cells, across two of them, or leaving the cell stays a row boundary")
    one_box = RowBorder(122.0, 124.0, 112.0, 178.0)
    convert.place_borders_on_cell_edges(Borders([one_box]), [fill((60.0, 100.0, 300.0, 148.0))])
    passed &= check(not one_box.is_reference,
                    "cell edges: lines inside a lone rectangle, which may be a panel of rows, keep their rows")

    # The table's own sides: the edge most rows of cells end on, not the first fill met.
    banner = [(0.0, 74.0, 76.8, 166.0), (76.8, 74.0, 612.0, 166.0)]
    glossary = [[(64.8, y, 172.8, y + 27.0), (172.8, y, 439.2, y + 27.0)] for y in (200.0, 254.0, 308.0)]
    fused_rows = [fill((0.0, 74.0, 612.0, 166.0), banner)]
    fused_rows += [fill((64.8, rows[0][1], 439.2, rows[0][3]), rows) for rows in glossary]
    right, left = Border("VR", 430.0, 620.0), Border("VL", 0.0, 72.0)
    convert.place_borders_on_cell_edges(Borders([right, left]), fused_rows)
    passed &= check(right.pinned == 439.2 and left.pinned == 64.8,
                    "cell edges: a band to the page edge above a table does not set the table's sides",
                    (left.pinned, right.pinned))
    lone_side = Border("VR", 430.0, 620.0)
    convert.place_borders_on_cell_edges(Borders([lone_side]), fused_rows[:2])
    passed &= check(lone_side.pinned is None,
                    "cell edges: a side one row of cells ends on, against one other, is left to pdf2docx")
    already = Border("VR", 430.0, 620.0, True)
    convert.place_borders_on_cell_edges(Borders([already]), fused_rows)
    passed &= check(already.pinned is None, "cell edges: a side ruled lines already placed is left alone")

    # A single rectangle, or fills that do not touch, change nothing.
    lone = Border("VI", 105.0, 120.0)
    convert.place_borders_on_cell_edges(Borders([lone]), [fill((60.0, 100.0, 300.0, 116.0))])
    passed &= check(lone.pinned is None, "cell edges: one rectangle across the table proves no inner edge")

    # The merge keeps what it fused, copied before the merge moves a box.
    class Piece:
        is_determined = False

        def __init__(self, color, box):
            self.color, self.bbox = color, box

    def fuse(shapes):
        whole = shapes[0]
        whole.bbox = (shapes[0].bbox[0], shapes[0].bbox[1], shapes[-1].bbox[2], shapes[-1].bbox[3])
        return [whole]

    kept = convert.merged_fills_keep_parts(fuse)
    pieces = [Piece(7, b) for b in row]
    merged = kept.__func__(pieces)
    passed &= check(isinstance(kept, staticmethod) and len(merged) == 1 and len(merged[0]._btct_parts) == 3
                    and merged[0]._btct_parts[0] == (60.0, 100.0, 110.0, 116.0),
                    "cell edges: a merged fill carries the boxes it was made from, first one as drawn",
                    getattr(merged[0], "_btct_parts", None))
    alone = kept.__func__([Piece(7, row[0])])
    passed &= check(not hasattr(alone[0], "_btct_parts"), "cell edges: an unmerged fill carries nothing")

    # The wrappers: ruled lines place borders first, the fills' edges second, and a failure is counted.
    order = []

    class Recording:
        _instances = []

        def finalize(self, strokes, fills):
            self._finalize_by_strokes(strokes)
            self._finalize_by_strokes(["from the fills"])

        def _finalize_by_strokes(self, strokes):
            order.append(strokes)

    counted = convert.Counter()
    finalize, by_strokes = convert.borders_follow_cell_edges(Recording.finalize, Recording._finalize_by_strokes, counted)
    placed = []
    original_place = convert.place_borders_on_cell_edges
    convert.place_borders_on_cell_edges = lambda borders, fills: placed.append(("placed", len(order)))
    try:
        Recording.finalize, Recording._finalize_by_strokes = finalize, by_strokes
        Recording().finalize("ruled", [])
        passed &= check(order == ["ruled", ["from the fills"]] and placed == [("placed", 1)],
                        "cell edges: the fills' edges are placed after the ruled lines and before the fills' own pass",
                        (order, placed))
        convert.place_borders_on_cell_edges = lambda borders, fills: (_ for _ in ()).throw(AttributeError("bbox"))
        order.clear()
        Recording().finalize("ruled", [])
        passed &= check(order == ["ruled", ["from the fills"]] and counted["cell-edges-fallthrough"] == 1,
                        "cell edges: a fill pdf2docx no longer describes is counted and the step carries on",
                        (order, dict(counted)))
    finally:
        convert.place_borders_on_cell_edges = original_place

    # Which rows of text carry on a cell of the row above: the second line of a cell.
    zebra = [(56.7, 89.7, 141.7, 119.4), (141.7, 89.7, 311.8, 119.4), (311.8, 89.7, 396.9, 119.4),
             (56.7, 119.4, 141.7, 136.0), (141.7, 119.4, 311.8, 136.0), (311.8, 119.4, 396.9, 136.0)]
    cells = convert.row_cell_boxes(zebra)
    passed &= check(sorted(cells) == sorted(zebra),
                    "cell rows: every cell of a row of equal-height cells is a cell", len(cells))
    sidebar = [(56.7, 89.7, 141.7, 400.0), (141.7, 89.7, 396.9, 150.0)]
    passed &= check(convert.row_cell_boxes(sidebar) == [],
                    "cell rows: two boxes side by side that are not the same height are not cells")
    passed &= check(convert.row_cell_boxes([(56.7, 89.7, 141.7, 119.4), (146.7, 89.7, 311.8, 119.4)]) == []
                    and convert.row_cell_boxes([(56.7, 89.7, 141.7, 119.4)]) == [],
                    "cell rows: a box with no neighbour touching it is not a cell")
    # Stripes with no fill on every other row: the rows between are bounded by the filled cells around them.
    columns = [(56.7, 141.7), (141.7, 311.8), (311.8, 396.9)]
    striped = [(x0, y0, x1, y1) for y0, y1 in ((73.1, 89.7), (119.4, 136.0), (165.6, 182.2)) for x0, x1 in columns]
    found = sorted(convert.row_cell_boxes(striped))
    gaps = sorted((x0, 89.7, x1, 119.4) for x0, x1 in columns) + sorted((x0, 136.0, x1, 165.6) for x0, x1 in columns)
    passed &= check(all(box in found for box in gaps) and len(found) == len(striped) + len(gaps),
                    "cell rows: the unfilled rows between filled stripes are cells too", len(found))
    one_column = [(56.7, 73.1, 141.7, 89.7), (56.7, 119.4, 141.7, 136.0)]
    passed &= check(convert.unfilled_row_cells(one_column) == [],
                    "cell rows: a gap in a single column is not a row of cells")
    split = [(56.7, 73.1, 141.7, 89.7), (141.7, 73.1, 311.8, 89.7),
             (56.7, 119.4, 141.7, 136.0), (141.7, 125.0, 311.8, 140.0)]
    passed &= check(convert.unfilled_row_cells(split) == [],
                    "cell rows: gaps that are not the same rows in two columns are not a row of cells")
    host, service, level = (61.7, 92.4, 91.5, 103.8), (146.7, 92.4, 179.9, 103.8), (316.8, 92.4, 350.7, 103.8)
    second = (146.7, 105.4, 181.7, 116.8)
    third = (146.7, 118.0, 181.7, 129.4)
    next_host = (61.7, 122.0, 91.5, 133.4)
    passed &= check(convert.rows_continuing_cells([[host, service, level], [second]], cells) == {1},
                    "cell rows: a lone line under a row of a table, inside the same cell, carries the cell on")
    # Three lines in one cell: the cell is tall enough for all of them.
    tall = convert.row_cell_boxes([(56.7, 89.7, 141.7, 132.0), (141.7, 89.7, 311.8, 132.0), (311.8, 89.7, 396.9, 132.0)])
    passed &= check(convert.rows_continuing_cells([[host, service, level], [second], [third]], tall) == {1, 2},
                    "cell rows: a third line carries on from the second")
    passed &= check(convert.rows_continuing_cells([[second], [third]], tall) == set(),
                    "cell rows: lines of text in a box with no row of a table above them are flowing text")
    passed &= check(convert.rows_continuing_cells([[host, service, level], [(146.7, 140.0, 181.7, 151.4)]], cells) == set(),
                    "cell rows: a lone line outside the cell stays the flowing text it was")
    passed &= check(convert.rows_continuing_cells([[host, service, level], [second, (200.0, 105.4, 250.0, 116.8)]], cells)
                    == set(),
                    "cell rows: a row of two blocks is its own row of a table and is left alone")
    passed &= check(convert.rows_continuing_cells([[host, service, level], [next_host]], cells) == set(),
                    "cell rows: the next row's first cell is not the cell above")
    # A card with an accent bar beside it is a pair of equal-height boxes, and the line under two pictures is in the card.
    card = convert.row_cell_boxes([(64.0, 265.2, 71.2, 651.3), (71.2, 265.2, 548.0, 651.3)])
    pictures = [(64.0, 416.6, 302.0, 550.5), (310.0, 416.6, 548.0, 550.5)]
    passed &= check(len(card) == 2 and convert.rows_continuing_cells([pictures, [(213.7, 553.8, 398.3, 566.4)]], card) == set(),
                    "cell rows: a row with one block in a card's cell is not a row of a table")
    # The note between two tables of the same columns lies in the gap that was taken for an unfilled row, with
    # the last row of the table above it (dns, 135.4). It stands on the column's edge (56.7), the cell text 5 pt
    # in, and 12.2 pt under that row.
    notes = convert.row_cell_boxes([(56.7, 89.7, 141.7, 119.4), (141.7, 89.7, 311.8, 119.4), (311.8, 89.7, 396.9, 119.4),
                                    (56.7, 165.0, 141.7, 181.0), (141.7, 165.0, 311.8, 181.0), (311.8, 165.0, 396.9, 181.0)])
    last_row = [(61.7, 124.0, 91.5, 135.4), (146.7, 124.0, 179.9, 135.4), (316.8, 124.0, 350.7, 135.4)]
    passed &= check(len(notes) == 9 and convert.rows_continuing_cells([last_row, [(61.7, 137.0, 100.0, 148.4)]], notes) == {1},
                    "cell rows: a line on the text's edge a line gap under the row above is that cell's next line")
    passed &= check(convert.rows_continuing_cells([last_row, [(56.7, 137.0, 100.0, 148.4)]], notes) == set(),
                    "cell rows: a line on the column's edge, 5 pt left of the cell's text, is not that cell's next line")
    # A gap of 12.2 pt (0.9 of a line) is also a paragraph break under par(spacing: 1.8em), so the gap alone
    # does not refuse it; the edge does, above. A line two lines below is not the cell's next line.
    passed &= check(convert.rows_continuing_cells([last_row, [(61.7, 147.6, 100.0, 159.0)]], notes) == {1},
                    "cell rows: a line on the cell's text edge 12.2 pt below (a wide paragraph break) carries on")
    passed &= check(convert.rows_continuing_cells([last_row, [(61.7, 163.0, 100.0, 174.4)]], notes) == set(),
                    "cell rows: a line more than a line and a half below is not that cell's next line")
    # A paragraph break inside a cell (a half line) still carries on; the edge, the right edge or the centre will do.
    passed &= check(convert.rows_continuing_cells([last_row, [(146.7, 141.4, 190.0, 152.8)]], notes) == {1},
                    "cell rows: a paragraph break inside a cell (6 pt, half a line) still carries on")
    # A cell's next line may be right-aligned, centred or indented under the line above (a reference, a
    # score, a nested list): only a line that starts left of the cell's text, on the column's edge, is refused.
    passed &= check(convert.rows_continuing_cells([last_row, [(165.0, 137.0, 179.9, 148.4)]], notes) == {1}
                    and convert.rows_continuing_cells([last_row, [(155.0, 137.0, 171.6, 148.4)]], notes) == {1}
                    and convert.rows_continuing_cells([last_row, [(160.0, 137.0, 190.0, 148.4)]], notes) == {1},
                    "cell rows: a right-aligned, centred or indented next line carries on")
    passed &= check(convert.rows_continuing_cells([last_row, [(144.0, 137.0, 190.0, 148.4)]], notes) == set(),
                    "cell rows: a line that starts left of the cell's text above is not its next line")

    # The walk's marking is installed for one call and taken off again, and a failure is counted.
    class Row(list):
        def is_flow_layout(self, *args, **kwargs):
            return len(self) <= 1

    class Walk:
        def __init__(self):
            self.rows = [Row([SimpleNamespace(bbox=box) for box in row]) for row in ([host, service, level], [second])]

        def group_by_rows(self, *args, **kwargs):
            return self.rows

    class Fill:
        def __init__(self, box):
            self.bbox = box

    layout = SimpleNamespace(_blocks=Walk(), _shapes=SimpleNamespace(table_fillings=[Fill(box) for box in zebra]))
    undo = convert.carry_cells_across_rows(layout)
    rows = layout._blocks.group_by_rows()
    passed &= check(rows[1].is_flow_layout() is False and rows[0].is_flow_layout() is False,
                    "cell rows: the walk sees the second line of a cell as part of a table")
    undo()
    passed &= check("group_by_rows" not in layout._blocks.__dict__,
                    "cell rows: the marking is taken off after the call")
    plain = SimpleNamespace(_blocks=Walk(), _shapes=SimpleNamespace(table_fillings=[Fill(zebra[0])]))
    convert.carry_cells_across_rows(plain)()
    passed &= check("group_by_rows" not in plain._blocks.__dict__,
                    "cell rows: a page with no cells is left to pdf2docx")
    lost = convert.Counter()
    convert.carry_cells_across_rows(SimpleNamespace(_blocks=Walk()), lost)()
    passed &= check(lost["cell-rows-fallthrough"] == 1,
                    "cell rows: a layout pdf2docx no longer describes is counted and the walk carries on", dict(lost))

    # The whole thing on a real page, converted both ways.
    workdir = tempfile.mkdtemp(prefix="btct-cell-edge-test-")
    try:
        widths = [52.1, 65.5, 174.6, 34.3, 47.9]
        pdf = os.path.join(workdir, "stripes.pdf")
        stripe_table_pdf(pdf, widths)
        declined = convert.Counter()
        convert.run_pdf2docx(pdf, os.path.join(workdir, "full.docx"), raw=False, declined=declined)
        convert.run_pdf2docx(pdf, os.path.join(workdir, "raw.docx"), raw=True, declined=convert.Counter())
        full = table_column_widths(os.path.join(workdir, "full.docx"), "mcd-fs1")
        raw = table_column_widths(os.path.join(workdir, "raw.docx"), "mcd-fs1")
        passed &= check(raw is not None and len(raw) == 5 and max(abs(a - b) for a, b in zip(raw, widths)) > 20.0,
                        "cell edges: raw pdf2docx still sets the columns off the words (if not, the patch can go)", raw)
        passed &= check(full is not None and len(full) == 5 and all(abs(a - b) < 1.5 for a, b in zip(full, widths)),
                        "cell edges: the converter's columns are the PDF's cells", full)
        passed &= check("cell-edges-fallthrough" not in declined, "cell edges: nothing fell through", dict(declined))
        # And pdf2docx is left as it was, so the next call is unpatched.
        again = convert.Counter()
        convert.run_pdf2docx(pdf, os.path.join(workdir, "raw2.docx"), raw=True, declined=again)
        passed &= check(table_column_widths(os.path.join(workdir, "raw2.docx"), "mcd-fs1") == raw,
                        "cell edges: the patch is taken off after the call")

        # A shaded stripe of two lines is one row with both lines in one cell. The
        # unshaded stripes between them prove no cell box and keep pdf2docx's rows.
        source = os.path.join(workdir, "wrapped.pdf")
        compile_case(WRAPPED_STRIPES, source)
        convert.run_pdf2docx(source, os.path.join(workdir, "wrapped.docx"), raw=False, declined=convert.Counter())
        convert.run_pdf2docx(source, os.path.join(workdir, "wrapped-raw.docx"), raw=True, declined=convert.Counter())
        rows, cell = table_rows_and_cell(os.path.join(workdir, "wrapped.docx"), "MCD-01", "Rotate")
        raw_rows, raw_cell = table_rows_and_cell(os.path.join(workdir, "wrapped-raw.docx"), "MCD-01", "Rotate")
        passed &= check(raw_cell is not None and "credentials" not in raw_cell and raw_rows == 13,
                        "cell edges: raw pdf2docx still cuts a two-line stripe into two rows (if not, the rule can go)",
                        (raw_rows, raw_cell))
        passed &= check(cell is not None and "credentials now" in cell and rows == 10,
                        "cell edges: a shaded two-line stripe stays one row, both lines in its cell", (rows, cell))
        wrapped_widths = table_column_widths(os.path.join(workdir, "wrapped.docx"), "MCD-01")
        passed &= check(wrapped_widths is not None
                        and all(abs(a - b) < 1.5 for a, b in zip(wrapped_widths, (40.0, 120.0, 50.0, 80.0, 52.0, 35.0))),
                        "cell edges: and the columns are still the PDF's", wrapped_widths)

        # A striped table whose cell has a line break in it, the second line alone in its row of text.
        # pdf2docx closed the table between the two lines, so the second service was a paragraph (or a
        # one-row table) between the host rows.
        source = os.path.join(workdir, "zebra.pdf")
        compile_case(ZEBRA_CELLS, source)
        counted = convert.Counter()
        convert.run_pdf2docx(source, os.path.join(workdir, "zebra.docx"), raw=False, declined=counted)
        convert.run_pdf2docx(source, os.path.join(workdir, "zebra-raw.docx"), raw=True, declined=convert.Counter())
        for name, whole in (("zebra", True), ("zebra-raw", False)):
            rows = {}
            # A second service in a paragraph of its own is outside every host's row, filled stripes
            # or unfilled ones.
            second_lines = {"rdp 3389", "winrm 5985", "http 8080", "vnc 5900", "pop3 110"}
            loose = [p.text.strip() for p in Document(os.path.join(workdir, name + ".docx")).paragraphs
                     if p.text.strip() in second_lines]
            passed &= check(bool(loose) != whole,
                            "zebra cells: no second service is left in a paragraph between the tables" if whole
                            else "zebra cells: raw pdf2docx still leaves second services in paragraphs (if not, the rule can go)",
                            loose)
            for table in Document(os.path.join(workdir, name + ".docx")).tables:
                for row in table.rows:
                    texts = [cell.text for cell in row.cells]
                    for host in ("bravo-02", "delta-04"):
                        if texts and texts[0].startswith(host):
                            rows.setdefault(host, []).append(" ".join(texts))
            together = ("rdp 3389" in " ".join(rows.get("bravo-02", [])) and "winrm 5985" in " ".join(rows.get("bravo-02", []))
                        and "http 8080" in " ".join(rows.get("delta-04", [])))
            passed &= check(together == whole,
                            "zebra cells: every service is in its own host's row" if whole
                            else "zebra cells: raw pdf2docx still puts a second line outside its host's row (if not, the rule can go)",
                            rows)
        passed &= check("cell-rows-fallthrough" not in counted, "zebra cells: nothing fell through", dict(counted))

        # A one-line note between two striped tables of the same columns is flowing text, and the tables stay two.
        for name, source_text, shapes, notes in (
                ("notes", NOTES_BETWEEN_TABLES, [4, 4, 5], ("Short note.", "After the last table.")),
                ("notes-columns", NOTES_BETWEEN_TABLES_OF_COLUMNS, [4, 4, 4, 2], ("Short bullet.", "Remediation:"))):
            source = os.path.join(workdir, name + ".pdf")
            compile_case(source_text, source)
            counted = convert.Counter()
            convert.run_pdf2docx(source, os.path.join(workdir, name + ".docx"), raw=False, declined=counted)
            document = Document(os.path.join(workdir, name + ".docx"))
            in_cells = [note for note in notes for table in document.tables for row in table.rows for cell in row.cells
                        if note in cell.text]
            in_body = [note for note in notes for paragraph in document.paragraphs if note in paragraph.text]
            passed &= check(not in_cells and sorted(in_body) == sorted(notes),
                            "note cells: %s stay paragraphs between the tables" % " and ".join(notes), (in_cells, in_body))
            passed &= check([len(table.rows) for table in document.tables] == shapes,
                            "note cells: the tables around %s are not joined" % notes[0], [len(t.rows) for t in document.tables])
            passed &= check("cell-rows-fallthrough" not in counted, "note cells: nothing fell through", dict(counted))
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return passed


def run_cell_checks():
    """Text inside a table, paired with the PDF by geometry rather than order."""
    passed = True

    # The rule the whole pass rests on, measured on Word's own render of the
    # reference report: a cell's content starts below its top border and its
    # top margin, and the first paragraph's spacing before is not dropped.
    document = Document()
    table = exact_table(document, [20.0], [100.0, 200.0])
    cell = table.rows[0].cells[0]
    passed &= check(convert.cell_content_top(cell._tc, 50.0) == 50.0,
                    "cells: a cell with no border and no margin starts at its row's top")
    cell_top_border(cell, 0.75)
    passed &= check(abs(convert.cell_content_top(cell._tc, 50.0) - 50.75) < 0.01,
                    "cells: a cell's top border insets its content",
                    convert.cell_content_top(cell._tc, 50.0))

    # The grid, which is the PDF's own: exact row heights and cell widths,
    # offset by the table's indent.
    grid = convert.table_grid(table._tbl, 100.0, 54.0)
    passed &= check(grid is not None and len(grid) == 1 and len(grid[0]["cells"]) == 2,
                    "cells: the grid is one rectangle per cell", grid and len(grid[0]["cells"]))
    passed &= check(abs(grid[0]["y1"] - 120.0) < 0.01 and abs(grid[0]["cells"][1]["x0"] - 154.0) < 0.01
                    and abs(grid[0]["cells"][1]["x1"] - 354.0) < 0.01,
                    "cells: rows and columns are measured from the table's corner",
                    grid and (grid[0]["y1"], grid[0]["cells"][1]["x0"]))
    indented = exact_table(Document(), [20.0], [100.0], indent=9.0)
    passed &= check(abs(convert.table_grid(indented._tbl, 0.0, 54.0)[0]["cells"][0]["x0"] - 63.0) < 0.01,
                    "cells: the table's indent moves the whole grid")
    loose = exact_table(Document(), [20.0], [100.0])
    loose.rows[0]._tr.find(convert.qn("w:trPr")).find(
        convert.qn("w:trHeight")).set(convert.qn("w:hRule"), "atLeast")
    passed &= check(convert.table_grid(loose._tbl, 0.0, 0.0) is None,
                    "cells: a row Word is free to size has no grid")

    # A cell whose whole text is too short to pair on is paired by being the
    # only thing in the cell. The number of a numbered heading squashes to "2",
    # which occurs all over a report.
    one = exact_table(Document(), [40.1], [24.8, 480.0])
    fill_cell(one.rows[0].cells[0], "2.", before=0.0, line=25.2)
    fill_cell(one.rows[0].cells[1], "INTRODUCTION", before=0.0, line=23.65)
    boxes = convert.table_grid(one._tbl, 69.4, 54.0)[0]["cells"]
    single = [pdf_line("2.", 88.29, 54.0, 66.0, ascent=12.6)]
    passed &= check(convert.sole_pairing(boxes[0]["cell"], single, 0) == (0, 1),
                    "cells: one paragraph and one line in a cell are each other's",
                    convert.sole_pairing(boxes[0]["cell"], single, 0))
    passed &= check(convert.sole_pairing(boxes[0]["cell"], single, 1) is None,
                    "cells: not when something in the cell is already placed")
    passed &= check(convert.sole_pairing(boxes[0]["cell"], single + single, 0) is None,
                    "cells: nor when two lines fell into the cell")
    two_paragraphs = exact_table(Document(), [40.1], [100.0])
    fill_cell(two_paragraphs.rows[0].cells[0], "first", before=0.0, line=12.0)
    second = two_paragraphs.rows[0].cells[0].add_paragraph()
    convert.set_spacing(second._p, before=0.0, after=0.0, line=12.0)
    second.add_run("second")
    passed &= check(convert.sole_pairing(two_paragraphs.rows[0].cells[0]._tc, single, 0) is None,
                    "cells: nor when the cell holds two paragraphs of text")

    # And end to end: the number lands on the title's baseline, which is what
    # the reader sees, rather than 1.25 pt under it.
    lines = [pdf_line("2.", 88.29, 54.0, 66.0, ascent=12.6),
             pdf_line("INTRODUCTION", 88.29, 87.35, 300.0, ascent=12.6)]
    moved, undo = convert.align_table_cells(one._tbl, 69.4, 54.0, lines)
    numbered = convert.paragraph_metrics(one.rows[0].cells[0].paragraphs[0]._p)
    passed &= check(moved >= 1 and numbered is not None
                    and abs(69.4 + numbered["before"] + 0.8 * numbered["line"] - 88.29) < 0.05,
                    "cells: a heading's number lands on the title's own baseline",
                    numbered and 69.4 + numbered["before"] + 0.8 * numbered["line"])

    # A cell holding two leadings is cut the way a top-level block is, and the
    # cut is undone if the table is given back.
    card = exact_table(Document(), [60.0], [200.0])
    body = fill_cell(card.rows[0].cells[0], "Heading one", before=0.0, line=15.8)
    for text in ("first body line", "Heading two", "second body line"):
        holder = convert.OxmlElement("w:r")
        holder.append(convert.OxmlElement("w:br"))
        body._p.append(holder)
        body.add_run(text)
    cell_lines = [pdf_line("Heading one", 113.5, 64.0, 150.0),
                  pdf_line("first body line", 127.0, 64.0, 200.0),
                  pdf_line("Heading two", 145.0, 64.0, 150.0),
                  pdf_line("second body line", 158.5, 64.0, 200.0)]
    boxes = convert.table_grid(card._tbl, 100.0, 54.0)[0]["cells"]
    made = []
    convert.cut_cell_pitches(boxes[0]["cell"], convert.merged_lines(cell_lines), made)
    pieces = boxes[0]["cell"].findall(convert.qn("w:p"))
    passed &= check(len(made) == 1 and len(pieces) == 2,
                    "cells: a cell paragraph is cut where the leading changes",
                    (len(made), len(pieces)))
    passed &= check("".join(convert.paragraph_text(p) for p in pieces)
                    == "Heading onefirst body lineHeading twosecond body line",
                    "cells: and not one character moved",
                    [convert.paragraph_text(p) for p in pieces])
    head = convert.paragraph_metrics(pieces[0])
    tail = convert.paragraph_metrics(pieces[1])
    passed &= check(abs((head["before"] + 2 * head["line"] + max(head["after"], tail["before"])
                         + 2 * tail["line"] + tail["after"]) - (0.0 + 4 * 15.8 + 0.0)) < 0.05,
                    "cells: the pieces are as tall together as the paragraph was",
                    (head, tail))
    convert.rejoin_unplaced(made, set())
    passed &= check(len(boxes[0]["cell"].findall(convert.qn("w:p"))) == 1,
                    "cells: a cut the pass never placed is joined back up",
                    len(boxes[0]["cell"].findall(convert.qn("w:p"))))

    # A list marker and the text beside it are two columns of one row and one
    # line of one cell.
    merged = convert.merged_lines([pdf_line("1.", 100.0, 60.0, 68.0),
                                   pdf_line("{{STEP}}", 100.0, 74.0, 140.0),
                                   pdf_line("and more", 113.0, 60.0, 150.0)])
    passed &= check(len(merged) == 2 and merged[0]["key"] == convert.compare_squash("1.{{STEP}}"),
                    "cells: two columns on one baseline read as one line", [m["key"] for m in merged])

    # End to end. Two cells of one row, each holding text the PDF put on its
    # own baseline, and the pair reads down one column and then the next in
    # the file while the PDF reads across.
    document = Document()
    table = exact_table(document, [34.5], [108.0, 396.0])
    left = fill_cell(table.rows[0].cells[0], "{{DATE}}", before=9.2, line=15.4)
    right = fill_cell(table.rows[0].cells[1], "The description runs to two lines here", before=2.4, line=14.8)
    right._p.find(convert.qn("w:r")).append(convert.OxmlElement("w:br"))
    right.add_run("second line of it")
    lines = [pdf_line("The description runs to two lines here", 263.4, 170.0, 500.0),
             pdf_line("{{DATE}}", 270.8, 60.0, 120.0),
             pdf_line("second line of it", 278.2, 170.0, 300.0)]
    moved, undo = convert.align_table_cells(table._tbl, 251.55, 54.0, lines)
    got_left = convert.paragraph_metrics(left._p)
    got_right = convert.paragraph_metrics(right._p)
    passed &= check(moved == 2 and len(undo) == 2, "cells: both cells of the row moved", (moved, len(undo)))
    passed &= check(abs(251.55 + got_left["before"] + 0.8 * got_left["line"] - 270.8) < 0.05,
                    "cells: the date column lands on the PDF's baseline",
                    got_left and 251.55 + got_left["before"] + 0.8 * got_left["line"])
    passed &= check(abs(251.55 + got_right["before"] + 0.8 * got_right["line"] - 263.4) < 0.05,
                    "cells: and so does the description beside it",
                    got_right and 251.55 + got_right["before"] + 0.8 * got_right["line"])
    passed &= check(abs(got_right["line"] - 14.8) < 0.05,
                    "cells: a cell of several lines is set at the PDF's own pitch", got_right["line"])
    height = table.rows[0]._tr.find(convert.qn("w:trPr")).find(convert.qn("w:trHeight"))
    passed &= check(height.get(convert.qn("w:val")) == "690",
                    "cells: and the row keeps its height, so no rule moves",
                    height.get(convert.qn("w:val")))
    convert.restore(undo)
    passed &= check(abs(convert.paragraph_metrics(left._p)["before"] - 9.2) < 0.05,
                    "cells: the undo puts the spacing back")

    # A line whose centre is in the next column belongs to that cell, whatever
    # the file's reading order says.
    document = Document()
    table = exact_table(document, [20.0], [108.0, 396.0])
    first = fill_cell(table.rows[0].cells[0], "Name", before=1.8, line=15.4)
    fill_cell(table.rows[0].cells[1], "{{NAME}}", before=1.8, line=15.4)
    moved, _ = convert.align_table_cells(
        table._tbl, 100.0, 54.0,
        [pdf_line("{{NAME}}", 112.0, 170.0, 260.0), pdf_line("Name", 112.0, 60.0, 90.0)])
    passed &= check(moved == 2 and abs(100.0 + convert.paragraph_metrics(first._p)["before"]
                                       + 0.8 * convert.paragraph_metrics(first._p)["line"] - 112.0) < 0.05,
                    "cells: a cell takes the line whose centre is inside it", moved)

    # The refusal: a correction that would push a cell's last line past the
    # bottom Word clips an exact row at gives the whole table back.
    document = Document()
    table = exact_table(document, [20.0, 20.0], [200.0])
    kept = fill_cell(table.rows[0].cells[0], "Row one", before=1.0, line=12.0)
    fill_cell(table.rows[1].cells[0], "Row two", before=1.0, line=12.0)
    moved, undo = convert.align_table_cells(
        table._tbl, 100.0, 54.0,
        [pdf_line("Row one", 120.0, 60.0, 120.0), pdf_line("Row two", 131.0, 60.0, 120.0)])
    passed &= check(moved == 0 and undo == []
                    and abs(convert.paragraph_metrics(kept._p)["before"] - 1.0) < 0.05,
                    "cells: a table whose text would be clipped is given back", (moved, undo))

    # A nested table inside a cell is placed inside that cell.
    document = Document()
    outer = exact_table(document, [40.0], [300.0])
    holder = outer.rows[0].cells[0]
    convert.set_spacing(holder.paragraphs[0]._p, before=0.0, after=0.0, line=1.0)
    inner = exact_table(document, [20.0], [150.0, 150.0], indent=6.0)
    holder._tc.append(inner._tbl)
    one = fill_cell(inner.rows[0].cells[0], "Techniques", before=1.9, line=11.7)
    two = fill_cell(inner.rows[0].cells[1], "Mitigations", before=1.9, line=11.7)
    moved, _ = convert.align_table_cells(
        outer._tbl, 100.0, 54.0,
        [pdf_line("Techniques", 113.0, 70.0, 130.0), pdf_line("Mitigations", 113.0, 230.0, 300.0)])
    passed &= check(moved == 2
                    and abs(101.0 + convert.paragraph_metrics(one._p)["before"]
                            + 0.8 * convert.paragraph_metrics(one._p)["line"] - 113.0) < 0.05
                    and abs(convert.paragraph_metrics(two._p)["before"]
                            - convert.paragraph_metrics(one._p)["before"]) < 0.05,
                    "cells: a nested table's own cells are placed too", moved)

    # A cell that opens with a paragraph Word sizes itself (an inline picture's
    # line has no exact height) used to be given up whole. The walk now has no
    # position until it meets a block it can pair with the PDF, which becomes
    # the anchor, and what follows is placed from there (shape-a, page 25: a
    # panel under a screenshot stood 7.9 pt above its rectangle).
    document = Document()
    table = exact_table(document, [200.0], [400.0])
    cell = table.rows[0].cells[0]
    picture = cell.paragraphs[0]
    picture.add_run()          # no spacing of its own: an unmodelled paragraph
    caption = cell.add_paragraph()
    convert.set_spacing(caption._p, before=8.0, after=0.0, line=12.0)
    caption.add_run("Figure caption text")
    body = cell.add_paragraph()
    convert.set_spacing(body._p, before=30.0, after=0.0, line=12.0)
    body.add_run("Body paragraph text here")
    moved, undo = convert.align_table_cells(
        table._tbl, 80.0, 54.0,
        [pdf_line("Figure caption text", 100.0, 60.0, 250.0), pdf_line("Body paragraph text here", 114.0, 60.0, 260.0)])
    passed &= check(moved == 1 and abs(convert.paragraph_metrics(body._p)["before"] - 2.0) < 0.05
                    and convert.paragraph_metrics(caption._p)["before"] == 8.0,
                    "cells: the walk anchors on the first paired block after a paragraph it cannot measure",
                    (moved, convert.paragraph_metrics(body._p)["before"], convert.paragraph_metrics(caption._p)["before"]))

    # A first line too short to pair on ("{") pairs with the cell's next line
    # when that is the same text, so it gets the PDF's place and the line
    # under it keeps its own height instead of shrinking to make up for the
    # gap above the block (page 61 of an 86-page report: "{" 8 pt low).
    document = Document()
    table = exact_table(document, [60.0], [400.0])
    cell = table.rows[0].cells[0]
    brace = fill_cell(cell, "{", before=14.5, line=9.0)
    host = cell.add_paragraph()
    convert.set_spacing(host._p, before=3.7, after=0.0, line=9.0)
    host.add_run('"host": "10.10.7.147",')
    moved, undo = convert.align_table_cells(
        table._tbl, 100.0, 54.0,
        [pdf_line("{", 109.0, 60.0, 66.0, ascent=7.0), pdf_line('"host": "10.10.7.147",', 121.7, 60.0, 200.0, ascent=7.0)])
    brace_metrics, host_metrics = convert.paragraph_metrics(brace._p), convert.paragraph_metrics(host._p)
    passed &= check(abs(brace_metrics["before"] - (109.0 - 0.8 * 9.0 - 100.0)) < 0.05 and host_metrics["line"] == 9.0
                    and abs(host_metrics["before"] - 3.7) < 0.3,
                    "cells: a short first line takes the PDF's place and the line under it keeps its height",
                    (brace_metrics["before"], host_metrics["before"], host_metrics["line"]))
    # Only the very next line: a short paragraph with a different line next in
    # the cell is not paired with a line further down.
    document = Document()
    table = exact_table(document, [60.0], [400.0])
    cell = table.rows[0].cells[0]
    brace = fill_cell(cell, "}", before=14.5, line=9.0)
    moved, undo = convert.align_table_cells(
        table._tbl, 100.0, 54.0,
        [pdf_line("alpha beta gamma", 109.0, 60.0, 160.0, ascent=7.0), pdf_line("}", 121.7, 60.0, 66.0, ascent=7.0)])
    passed &= check(convert.paragraph_metrics(brace._p)["before"] == 14.5,
                    "cells: a short paragraph is not paired with a line further down than the next",
                    convert.paragraph_metrics(brace._p)["before"])

    # A cell of nothing but such paragraphs has no anchor and is left alone.
    document = Document()
    table = exact_table(document, [200.0], [400.0])
    table.rows[0].cells[0].paragraphs[0].add_run()
    stray = table.rows[0].cells[0].add_paragraph()
    convert.set_spacing(stray._p, before=30.0, after=0.0, line=12.0)
    stray.add_run("Only text")
    moved, undo = convert.align_table_cells(table._tbl, 80.0, 54.0, [])
    passed &= check(moved == 0 and convert.paragraph_metrics(stray._p)["before"] == 30.0,
                    "cells: a cell whose text has no line of the PDF to pair with stays as written", moved)
    return passed


def anchor_row(text, base, y0, y1, x0=54.0, x1=300.0):
    """One page row in the shape `page_anchor_rows` hands the vertical pass."""
    return {"key": convert.compare_squash(text), "base": base, "y0": y0, "y1": y1,
            "x0": x0, "x1": x1, "text": text,
            "columns": [{"x0": x0, "x1": x1, "y0": y0, "y1": y1, "base": base, "text": text}]}


def code_document(lines, indent=25.8, before=14.5, line=9.0, after=19.8, shaded=False):
    """A paragraph of body text and then a shaded code block under it."""
    document = Document()
    document.sections[0].top_margin = Pt(72)
    lead = document.add_paragraph()
    convert.set_spacing(lead._p, before=0.0, after=8.3, line=13.5)
    lead.add_run("Step title")
    made = []
    for order, text in enumerate(lines):
        paragraph = document.add_paragraph()
        convert.set_spacing(paragraph._p, before=before if not order else 3.7,
                            after=after if order == len(lines) - 1 else 0.0, line=line)
        paragraph.add_run(text)
        paragraph.paragraph_format.left_indent = Pt(indent)
        if shaded:
            convert.shade_paragraph(paragraph, (0.93, 0.93, 0.93))
        made.append(paragraph._p)
    return document, made


def code_panel(top, height, x0=71.2, x1=548.0, text_x=80.2, bar=7.2, lines=()):
    """A panel as `find_decoration` describes one, with an accent bar."""
    return {
        "fill": (0.93, 0.93, 0.93),
        "rect": pymupdf.Rect(x0, top, x1, top + height),
        "page": 0,
        "bar": {"fill": (0.1, 0.14, 0.22), "width": bar, "space": text_x - x0},
        "rows": [{"x0": text_x, "x1": text_x + 200.0, "y0": base - 7.0, "y1": base + 2.0,
                  "base": base, "text": text,
                  "columns": [{"x0": text_x, "x1": text_x + 200.0, "y0": base - 7.0,
                               "y1": base + 2.0, "base": base, "text": text}]}
                 for text, base in lines],
    }


def run_table_place_checks():
    """Moving a table onto the rules the PDF drew for its own rows."""
    passed = True

    # A rule is one edge at its centre; a filled block is not an edge at all,
    # because Word paints a cell's fill below its top border and matching a
    # row boundary to one took four pages backwards.
    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    page.draw_line((LEFT, 200.0), (RIGHT, 200.0), width=0.8)
    page.draw_rect(pymupdf.Rect(LEFT, 300.0, RIGHT, 400.0), color=None, fill=(0.9, 0.9, 0.9))
    page.draw_line((LEFT, 500.0), (LEFT + 20.0, 500.0), width=0.8)
    edges = convert.page_edges(doc)[0]
    doc.close()
    passed &= check(len(edges) == 1 and abs(edges[0]["y"] - 200.0) < 0.2,
                    "place: a rule is an edge, a filled block and a stub are not",
                    [round(edge["y"], 2) for edge in edges])

    heights = [19.7, 34.6, 34.5, 34.5]
    table = exact_table(Document(), heights, [108.0, 396.0])
    drawn = [231.8, 251.55, 286.06, 320.56, 355.07]
    across = [{"y": y, "x0": 53.9, "x1": 558.1} for y in drawn]
    offset = convert.table_top_offset(across, table._tbl, LEFT, 232.9)
    passed &= check(abs(offset + 1.1) < 0.05,
                    "place: the table moves onto the rules the PDF drew for it", offset)
    passed &= check(convert.table_top_offset(across, table._tbl, LEFT, 231.8) == 0.0,
                    "place: and stays where it is when they already agree")

    # Downwards is never a correction: every error the converter's spacing
    # makes points the same way, so a table is drawn below the PDF's rules
    # and never above them.
    below = [{"y": y + 2.2, "x0": 53.9, "x1": 558.1} for y in drawn]
    passed &= check(convert.table_top_offset(below, table._tbl, LEFT, 232.9) == 0.0,
                    "place: an edge below the table is not a correction")

    # A table whose rows do not add up to the PDF's own cannot be moved: only
    # its top would land, and one boundary is a coincidence.
    stretched = [{"y": y, "x0": 53.9, "x1": 558.1} for y in (231.8, 253.0, 290.0, 330.0)]
    passed &= check(convert.table_top_offset(stretched, table._tbl, LEFT, 232.9) == 0.0,
                    "place: a grid that does not match the PDF's keeps its place")
    # Nor one whose rules belong to something narrower, such as a table
    # nested inside it.
    inside = [{"y": y, "x0": 120.0, "x1": 400.0} for y in drawn]
    passed &= check(convert.table_top_offset(inside, table._tbl, LEFT, 232.9) == 0.0,
                    "place: a rule that does not span the table is not its rule")
    return passed


def small_picture():
    """The bytes of a tiny PNG, for a picture whose size is all that matters."""
    return pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 8, 8), False).tobytes("png")


def picture_paragraph(container, height, before, after=0.0):
    """A paragraph holding one inline picture the way the picture pass writes it: one automatic line."""
    import io
    paragraph = container.add_paragraph()
    paragraph.add_run().add_picture(io.BytesIO(small_picture()), width=Pt(height), height=Pt(height))
    convert.set_spacing(paragraph._p, before=before, after=after)
    spacing = paragraph._p.get_or_add_pPr().find(convert.qn("w:spacing"))
    spacing.set(convert.qn("w:line"), "240")
    spacing.set(convert.qn("w:lineRule"), "auto")
    return paragraph


def exact_paragraph(container, text, before, line=15.0, after=0.0):
    paragraph = container.add_paragraph()
    convert.set_spacing(paragraph._p, before=before, after=after, line=line)
    paragraph.add_run(text)
    return paragraph


def run_picture_walk_checks():
    """The vertical pass counts a picture and keeps it where the converter had it."""
    import io
    passed = True

    document = Document()
    picture = picture_paragraph(document, 60.0, 10.0, 3.0)
    found = convert.picture_line(picture._p)
    passed &= check(found is not None and abs(found[0] - 10.0) < 0.05 and abs(found[1] - 60.0) < 0.05
                    and abs(found[2] - 3.0) < 0.05,
                    "picture: a paragraph of one inline picture is spacing, the picture's height, spacing", found)
    captioned = picture_paragraph(document, 60.0, 0.0)
    captioned.add_run("a caption")
    passed &= check(convert.picture_line(captioned._p) is None,
                    "picture: one with text beside the picture is not modelled")
    two = picture_paragraph(document, 60.0, 0.0)
    two.add_run().add_picture(io.BytesIO(small_picture()), width=Pt(10.0), height=Pt(10.0))
    passed &= check(convert.picture_line(two._p) is None,
                    "picture: nor one holding two pictures")
    fixed = picture_paragraph(document, 60.0, 0.0)
    convert.set_spacing(fixed._p, line=20.0)
    passed &= check(convert.picture_line(fixed._p) is None,
                    "picture: nor one whose line is exact, which is a line of text")
    passed &= check(convert.picture_line(exact_paragraph(document, "words", 0.0)._p) is None,
                    "picture: nor a paragraph of text")

    def build(second_before, picture_before):
        document = Document()
        document.sections[0].top_margin = Pt(72)
        document.sections[0].bottom_margin = Pt(72)
        exact_paragraph(document, "First block here", 0.0)
        exact_paragraph(document, "Second block here", second_before)
        picture_paragraph(document, 60.0, picture_before)
        exact_paragraph(document, "Fourth block here", 2.0)
        return document

    def rows_for(second, fourth):
        return [[anchor_row("First block here", 84.0, 76.0, 87.0),
                 anchor_row("Second block here", second, second - 8.0, second + 3.0),
                 anchor_row("Fourth block here", fourth, fourth - 8.0, fourth + 3.0)]]

    # The block under a picture is placed on its baseline: before this the
    # walk stopped at the picture and left it where the converter had it.
    # The first block ends at 87, the second keeps its place (89 to 104), the
    # picture stands at 114 to 174 and the fourth wants its baseline at 191.
    document = build(2.0, 10.0)
    convert.align_vertical_rhythm(document, rows_for(101.0, 191.0), 792.0)
    fourth = convert.paragraph_metrics(document.paragraphs[3]._p)
    passed &= check(fourth is not None and abs(fourth["before"] - 5.0) < 0.3,
                    "picture: the block under it is placed on its own baseline", fourth and fourth["before"])

    # A picture keeps the place it had when the block above it moved: the
    # second block comes 8 pt lower, so the picture's spacing loses 8 pt.
    document = build(2.0, 10.0)
    convert.align_vertical_rhythm(document, rows_for(109.0, 191.0), 792.0)
    second = convert.paragraph_metrics(document.paragraphs[1]._p)
    held = convert.picture_line(document.paragraphs[2]._p)
    fourth = convert.paragraph_metrics(document.paragraphs[3]._p)
    passed &= check(second is not None and abs(second["before"] - 10.0) < 0.3,
                    "picture: the block above moved down 8 pt", second and second["before"])
    passed &= check(held is not None and abs(held[0] - 2.0) < 0.3,
                    "picture: and the picture's spacing before took the 8 pt back", held)
    passed &= check(fourth is not None and abs(fourth["before"] - 5.0) < 0.3,
                    "picture: so what is under it is measured from where it stands", fourth and fourth["before"])
    spacing = document.paragraphs[2]._p.find(convert.qn("w:pPr")).find(convert.qn("w:spacing"))
    passed &= check(spacing.get(convert.qn("w:lineRule")) == "auto" and spacing.get(convert.qn("w:line")) == "240",
                    "picture: its line is still the automatic one, which is what holds the picture")

    # A page that would run past its bottom margin gives all of it back.
    document = build(2.0, 10.0)
    convert.align_vertical_rhythm(document, rows_for(109.0, 900.0), 792.0)
    held = convert.picture_line(document.paragraphs[2]._p)
    spacing = document.paragraphs[2]._p.find(convert.qn("w:pPr")).find(convert.qn("w:spacing"))
    passed &= check(held is not None and abs(held[0] - 10.0) < 0.05 and spacing.get(convert.qn("w:lineRule")) == "auto",
                    "picture: a refused page gives the picture its spacing back", held)

    # Inside a cell the walk is the same: it used to stop at the picture,
    # which left a panel and the table under it where the converter's flow
    # had them.
    table = exact_table(Document(), [300.0], [400.0])
    cell = table.rows[0].cells[0]
    cell._tc.remove(cell.paragraphs[0]._p)
    exact_paragraph(cell, "Intro line text", 0.0)
    exact_paragraph(cell, "Lead line text", 2.0)
    picture_paragraph(cell, 60.0, 10.0)
    exact_paragraph(cell, "Tail line text", 2.0)
    lines = [pdf_line("Intro line text", 112.0, 60.0, 200.0, ascent=12.0),
             pdf_line("Lead line text", 137.0, 60.0, 200.0, ascent=12.0),
             pdf_line("Tail line text", 215.0, 60.0, 200.0, ascent=12.0)]
    moved, undo = convert.align_table_cells(table._tbl, 100.0, 54.0, lines)
    parts = cell._tc.findall(convert.qn("w:p"))
    lead_before = convert.paragraph_metrics(parts[1])["before"]
    held = convert.picture_line(parts[2])
    tail = convert.paragraph_metrics(parts[3])
    # Lead moves from 2 to 10 (100 + 15 + 10 + 12 = 137), the picture takes
    # the 8 back and stands at 142 to 202, and the tail then lands
    # 215 - 12 - 202 = 1 under it.
    passed &= check(abs(lead_before - 10.0) < 0.3 and held is not None and abs(held[0] - 2.0) < 0.3
                    and tail is not None and abs(tail["before"] - 1.0) < 0.3,
                    "picture: a cell's walk counts the picture and keeps it where it was",
                    (lead_before, held, tail and tail["before"]))
    convert.restore(undo)
    held = convert.picture_line(cell._tc.findall(convert.qn("w:p"))[2])
    passed &= check(held is not None and abs(held[0] - 10.0) < 0.05,
                    "picture: and the table given back has its picture spacing back", held)
    return passed


def run_table_box_checks():
    """A table that is a rectangle the PDF drew belongs on that rectangle."""
    passed = True

    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    page.draw_rect(pymupdf.Rect(79.2, 151.45, 532.8, 475.45), color=(0, 0, 0), fill=(0.9, 0.9, 0.9))
    page.draw_line((LEFT, 600.0), (RIGHT, 600.0), width=0.8)
    page.draw_rect(pymupdf.Rect(LEFT, 650.0, LEFT + 30.0, 700.0), color=(0, 0, 0))
    boxes = convert.page_boxes(doc)[0]
    doc.close()
    passed &= check(len(boxes) == 1 and abs(boxes[0]["height"] - 324.0) < 0.5,
                    "box: a drawn rectangle is a box, a rule and a stub are not",
                    [(round(b["x0"], 1), round(b["y0"], 2), round(b["height"], 2)) for b in boxes])

    table = exact_table(Document(), [324.0], [453.6], indent=25.2)
    passed &= check(abs(convert.table_box_offset(boxes, table._tbl, LEFT, 155.28) + 3.83) < 0.05,
                    "box: the table moves down onto the rectangle it is",
                    convert.table_box_offset(boxes, table._tbl, LEFT, 155.28))
    passed &= check(abs(convert.table_box_offset(boxes, table._tbl, LEFT, 147.6) - 3.85) < 0.05,
                    "box: and up onto it, which the fit against rules may not",
                    convert.table_box_offset(boxes, table._tbl, LEFT, 147.6))
    passed &= check(convert.table_box_offset(boxes, table._tbl, LEFT, 151.45) == 0.0,
                    "box: a table already on its rectangle keeps its place")

    taller = exact_table(Document(), [340.0], [453.6], indent=25.2)
    passed &= check(convert.table_box_offset(boxes, taller._tbl, LEFT, 155.28) == 0.0,
                    "box: a table of another height is not that rectangle")
    narrow = exact_table(Document(), [324.0], [300.0], indent=25.2)
    passed &= check(convert.table_box_offset(boxes, narrow._tbl, LEFT, 155.28) == 0.0,
                    "box: nor is one of another width")
    passed &= check(convert.table_box_offset(boxes, table._tbl, LEFT, 200.0) == 0.0,
                    "box: nor one the rectangle is nowhere near")

    two = boxes + [dict(boxes[0], y0=boxes[0]["y0"] + 3.0)]
    passed &= check(convert.table_box_offset(two, table._tbl, LEFT, 155.28) == 0.0,
                    "box: two rectangles that disagree are no evidence")

    # A header row is drawn one rectangle per cell: 108 and 396 pt for a
    # 504 pt grid, which is how a host table under a code panel stood 19.8 pt
    # below the PDF's place, neither piece being as wide as the table.
    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    header = (0.2627, 0.2628, 0.2627)
    page.draw_rect(pymupdf.Rect(LEFT, 618.4, 162.0, 638.2), color=None, fill=header)
    page.draw_rect(pymupdf.Rect(162.0, 618.4, RIGHT, 638.2), color=None, fill=(0.2628, 0.2627, 0.2628))
    # A panel's accent bar touches its body the same way and is not a header.
    page.draw_rect(pymupdf.Rect(LEFT, 200.0, LEFT + 7.2, 271.6), color=None, fill=(0.35, 0.35, 0.35))
    page.draw_rect(pymupdf.Rect(LEFT + 7.2, 200.0, RIGHT, 271.6), color=None, fill=(0.93, 0.93, 0.93))
    # Two cells with a gap between them are two boxes.
    page.draw_rect(pymupdf.Rect(LEFT, 400.0, 162.0, 419.8), color=None, fill=header)
    page.draw_rect(pymupdf.Rect(170.0, 400.0, RIGHT, 419.8), color=None, fill=header)
    boxes = convert.page_boxes(doc)[0]
    doc.close()

    def spans(y0):
        return [(round(b["x0"], 1), round(b["x1"], 1)) for b in boxes if abs(b["y0"] - y0) < 0.5]

    passed &= check((LEFT, RIGHT) in spans(618.4),
                    "box: the cells of a filled row are one box beside their parts", spans(618.4))
    passed &= check(spans(200.0) == [(LEFT + 7.2, RIGHT)],
                    "box: a bar and a body of another colour are not joined", spans(200.0))
    passed &= check((LEFT, RIGHT) not in spans(400.0),
                    "box: cells that do not touch are not joined", spans(400.0))

    grid = exact_table(Document(), [19.8, 19.7, 19.7, 19.7, 19.7, 19.7], [108.0, 396.0])
    moved = convert.table_box_offset(boxes, grid._tbl, LEFT, 638.2, [632.2])
    passed &= check(abs(moved + 19.8) < 0.05,
                    "box: a table moves up onto its header row's fill", moved)
    parts = [b for b in boxes if not (b["x0"] <= LEFT + 0.5 and b["x1"] >= RIGHT - 0.5)]
    passed &= check(convert.table_box_offset(parts, grid._tbl, LEFT, 638.2, [632.2]) == 0.0,
                    "box: the pieces alone never matched it")
    passed &= check(convert.table_box_offset(boxes, grid._tbl, LEFT, 618.4, [632.2]) == 0.0,
                    "box: a table already under its header fill keeps its place")
    passed &= check(convert.table_box_offset(boxes, grid._tbl, LEFT, 638.2) == 0.0,
                    "box: a header fill is not believed without the PDF's first row of the table")

    # A one-line finding bar is as tall as a header row, with typst's default
    # 5 pt insets, and stands about 34 pt above a table that is where it
    # belongs (the reviewer's probe: 16.95 pt both; here the bar is at 72.0
    # and the table at 104.0). The bar is the only box on the page, so it
    # was the table's header and a correctly placed table was pulled up
    # onto it.
    bar = [{"x0": LEFT, "x1": RIGHT, "y0": 72.0, "height": 16.95}]
    probe = exact_table(Document(), [16.95, 20.0], [108.0, 396.0])
    passed &= check(convert.table_box_offset(bar, probe._tbl, LEFT, 104.0, [116.0]) == 0.0,
                    "box: a bar the table's own text is not in is not its header")
    passed &= check(abs(convert.table_box_offset(bar, probe._tbl, LEFT, 104.0, [81.0]) + 32.0) < 0.05,
                    "box: a box the first row stands in is still believed")

    # End to end, where the walk gives the table the PDF row it consumed. The
    # paragraph above it has 6 pt after it, which is what a pull up spends.
    def probe_page():
        document = Document()
        document.sections[0].top_margin = Pt(72)
        document.sections[0].bottom_margin = Pt(72)
        document.sections[0].left_margin = Pt(LEFT)
        exact_paragraph(document, "Finding bar text here", 0.0, line=12.0)
        exact_paragraph(document, "Between the two here", 2.0, line=12.0, after=6.0)
        table = exact_table(document, [16.95, 20.0], [108.0, 396.0])
        for cell, text in zip(table.rows[0].cells + table.rows[1].cells,
                              ("Host name", "Result value", "ten host name", "fine value here")):
            fill_cell(cell, text, before=2.0, line=12.0)
        return document

    page_rows = [[anchor_row("Finding bar text here", 81.6, 73.0, 85.0),
                  anchor_row("Between the two here", 95.6, 87.0, 99.0),
                  anchor_row("Host name", 116.0, 108.0, 119.0, x0=60.0, x1=100.0),
                  anchor_row("Result value", 116.0, 108.0, 119.0, x0=170.0, x1=230.0),
                  anchor_row("ten host name", 136.0, 128.0, 139.0, x0=60.0, x1=130.0),
                  anchor_row("fine value here", 136.0, 128.0, 139.0, x0=170.0, x1=250.0)]]
    document = probe_page()
    convert.align_vertical_rhythm(document, page_rows, 792.0, boxes=[bar])
    between = convert.paragraph_metrics(document.paragraphs[1]._p)
    passed &= check(between is not None and abs(between["after"] - 6.0) < 0.05,
                    "box: a finding bar above a table does not pull the table up",
                    between and between["after"])

    # The first row's words are found by text and column, so a table whose
    # rows do not pair with the page's in reading order (a cell of many lines
    # beside a cell of one, read across by the PDF and down by the converter)
    # still has its own header confirmed. Here the PDF's rows come in an
    # order the table's text cannot join, and the page's one box is the
    # table's true header, 10 pt above where the table is modelled.
    columns = convert.text_columns(page_rows[0])
    found = convert.first_row_baselines(probe_page().tables[0]._tbl, LEFT, columns)
    passed &= check(found == [116.0, 116.0],
                    "box: the first row's words are found by text and column", found)
    wrong_column = [dict(column, x0=column["x0"] + 600.0, x1=column["x1"] + 600.0) for column in columns]
    passed &= check(convert.first_row_baselines(probe_page().tables[0]._tbl, LEFT, wrong_column) == [],
                    "box: and not in another column")
    # A label above the table that opens like its first cell is not the
    # cell's line: "Host" over a table whose first cell reads "Host name".
    # Its neighbour cell has no line of its own near it. The table's own lines
    # are corroborated by the other cell and kept, and a lone label is not.
    cell_a = next(c for c in columns if c["key"] == "hostname")
    cell_b = next(c for c in columns if c["key"] == "resultvalue")
    label = dict(cell_a, key="host", base=81.0, y0=72.0, y1=84.0, x1=cell_a["x0"] + 24.0)
    found = convert.first_row_baselines(probe_page().tables[0]._tbl, LEFT, [label, cell_a, cell_b])
    passed &= check(found == [116.0, 116.0],
                    "box: a label above the table that opens like its first cell is not its first line", found)
    found = convert.first_row_baselines(probe_page().tables[0]._tbl, LEFT, [label, cell_b])
    passed &= check(found == [], "box: a label with no neighbour cell's line near it, and a lone cell line, prove nothing", found)
    found = convert.first_row_baselines(probe_page().tables[0]._tbl, LEFT, [label])
    passed &= check(found == [], "box: a label alone is no evidence of the table's first row", found)
    short = dict(cell_a, key="ho", base=81.0)
    found = convert.first_row_baselines(probe_page().tables[0]._tbl, LEFT, [short])
    passed &= check(found == [], "box: a key shorter than the comparison floor opens nothing", found)
    header = [{"x0": LEFT, "x1": RIGHT, "y0": 104.0, "height": 16.95}]
    shuffled = [[page_rows[0][0], page_rows[0][1], page_rows[0][3], page_rows[0][2],
                 page_rows[0][5], page_rows[0][4]]]
    document = probe_page()
    convert.set_spacing(document.paragraphs[1]._p, after=16.0)
    convert.align_vertical_rhythm(document, shuffled, 792.0, boxes=[header])
    between = convert.paragraph_metrics(document.paragraphs[1]._p)
    passed &= check(between is not None and abs(between["after"] - 6.0) < 0.05,
                    "box: a table whose rows do not pair is still moved onto its own header",
                    between and between["after"])
    return passed


def split_document(text_lines, before=6.0, after=9.0, line=19.7):
    """One paragraph holding several lines, broken the way the converter breaks
    them: a run of text, then a run holding nothing but the break."""
    document = Document()
    document.sections[0].top_margin = Pt(72)
    paragraph = document.add_paragraph()
    convert.set_spacing(paragraph._p, before=before, after=after, line=line)
    for order, text in enumerate(text_lines):
        if order:
            holder = convert.OxmlElement("w:r")
            holder.append(convert.OxmlElement("w:br"))
            paragraph._p.append(holder)
        paragraph.add_run(text)
    return document, paragraph._p


def paragraph_shape(element):
    """What a paragraph holds, for a before-and-after comparison."""
    return (convert.paragraph_text(element),
            len([node for node in element.iter(convert.qn("w:br"))]),
            len([node for node in element.iter(convert.qn("w:r"))]))


def run_pitch_split_checks():
    """Cutting a paragraph where the PDF changed the leading inside it."""
    passed = True

    def rows(bases, ascent=11.6):
        return [anchor_row("Line %d here" % order, base, base - ascent, base + 3.0)
                for order, base in enumerate(bases)]

    passed &= check(convert.pitch_cuts(rows([100.0, 114.8, 129.6, 144.4])) == [],
                    "split: one leading all the way down is not a cut")
    passed &= check(convert.pitch_cuts(rows([100.0])) == [] and convert.pitch_cuts(rows([100.0, 120.0])) == [],
                    "split: one gap has nothing to be read against")
    passed &= check(convert.pitch_cuts(rows([100.0, 119.7, 134.5, 149.3])) == [1],
                    "split: a heading and the body under it are two blocks",
                    convert.pitch_cuts(rows([100.0, 119.7, 134.5, 149.3])))
    passed &= check(convert.pitch_cuts(rows([100.0, 122.8, 138.5, 161.3, 177.0])) == [1, 3],
                    "split: and a list of them is as many as it has",
                    convert.pitch_cuts(rows([100.0, 122.8, 138.5, 161.3, 177.0])))
    passed &= check(convert.pitch_cuts(rows([100.0, 114.8, 129.9, 144.7])) == [],
                    "split: a difference under half a point is rounding, not a block")

    # The breaks have to be replaceable one for one.
    document, element = split_document(["Line 0 here", "Line 1 here", "Line 2 here"])
    holders = convert.break_runs(element)
    passed &= check(holders is not None and len(holders) == 2,
                    "split: the paragraph's own line breaks are found", holders and len(holders))
    page_break = convert.OxmlElement("w:r")
    node = convert.OxmlElement("w:br")
    node.set(convert.qn("w:type"), "page")
    page_break.append(node)
    element.append(page_break)
    passed &= check(convert.break_runs(element) is None,
                    "split: a page break is not a line break and stops the cut")
    element.remove(page_break)
    mixed = convert.OxmlElement("w:r")
    mixed.append(convert.OxmlElement("w:br"))
    text = convert.OxmlElement("w:t")
    text.text = "tail"
    mixed.append(text)
    element.append(mixed)
    passed &= check(convert.break_runs(element) is None,
                    "split: a break sharing its run with text stops it too")
    element.remove(mixed)

    passed &= check(convert.splittable(element, ()), "split: an ordinary paragraph may be cut")
    numbered = convert.OxmlElement("w:numPr")
    element.get_or_add_pPr().insert(0, numbered)
    passed &= check(not convert.splittable(element, ()),
                    "split: a numbered paragraph would take a second number")
    element.get_or_add_pPr().remove(numbered)
    passed &= check(not convert.splittable(element, (element,)),
                    "split: code has to come out as it went in")

    # A section break lives in the properties of the paragraph it ends, and
    # the properties are what a cut copies.
    document, element = split_document(["Line 0 here", "Line 1 here", "Line 2 here"])
    section = convert.OxmlElement("w:sectPr")
    convert.insert_ordered(element.get_or_add_pPr(), section, convert.PPR_ORDER)
    passed &= check(not convert.splittable(element, ()),
                    "split: a paragraph that ends a section is not cut")
    # And if one is cut anyway, the break lands on the last piece and nowhere
    # else, then comes home when the cut is joined back up.
    holders = convert.break_runs(element)
    seams = convert.cut_paragraph(element, holders)
    breaks = [len(p.findall(convert.qn("w:sectPr")))
              for p in [element.find(convert.qn("w:pPr"))]
              + [s[1].find(convert.qn("w:pPr")) for s in seams]]
    passed &= check(breaks == [0, 0, 1], "split: the section break is on the last piece only", breaks)
    for seam in reversed(seams):
        convert.rejoin_paragraph(*seam)
    passed &= check(len(element.find(convert.qn("w:pPr")).findall(convert.qn("w:sectPr"))) == 1,
                    "split: and joining back up brings the break home",
                    len(element.find(convert.qn("w:pPr")).findall(convert.qn("w:sectPr"))))
    passed &= check(len(document.element.body.findall(convert.qn("w:p"))) == 1,
                    "split: with one paragraph left",
                    len(document.element.body.findall(convert.qn("w:p"))))

    # The cut itself: the same characters, one break fewer, the same height.
    document, element = split_document(["Line 0 here", "Line 1 here", "Line 2 here"])
    before_text = convert.paragraph_text(element)
    before_breaks = paragraph_shape(element)[1]
    holders = convert.break_runs(element)
    seams = convert.cut_paragraph(element, [holders[0]])
    tail = seams[0][1]
    passed &= check(len(seams) == 1, "split: one cut makes one seam", len(seams))
    passed &= check(convert.paragraph_text(element) + convert.paragraph_text(tail) == before_text,
                    "split: not one character was added or lost",
                    (convert.paragraph_text(element), convert.paragraph_text(tail)))
    passed &= check(paragraph_shape(element)[1] + paragraph_shape(tail)[1] == before_breaks - 1,
                    "split: a paragraph break replaced exactly one line break")
    head_metrics = convert.paragraph_metrics(element)
    tail_metrics = convert.paragraph_metrics(tail)
    passed &= check(head_metrics["before"] == 6.0 and head_metrics["after"] == 0.0,
                    "split: the first piece keeps the spacing before", head_metrics)
    passed &= check(tail_metrics["before"] == 0.0 and tail_metrics["after"] == 9.0,
                    "split: the last keeps the spacing after", tail_metrics)
    was = 6.0 + 3 * 19.7 + 9.0
    now = (head_metrics["before"] + head_metrics["lines"] * head_metrics["line"]
           + max(head_metrics["after"], tail_metrics["before"])
           + tail_metrics["lines"] * tail_metrics["line"] + tail_metrics["after"])
    passed &= check(abs(now - was) < 0.05,
                    "split: the pieces are as tall together as the paragraph was", (was, now))

    convert.rejoin_paragraph(*seams[0])
    passed &= check(convert.paragraph_text(element) == before_text
                    and paragraph_shape(element)[1] == before_breaks,
                    "split: and joining it back up gives the paragraph it came from",
                    paragraph_shape(element))
    passed &= check(convert.paragraph_metrics(element)["after"] == 9.0,
                    "split: with the spacing after it had")

    # End to end, against a page whose rows say where the leading changed.
    document, element = split_document(["Line 0 here", "Line 1 here", "Line 2 here"])
    pdf_rows = [rows([100.0, 119.7, 134.5])]
    seams = convert.split_pitch_changes(document, pdf_rows)
    passed &= check(len(seams) == 1 and len(document.paragraphs) == 2,
                    "split: the pass finds the change and cuts there",
                    (len(seams), len(document.paragraphs)))
    convert.align_vertical_rhythm(document, pdf_rows, 792.0, seams=seams)
    passed &= check(len(document.paragraphs) == 2,
                    "split: a piece the layout pass placed stays a paragraph",
                    len(document.paragraphs))
    heights = [convert.paragraph_metrics(p._p)["line"] for p in document.paragraphs]
    passed &= check(abs(heights[1] - 14.8) < 0.1,
                    "split: and the body under it gets the leading the PDF set", heights)

    # A piece nothing placed is joined back up rather than left with no space.
    document, element = split_document(["Line 0 here", "Line 1 here", "Line 2 here"])
    seams = convert.split_pitch_changes(document, pdf_rows)
    passed &= check(convert.rejoin_unplaced(seams, set()) == 1 and len(document.paragraphs) == 1,
                    "split: a piece the layout pass never reached is joined back up",
                    len(document.paragraphs))
    return passed


def run_panel_checks():
    """A code panel rebuilt as a table, so the grey gets the PDF's padding."""
    passed = True
    text = ["import requests", "r = requests.get(url)", "print(r.status_code)"]
    bases = [128.97, 154.34, 167.03]
    document, made = code_document(text, shaded=True)
    panel = code_panel(115.13, 58.9, lines=list(zip(text, bases)))
    built = convert.build_panel_table(panel, made)
    passed &= check(built is not None, "panel: the panel was rebuilt as a table")
    table, top, was = built
    passed &= check(abs(top - 115.13) < 0.01 and was > 40.0,
                    "panel: it reports the PDF's top edge and the height it replaced", (top, was))
    row = table.findall(convert.qn("w:tr"))[0]
    cells = row.findall(convert.qn("w:tc"))
    height = row.find(convert.qn("w:trPr")).find(convert.qn("w:trHeight"))
    passed &= check(height.get(convert.qn("w:hRule")) == "exact"
                    and abs(convert.twips_of(height.get(convert.qn("w:val"))) - 58.9) < 0.05,
                    "panel: the row is exactly as tall as the PDF's rectangle",
                    height.get(convert.qn("w:val")))
    passed &= check(len(cells) == 2, "panel: the accent bar is a cell of its own", len(cells))
    fills = [cell.find(convert.qn("w:tcPr")).find(convert.qn("w:shd")).get(convert.qn("w:fill"))
             for cell in cells]
    passed &= check(fills == ["1A2438", "EDEDED"], "panel: each cell carries its own colour", fills)
    widths = [convert.twips_of(cell.find(convert.qn("w:tcPr")).find(
        convert.qn("w:tcW")).get(convert.qn("w:w"))) for cell in cells]
    passed &= check(abs(widths[0] - 7.2) < 0.05 and abs(widths[1] - 476.8) < 0.05,
                    "panel: the bar and the panel are the PDF's own widths", widths)
    margins = cells[1].find(convert.qn("w:tcPr")).find(convert.qn("w:tcMar"))
    sides = {name: convert.twips_of(margins.find(convert.qn("w:" + name)).get(convert.qn("w:w")))
             for name in ("top", "start", "bottom", "end")}
    passed &= check(abs(sides["start"] - 9.0) < 0.05 and sides["top"] == 0.0 and sides["bottom"] == 0.0,
                    "panel: the left padding is the PDF's, and Word is left to fill the rest", sides)
    indent = convert.twips_of(table.find(convert.qn("w:tblPr")).find(
        convert.qn("w:tblInd")).get(convert.qn("w:w")))
    passed &= check(abs(indent - (25.8 - 7.2 - 9.0)) < 0.05,
                    "panel: the table's indent leaves the code at the x it already had", indent)
    passed &= check([convert.paragraph_text(node) for node in cells[1].findall(convert.qn("w:p"))] == text,
                    "panel: the code moved in unchanged, byte for byte")
    properties = made[0].find(convert.qn("w:pPr"))
    passed &= check(properties.find(convert.qn("w:shd")) is None
                    and properties.find(convert.qn("w:ind")) is None,
                    "panel: the paragraph's own fill and indent are gone, the cell has them now")
    passed &= check(table.getprevious() is not None and table.getprevious().tag == convert.qn("w:p")
                    and table.getnext() is not None and table.getnext().tag == convert.qn("w:p"),
                    "panel: a paragraph on either side, so Word does not join two tables")

    # End to end: the vertical pass puts the panel on the PDF's own top edge.
    rows = [[anchor_row("Step title", 100.0, 92.0, 103.0)]
            + [anchor_row(line, base, base - 7.0, base + 2.0, x0=80.2)
               for line, base in zip(text, bases)]]
    convert.align_vertical_rhythm(document, rows, 792.0, [built])
    spacing = convert.paragraph_metrics(table.getprevious())
    passed &= check(spacing is not None and abs(spacing["after"] - (115.13 - 72.0 - 13.5)) < 0.3,
                    "panel: the gap above it lands the panel on the PDF's top edge",
                    spacing and spacing["after"])
    first = convert.paragraph_metrics(made[0])
    passed &= check(abs(115.13 + first["before"] + convert.EXACT_BASELINE_RATIO * first["line"] - 128.97) < 0.1,
                    "panel: and the first line of code lands on the PDF's baseline",
                    first and first["before"])

    # pdf2docx's own shape for a code block: one table per line, each one row
    # of exact height, a 1 pt spacer between them. Taken apart they become
    # one panel table, the spacers go, and every line keeps its pitch.
    document = Document()
    document.sections[0].top_margin = Pt(72)
    lead = document.add_paragraph("Step title")
    convert.set_spacing(lead._p, before=0.0, after=8.3, line=13.5)
    heights = [17.8, 12.6, 15.7]
    befores = [6.6, 0.0, 0.0]
    wrapped = []
    for order, (line_text, height, before) in enumerate(zip(text, heights, befores)):
        if order:
            convert.set_spacing(document.add_paragraph()._p, before=0.0, after=0.0, line=1.0)
        table = document.add_table(rows=1, cols=1)
        row = table.rows[0]._tr
        span = convert.OxmlElement("w:trHeight")
        span.set(convert.qn("w:hRule"), "exact")
        span.set(convert.qn("w:val"), convert.twips_attribute(height))
        row.get_or_add_trPr().append(span)
        cell_paragraph = table.rows[0].cells[0].paragraphs[0]
        cell_paragraph.add_run(line_text)
        convert.set_spacing(cell_paragraph._p, before=before, after=0.0, line=9.0)
        cell_paragraph.paragraph_format.left_indent = Pt(16.2)
        wrapped.append(cell_paragraph._p)
    panel = code_panel(115.13, sum(heights), x0=71.2, text_x=80.2 + 7.2 + 5.4,
                       lines=list(zip(text, [128.97, 141.4, 154.0])))
    painted, anchors = convert.shade_run_of_paragraphs(document, [panel])
    tables = document.element.body.findall(convert.qn("w:tbl"))
    passed &= check(len(tables) == 1 and len(anchors) == 1 and painted == wrapped,
                    "unwrap: three one-line tables became one panel table", (len(tables), len(anchors)))
    if tables:
        row = tables[0].find(convert.qn("w:tr"))
        height = convert.twips_of(row.find(convert.qn("w:trPr")).find(convert.qn("w:trHeight")).get(convert.qn("w:val")))
        inside = [convert.paragraph_text(p) for p in row.findall(convert.qn("w:tc"))[-1].findall(convert.qn("w:p"))]
        passed &= check(abs(height - sum(heights)) < 0.06 and inside == text,
                        "unwrap: the row is as tall as the rectangle and holds every line", (height, inside))
        spacing = [convert.paragraph_metrics(p)["before"] for p in wrapped]
        # before_k = before_k + (h_{k-1} - before_{k-1} - line) + spacer
        passed &= check([round(s, 2) for s in spacing] == [6.6, 3.2, 1.4],
                        "unwrap: each line's pitch survives as spacing before the next", spacing)
        passed &= check(not [p for p in document.element.body.iterchildren(convert.qn("w:p"))
                             if convert.blank_spacer(p) and convert.paragraph_metrics(p)["line"] >= 0.9],
                        "unwrap: the 1 pt spacers are gone")

    # Lines pdf2docx joined are broken where the PDF broke them, and the
    # leading spaces it turned into an indent come back as spaces.
    document = Document()
    joined = document.add_paragraph()
    joined.add_run("Sharename       Type      Comment delta$          Disk      Sit amet")
    indented = document.add_paragraph()
    indented.add_run("golf            Disk      Quis nostrud")
    indented.paragraph_format.left_indent = Pt(37.9 + 16.2)
    width = 5.42 * len("Sharename       Type      Comment")
    rows = [{"x0": 80.2, "x1": 80.2 + width, "text": "Sharename       Type      Comment"},
            {"x0": 80.2, "x1": 80.2 + 5.42 * len("delta$          Disk      Sit amet"), "text": "delta$          Disk      Sit amet"},
            {"x0": 80.2 + 7 * 5.42, "x1": 80.2 + 7 * 5.42 + 5.42 * len("golf            Disk      Quis nostrud"),
             "text": "golf            Disk      Quis nostrud"}]
    panel = {"rows": rows, "rect": pymupdf.Rect(61.2, 100.0, 558.0, 150.0), "fill": (0.93, 0.93, 0.93), "bar": None}
    elements = [joined._p, indented._p]
    passed &= check(convert.split_joined_code(elements, panel) == 1 and joined._p.find(convert.qn("w:r")) is not None
                    and len(list(joined._p.iter(convert.qn("w:br")))) == 1,
                    "code: two lines pdf2docx joined get the PDF's break between them")
    # A rule of dashes under a listing's header squashes to nothing in the
    # comparison, so its break is cut from the paragraph's own text.
    ruled = document.add_paragraph()
    ruled.add_run("    Sharename       Type      Comment---------       ----      -------charlie$        Disk      Aute irure dolor")
    ruled_rows = [{"x0": 80.2, "x1": 80.2 + 5.42 * len(line), "text": line}
                  for line in ("Sharename       Type      Comment", "---------       ----      -------",
                               "charlie$        Disk      Aute irure dolor")]
    ruled_panel = {"rows": ruled_rows, "rect": pymupdf.Rect(61.2, 100.0, 558.0, 150.0), "fill": (0.93, 0.93, 0.93), "bar": None}
    made = convert.split_joined_code([ruled._p], ruled_panel)
    ruled_lines = [""]
    for run in ruled._p.iter(convert.qn("w:r")):
        for child in run:
            if child.tag == convert.qn("w:t"):
                ruled_lines[-1] += child.text or ""
            elif child.tag == convert.qn("w:br"):
                ruled_lines.append("")
    passed &= check(made == 2 and [line.strip() for line in ruled_lines] == [row["text"] for row in ruled_rows],
                    "code: a rule of dashes keeps a line of its own", (made, ruled_lines))
    # The same listing cut by a page right under its rule: the paragraph ENDS
    # with the rule, no held row follows it, and the break before it was lost.
    closed = document.add_paragraph()
    closed.add_run("$ smbclient -L //10.10.7.229 -N     Sharename       Type      Comment---------       ----      -------")
    closed_rows = [{"x0": 80.2, "x1": 80.2 + 5.42 * len(line), "text": line}
                   for line in ("$ smbclient -L //10.10.7.229 -N", "    Sharename       Type      Comment",
                                "    ---------       ----      -------", "    alpha           Disk      Ipsum dolor sit")]
    closed_panel = {"rows": closed_rows[:3], "rect": pymupdf.Rect(61.2, 100.0, 558.0, 150.0), "fill": (0.93, 0.93, 0.93), "bar": None}
    made = convert.split_joined_code([closed._p], closed_panel)
    closed_lines = [""]
    for run in closed._p.iter(convert.qn("w:r")):
        for child in run:
            if child.tag == convert.qn("w:t"):
                closed_lines[-1] += child.text or ""
            elif child.tag == convert.qn("w:br"):
                closed_lines.append("")
    passed &= check(made == 2 and [line.strip() for line in closed_lines] == [row["text"].strip() for row in closed_rows[:3]],
                    "code: a rule that is the last row of its page keeps a line of its own", (made, closed_lines))
    # A rule the paragraph does not end with is not claimed by the extension.
    open_rows = closed_rows[:3] + [{"x0": 80.2, "x1": 160.0, "text": "    -------- ------"}]
    open_panel = dict(closed_panel, rows=open_rows)
    stray = document.add_paragraph()
    stray.add_run("$ smbclient -L //10.10.7.229 -N     Sharename       Type      Comment")
    passed &= check(convert.raw_row_ends(stray._p, open_rows[:2] + [open_rows[3]], to_end=True) is None
                    and convert.raw_row_ends(stray._p, open_rows[:2], to_end=True) == [31],
                    "code: raw_row_ends(to_end) refuses a rule the paragraph does not hold")
    passed &= check(convert.restore_leading_spaces(elements, panel) == 1
                    and convert.paragraph_text(indented._p).startswith("       golf")
                    and abs(convert.paragraph_indent(indented._p) - 16.2) < 0.1,
                    "code: seven spaces come back and the indent returns to the panel's",
                    (repr(convert.paragraph_text(indented._p)[:12]), convert.paragraph_indent(indented._p)))

    # A listing's continuation panel whose rows are all four characters in
    # takes the edge the document's other panels share.
    def edge_panel(x0, offsets):
        return {"rect": pymupdf.Rect(61.2, 100.0, 558.0, 150.0), "fill": (0.93, 0.93, 0.93), "bar": None,
                "rows": [{"x0": x0 + k * 5.42, "x1": x0 + k * 5.42 + 5.42 * 20, "y0": 104.0, "y1": 113.0, "text": "a" * 20} for k in offsets]}
    shared = [edge_panel(70.2, [0, 4]), edge_panel(70.2, [0]), edge_panel(70.2, [0, 2])]
    continued = edge_panel(70.2, [4, 4, 4])
    odd = edge_panel(70.2, [0.5])
    settled = convert.settle_panel_text_left(shared + [continued, odd])
    passed &= check(settled == 1 and abs(convert.panel_text_left(continued) - 70.2) < 0.05
                    and abs(convert.panel_padding(continued)[0] - 9.0) < 0.05
                    and abs(convert.panel_text_left(shared[0]) - 70.2) < 0.05,
                    "code: a panel whose rows are all indented takes the shared text edge",
                    (settled, convert.panel_text_left(continued)))

    # YAML whose nested lines pdf2docx indented with tab characters.
    yaml_doc = Document()
    yaml = yaml_doc.add_paragraph()
    yaml.add_run("server:")
    yaml.add_run().add_break()
    tab_run = yaml.add_run()
    tab_run.add_tab()
    yaml.add_run("listen: 10.10.6.176:8080")
    yaml.add_run().add_break()
    deep = yaml.add_run()
    deep.add_tab()
    deep.add_tab()
    yaml.add_run("enabled: false")
    yaml_rows = [{"x0": 80.2, "x1": 80.2 + 5.42 * 7, "y0": 104.0, "y1": 113.0, "text": "server:"},
                 {"x0": 80.2 + 2 * 5.42, "x1": 80.2 + 2 * 5.42 + 5.42 * 24, "y0": 117.0, "y1": 126.0, "text": "listen: 10.10.6.176:8080"},
                 {"x0": 80.2, "inkX0": 80.2 + 4 * 5.42, "x1": 80.2 + 4 * 5.42 + 5.42 * 14, "y0": 130.0, "y1": 139.0, "text": "enabled: false"}]
    yaml_panel = {"rows": yaml_rows, "rect": pymupdf.Rect(71.2, 100.0, 548.0, 145.0), "fill": (0.93, 0.93, 0.93), "bar": None}
    convert.restore_leading_spaces([yaml._p], yaml_panel)
    yaml_lines = [""]
    for run in yaml._p.iter(convert.qn("w:r")):
        for child in run:
            if child.tag == convert.qn("w:t"):
                yaml_lines[-1] += child.text or ""
            elif child.tag == convert.qn("w:br"):
                yaml_lines.append("")
            elif child.tag == convert.qn("w:tab"):
                yaml_lines[-1] += "<TAB>"
    passed &= check(yaml_lines == ["server:", "  listen: 10.10.6.176:8080", "    enabled: false"],
                    "code: tabs in front of a nested line become the PDF's count of spaces, from where its ink starts",
                    yaml_lines)

    # A nested line whose spaces pdf2docx kept as text, indented at the panel's edge: the indent stays.
    kept_doc = Document()
    kept = kept_doc.add_paragraph()
    kept.add_run("  listen: 10.10.6.176:8080")
    kept.paragraph_format.left_indent = Pt(25.8)
    kept_rows = [{"x0": 80.2, "x1": 80.2 + 5.42 * 7, "y0": 104.0, "y1": 113.0, "text": "server:"},
                 {"x0": 80.2, "inkX0": 80.2 + 2 * 5.42, "x1": 80.2 + 2 * 5.42 + 5.42 * 24, "y0": 117.0, "y1": 126.0,
                  "text": "listen: 10.10.6.176:8080"}]
    kept_panel = {"rows": kept_rows, "rect": pymupdf.Rect(71.2, 100.0, 548.0, 145.0), "fill": (0.93, 0.93, 0.93), "bar": None}
    kept_rows_only = dict(kept_panel, rows=kept_rows)
    convert.restore_leading_spaces([kept._p], kept_rows_only)
    passed &= check(convert.paragraph_text(kept._p) == "  listen: 10.10.6.176:8080"
                    and abs(convert.paragraph_indent(kept._p) - 25.8) < 0.1,
                    "code: a line whose spaces are already text keeps its spaces and its indent",
                    (repr(convert.paragraph_text(kept._p)), convert.paragraph_indent(kept._p)))

    # A JSON block joined into one paragraph whose first line, "{", is too short to measure.
    json_doc = Document()
    json_para = json_doc.add_paragraph()
    json_para.add_run("{")
    json_para.add_run().add_break()
    json_para.add_run('"host": "10.10.9.69",')
    json_para.paragraph_format.left_indent = Pt(25.8)
    json_rows = [{"x0": 80.2, "x1": 80.2 + 5.42, "y0": 104.0, "y1": 113.0, "text": "{"},
                 {"x0": 80.2 + 2 * 5.42, "x1": 80.2 + 2 * 5.42 + 5.42 * 21, "y0": 117.0, "y1": 126.0, "text": '"host": "10.10.9.69",'},
                 {"x0": 80.2 + 2 * 5.42, "x1": 80.2 + 2 * 5.42 + 5.42 * 23, "y0": 130.0, "y1": 139.0, "text": '"ports": [22, 80, 443],'}]
    json_panel = {"rows": json_rows, "rect": pymupdf.Rect(71.2, 100.0, 548.0, 145.0), "fill": (0.93, 0.93, 0.93), "bar": None}
    convert.restore_leading_spaces([json_para._p], json_panel)
    json_lines = [""]
    for run in json_para._p.iter(convert.qn("w:r")):
        for child in run:
            if child.tag == convert.qn("w:t"):
                json_lines[-1] += child.text or ""
            elif child.tag == convert.qn("w:br"):
                json_lines.append("")
    passed &= check(json_lines == ["{", '  "host": "10.10.9.69",'] and abs(convert.paragraph_indent(json_para._p) - 25.8) < 0.1,
                    "code: the first line decides the indent, the next line only gets its spaces",
                    (json_lines, convert.paragraph_indent(json_para._p)))

    # A closing brace pdf2docx glued onto the end of the last value. The row
    # is too short to compare and it is the one at the panel's edge; unseen,
    # the block's edge was taken from the "}," row (one space kept of two) and
    # every line of the block came out one character left of the PDF.
    glue_advance = 5.42
    glue_doc = Document()
    glue_para = glue_doc.add_paragraph()
    glue_para.paragraph_format.left_indent = Pt(25.8)
    for number, glued_line in enumerate(['"findings": {', '"critical": 2,', " }, ", '"notes": "Dolor sit amet" }']):
        if number:
            glue_para.add_run().add_break()
        glue_para.add_run(glued_line)

    def glue_row(row_text, offset, ink=None):
        row = {"x0": 80.2 + offset * glue_advance, "x1": 80.2 + (offset if ink is None else ink) * glue_advance + glue_advance * len(row_text),
               "y0": 104.0, "y1": 113.0, "text": row_text}
        if ink is not None:
            row["inkX0"] = 80.2 + ink * glue_advance
        return row

    glue_rows = [glue_row('"findings": {', 2), glue_row('"critical": 2,', 4), glue_row("},", 0, 2),
                 glue_row('"notes": "Dolor sit amet"', 2), glue_row("}", 0)]
    glue_panel = {"rows": glue_rows, "rect": pymupdf.Rect(71.2, 100.0, 548.0, 145.0), "fill": (0.93, 0.93, 0.93), "bar": None}
    cut = convert.split_joined_code([glue_para._p], glue_panel)
    convert.restore_leading_spaces([glue_para._p], glue_panel)
    glue_lines = [""]
    for run in glue_para._p.iter(convert.qn("w:r")):
        for child in run:
            if child.tag == convert.qn("w:t"):
                glue_lines[-1] += child.text or ""
            elif child.tag == convert.qn("w:br"):
                glue_lines.append("")
    passed &= check(cut == 1 and [line.rstrip() for line in glue_lines]
                    == ['  "findings": {', '    "critical": 2,', "  },", '  "notes": "Dolor sit amet"', "}"]
                    and abs(convert.paragraph_indent(glue_para._p) - 25.8) < 0.1,
                    "code: a short row glued to the end of a line gets a line of its own and the block keeps the panel's edge",
                    (cut, glue_lines, convert.paragraph_indent(glue_para._p)))
    squashed_rows = [convert.compare_squash(row["text"]) for row in glue_rows]
    passed &= check(convert.glued_tail(convert.compare_squash('"notes": "Dolor sit amet" }'), squashed_rows) == (3, [4])
                    and convert.glued_tail(convert.compare_squash('"notes": "Dolor sit amet" ]'), squashed_rows) is None
                    and convert.glued_tail(convert.compare_squash('"notes": "Dolor sit amet"'), squashed_rows) is None,
                    "code: a tail is only the short rows that follow the line's own row, in the panel's order")

    # A listing pdf2docx wrote as one paragraph with a hanging indent: the
    # first line at the panel's edge, the rest four characters in.
    listing = document.add_paragraph()
    listing.add_run("$ smbclient -L //10.10.4.10 -N")
    listing.add_run().add_break()
    listing.add_run("Sharename       Type      Comment")
    listing.add_run().add_break()
    listing.add_run("alpha           Disk      Velit esse")
    listing.paragraph_format.left_indent = Pt(47.5)
    listing.paragraph_format.first_line_indent = Pt(-21.7)
    advance = 5.42
    listing_rows = [{"x0": 70.2, "x1": 70.2 + advance * 30, "text": "$ smbclient -L //10.10.4.10 -N"},
                    {"x0": 70.2 + 4 * advance, "x1": 70.2 + 4 * advance + advance * 33, "text": "Sharename       Type      Comment"},
                    {"x0": 70.2 + 4 * advance, "x1": 70.2 + 4 * advance + advance * 36, "text": "alpha           Disk      Velit esse"}]
    listing_panel = {"rows": listing_rows, "rect": pymupdf.Rect(61.2, 100.0, 558.0, 150.0), "fill": (0.93, 0.93, 0.93), "bar": None}
    convert.restore_leading_spaces([listing._p], listing_panel)
    listing_lines = [""]
    for run in listing._p.iter(convert.qn("w:r")):
        for child in run:
            if child.tag == convert.qn("w:t"):
                listing_lines[-1] += child.text or ""
            elif child.tag == convert.qn("w:br"):
                listing_lines.append("")
    ind = listing._p.find(convert.qn("w:pPr")).find(convert.qn("w:ind"))
    passed &= check(listing_lines == ["$ smbclient -L //10.10.4.10 -N", "    Sharename       Type      Comment",
                                      "    alpha           Disk      Velit esse"]
                    and ind.get(convert.qn("w:hanging")) is None and ind.get(convert.qn("w:firstLine")) is None
                    and abs(convert.paragraph_indent(listing._p) - 25.8) < 0.1,
                    "code: each line of a joined listing gets its own spaces and the hanging indent goes",
                    (listing_lines, dict(ind.attrib)))

    # A column-mate counts as a grid neighbour only when it touches.
    here = pymupdf.Rect(61.2, 300.0, 558.0, 400.0)
    far = {"rect": pymupdf.Rect(61.2, 520.0, 558.0, 620.0)}
    touching = {"rect": pymupdf.Rect(61.2, 400.4, 558.0, 500.0)}
    beside = {"rect": pymupdf.Rect(300.0, 300.0, 558.0, 400.0)}
    passed &= check(not convert.has_grid_neighbour(here, [far]) and convert.has_grid_neighbour(here, [touching])
                    and convert.has_grid_neighbour(here, [beside]),
                    "neighbour: a second panel down the page is not a table cell, a touching one is")

    # A panel taller than the paragraphs it would replace is left alone, or
    # the page it is on could gain a line.
    document, made = code_document(text, before=0.0, after=0.0, line=9.0)
    tall = code_panel(115.13, 200.0, lines=list(zip(text, bases)))
    passed &= check(convert.build_panel_table(tall, made) is None,
                    "panel: a panel taller than the text it replaces is refused")

    # So is one whose paragraphs are not next to each other.
    document, made = code_document(text)
    stray = document.add_paragraph("something else")
    made[1].addprevious(stray._p)
    passed &= check(convert.build_panel_table(panel, made) is None,
                    "panel: and one with a stranger in among its lines")

    # A panel whose rows pdf2docx joined into one paragraph declares fewer
    # lines than it has rows; it is not too tall, and is rebuilt.
    joined_rows = ["Sharename       Type      Comment", "alpha           Disk      Velit esse cillum",
                   "bravo           Disk      Dolor sit amet", "charlie         Disk      Nisi ut aliquip"]
    document, made = code_document([" ".join(joined_rows)], before=0.0, after=0.0, line=9.0)
    joined_panel = code_panel(115.13, 120.0, lines=[(text, 128.97 + 12.7 * k) for k, text in enumerate(joined_rows)])
    painted, anchors = convert.shade_run_of_paragraphs(document, [joined_panel])
    passed &= check(len(anchors) == 1 and abs(anchors[0][1] - 115.13) < 0.01,
                    "panel: rows pdf2docx joined into one paragraph do not make the panel too tall", (len(painted), len(anchors)))
    # The same height over separate paragraphs that really are shorter stays refused (below).

    # The caller then shades them the way it used to.
    document, made = code_document(text)
    tall = code_panel(115.13, 200.0, lines=list(zip(text, bases)))
    painted, anchors = convert.shade_run_of_paragraphs(document, [tall])
    properties = painted[0].find(convert.qn("w:pPr")) if painted else None
    passed &= check(len(painted) == 3 and anchors == []
                    and properties is not None and properties.find(convert.qn("w:shd")) is not None,
                    "panel: a panel that cannot be a table is still shaded", (len(painted), anchors))

    # A "{" above the first line and a "}" under the last are too short to
    # match anything, so they were left outside the table the panel became:
    # the table took the whole rectangle's height, the "}" was pushed below
    # it and the page under it fell 31 pt (shape-f, page 68).
    framed = ["{", '"host": "10.10.9.85",', '"ports": [22, 80, 443],', "}"]
    framed_bases = [128.97, 141.67, 154.37, 167.07]
    def framed_panel(rows):
        return code_panel(115.13, 12.7 * 4 + 6.0, lines=list(zip(framed, framed_bases))[:rows] if rows else [])
    document, made = code_document(framed, before=14.5, after=0.0, line=9.0)
    painted, anchors = convert.shade_run_of_paragraphs(document, [framed_panel(4)])
    tables = document.element.body.findall(convert.qn("w:tbl"))
    inside = ([convert.paragraph_text(p) for p in tables[0].find(convert.qn("w:tr")).findall(convert.qn("w:tc"))[-1]
              .findall(convert.qn("w:p"))] if len(tables) == 1 else [])
    passed &= check(inside == framed and len(anchors) == 1,
                    "panel: the short first and last rows go inside the table with the others", inside)
    outside = [convert.paragraph_text(p) for p in document.element.body.findall(convert.qn("w:p"))]
    passed &= check(outside[0] == "Step title" and not any(t in ("{", "}") for t in outside),
                    "panel: and nothing of the panel is left loose beside it", outside)

    # A listing pdf2docx wrote as one paragraph of one declared line holds
    # three rows of the panel. Word draws three lines, so that is the flow the
    # blocks below were placed after: the height the pass is told the panel
    # replaced counts a line per row, not just the rectangle (shape-a, page
    # 25: the figure under a 6-row listing stood 18.1 pt too high).
    document, made = code_document(["alpha one two beta three four gamma five six"], before=10.8, after=15.4, line=12.7)
    rows = [("alpha one two", 128.97), ("beta three four", 141.67), ("gamma five six", 154.37)]
    painted, anchors = convert.shade_run_of_paragraphs(document, [code_panel(115.13, 12.7 * 3 + 6.0, lines=rows)])
    passed &= check(len(anchors) == 1 and abs(anchors[0][2] - (10.8 + 3 * 12.7 + 15.4)) < 0.3,
                    "panel: the flow it replaced counts a line per row the PDF has",
                    anchors and anchors[0][2])

    # Only a paragraph that IS the row's text is taken: a body paragraph just
    # above the panel is not, whatever it says.
    document, made = code_document(["Next steps", "alpha beta gamma", "delta epsilon zeta", "}"], before=14.5, after=0.0, line=9.0)
    rows = [("alpha beta gamma", 128.97), ("delta epsilon zeta", 141.67), ("}", 154.37)]
    painted, anchors = convert.shade_run_of_paragraphs(document, [code_panel(115.13, 12.7 * 3 + 6.0, lines=rows)])
    tables = document.element.body.findall(convert.qn("w:tbl"))
    inside = ([convert.paragraph_text(p) for p in tables[0].find(convert.qn("w:tr")).findall(convert.qn("w:tc"))[-1]
              .findall(convert.qn("w:p"))] if len(tables) == 1 else [])
    passed &= check(inside == ["alpha beta gamma", "delta epsilon zeta", "}"],
                    "panel: a body line above the panel stays outside it", inside)


    # A finding's title banner is a panel too, and its number is followed by a
    # tab to a stop of the paragraph's own. The indent that stop was counted
    # from goes into the table, so the stop loses it, and the text keeps its x.
    titles = ["import requests", "r = requests.get(url)", "print(r.status_code)"]
    document, made = code_document(titles)
    convert.Paragraph(made[0], document).paragraph_format.tab_stops.add_tab_stop(Pt(100.0))
    built = convert.build_panel_table(code_panel(115.13, 58.9, lines=list(zip(titles, [128.97, 154.34, 167.03]))), made)
    stop = made[0].find(convert.qn("w:pPr")).find(convert.qn("w:tabs")).find(convert.qn("w:tab"))
    passed &= check(built is not None and abs(convert.twips_of(stop.get(convert.qn("w:pos"))) - (100.0 - 25.8)) < 0.1,
                    "panel: a paragraph's tab stop is counted from the cell's text edge once the indent is gone",
                    convert.twips_of(stop.get(convert.qn("w:pos"))))
    passed &= run_numbered_panel_checks()
    return passed



def numbered_row(number, text, base, number_x0, number_x1, code_x0, code_x1):
    """A page row of a numbered listing as `page_rows` reads it: the number, then the code."""
    columns = [{"x0": number_x0, "x1": number_x1, "inkX0": number_x0, "y0": base - 7.0, "y1": base + 2.0,
                "base": base, "text": str(number), "font": "DejaVu Sans Mono"},
               {"x0": code_x0, "x1": code_x1, "inkX0": code_x0, "y0": base - 7.0, "y1": base + 2.0,
                "base": base, "text": text, "font": "DejaVu Sans Mono"}]
    return {"x0": number_x0, "inkX0": number_x0, "x1": code_x1, "y0": base - 7.0, "y1": base + 2.0,
            "base": base, "columns": columns, "text": str(number) + " " + text}


def numbered_listing_document(texts, quote=None):
    """What pdf2docx wrote for a numbered listing: one row, a cell of numbers and a cell of code."""
    document = Document()
    document.sections[0].top_margin = Pt(72)
    lead = document.add_paragraph()
    convert.set_spacing(lead._p, before=0.0, after=8.3, line=13.5)
    lead.add_run(quote or "Step title")
    table = document.add_table(rows=1, cols=2)
    first, second = table.rows[0].cells
    first.width, second.width = Pt(22.6), Pt(324.0)
    for cell, lines, hand in ((first, [str(n + 1) + " " for n in range(len(texts))], "right"),
                              (second, [text + " " for text in texts], "left")):
        paragraph = cell.paragraphs[0]
        convert.set_spacing(paragraph._p, before=3.6, after=0.0, line=10.3)
        paragraph.paragraph_format.left_indent = Pt(4.3)
        for index, line in enumerate(lines):
            run = paragraph.add_run(line)
            run.font.name = "DejaVu Sans Mono"
            if index < len(lines) - 1:
                paragraph.add_run().add_break()
    tail = document.add_paragraph()
    convert.set_spacing(tail._p, before=0.0, after=0.0, line=10.3)
    tail.add_run("Next block")
    return document, table


def run_numbered_panel_checks():
    """A listing with line numbers is one panel again, and a sentence quoting its first line is not its anchor."""
    from collections import Counter

    passed = True
    texts = ["$ nmap -sV -p 22,80 10.10.6.236", "Starting Nmap 7.94 ( https://nmap.org )",
             "PORT    STATE SERVICE  VERSION"]
    bases = [128.97, 139.27, 149.57]
    rows = [numbered_row(index + 1, text, base, 70.0, 74.5, 80.0, 80.0 + 5.0 * len(text))
            for index, (text, base) in enumerate(zip(texts, bases))]
    panel = {"fill": (0.97, 0.97, 0.97), "rect": pymupdf.Rect(60.0, 120.0, 440.0, 160.0), "page": 0,
             "bar": None, "rows": rows}
    passed &= check(convert.split_line_numbers(panel) is not None
                    and convert.split_line_numbers(panel)["gutter"]["numbers"] == ["1", "2", "3"],
                    "numbers: rows that start with 1, 2, 3 in the code's font are a numbered listing")
    skipped = [dict(rows[0]), dict(rows[1], columns=[dict(rows[1]["columns"][0], text="5")] + rows[1]["columns"][1:]),
               dict(rows[2])]
    passed &= check(convert.split_line_numbers(dict(panel, rows=skipped)) is None,
                    "numbers: a column that does not count up by one is not line numbers")
    other = [dict(row, columns=[dict(row["columns"][0], font="Libertinus Serif")] + row["columns"][1:])
             for row in rows]
    passed &= check(convert.split_line_numbers(dict(panel, rows=other)) is None,
                    "numbers: a number in another font than the text beside it is a list of prose")

    document, table = numbered_listing_document(texts)
    declined = Counter()
    painted, anchors = convert.shade_run_of_paragraphs(document, [panel], plain=False, declined=declined)
    passed &= check(not declined and len(anchors) == 1 and len(painted) == 1,
                    "numbers: the listing was matched and its panel built", (dict(declined), len(anchors)))
    tables = document.element.body.findall(convert.qn("w:tbl"))
    passed &= check(len(tables) == 1 and tables[0] is anchors[0][0],
                    "numbers: pdf2docx's two-cell table is gone and the panel's is the only one")
    cells = tables[0].find(convert.qn("w:tr")).findall(convert.qn("w:tc"))
    holder = [child for child in cells[0] if child.tag == convert.qn("w:p")]
    passed &= check(len(cells) == 1 and len(holder) == 1,
                    "numbers: one cell the size of the rectangle holds one paragraph of lines", len(cells))
    lines = ["".join(node.text or "" for node in line).rstrip() for line in convert.code_line_nodes(holder[0])]
    passed &= check(lines == [str(n + 1) + text for n, text in enumerate(texts)],
                    "numbers: every line is its number, then its code, with nothing lost", lines)
    stops = holder[0].find(convert.qn("w:pPr")).find(convert.qn("w:tabs")).findall(convert.qn("w:tab"))
    kinds = [stop.get(convert.qn("w:val")) for stop in stops]
    places = [convert.twips_of(stop.get(convert.qn("w:pos"))) for stop in stops]
    passed &= check(kinds == ["right", "left"] and abs(places[1] - 10.0) < 0.1 and abs(places[1] - places[0] - 5.5) < 0.1,
                    "numbers: a right stop where the numbers end and a left stop where the code starts, as the PDF has them",
                    (kinds, places))
    passed &= check(sum(1 for run in holder[0].iter(convert.qn("w:r")) if run.find(convert.qn("w:tab")) is not None)
                    == 2 * len(texts), "numbers: two tabs on every line")
    indent = convert.twips_of(tables[0].find(convert.qn("w:tblPr")).find(convert.qn("w:tblInd")).get(convert.qn("w:w")))
    # No accent bar, so the text cell is the first cell and Word hangs the
    # table's edge its 10 pt margin before tblInd: the indent is where the
    # numbers start, and the rectangle's edge falls a padding to the left of it.
    passed &= check(abs(indent - (4.3 + 22.6 + convert.DEFAULT_CELL_MARGIN_PT - 10.0)) < 0.1,
                    "numbers: tblInd is where the numbers start, the edge falls the cell's margin before it", indent)
    # A bar-less panel in the body gets the first cell's margin back: Word
    # hangs a body table's edge that margin before tblInd whatever the panel
    # holds (design-b1's and design-b2's terminal windows stood 10 and 9 pt left).
    def table_indent(built):
        return convert.twips_of(built[0].find(convert.qn("w:tblPr")).find(convert.qn("w:tblInd")).get(convert.qn("w:w")))

    def set_alignment(node, hand, indent):
        jc = convert.OxmlElement("w:jc")
        jc.set(convert.qn("w:val"), hand)
        node.get_or_add_pPr().append(jc)
        node.find(convert.qn("w:pPr")).find(convert.qn("w:ind")).set(convert.qn("w:left"), convert.twips_attribute(indent))

    def hands_of(nodes):
        found = [node.find(convert.qn("w:pPr")).find(convert.qn("w:jc")) for node in nodes]
        return [hand.get(convert.qn("w:val")) if hand is not None else None for hand in found]

    document, made = code_document(texts[:2], shaded=True)
    plain_panel = code_panel(115.13, 58.9, lines=list(zip(texts[:2], bases[:2])))
    plain_panel["bar"] = None
    passed &= check(abs(table_indent(convert.build_panel_table(plain_panel, made)) - 25.8) < 0.05,
                    "terminal: a bar-less panel's tblInd is where its text starts, the margin added back")
    # A panel with a bar keeps the indent it had: the bar's cell has no margin.
    document, made = code_document(texts[:2], shaded=True)
    barred = code_panel(115.13, 58.9, lines=list(zip(texts[:2], bases[:2])))
    passed &= check(abs(table_indent(convert.build_panel_table(barred, made)) - (25.8 - 7.2 - 9.0)) < 0.05,
                    "terminal: a panel with an accent bar keeps its indent")
    # A bar-less panel inside a table's cell does not: design-a2's nine stood
    # 8.9 pt right of the PDF's with the margin added.
    document = Document()
    document.sections[0].top_margin = Pt(72)
    outer = document.add_table(rows=1, cols=1).rows[0].cells[0]
    nested = []
    for text in texts[:2]:
        paragraph = outer.add_paragraph()
        convert.set_spacing(paragraph._p, before=3.7, after=0.0, line=9.0)
        paragraph.add_run(text)
        paragraph.paragraph_format.left_indent = Pt(25.8)
        nested.append(paragraph._p)
    cell_panel = code_panel(115.13, 58.9, lines=list(zip(texts[:2], bases[:2])))
    cell_panel["bar"] = None
    passed &= check(abs(table_indent(convert.build_panel_table(cell_panel, nested)) - (25.8 - 9.0)) < 0.05,
                    "terminal: a bar-less panel inside a cell keeps the indent it had")

    # A terminal window pdf2docx read as three blocks: the first line centred
    # with an indent 10 pt short (design-b1 page 20), a wrapped tail
    # right-aligned. The PDF starts both at the panel's text edge. The indent
    # comes from a line that kept its alignment, and neither line is left
    # centred or right-aligned in the cell.
    document, made = code_document(texts[:3], shaded=True)
    set_alignment(made[0], "center", 15.8)
    set_alignment(made[2], "right", 5.0)
    window = code_panel(115.13, 58.9, lines=list(zip(texts[:3], bases[:3])))
    window["bar"] = None
    passed &= check(abs(table_indent(convert.build_panel_table(window, made)) - 25.8) < 0.05,
                    "terminal: a centred first line does not give the table its indent")
    passed &= check(hands_of(made) == ["left", None, "left"],
                    "terminal: the centred and the right-aligned line are left-aligned in the cell", hands_of(made))

    # A shaded box whose title the PDF centres, over left-aligned lines, a
    # tail at the edge and an indented line pdf2docx centred: the title stays
    # centred, the other two go left. The panel's text area is 80.2 to 539.0,
    # so the title's middle is 309.6; the indented line starts four
    # characters in with no symmetry.
    box_texts = [texts[0], texts[1], "SMB 10.10.4.33 445", "    indented line of code"]
    box_bases = [128.97, 139.27, 149.57, 159.87]
    document, made = code_document(box_texts, shaded=True)
    set_alignment(made[0], "center", 190.0)
    set_alignment(made[2], "right", 5.0)
    set_alignment(made[3], "center", 40.0)
    box = code_panel(115.13, 58.9, lines=list(zip(box_texts, box_bases)))
    box["bar"] = None
    for row, (start, end) in zip(box["rows"], ((250.0, 369.2), (80.2, 280.0), (80.2, 200.0), (98.2, 250.0))):
        row["x0"], row["x1"] = start, end
        row["inkX0"] = start
    box["rows"][1]["inkX0"] = 80.2
    passed &= check(convert.panel_text_left(box) == 80.2, "terminal: the box's text edge is its leftmost row")
    passed &= check([convert.drifted_alignment(node, box) for node in made] == [False, False, True, True],
                    "terminal: only the tail at the edge and the unsymmetric indented line are drifted",
                    [convert.drifted_alignment(node, box) for node in made])
    built = convert.build_panel_table(box, made)
    passed &= check(built is not None and hands_of(made) == ["center", None, "left", "left"],
                    "terminal: the title the PDF centres stays centred in the cell", hands_of(made))
    passed &= check(built is not None and abs(table_indent(built) - 25.8) < 0.05,
                    "terminal: and does not give the box its indent either")

    right_to_left = convert.OxmlElement("w:p")
    right_to_left.get_or_add_pPr().append(convert.OxmlElement("w:bidi"))
    jc = convert.OxmlElement("w:jc")
    jc.set(convert.qn("w:val"), "right")
    right_to_left.get_or_add_pPr().append(jc)
    passed &= check(not convert.pushed_alignment(right_to_left) and not convert.pushed_alignment(made[1]),
                    "terminal: right-to-left 'right' is the start edge and a plain line is not pushed")
    unmatched = code_document(["no row has this text"], shaded=True)[1][0]
    set_alignment(unmatched, "center", 30.0)
    passed &= check(not convert.drifted_alignment(unmatched, box),
                    "terminal: a centred line no PDF row matches keeps its alignment")

    # On the rung that moves nothing the listing stays pdf2docx's table.
    document, table = numbered_listing_document(texts)
    declined = Counter()
    painted, anchors = convert.shade_run_of_paragraphs(document, [panel], plain=True, declined=declined)
    passed &= check(not painted and dict(declined) == {"panels-unmatched": 1}
                    and document.element.body.find(convert.qn("w:tbl")) is table._tbl,
                    "numbers: the plain rung leaves the table and counts the panel unmatched", dict(declined))

    # A sentence that names the first command is not where the panel begins.
    text = ["import requests", "r = requests.get(url)", "print(r.status_code)"]
    bases = [128.97, 154.34, 167.03]
    document, made = code_document(text, shaded=False)
    prose = document.add_paragraph()
    convert.set_spacing(prose._p, before=0.0, after=6.0, line=13.5)
    prose.add_run("First we ran import requests in a shell and read the reply.")
    made[0].addprevious(prose._p)
    panel = code_panel(115.13, 58.9, lines=list(zip(text, bases)))
    declined = Counter()
    painted, anchors = convert.shade_run_of_paragraphs(document, [panel], plain=False, declined=declined)
    passed &= check(prose._p.getparent() is document.element.body and len(anchors) == 1
                    and [convert.paragraph_text(node) for node in anchors[0][0].iter(convert.qn("w:p"))
                         if convert.paragraph_text(node)] == text,
                    "quote: the sentence stays outside the panel and the panel holds the three code lines",
                    (dict(declined), len(anchors)))
    passed &= run_numbered_gap_checks()

    # Steps in a prose font pass every test split_line_numbers has. When the
    # fold cannot be made (here pdf2docx kept each step in one paragraph, so
    # there are no number cells to take apart) the callout used to stay
    # unbuilt and unfilled; it now goes down the ordinary path, with its
    # rows as read, numbers and all, and the text is the converter's own.
    steps = ["Disable SMBv1 on every file server", "Enforce LDAP signing on all domain controllers",
             "Rotate every service account password"]
    prose_rows = []
    for index, (step, base) in enumerate(zip(steps, [128.97, 139.27, 149.57])):
        row = numbered_row(index + 1, step, base, 70.0, 74.5, 80.0, 80.0 + 4.0 * len(step))
        row["columns"] = [dict(column, font="Libertinus Serif") for column in row["columns"]]
        prose_rows.append(row)
    callout = {"fill": (0.93, 0.95, 0.97), "rect": pymupdf.Rect(60.0, 120.0, 440.0, 160.0), "page": 0,
               "bar": None, "rows": prose_rows}
    passed &= check(convert.split_line_numbers(callout) is not None,
                    "numbers: steps in a prose font are read as a numbered listing, the case the fallback is for")
    document, made = code_document(["%d %s" % (index + 1, step) for index, step in enumerate(steps)], shaded=False)
    declined = Counter()
    painted, anchors = convert.shade_run_of_paragraphs(document, [callout], plain=False, declined=declined)
    texts = [convert.paragraph_text(node) for anchor in anchors for node in anchor[0].iter(convert.qn("w:p"))]
    passed &= check(not declined and len(anchors) == 1 and texts == ["%d %s" % (index + 1, step) for index, step in enumerate(steps)]
                    and not any(run.find(convert.qn("w:tab")) is not None for run in document.element.body.iter(convert.qn("w:r"))),
                    "numbers: a fold that cannot be made builds the callout the ordinary way, text intact and no stops added",
                    (dict(declined), len(anchors), texts))
    # The second attempt does not leave the first one's claim on the paragraphs:
    # a later panel on the same page still finds its own.
    document, made = code_document(["%d %s" % (index + 1, step) for index, step in enumerate(steps)]
                                   + ["$ nmap -sV 10.10.6.236", "PORT STATE SERVICE"], shaded=False)
    later = code_panel(165.0, 30.0, lines=[("$ nmap -sV 10.10.6.236", 175.0), ("PORT STATE SERVICE", 185.0)])
    later["page"] = 0
    declined = Counter()
    painted, anchors = convert.shade_run_of_paragraphs(document, [callout, later], plain=False, declined=declined)
    passed &= check(not declined and len(anchors) == 2,
                    "numbers: and the panel after it is still built", (dict(declined), len(anchors)))
    return passed

def blank_numbered_row(number, base, number_x0, number_x1):
    """The row of a blank code line: its number and nothing else."""
    column = {"x0": number_x0, "x1": number_x1, "inkX0": number_x0, "y0": base - 7.0, "y1": base + 2.0,
              "base": base, "text": str(number), "font": "DejaVu Sans Mono"}
    return {"x0": number_x0, "inkX0": number_x0, "x1": number_x1, "y0": base - 7.0, "y1": base + 2.0,
            "base": base, "columns": [column], "text": str(number)}


def wrapped_row(text, base, x0):
    """The row a wrapped line continues on: no number, and it starts at the block's edge."""
    column = {"x0": x0, "x1": x0 + 5.0 * len(text), "inkX0": x0, "y0": base - 7.0, "y1": base + 2.0,
              "base": base, "text": text, "font": "DejaVu Sans Mono"}
    return {"x0": x0, "inkX0": x0, "x1": x0 + 5.0 * len(text), "y0": base - 7.0, "y1": base + 2.0,
            "base": base, "columns": [column], "text": text}


def numbered_tables_document(tables):
    """What pdf2docx wrote for a numbered listing it cut into tables.

    `tables` is a list of tables, each a list of rows, each row a pair of
    the number cell's lines (None for a row of one cell, code alone) and the
    code cell's lines.
    """
    document = Document()
    document.sections[0].top_margin = Pt(72)
    lead = document.add_paragraph()
    convert.set_spacing(lead._p, before=0.0, after=8.3, line=13.5)
    lead.add_run("Step title")
    for rows in tables:
        table = document.add_table(rows=len(rows), cols=2)
        for row, (marks, texts) in zip(table.rows, rows):
            first, second = row.cells
            first.width, second.width = Pt(22.6), Pt(324.0)
            if marks is None:
                row._tr.remove(first._tc)
                holders = ((second, texts),)
            else:
                holders = ((first, marks), (second, texts))
            for cell, lines in holders:
                paragraph = cell.paragraphs[0]
                convert.set_spacing(paragraph._p, before=3.6, after=0.0, line=10.3)
                paragraph.paragraph_format.left_indent = Pt(4.3)
                for index, line in enumerate(lines):
                    run = paragraph.add_run(line)
                    run.font.name = "DejaVu Sans Mono"
                    if index < len(lines) - 1:
                        paragraph.add_run().add_break()
    tail = document.add_paragraph()
    convert.set_spacing(tail._p, before=0.0, after=0.0, line=10.3)
    tail.add_run("Next block")
    return document


def run_numbered_gap_checks():
    """A listing with a blank line or a wrapped line is folded, each number once and beside its own line."""
    from collections import Counter
    from lxml import etree

    passed = True
    base = [128.97, 139.27, 149.57, 159.87, 170.17]

    def panel_for(rows):
        return {"fill": (0.97, 0.97, 0.97), "rect": pymupdf.Rect(60.0, 120.0, 440.0, 190.0),
                "page": 0, "bar": None, "rows": rows}

    def request_rows():
        made = [numbered_row(index + 1, text, base[index], 70.0, 74.5, 80.0, 80.0 + 5.0 * len(text))
                for index, text in enumerate(http)]
        return made + [blank_numbered_row(4, base[3], 70.0, 74.5),
                       numbered_row(5, body, base[4], 70.0, 74.5, 80.0, 165.0)]

    http = ["POST /login HTTP/1.1", "Host: example.test", "Content-Length: 17"]
    body = '{"user": "admin"}'
    rows = request_rows()
    split = convert.split_line_numbers(panel_for(rows))
    passed &= check(split is not None and split["gutter"]["numbers"] == ["1", "2", "3", "4", "5"]
                    and [line["blank"] for line in split["gutter"]["lines"]] == [False, False, False, True, False]
                    and len(split["rows"]) == 4,
                    "numbers: a row with its number alone is a blank code line, kept out of the code rows",
                    split and split["gutter"]["lines"])
    skipped = list(rows[:4]) + [dict(rows[4], columns=[dict(rows[4]["columns"][0], text="7")] + rows[4]["columns"][1:])]
    passed &= check(convert.split_line_numbers(panel_for(skipped)) is None,
                    "numbers: a column that skips a number past a blank line is not line numbers")
    stray = list(rows[:4]) + [blank_numbered_row(9, base[4], 70.0, 74.5)]
    passed &= check(convert.split_line_numbers(panel_for(stray)) is None,
                    "numbers: a lone number that is not the next one is not a blank line")
    passed &= check(convert.split_line_numbers(panel_for([blank_numbered_row(1, base[0], 70.0, 74.5)] + rows[1:])) is None,
                    "numbers: the first row of a listing has code")
    indented = list(rows[:3]) + [wrapped_row("stray words", base[3], 120.0)]
    passed &= check(convert.split_line_numbers(panel_for(indented)) is None,
                    "numbers: a row without a number that starts right of the code column is not a continuation")

    # The request as pdf2docx wrote it: the rows before the blank line in one
    # table, the blank line's number and the body in the next.
    texts = [text + " " for text in http]
    tables = [[(["1 ", "2 ", "3 "], texts)], [(["4 ", "5"], [body])]]
    document = numbered_tables_document(tables)
    declined = Counter()
    painted, anchors = convert.shade_run_of_paragraphs(document, [panel_for(rows)], plain=False, declined=declined)
    found = document.element.body.findall(convert.qn("w:tbl"))
    passed &= check(not declined and len(anchors) == 1 and len(found) == 1 and found[0] is anchors[0][0],
                    "numbers: the request's two tables are one panel table", (dict(declined), len(found)))
    holder = [node for node in found[0].iter(convert.qn("w:p"))] if found else []
    lines = ["".join(node.text or "" for node in line).rstrip() for line in convert.code_line_nodes(holder[0])] if holder else []
    passed &= check(len(holder) == 1 and lines == ["1" + http[0], "2" + http[1], "3" + http[2], "4", "5" + body],
                    "numbers: each number is there once and the body stands beside its own", lines)

    # A line typst wrapped: the first piece, the rest in a table of its own, then
    # a row of code alone above two rows of a number and its code.
    long_line = "curl -sk https://example.test/api/v2/export?format=csv&"
    rest = "include=all_fields -H 'X: y'"
    rows = [numbered_row(1, long_line, base[0], 70.0, 74.5, 80.0, 80.0 + 5.0 * len(long_line)),
            wrapped_row(rest, base[1], 60.0),
            numbered_row(2, "id", base[2], 70.0, 74.5, 80.0, 90.0),
            numbered_row(3, "uname -a", base[3], 70.0, 74.5, 80.0, 120.0)]
    split = convert.split_line_numbers(panel_for(rows))
    passed &= check(split is not None and [line["number"] for line in split["gutter"]["lines"]] == ["1", None, "2", "3"]
                    and abs(split["gutter"]["reach"] - 60.0) < 0.01 and len(split["gutter"]["continued"]) == 1,
                    "numbers: a row without a number, left of the code, continues the line above",
                    split and split["gutter"]["lines"])
    tables = [[(["1 "], [long_line])], [(None, [rest])], [(["2 "], ["id"]), (["3 "], ["uname -a"])]]
    document = numbered_tables_document(tables)
    declined = Counter()
    painted, anchors = convert.shade_run_of_paragraphs(document, [panel_for(rows)], plain=False, declined=declined)
    found = document.element.body.findall(convert.qn("w:tbl"))
    holder = [node for node in found[0].iter(convert.qn("w:p"))] if len(found) == 1 else []
    lines = ["".join(node.text or "" for node in line).rstrip() for line in convert.code_line_nodes(holder[0])] if holder else []
    passed &= check(not declined and len(holder) == 1
                    and lines == ["1" + long_line, rest, "2id", "3uname -a"],
                    "numbers: a wrapped line is folded with its rest under it and no number of its own", (dict(declined), lines))
    if holder:
        tabs = sum(1 for run in holder[0].iter(convert.qn("w:r")) if run.find(convert.qn("w:tab")) is not None)
        indent = convert.paragraph_indent(holder[0])
        stops = holder[0].find(convert.qn("w:pPr")).find(convert.qn("w:tabs")).findall(convert.qn("w:tab"))
        places = [convert.twips_of(stop.get(convert.qn("w:pos"))) for stop in stops]
        # The wrapped line starts at x 60 and the panel's table puts its text edge there, so
        # the paragraph has no indent left and the stops are the PDF's: numbers end at 74.5, code at 80.
        passed &= check(tabs == 6 and abs(indent) < 0.1 and len(places) == 2
                        and abs(places[0] - 14.5) < 0.1 and abs(places[1] - 20.0) < 0.1,
                        "numbers: the text edge is where the wrapped line starts and the stops are the PDF's, in order",
                        (tabs, indent, places))

    # The same line with its numbers in a column of their own (a grid): the
    # rest of the wrapped line starts in the code column, 80, not at the
    # block's edge, and the fold gives it two tabs (the right stop over
    # nothing, then the code's stop) where it used to write it at the paragraph's
    # edge, 10 pt left of the PDF, under the numbers (numbers2.typ: 12.6).
    rows = [numbered_row(1, long_line, base[0], 70.0, 74.5, 80.0, 80.0 + 5.0 * len(long_line)),
            wrapped_row(rest, base[1], 80.05),
            numbered_row(2, "id", base[2], 70.0, 74.5, 80.0, 90.0),
            numbered_row(3, "uname -a", base[3], 70.0, 74.5, 80.0, 120.0)]
    split = convert.split_line_numbers(panel_for(rows))
    passed &= check(split is not None and [line.get("atCode") for line in split["gutter"]["lines"]] == [None, True, None, None]
                    and abs(split["gutter"]["reach"] - 80.05) < 0.01,
                    "numbers: a row without a number that starts in the code column is a grid listing's continuation",
                    split and split["gutter"]["lines"])
    box_rows = list(rows[:1]) + [wrapped_row(rest, base[1], 60.0)] + list(rows[2:])
    split = convert.split_line_numbers(panel_for(box_rows))
    passed &= check(split is not None and not any(line.get("atCode") for line in split["gutter"]["lines"]),
                    "numbers: and one at the block's edge is a boxed number's, which has no code-column start")
    tables = [[(["1 "], [long_line, rest]), (["2 ", "3 "], ["id", "uname -a"])]]
    document = numbered_tables_document(tables)
    declined = Counter()
    painted, anchors = convert.shade_run_of_paragraphs(document, [panel_for(rows)], plain=False, declined=declined)
    found = document.element.body.findall(convert.qn("w:tbl"))
    holder = [node for node in found[0].iter(convert.qn("w:p"))] if len(found) == 1 else []
    lines = ["".join(node.text or "" for node in line).rstrip() for line in convert.code_line_nodes(holder[0])] if holder else []
    passed &= check(not declined and len(holder) == 1 and lines == ["1" + long_line, rest, "2id", "3uname -a"],
                    "numbers: a grid listing's wrapped line is folded with its rest under it", (dict(declined), lines))
    if holder:
        runs = list(holder[0].findall(convert.qn("w:r")))
        stop_at = [i for i, run in enumerate(runs) if run.find(convert.qn("w:br")) is not None]
        after_break = runs[stop_at[0] + 1:stop_at[0] + 4] if stop_at else []
        starts = [run.find(convert.qn("w:tab")) is not None for run in after_break]
        tabs = sum(1 for run in runs if run.find(convert.qn("w:tab")) is not None)
        stops = holder[0].find(convert.qn("w:pPr")).find(convert.qn("w:tabs")).findall(convert.qn("w:tab"))
        places = [convert.twips_of(stop.get(convert.qn("w:pos"))) for stop in stops]
        # Two tabs open the continuation. The table's text edge is the numbers' (70), so the
        # stops are the PDF's 74.5 and 80 measured from it.
        passed &= check(starts == [True, True, False] and tabs == 8 and len(places) == 2
                        and abs(places[0] - 4.5) < 0.1 and abs(places[1] - 10.0) < 0.1,
                        "numbers: the continuation opens with two tabs, so it stands at the code stop and not under the numbers",
                        (starts, tabs, places))

    # A refusal leaves the document as it was found: numbers that are not the
    # PDF's, or code with a line fewer than the PDF has.
    for name, tables in (("numbers that are not the PDF's", [[(["1 ", "2 ", "9 "], texts)], [(["4 ", "5"], [body])]]),
                         ("a line fewer than the PDF has", [[(["1 ", "2 ", "3 "], texts[:2])], [(["4 ", "5"], [body])]])):
        document = numbered_tables_document(tables)
        before = etree.tostring(document.element.body)
        declined = Counter()
        painted, anchors = convert.shade_run_of_paragraphs(document, [panel_for(request_rows())], plain=False,
                                                           declined=declined)
        passed &= check(dict(declined) == {"panels-unmatched": 1} and etree.tostring(document.element.body) == before,
                        "numbers: " + name + " leaves the tables exactly as they were", dict(declined))
    return passed


def run_rhythm_checks():
    """The vertical pass: its layout model, and what it refuses to do."""
    passed = True

    # The model itself, which is the whole of the correction.
    metrics = convert.paragraph_metrics(rhythm_document(6.0, 3.0, 15.0).paragraphs[1]._p)
    passed &= check(metrics is not None and abs(metrics["before"] - 6.0) < 0.05
                    and abs(metrics["after"] - 3.0) < 0.05 and abs(metrics["line"] - 15.0) < 0.05,
                    "rhythm: spacing is read back in points", metrics)
    inherited = Document().add_paragraph("no spacing of its own")
    passed &= check(convert.paragraph_metrics(inherited._p) is None,
                    "rhythm: a paragraph without an exact line height is not modelled")

    rows = [{"base": 100.0, "y0": 92.0, "y1": 103.0, "key": "a"},
            {"base": 116.0, "y0": 108.0, "y1": 119.0, "key": "b"}]
    line, floor = convert.wanted_line_height(rows, 19.6)
    passed &= check(abs(line - 16.0) < 0.01, "rhythm: the line height is the PDF's own pitch", line)
    passed &= check(abs(floor - 10.0) < 0.01, "rhythm: the floor holds the row's ascent", floor)
    single, single_floor = convert.wanted_line_height(rows[:1], 19.6)
    passed &= check(abs(single - 19.6) < 0.01, "rhythm: one row keeps the converter's height", single)
    passed &= check(single_floor <= single, "rhythm: the floor never raises a height", single_floor)
    passed &= check(abs(convert.first_baseline_target(rows, 16.0) - 100.0) < 0.01,
                    "rhythm: an even block is placed on its first line")
    staggered = [dict(rows[0]), dict(rows[1], base=120.0)]
    passed &= check(convert.first_baseline_target(staggered, 16.0) > 100.0,
                    "rhythm: an uneven block straddles its lines")

    # The correction, end to end on a document whose PDF says where to put it.
    pdf_rows = [[anchor_row("First block here", 100.0, 92.0, 103.0),
                 anchor_row("Second block here", 160.0, 152.0, 163.0)]]
    document = rhythm_document(0.0, 0.0, 15.0)
    moved = convert.align_vertical_rhythm(document, pdf_rows, 792.0)
    after = convert.paragraph_metrics(document.paragraphs[1]._p)
    # 72 top margin, a 12 pt first box, then the gap that puts the second
    # baseline at 160: 160 - 0.8 x 15 - 84.
    passed &= check(moved == 1, "rhythm: one block was moved", moved)
    passed &= check(after is not None and abs(after["before"] - 64.0) < 0.3,
                    "rhythm: the spacing is what the PDF's baseline asks for", after and after["before"])

    # A gap the PDF makes negative comes out of the block above first.
    tight = [[anchor_row("First block here", 100.0, 92.0, 103.0),
              anchor_row("Second block here", 92.0, 84.0, 95.0)]]
    document = rhythm_document(0.0, 20.0, 15.0)
    convert.align_vertical_rhythm(document, tight, 792.0)
    above = convert.paragraph_metrics(document.paragraphs[0]._p)
    below = convert.paragraph_metrics(document.paragraphs[1]._p)
    passed &= check(above is not None and above["after"] == 0.0,
                    "rhythm: the block above gives its spacing back first", above and above["after"])
    passed &= check(below is not None and below["before"] == 0.0 and below["line"] < 15.0,
                    "rhythm: a box too tall for its place is shortened", below and below["line"])

    # And a page that would end below its bottom margin keeps what it had.
    far = [[anchor_row("First block here", 100.0, 92.0, 103.0),
            anchor_row("Second block here", 900.0, 892.0, 903.0)]]
    document = rhythm_document(2.0, 0.0, 15.0)
    moved = convert.align_vertical_rhythm(document, far, 792.0)
    kept = convert.paragraph_metrics(document.paragraphs[1]._p)
    passed &= check(moved == 0 and kept is not None and abs(kept["before"] - 2.0) < 0.05,
                    "rhythm: a correction past the bottom margin is refused", kept and kept["before"])

    # A table stands as tall as its rows plus the border under the last one,
    # and the converter writes that width as "12.0", which is not an integer.
    document = Document()
    table = document.add_table(rows=2, cols=1)
    for order, row in enumerate(table.rows):
        properties = row._tr.get_or_add_trPr()
        node = convert.OxmlElement("w:trHeight")
        node.set(convert.qn("w:val"), "400")
        node.set(convert.qn("w:hRule"), "exact")
        properties.append(node)
        if order:
            borders = convert.OxmlElement("w:tcBorders")
            bottom = convert.OxmlElement("w:bottom")
            bottom.set(convert.qn("w:val"), "single")
            bottom.set(convert.qn("w:sz"), "12.0")
            borders.append(bottom)
            row.cells[0]._tc.get_or_add_tcPr().append(borders)
    metrics = convert.table_metrics(table._tbl)
    passed &= check(metrics is not None and abs(metrics["height"] - 41.5) < 0.05,
                    "rhythm: a table is its rows plus its last border", metrics and metrics["height"])
    for row in table.rows:
        row._tr.find(convert.qn("w:trPr")).find(convert.qn("w:trHeight")).set(convert.qn("w:hRule"), "atLeast")
    passed &= check(convert.table_metrics(table._tbl) is None,
                    "rhythm: a table Word is free to size is not modelled")

    # Word keeps the larger of the spacing after one paragraph and the
    # spacing before the next, so the one above has to give its spacing back
    # before the one below can come any higher, even when the gap it wants is
    # positive.
    collapsing = [[anchor_row("First block here", 100.0, 92.0, 103.0),
                   anchor_row("Second block here", 112.0, 104.0, 115.0)]]
    document = rhythm_document(0.0, 0.0, 15.0)
    convert.set_spacing(document.paragraphs[0]._p, after=20.0)
    convert.align_vertical_rhythm(document, collapsing, 792.0)
    above = convert.paragraph_metrics(document.paragraphs[0]._p)
    below = convert.paragraph_metrics(document.paragraphs[1]._p)
    # 72 top margin, a 12 pt box, then 112 - 0.8 x 15 - 84 = 16 of gap, which
    # is less than the 20 above it and therefore unreachable until that goes.
    passed &= check(above is not None and above["after"] == 0.0,
                    "rhythm: a larger spacing above is given back, not added to",
                    above and above["after"])
    passed &= check(below is not None and abs(below["before"] - 16.0) < 0.3,
                    "rhythm: and the block below asks for the whole gap itself",
                    below and below["before"])

    passed &= run_cell_checks()
    passed &= run_table_place_checks()
    passed &= run_table_box_checks()
    passed &= run_picture_walk_checks()
    passed &= run_pitch_split_checks()
    passed &= run_panel_checks()

    # A word the typesetter broke keeps its hyphen where the line still ends.
    document = Document()
    first = document.add_paragraph()
    first.add_run("Therefore, the company also em")
    second = document.add_paragraph()
    second.add_run("ployed a custom system")
    broken = {(convert.compare_squash("Therefore, the company also em")[-convert.PAIR_WINDOW:],
               convert.compare_squash("ployed a custom system")[:convert.PAIR_WINDOW])}
    added = convert.restore_break_hyphens(document, broken)
    ends = list(first._p.iter(convert.qn("w:t")))[-1]
    passed &= check(added == 1 and ends.text.endswith("em-"),
                    "rhythm: the break hyphen came back", (added, ends.text[-6:]))
    convert.restore_break_hyphens(document, broken)
    passed &= check(ends.text.endswith("em-") and not ends.text.endswith("em--"),
                    "rhythm: and only once", ends.text[-6:])
    elsewhere = Document()
    one = elsewhere.add_paragraph()
    one.add_run("a line that was not broken")
    elsewhere.add_paragraph().add_run("and the next one")
    passed &= check(convert.restore_break_hyphens(elsewhere, broken) == 0,
                    "rhythm: a line end the PDF did not break gets nothing")

    # A code panel's accent bar stands off the text by the panel's padding.
    document = Document()
    paragraph = document.add_paragraph("nmap -sV")
    convert.left_border(paragraph, {"fill": (0.4, 0.4, 0.4), "width": 7.2, "space": 9.0})
    left = paragraph._p.find(convert.qn("w:pPr")).find(convert.qn("w:pBdr")).find(convert.qn("w:left"))
    passed &= check(left.get(convert.qn("w:sz")) == "48" and left.get(convert.qn("w:space")) == "9",
                    "rhythm: the accent bar's width and stand-off come from the PDF",
                    (left.get(convert.qn("w:sz")), left.get(convert.qn("w:space"))))
    convert.left_border(paragraph, {"fill": (0.4, 0.4, 0.4), "width": 7.2, "space": 90.0})
    left = paragraph._p.find(convert.qn("w:pPr")).find(convert.qn("w:pBdr")).find(convert.qn("w:left"))
    passed &= check(left.get(convert.qn("w:space")) == "31",
                    "rhythm: a stand-off stays inside what Word accepts", left.get(convert.qn("w:space")))

    # A tab is already a gap, so the space repair may not add one across it.
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("2.1")
    paragraph.add_run("\t")
    paragraph.add_run("NON-DISCLOSURE STATEMENT")
    nodes = list(paragraph._p.iter(convert.qn("w:t")))
    passed &= check(convert.adjacent_text(paragraph._p) == [],
                    "rhythm: a tab between two runs is not a run boundary",
                    len(convert.adjacent_text(paragraph._p)))
    convert.repair_text(document, {(convert.compare_squash("2.1")[-8:],
                                    convert.compare_squash("NON-DISC")[:8])})
    passed &= check(nodes[1].text == "NON-DISCLOSURE STATEMENT",
                    "rhythm: the title after a tab keeps its place", nodes[1].text[:12])
    joined = document.add_paragraph()
    joined.add_run("1.")
    joined.add_run("{{STEP}}")
    pieces = list(joined._p.iter(convert.qn("w:t")))
    convert.repair_text(document, {(convert.compare_squash("1.")[-8:],
                                    convert.compare_squash("{{STEP}}")[:8])})
    passed &= check(pieces[1].text.startswith(" "),
                    "rhythm: a run boundary with no tab still gets its space", pieces[1].text)

    # A page count the converter did not write one section per is left alone.
    document = rhythm_document(2.0, 0.0, 15.0)
    passed &= check(convert.align_vertical_rhythm(document, pdf_rows + pdf_rows, 792.0) == 0,
                    "rhythm: sections and pages have to line up")
    return passed


def run_lost_walk_checks():
    """The vertical pass does not pay for lines it never counted, nor let a footnote past the margin."""
    passed = True

    def build(second_text):
        document = Document()
        document.sections[0].top_margin = Pt(72)
        document.sections[0].bottom_margin = Pt(72)
        exact_paragraph(document, "First block here", 0.0, line=12.0)
        exact_paragraph(document, second_text, 5.0)
        exact_paragraph(document, "Third block here", 3.0)
        return document

    def rows(with_second):
        page = [anchor_row("First block here", 100.0, 92.0, 103.0)]
        if with_second:
            page.append(anchor_row("A paragraph the page holds", 120.0, 112.0, 123.0))
        page.append(anchor_row("Third block here", 200.0, 192.0, 203.0))
        return [page]

    # The middle paragraph's rows are there: the third block is placed on its
    # own baseline, 200 - 0.8 x 15 less the 123 where the middle one ends.
    document = build("A paragraph the page holds")
    declined = convert.Counter()
    convert.align_vertical_rhythm(document, rows(True), 792.0, declined=declined)
    third = convert.paragraph_metrics(document.paragraphs[2]._p)
    passed &= check(third is not None and abs(third["before"] - 65.0) < 0.3 and not declined["rhythm-lost-walk"],
                    "lost walk: with its rows found the block under is placed on its baseline", third and third["before"])

    # The middle paragraph is text the page's rows do not hold (Word wraps it
    # into lines the converter counted as one). The cursor is behind by those
    # lines, so the 84 pt gap the PDF's baseline asks for would be too wide;
    # the block under it keeps the converter's 3 pt and the pass says so.
    document = build("Some prose the rows never listed")
    declined = convert.Counter()
    convert.align_vertical_rhythm(document, rows(False), 792.0, declined=declined)
    third = convert.paragraph_metrics(document.paragraphs[2]._p)
    passed &= check(third is not None and abs(third["before"] - 3.0) < 0.05 and declined["rhythm-lost-walk"] == 1,
                    "lost walk: a block found again after an unmatched one keeps the converter's gap",
                    (third and third["before"], dict(declined)))

    # The last block with text on a page carries the page's slack as spacing
    # before. Its bottom is held inside the page's own guard, so drift above
    # it does not push it onto a page of its own.
    document = Document()
    document.sections[0].top_margin = Pt(72)
    document.sections[0].bottom_margin = Pt(72)
    exact_paragraph(document, "First block here", 0.0, line=12.0)
    exact_paragraph(document, "Footnote line here", 0.0)
    page = [[anchor_row("First block here", 100.0, 92.0, 103.0),
             anchor_row("Footnote line here", 720.0, 712.0, 723.0)]]
    declined = convert.Counter()
    convert.align_vertical_rhythm(document, page, 792.0, declined=declined)
    note = convert.paragraph_metrics(document.paragraphs[1]._p)
    # 84 + gap + 15 = 792 - 72 - 6 puts the gap at 615; the PDF asked for 624.
    passed &= check(note is not None and abs(note["before"] - 615.0) < 0.3 and declined["rhythm-last-block-capped"] == 1,
                    "last block: a footnote's spacing before stops at the guard above the margin",
                    (note and note["before"], dict(declined)))
    document = Document()
    document.sections[0].top_margin = Pt(72)
    document.sections[0].bottom_margin = Pt(72)
    exact_paragraph(document, "First block here", 0.0, line=12.0)
    exact_paragraph(document, "Footnote line here", 0.0)
    page = [[anchor_row("First block here", 100.0, 92.0, 103.0),
             anchor_row("Footnote line here", 700.0, 692.0, 703.0)]]
    declined = convert.Counter()
    convert.align_vertical_rhythm(document, page, 792.0, declined=declined)
    note = convert.paragraph_metrics(document.paragraphs[1]._p)
    passed &= check(note is not None and abs(note["before"] - 604.0) < 0.3 and not declined["rhythm-last-block-capped"],
                    "last block: a footnote that already ends inside the guard is not moved",
                    (note and note["before"], dict(declined)))

    # A panel's "}" that pdf2docx made the first line of the step's text cell
    # under it, behind the cell holding the step's number: the panel takes it
    # back, and the step keeps its own text. Left there it stood 30 pt below
    # the panel it closed and pushed the step 11.5 pt late (design-b2, page 13).
    rows = ["{", '"host": "10.10.9.85",', '"ports": [22, 80, 443],', "}"]
    bases = [128.97, 141.67, 154.37, 167.07]
    def stray_document(first_cell):
        document, made = code_document(rows[:3], before=14.5, after=0.0, line=9.0)
        step = document.add_table(rows=1, cols=2)
        step.cell(0, 0).paragraphs[0].add_run("4")
        cell = step.cell(0, 1)
        cell.paragraphs[0].add_run(first_cell)
        convert.set_spacing(cell.paragraphs[0]._p, before=1.9, after=0.0, line=9.0)
        cell.add_paragraph("Detraxisse ut aliquam ex eo est")
        return document, step
    document, step = stray_document("}")
    declined = convert.Counter()
    convert.shade_run_of_paragraphs(document, [code_panel(115.13, 12.7 * 4 + 6.0, lines=list(zip(rows, bases)))],
                                    declined=declined)
    panel_table = document.element.body.findall(convert.qn("w:tbl"))[0]
    inside = [convert.paragraph_text(p) for p in panel_table.iter(convert.qn("w:p"))]
    left = [convert.paragraph_text(p) for p in step.cell(0, 1)._tc.findall(convert.qn("w:p"))]
    passed &= check("}" in inside and inside.index("}") == len(inside) - 1 and left == ["Detraxisse ut aliquam ex eo est"]
                    and declined["panels-stray-tail"] == 1,
                    "panel: a closing row in the next step's cell goes back into the panel",
                    (inside, left, dict(declined)))
    # Only that shape: a "}" that is the step's own text is not the panel's.
    document, step = stray_document("{")
    declined = convert.Counter()
    convert.shade_run_of_paragraphs(document, [code_panel(115.13, 12.7 * 4 + 6.0, lines=list(zip(rows, bases)))],
                                    declined=declined)
    passed &= check(declined["panels-stray-tail"] == 0
                    and [convert.paragraph_text(p) for p in step.cell(0, 1)._tc.findall(convert.qn("w:p"))][0] == "{",
                    "panel: and a different short line there stays where it is", dict(declined))
    return passed


def run_comparison_checks():
    """The text check, on text rather than on documents.

    These are the shapes it has to be blind to (a re-flow) and the one shape
    it has to see (a line that is gone).
    """
    passed = True
    line = lambda text: {"text": text, "pieces": [text]}  # noqa: E731

    joined = "The quick brown fox\njumped over the lazy dog"
    passed &= check(
        convert.missing_from_docx([line("The quick brown"), line("fox jumped over")], joined) == [],
        "a line joined to its neighbour is not missing",
    )
    passed &= check(
        convert.missing_from_docx([line("em-"), line("ployed a custom system")], "employed a custom system") == [],
        "a word broken across two lines is not missing",
    )
    passed &= check(
        convert.missing_from_docx([line("Title . . . . . . . . 12")], "Title\t12") == [],
        "a dot leader drawn from a tab stop is not missing",
    )
    passed &= check(
        convert.missing_from_docx([line("1. {{RECOMMENDATION}}")], "1.{{RECOMMENDATION}}") == [],
        "a space the converter did not write is not missing",
    )
    gone = convert.missing_from_docx(
        [line("Host-111 is vulnerable"), line("Host-112 is vulnerable")],
        "Host-112 is vulnerable",
    )
    passed &= check([g for g in gone] == ["Host-111 is vulnerable"], "a dropped paragraph is caught", gone)
    passed &= check(
        convert.missing_from_docx([{"text": "TECHNIQUES MITIGATIONS", "pieces": ["TECHNIQUES", "MITIGATIONS"]}],
                                  "TECHNIQUES\n{{T}}\nMITIGATIONS\n{{M}}") == [],
        "two columns written one after the other are not missing",
    )
    passed &= check(
        convert.missing_from_docx([{"text": "LOW", "pieces": ["LOW"]}], "nothing here") == [],
        "a piece too short to be evidence is not reported",
    )
    # Presence is not enough: three identical rows against one is two rows
    # lost, and a substring test calls that a pass.
    thrice = [line("Host-200 is vulnerable")] * 3
    passed &= check(convert.missing_from_docx(thrice, "Host-200 is vulnerable") != [],
                    "three identical lines against one copy are reported")
    passed &= check(
        convert.missing_from_docx(thrice, "Host-200 is vulnerable " * 3) == [],
        "three identical lines against three copies are not",
    )
    # A soft hyphen is where typst broke the word, so the PDF's two lines have
    # to match the one word the Word file writes.
    passed &= check(
        convert.missing_from_docx([line("em­"), line("ploying a custom system")],
                                  "employing a custom system") == [],
        "a word broken at a soft hyphen is not missing",
    )
    passed &= check(convert.compare_squash("em­ploying") == "employing",
                    "a soft hyphen counts as nothing", convert.compare_squash("em­ploying"))

    # The ladder puts back the highest rung whose only losses are text the PDF
    # overprints and every rung below lost too.
    piled = [{"text": "Row A", "overprinted": True, "page": 2},
             {"text": "Row A", "overprinted": True, "page": 2},
             {"text": "Row B", "overprinted": True, "page": 2},
             {"text": "Row C", "overprinted": False, "page": 2}]
    top = {"missing": ["Row A", "Row B"], "lines": piled}
    middle = {"missing": ["Row A"], "lines": piled}
    passed &= check(convert.better_rung([top, middle], ["Row A", "Row B"]) is top,
                    "ladder: a rung whose overprinted losses recur below it is put back")
    passed &= check(convert.better_rung([top, middle], ["Row A"]) is middle,
                    "ladder: a rung that lost more than the last is passed over")
    passed &= check(convert.better_rung([top], ["Row A"]) is None,
                    "ladder: nothing is put back when every rung above lost more")
    passed &= check(convert.better_rung([{"missing": ["Row A", "Row A"], "lines": piled}], ["Row A"]) is None,
                    "ladder: losing a line twice is not losing it once")
    passed &= check(convert.better_rung([{"missing": ["Row C"], "lines": piled}], ["Row C"]) is None,
                    "ladder: a legible line lost on every rung does not put a rung back")

    # Only the missing lines the PDF overprints are named as the PDF's fault.
    lines = [{"text": "Row A", "overprinted": True, "page": 3},
             {"text": "Row B", "overprinted": False, "page": 3},
             {"text": "Row A", "overprinted": True, "page": 4}]
    named = convert.overprinted_missing(lines, ["Row A", "Row B"])
    passed &= check([(l["text"], l["page"]) for l in named] == [("Row A", 3)],
                    "overprint: a missing line is named once, from the first page drawing it", named)
    return passed


def run_tail_checks():
    """The section tail is a 1 pt stub and the padding before it is gone."""
    passed = True
    document = Document()
    body = document.element.body
    document.add_table(rows=1, cols=1)
    for _ in range(3):
        convert.set_spacing(document.add_paragraph()._p, before=0.0, after=0.0, line=0.7)
    document.add_section()
    document.add_table(rows=1, cols=1)
    convert.set_spacing(document.add_paragraph()._p, before=0.0, after=0.0, line=convert.SPACER_LINE_PT)
    document.add_table(rows=1, cols=1)
    convert.set_spacing(document.add_paragraph()._p, before=0.0, after=0.0, line=0.7)
    document.add_paragraph("A last line of text")
    convert.set_spacing(document.add_paragraph()._p, before=0.0, after=0.0, line=0.7)
    document.add_section()
    removed = convert.trim_section_tails(document)
    tags = [child.tag.split("}")[1] for child in body]
    passed &= check(removed == 4 and tags == ["tbl", "p", "tbl", "p", "tbl", "p", "p", "p", "sectPr"],
                    "tail: padding paragraphs go, the spacer between two tables stays", (removed, tags))
    breaks = [p for p in body.iterchildren(convert.qn("w:p")) if p.find(convert.qn("w:pPr")) is not None
              and p.find(convert.qn("w:pPr")).find(convert.qn("w:sectPr")) is not None]
    metrics = [convert.paragraph_metrics(p) for p in breaks]
    sizes = [p.find(convert.qn("w:pPr")).find(convert.qn("w:rPr")).find(convert.qn("w:sz")).get(convert.qn("w:val"))
             for p in breaks]
    passed &= check(len(breaks) == 2 and all(m and m["line"] == 1.0 and m["before"] == 0 and m["after"] == 0 for m in metrics)
                    and sizes == ["2", "2"],
                    "tail: each break paragraph is a 1 pt exact line in a 1 pt font", (metrics, sizes))
    return passed


def run_column_checks():
    """pdf2docx's side-by-side columns folded into a table row, and pages counted past them."""
    passed = True
    qn = convert.qn
    document = Document()
    body = document.element.body

    def section_paragraph(kind=None, columns=None):
        paragraph = document.add_paragraph()._p
        section = convert.OxmlElement("w:sectPr")
        if kind:
            node = convert.OxmlElement("w:type")
            node.set(qn("w:val"), kind)
            section.append(node)
        size = convert.OxmlElement("w:pgSz")
        size.set(qn("w:w"), "12240")
        size.set(qn("w:h"), "15840")
        section.append(size)
        margins = convert.OxmlElement("w:pgMar")
        for name, value in (("w:top", "1088"), ("w:right", "1002"), ("w:bottom", "1440"), ("w:left", "1080")):
            margins.set(qn(name), value)
        section.append(margins)
        cols = convert.OxmlElement("w:cols")
        if columns:
            cols.set(qn("w:num"), str(len(columns)))
            cols.set(qn("w:equalWidth"), "0")
            for width in columns:
                col = convert.OxmlElement("w:col")
                col.set(qn("w:w"), str(width))
                col.set(qn("w:space"), "0")
                cols.append(col)
        section.append(cols)
        paragraph.get_or_add_pPr().append(section)
        return paragraph

    def picture_paragraph(height):
        paragraph = document.add_paragraph()._p
        run = convert.OxmlElement("w:r")
        drawing = convert.OxmlElement("w:drawing")
        inline = convert.OxmlElement("wp:inline")
        extent = convert.OxmlElement("wp:extent")
        extent.set("cx", str(int(248 * convert.EMU_PER_PT)))
        extent.set("cy", str(int(height * convert.EMU_PER_PT)))
        inline.append(extent)
        drawing.append(inline)
        run.append(drawing)
        paragraph.append(run)
        return paragraph

    def sized(paragraph, size):
        for run in paragraph.findall(qn("w:r")):
            node = convert.OxmlElement("w:sz")
            node.set(qn("w:val"), str(size))
            run.get_or_add_rPr().append(node)

    intro = document.add_paragraph("Before the figure")._p
    section_paragraph()
    first = picture_paragraph(139.5)
    stacked = picture_paragraph(50.0)
    section_paragraph("continuous", [5040, 5118])
    second = picture_paragraph(120.0)
    closing_column = section_paragraph("nextColumn", [5040, 5118])
    convert.set_spacing(closing_column, after=9.6)
    caption = document.add_paragraph("Figure 6. Two captures")._p
    convert.set_spacing(caption, before=0.0, after=0.0, line=1.0)
    sized(caption, 18)
    text = document.add_paragraph("In voluptate velit esse")._p
    convert.set_spacing(text, before=9.6, after=0.0, line=14.8)
    sized(text, 18)
    closing = body.find(qn("w:sectPr"))
    kind = convert.OxmlElement("w:type")
    kind.set(qn("w:val"), "continuous")
    closing.insert(0, kind)

    groups = convert.page_blocks(body)
    pages = convert.paragraph_pages(document, 1)
    passed &= check(len(groups) == 1 and len(document.sections) == 4 and pages is not None and pages.get(id(text)) == 0,
                    "columns: four sections of one page count as one page", (len(groups), len(document.sections), pages is not None))
    passed &= check(len(convert.page_sections(document)) == 1, "columns: one section starts the page",
                    len(convert.page_sections(document)))

    folded = convert.fold_column_sections(document)
    table = body.find(qn("w:tbl"))
    cells = table.find(qn("w:tr")).findall(qn("w:tc")) if table is not None else []
    widths = [cell.find(qn("w:tcPr")).find(qn("w:tcW")).get(qn("w:w")) for cell in cells]
    height = table.find(qn("w:tr")).find(qn("w:trPr")).find(qn("w:trHeight")).get(qn("w:val")) if table is not None else None
    passed &= check(folded == 1 and widths == ["5040", "5118"] and height == str(int(round(189.5 * 20))),
                    "columns: the pictures become one row of two cells at the columns' widths, as tall as the taller column",
                    (folded, widths, height))
    passed &= check(stacked.getparent() is cells[0] if cells else False, "columns: a column's second picture is in its cell too")
    passed &= check(abs(convert.spacing_pt(caption, "w:before") - 9.6) < 0.05,
                    "columns: the gap pdf2docx wrote after the columns is carried to the caption",
                    convert.spacing_pt(caption, "w:before"))
    passed &= check(convert.strip_column_marks(document) == 1 and table.find(qn("w:tblPr")).find(qn("w:tblDescription")) is None,
                    "columns: the fold's own marker never reaches the file a reader opens")
    passed &= check(len(document.sections) == 1 and document.sections[0]._sectPr.find(qn("w:type")) is None,
                    "columns: the page is one section again, and it starts a page", len(document.sections))
    passed &= check(cells and first.getparent() is cells[0] and second.getparent() is cells[1],
                    "columns: each picture sits in its column's cell")
    order = list(body)
    passed &= check(table is not None and order.index(intro) < order.index(table) < order.index(caption),
                    "columns: the row stands where the pictures were")
    metrics = convert.paragraph_metrics(caption)
    passed &= check(metrics is not None and abs(metrics["line"] - 14.8) < 0.05,
                    "columns: the caption's 1 pt line is taken from its neighbour at the same size", metrics)

    # One centred line in a cell keeps its centring; two lines at one x in a cell are flush left.
    dates = document.add_table(rows=1, cols=2)
    date_cell = dates.rows[0].cells[0]
    date = date_cell.paragraphs[0]
    date.add_run("August 18, 2026")
    date.alignment = convert.WD_ALIGN_PARAGRAPH.CENTER
    lines = dates.rows[0].cells[1].paragraphs[0]
    lines.add_run("ACME contracted Blue Harbor Security to perform")
    lines.add_run().add_break()
    lines.add_run("a penetration test of the internal network.")
    lines.alignment = convert.WD_ALIGN_PARAGRAPH.CENTER
    runs = [{"lines": ["august18,2026"], "short": [False], "justified": False, "indented": True, "page": None},
            {"lines": ["acmecontractedblueharborsecuritytoperform", "apenetrationtestoftheinternalnetwork."],
             "short": [False, True], "justified": False, "indented": True, "page": None}]
    convert.align_paragraphs(document, runs)
    passed &= check(date.alignment == convert.WD_ALIGN_PARAGRAPH.CENTER and lines.alignment == convert.WD_ALIGN_PARAGRAPH.LEFT,
                    "columns: a cell's one centred line stays centred, its two-line run goes flush left",
                    (date.alignment, lines.alignment))

    # Columns that end their page: the break after them is the next page's, so
    # the page's own break moves after the row and no two pages become one.
    end_doc = Document()
    end_body = end_doc.element.body

    def end_section(kind=None, columns=None):
        paragraph = end_doc.add_paragraph()._p
        section = convert.OxmlElement("w:sectPr")
        if kind:
            node = convert.OxmlElement("w:type")
            node.set(qn("w:val"), kind)
            section.append(node)
        size = convert.OxmlElement("w:pgSz")
        size.set(qn("w:w"), "12240")
        size.set(qn("w:h"), "15840")
        section.append(size)
        margins = convert.OxmlElement("w:pgMar")
        for name, value in (("w:top", "1088"), ("w:right", "1002"), ("w:bottom", "1440"), ("w:left", "1080")):
            margins.set(qn(name), value)
        section.append(margins)
        cols = convert.OxmlElement("w:cols")
        if columns:
            cols.set(qn("w:num"), str(len(columns)))
            for width in columns:
                col = convert.OxmlElement("w:col")
                col.set(qn("w:w"), str(width))
                cols.append(col)
        section.append(cols)
        paragraph.get_or_add_pPr().append(section)
        return paragraph

    def end_picture(height):
        paragraph = end_doc.add_paragraph()._p
        run = convert.OxmlElement("w:r")
        drawing = convert.OxmlElement("w:drawing")
        inline = convert.OxmlElement("wp:inline")
        extent = convert.OxmlElement("wp:extent")
        extent.set("cx", str(int(248 * convert.EMU_PER_PT)))
        extent.set("cy", str(int(height * convert.EMU_PER_PT)))
        inline.append(extent)
        drawing.append(inline)
        run.append(drawing)
        paragraph.append(run)

    end_doc.add_paragraph("Page one text")
    end_section()
    end_doc.add_paragraph("Page two text")
    end_section()
    end_picture(120.0)
    end_section("continuous", [5040, 5118])
    end_picture(120.0)
    end_section("nextColumn", [5040, 5118])
    end_doc.add_paragraph("Page three text")
    before_pages = len(convert.page_sections(end_doc))
    convert.fold_column_sections(end_doc)
    passed &= check(before_pages == 3 and len(convert.page_sections(end_doc)) == 3 and len(convert.page_blocks(end_body)) == 3,
                    "columns: a picture row that ends its page keeps the page count",
                    (before_pages, len(convert.page_sections(end_doc)), len(convert.page_blocks(end_body))))

    # A tab before a break reaches a stop a point inside the column's edge.
    plain = document.add_paragraph("Short line")._p
    convert.Paragraph(plain, document).paragraph_format.right_indent = Pt(10)
    convert.edge_tab_stop(document, plain)
    stops = plain.find(qn("w:pPr")).find(qn("w:tabs"))
    stop = stops.findall(qn("w:tab"))[-1] if stops is not None else None
    section = document.sections[0]
    width = section.page_width.pt - section.left_margin.pt - section.right_margin.pt
    passed &= check(stop is not None and stop.get(qn("w:val")) == "right"
                    and abs(int(stop.get(qn("w:pos"))) / 20.0 - (width - 11.0)) < 0.1,
                    "columns: a body paragraph's edge stop is its text width less the indent and a point",
                    (None if stop is None else stop.attrib, width))
    cell_text = cells[0].find(qn("w:p"))
    convert.edge_tab_stop(document, cell_text)
    stops = cell_text.find(qn("w:pPr")).find(qn("w:tabs"))
    stop = stops.findall(qn("w:tab"))[-1] if stops is not None else None
    passed &= check(stop is not None and abs(int(stop.get(qn("w:pos"))) / 20.0 - (5040 / 20.0 - 1.0)) < 0.1,
                    "columns: a cell paragraph's edge stop is its cell's width less a point",
                    None if stop is None else stop.attrib)

    # The right indent of a justified paragraph is its section's text edge
    # less the column its run ended at: a landscape page in a portrait report
    # got 248 pt from the document's one edge, and every forced line wrapped.
    wide = Document()
    wide.sections[0].page_width, wide.sections[0].page_height = Pt(612.0), Pt(792.0)
    wide.sections[0].left_margin = wide.sections[0].right_margin = Pt(54.0)
    portrait = wide.add_paragraph("Quis nostrud exercitation ullamco laboris nisi ut aliquip")
    landscape_section = wide.add_section()
    landscape_section.page_width, landscape_section.page_height = Pt(792.0), Pt(612.0)
    landscape_section.left_margin = landscape_section.right_margin = Pt(54.0)
    landscape = wide.add_paragraph("Duis aute irure dolor in reprehenderit in voluptate velit")
    try:
        convert.RUN_COLUMNS.reset([], 558.0, {792: 180.0})
        got = (convert.column_right_indent(wide, portrait._p, 558.0), convert.column_right_indent(wide, landscape._p, 558.0))
        passed &= check(got == (0.0, 0.0),
                        "indent: a paragraph no run claims is measured against the edge of pages its size", got)
        got = (convert.column_right_indent(wide, landscape._p, 558.0, column=700.0),
               convert.column_right_indent(wide, portrait._p, 558.0, column=290.0))
        passed &= check(got == (38.0, 268.0),
                        "indent: a column handed in is the paragraph's text edge less that column", got)
        got = (convert.column_right_indent(wide, portrait._p, 558.0, column=555.2),
               convert.column_right_indent(wide, portrait._p, 558.0, column=549.0))
        passed &= check(got == (0.0, 9.0),
                        "indent: a column within a short line's reach of the page's edge is the page's edge, one past it is not", got)
        # Short of the edge the reach is a mark's overhang in the paragraph's own type (0.3 em, 3.3 pt at 11 pt), not a
        # flat 6 pt: a block inset 5 pt from the right is a column of its own.
        portrait.runs[0].font.size = Pt(11.0)
        got = (convert.column_right_indent(wide, portrait._p, 558.0, column=555.2),
               convert.column_right_indent(wide, portrait._p, 558.0, column=553.0),
               convert.column_right_indent(wide, portrait._p, 558.0, column=560.5))
        passed &= check(got == (0.0, 5.0, 0.0),
                        "indent: an inset block 5 pt short of the edge keeps its own column, a hanging mark's 2.8 pt does not", got)
        # Small type is measured with the body's window: its rows end 0.25 em of the difference short of the
        # body's, so a 6 pt disclaimer in an 11 pt report 2.5 pt short of the edge is on the edge.
        small_doc = Document()
        small_doc.sections[0].page_width, small_doc.sections[0].page_height = Pt(612.0), Pt(792.0)
        small_doc.sections[0].left_margin = small_doc.sections[0].right_margin = Pt(54.0)
        for _ in range(4):
            small_doc.add_paragraph().add_run("Body text set in the report's own eleven point type.").font.size = Pt(11.0)
        notice = small_doc.add_paragraph()
        notice.add_run("Disclaimer: this report is provided as is.").font.size = Pt(6.0)
        got = (convert.body_font_size(small_doc),
               convert.column_right_indent(small_doc, notice._p, 558.0, column=555.5),
               convert.column_right_indent(small_doc, notice._p, 558.0, column=553.0))
        passed &= check(got == (11.0, 0.0, 5.0),
                        "indent: small type snaps to the page's edge within the body's window, an inset block does not", got)
        convert.RUN_COLUMNS.reset([], None, {})
        passed &= check(convert.column_right_indent(wide, landscape._p, 558.0) == 180.0,
                        "indent: with no run and no page sizes known the document's own edge is the column")
    finally:
        convert.RUN_COLUMNS.reset([], None, {})
    # A stop at the text edge follows the indent the column gave it.
    convert.Paragraph(landscape._p, wide).paragraph_format.right_indent = Pt(2.0)
    convert.edge_tab_stop(wide, landscape._p)
    stop = landscape._p.find(qn("w:pPr")).find(qn("w:tabs")).find(qn("w:tab"))
    passed &= check(stop is not None and abs(int(stop.get(qn("w:pos"))) / 20.0 - (792.0 - 108.0 - 2.0 - 1.0)) < 0.1,
                    "columns: a landscape paragraph's edge stop is its own section's width less the indent and a point",
                    None if stop is None else stop.attrib)
    return passed


def run_column_page_end_checks():
    """Columns that end their page leave the next page's blank run-up with the next page."""
    passed = True
    qn = convert.qn
    document = Document()
    body = document.element.body

    def section_paragraph(kind=None, columns=None):
        paragraph = document.add_paragraph()._p
        section = convert.OxmlElement("w:sectPr")
        if kind:
            node = convert.OxmlElement("w:type")
            node.set(qn("w:val"), kind)
            section.append(node)
        size = convert.OxmlElement("w:pgSz")
        size.set(qn("w:w"), "12240")
        size.set(qn("w:h"), "15840")
        section.append(size)
        margins = convert.OxmlElement("w:pgMar")
        for name, value in (("w:top", "720"), ("w:right", "1002"), ("w:bottom", "1440"), ("w:left", "1080")):
            margins.set(qn(name), value)
        section.append(margins)
        cols = convert.OxmlElement("w:cols")
        if columns:
            cols.set(qn("w:num"), str(len(columns)))
            cols.set(qn("w:equalWidth"), "0")
            for width in columns:
                col = convert.OxmlElement("w:col")
                col.set(qn("w:w"), str(width))
                col.set(qn("w:space"), "0")
                cols.append(col)
        section.append(cols)
        paragraph.get_or_add_pPr().append(section)
        return paragraph

    def picture_paragraph(height):
        paragraph = document.add_paragraph()._p
        run = convert.OxmlElement("w:r")
        drawing = convert.OxmlElement("w:drawing")
        inline = convert.OxmlElement("wp:inline")
        extent = convert.OxmlElement("wp:extent")
        extent.set("cx", str(int(248 * convert.EMU_PER_PT)))
        extent.set("cy", str(int(height * convert.EMU_PER_PT)))
        inline.append(extent)
        drawing.append(inline)
        run.append(drawing)
        paragraph.append(run)
        return paragraph

    intro = document.add_paragraph("Before the figure")._p
    section_paragraph()
    picture_paragraph(139.5)
    section_paragraph("continuous", [5040, 5118])
    picture_paragraph(120.0)
    section_paragraph("nextColumn", [5040, 5118])
    # pdf2docx's blank line for the next page's header, then that page's first text.
    run_up = document.add_paragraph()._p
    convert.set_spacing(run_up, before=0.0, after=25.2, line=11.0)
    caption = document.add_paragraph("Figure 6. Two captures from day 4")._p
    convert.set_spacing(caption, before=0.0, after=0.0, line=16.7)
    for node in caption.findall(qn("w:r")):
        size = convert.OxmlElement("w:sz")
        size.set(qn("w:val"), "22")
        node.get_or_add_rPr().append(size)

    folded = convert.fold_column_sections(document)
    table = body.find(qn("w:tbl"))
    groups = convert.section_blocks(body)
    stub = table.getnext() if table is not None else None
    passed &= check(folded == 1 and stub is not None and convert.section_break_of(stub) is not None
                    and len(groups) == 2 and groups[0][-1] is stub and groups[1][0] is run_up and groups[1][1] is caption,
                    "columns: the page's break stands straight after the row, so the next page keeps its run-up",
                    (folded, [len(group) for group in groups]))
    left = convert.shave_leading_space(groups[1], 18.2)
    metrics = convert.paragraph_metrics(caption)
    passed &= check(left < 0.01 and metrics is not None and abs(metrics["line"] - 16.7) < 0.05,
                    "columns: the header's margin comes off the run-up, and the caption keeps its line",
                    (left, metrics))
    passed &= check(children_in_order(body, intro, table, run_up, caption),
                    "columns: the row stands where the pictures were, the run-up and the caption follow it")
    return passed


def children_in_order(body, *elements):
    """Are these children of the body in the order given?"""
    order = list(body)
    positions = [order.index(element) for element in elements]
    return positions == sorted(positions)


def run_join_checks():
    """Lines pdf2docx joined: a link list's break reused, a nested item's run placed, a cut line let run off the page."""
    passed = True
    qn = convert.qn
    document = Document()

    def hyperlink(paragraph, text):
        link = convert.OxmlElement("w:hyperlink")
        run = convert.OxmlElement("w:r")
        node = convert.OxmlElement("w:t")
        node.text = text
        run.append(node)
        link.append(run)
        paragraph._p.append(link)

    # A resource list: a bullet, a link, the converter's break, the next bullet and link.
    listed = document.add_paragraph()
    listed.add_run("• ")
    hyperlink(listed, "Nostrud exercitation ullamco laboris")
    listed.add_run().add_break()
    listed.add_run("• ")
    hyperlink(listed, "Ad minim veniam quis")
    added = convert.split_paragraph_at(listed._p, ["•nostrudexercitationullamcolaboris", "•adminimveniamquis"])
    breaks = sum(1 for _ in listed._p.iter(qn("w:br")))
    passed &= check(added == 1 and breaks == 1,
                    "join: a cut at the end of a link reuses the break after the link", (added, breaks))

    # A run whose last line is a table row read across two cells: the line is dropped, the rest placed.
    card = document.add_paragraph("Fugiat nulla pariatur Lorem ipsum dolor sit amet consectetur adipiscing elit Sed")
    grid = document.add_table(rows=1, cols=2)
    grid.rows[0].cells[0].paragraphs[0].add_run("TECHNIQUES")
    grid.rows[0].cells[1].paragraphs[0].add_run("MITIGATIONS")
    across = {"lines": ["fugiatnullapariaturloremipsumdolor", "sitametconsecteturadipiscingelitsed", "techniquesmitigations"],
              "short": [False, True, True], "widths": [0.0, 0.0, 0.0], "spaces": [0, 0, 0], "justified": False, "page": None}
    broken = convert.force_line_breaks(document, [across], set())
    passed &= check(broken == 1 and sum(1 for _ in card._p.iter(qn("w:br"))) == 1,
                    "join: a last line no paragraph holds is dropped and the paragraph above still gets its break", broken)

    # A cell paragraph pdf2docx joined from two of the PDF's lines in that cell.
    pages_doc = Document()
    steps = pages_doc.add_table(rows=1, cols=2)
    steps.rows[0].cells[0].paragraphs[0].add_run("Steps taken")
    joined_steps = steps.rows[0].cells[1].paragraphs[0]
    joined_steps.add_run("2. Ea commodo consequat Duis aute irure dolor. ")
    joined_steps.add_run("3. Exercitation ullamco laboris nisi ut.")
    joined_steps.paragraph_format.right_indent = Pt(144)
    cell_lines = [{"text": "Steps taken 2. Ea commodo consequat Duis aute irure dolor.", "page": 1,
                   "pieces": ["Steps taken", "2. Ea commodo consequat Duis aute irure dolor."]},
                  {"text": "3. Exercitation ullamco laboris nisi ut.", "page": 1,
                   "pieces": ["3. Exercitation ullamco laboris nisi ut."]}]
    done = convert.break_cell_lines(pages_doc, cell_lines, set(), page_count=1)
    breaks = sum(1 for _ in joined_steps._p.iter(qn("w:br")))
    passed &= check(done == 1 and breaks == 1 and joined_steps.paragraph_format.right_indent == 0,
                    "join: a cell's joined steps break where the PDF's lines in that cell end", (done, breaks))

    # Two paragraphs end with the same words; the run's tail is the later one's.
    tie_doc = Document()
    earlier = tie_doc.add_paragraph("Heading words. Alpha beta gamma delta epsilon zeta")
    later = tie_doc.add_paragraph("epsilon zeta")
    tie_run = {"lines": ["headingwords.", "alphabetagammadeltaepsilonzeta", "epsilonzeta"], "short": [True, False, True],
               "widths": [0.0, 0.0, 0.0], "spaces": [0, 0, 0], "justified": False, "page": None}
    convert.force_line_breaks(tie_doc, [tie_run], set())
    passed &= check(sum(1 for _ in earlier._p.iter(qn("w:br"))) == 1 and sum(1 for _ in later._p.iter(qn("w:br"))) == 0,
                    "join: a tail two paragraphs end with goes to the later one",
                    convert.paragraph_text(earlier._p))

    # A paragraph whose last line is two letters and a stop still gets the break before it.
    ending = document.add_paragraph("Alpha one two three four five six seven in.")
    short_run = {"lines": ["alphaonetwothreefour", "fivesixseven", "in."], "short": [False, False, True],
                 "widths": [0.0, 0.0, 0.0], "spaces": [0, 0, 0], "justified": False, "page": None}
    broken = convert.force_line_breaks(document, [short_run], set())
    passed &= check(broken == 2 and sum(1 for _ in ending._p.iter(qn("w:br"))) == 2,
                    "join: the run's short last line keeps the break before it", broken)

    # pdf2docx wraps a link in a run of its own, with its break after the wrapper.
    wrapped_list = document.add_paragraph()
    wrapped_list.add_run("•")
    for text, last in (("Labore et dolore magna", False), ("Aliqua Ut enim ad", True)):
        outer = convert.OxmlElement("w:r")
        link = convert.OxmlElement("w:hyperlink")
        inner = convert.OxmlElement("w:r")
        node = convert.OxmlElement("w:t")
        node.text = " " + text
        inner.append(node)
        link.append(inner)
        outer.append(link)
        wrapped_list._p.append(outer)
        if not last:
            wrapped_list.add_run().add_break()
            wrapped_list.add_run("•")
    added = convert.split_paragraph_at(wrapped_list._p, ["•laboreetdoloremagna", "•aliquautenimad"])
    breaks = sum(1 for _ in wrapped_list._p.iter(qn("w:br")))
    passed &= check(added == 1 and breaks == 1,
                    "join: a cut at the end of a wrapped link reuses the break after the wrapper", (added, breaks))

    # A row of exact height whose declared content is taller grows instead of clipping.
    grid = document.add_table(rows=2, cols=1)
    for index, row in enumerate(grid.rows):
        properties = row._tr.get_or_add_trPr()
        height = convert.OxmlElement("w:trHeight")
        height.set(qn("w:hRule"), "exact")
        height.set(qn("w:val"), "400")
        properties.append(height)
        paragraph = row.cells[0].paragraphs[0]
        paragraph.add_run("one")
        if index == 0:
            paragraph.add_run().add_break()
            paragraph.add_run("two")
            paragraph.add_run().add_break()
            paragraph.add_run("three")
        convert.set_spacing(paragraph._p, before=0.0, after=0.0, line=13.5)
    relaxed = convert.relax_overfull_rows(document)
    rules = [row._tr.find(qn("w:trPr")).find(qn("w:trHeight")).get(qn("w:hRule")) for row in grid.rows]
    # A vertically merged cell's content spans the rows below; it never grows its first row.
    merged = document.add_table(rows=2, cols=2)
    for row in merged.rows:
        properties = row._tr.get_or_add_trPr()
        height = convert.OxmlElement("w:trHeight")
        height.set(qn("w:hRule"), "exact")
        height.set(qn("w:val"), "400")
        properties.append(height)
    tall = merged.rows[0].cells[0].paragraphs[0]
    tall.add_run("one")
    for _ in range(3):
        tall.add_run().add_break()
        tall.add_run("more")
    convert.set_spacing(tall._p, before=0.0, after=0.0, line=13.5)
    for k in range(2):
        merged.rows[k].cells[1].paragraphs[0].add_run("x")
        convert.set_spacing(merged.rows[k].cells[1].paragraphs[0]._p, before=0.0, after=0.0, line=13.5)
    for k, value in enumerate(("restart", None)):
        node = convert.OxmlElement("w:vMerge")
        if value:
            node.set(qn("w:val"), value)
        merged.rows[k].cells[0]._tc.get_or_add_tcPr().append(node)
    convert.relax_overfull_rows(document)
    merged_rules = [row._tr.find(qn("w:trPr")).find(qn("w:trHeight")).get(qn("w:hRule")) for row in merged.rows]
    passed &= check(merged_rules == ["exact", "exact"],
                    "join: a vertically merged cell does not grow its first row", merged_rules)
    passed &= check(relaxed >= 1 and rules == ["atLeast", "exact"],
                    "join: a 20 pt row holding three 13.5 pt lines may grow, a row that fits stays exact", (relaxed, rules))

    # A run that ends inside a paragraph: the nested item pdf2docx joined to its bullet.
    prose = document.add_paragraph("Alpha one two three four five six seven eight")
    bullet = document.add_paragraph("• Bullet line text ‣ Sub item joined")
    run = {"lines": ["alphaonetwothreefour", "fivesixseveneight", "•bulletlinetext"], "short": [False, True, True],
           "widths": [0.0, 0.0, 0.0], "spaces": [0, 0, 0], "justified": False, "page": None}
    broken = convert.force_line_breaks(document, [run], set())
    prose_breaks = sum(1 for _ in prose._p.iter(qn("w:br")))
    bullet_breaks = sum(1 for _ in bullet._p.iter(qn("w:br")))
    passed &= check(broken == 2 and prose_breaks == 1 and bullet_breaks == 1,
                    "join: the paragraph above gets its break, and the run's last line is cut from the item joined to it",
                    (broken, prose_breaks, bullet_breaks))

    # Two bullets, the second the first one's first line with a stop after it.
    # A run ending inside a paragraph used to go to the later of two such
    # paragraphs, and the stop was wrapped onto a line of its own (page 168 of
    # a 291-page report). What follows the part has to start a line of the
    # page: "tempor." does, "." does not.
    twin_doc = Document()
    twin_head = twin_doc.add_paragraph("• Alpha one two")
    long_item = twin_doc.add_paragraph("• Nulla pariatur sit amet tempor.")
    short_item = twin_doc.add_paragraph("• Nulla pariatur sit amet.")
    twin_runs = [{"lines": ["•alphaonetwo", "•nullapariatursitamet"], "short": [False, True],
                  "widths": [0.0, 0.0], "spaces": [0, 0], "justified": False, "page": 0,
                  "pageKeys": ["•alphaonetwo", "•nullapariatursitamet", "tempor.", "•nullapariatursitamet."]}]
    broken = convert.force_line_breaks(twin_doc, twin_runs, set(), page_count=1)
    long_breaks = sum(1 for _ in long_item._p.iter(qn("w:br")))
    short_breaks = sum(1 for _ in short_item._p.iter(qn("w:br")))
    passed &= check(long_breaks == 1 and short_breaks == 0 and broken == 1,
                    "join: a part that ends inside a paragraph goes to the one whose rest starts a line of the page",
                    (broken, long_breaks, short_breaks))
    # The rest may be a one-word last line that is no run of its own: it is a
    # row of the page all the same (page 238 of the same report lost the break
    # before "ut." when only the runs' lines were looked in).
    last_doc = Document()
    last_doc.add_paragraph("• Alpha one two")
    ends_short = last_doc.add_paragraph("• Lorem ipsum dolor sit amet ut.")
    last_runs = [{"lines": ["•alphaonetwo", "•loremipsumdolorsitamet"], "short": [False, True],
                  "widths": [0.0, 0.0], "spaces": [0, 0], "justified": False, "page": 0,
                  "pageKeys": ["•alphaonetwo", "•loremipsumdolorsitamet", "ut."]}]
    convert.force_line_breaks(last_doc, last_runs, set(), page_count=1)
    passed &= check(sum(1 for _ in ends_short._p.iter(qn("w:br"))) == 1,
                    "join: a one-word last line the runs do not hold still gets its break", convert.paragraph_text(ends_short._p))
    passed &= check(convert.begins_a_line("", ["abc"]) and convert.begins_a_line("tempor.", ["tempor."])
                    and convert.begins_a_line("tempor.andmore", ["tempor."]) and not convert.begins_a_line(".", ["tempor."]),
                    "join: a rest begins a line when it is empty, a line's prefix or a line's start")

    # A code line longer than its cell: the right indent lets it run past the edge, on one line.
    table = document.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    width = cell._tc.get_or_add_tcPr()
    size = width.find(qn("w:tcW"))
    if size is None:
        size = convert.OxmlElement("w:tcW")
        width.append(size)
    size.set(qn("w:w"), "3000")
    size.set(qn("w:type"), "dxa")
    code = cell.paragraphs[0]
    long_run = code.add_run("https://10.10.9.70/downloads/releases/" + "a" * 60)
    long_run.font.size = Pt(9)
    changed = convert.let_line_overflow(document, code._p)
    indent = code._p.find(qn("w:pPr")).find(qn("w:ind"))
    right = convert.twips_of(indent.get(qn("w:right"))) if indent is not None else None
    expected = -(98 * convert.MONO_ADVANCE_EM * 9.0 * (1.0 + convert.WORD_WIDTH_EXCESS) - (150.0 - 2 * convert.DEFAULT_CELL_MARGIN_PT) + 2.0)
    passed &= check(changed and right is not None and abs(right - expected) < 0.2,
                    "join: a cut line wider than its cell gets the right indent that keeps it on one line", (right, expected))
    short = cell.add_paragraph()
    short.add_run("ok").font.size = Pt(9)
    passed &= check(not convert.let_line_overflow(document, short._p), "join: a line that fits is left alone")
    passed &= run_overflow_row_checks()
    passed &= run_checkbox_checks()

    # A wrapper table whose second row is one cell holding one picture: the
    # line comes out, the picture row stays a table of its own.
    single = document.add_table(rows=2, cols=1)
    for row in single.rows:
        properties = row._tr.get_or_add_trPr()
        height = convert.OxmlElement("w:trHeight")
        height.set(qn("w:hRule"), "exact")
        height.set(qn("w:val"), "2680")
        properties.append(height)
    single.rows[0].cells[0].paragraphs[0].add_run("$ whoami")
    lone = single.rows[1].cells[0].paragraphs[0]._p
    run = convert.OxmlElement("w:r")
    drawing = convert.OxmlElement("w:drawing")
    inline = convert.OxmlElement("wp:inline")
    extent = convert.OxmlElement("wp:extent")
    extent.set("cx", str(int(238 * convert.EMU_PER_PT)))
    extent.set("cy", str(int(134 * convert.EMU_PER_PT)))
    inline.append(extent)
    drawing.append(inline)
    run.append(drawing)
    lone.append(run)
    lifted, kept_height = convert.unwrap_panel_blocks([single._tbl])
    passed &= check(len(lifted) == 1 and len(single._tbl.findall(qn("w:tr"))) == 1
                    and sum(1 for _ in single._tbl.iter(qn("w:drawing"))) == 1 and abs(kept_height - 134.0) < 0.1,
                    "join: a lone picture in the column's own cell keeps its row", (len(lifted), kept_height))

    # A listing whose first row is a short prompt: pdf2docx's break after it is kept.
    prompt = document.add_paragraph()
    prompt.add_run("$ id")
    prompt.add_run().add_break()
    prompt.add_run("uid=0(root) gid=0(root) groups=0(root)$ whoami")
    prompt.add_run().add_break()
    prompt.add_run("root")
    prompt_rows = [{"x0": 80.2, "x1": 80.2 + 5.42 * len(line), "text": line}
                   for line in ("$ id", "uid=0(root) gid=0(root) groups=0(root)", "$ whoami", "root")]
    prompt_panel = {"rows": prompt_rows, "rect": pymupdf.Rect(61.2, 100.0, 558.0, 160.0), "fill": (0.93, 0.93, 0.93), "bar": None}
    made = convert.split_joined_code([prompt._p], prompt_panel)
    prompt_lines = [""]
    for run in prompt._p.iter(qn("w:r")):
        for child in run:
            if child.tag == qn("w:t"):
                prompt_lines[-1] += child.text or ""
            elif child.tag == qn("w:br"):
                prompt_lines.append("")
    passed &= check(prompt_lines == [row["text"] for row in prompt_rows],
                    "join: a short first row keeps the break after it", (made, prompt_lines))

    # A wrapper table whose second row holds two pictures: the line comes out,
    # the picture row stays a table of its own.
    wrapped = document.add_table(rows=2, cols=2)
    for row in wrapped.rows:
        properties = row._tr.get_or_add_trPr()
        height = convert.OxmlElement("w:trHeight")
        height.set(qn("w:hRule"), "exact")
        height.set(qn("w:val"), "2680")
        properties.append(height)
    line = wrapped.rows[0].cells[0].paragraphs[0]
    line.add_run("print(port, probe(host))")
    pictures = []
    for cell in wrapped.rows[1].cells:
        paragraph = cell.paragraphs[0]._p
        run = convert.OxmlElement("w:r")
        drawing = convert.OxmlElement("w:drawing")
        inline = convert.OxmlElement("wp:inline")
        extent = convert.OxmlElement("wp:extent")
        extent.set("cx", str(int(238 * convert.EMU_PER_PT)))
        extent.set("cy", str(int(134 * convert.EMU_PER_PT)))
        inline.append(extent)
        drawing.append(inline)
        run.append(drawing)
        paragraph.append(run)
        pictures.append(paragraph)
    table = wrapped._tbl
    lifted, kept_height = convert.unwrap_panel_blocks([table])
    rows_left = table.findall(qn("w:tr"))
    drawings = sum(1 for _ in table.iter(qn("w:drawing")))
    passed &= check(lifted == [line._p] and line._p.getnext() is table and len(rows_left) == 1 and drawings == 2
                    and abs(kept_height - 134.0) < 0.1 and table.getparent() is document.element.body,
                    "join: a picture row under a panel's last line stays a table of its own, after the line",
                    (len(lifted), len(rows_left), drawings, kept_height))
    return passed


def run_align_checks():
    """Centred lines stay centred, and a break lands after the dash or space the PDF's line ended on."""
    passed = True
    qn = convert.qn
    centre = convert.WD_ALIGN_PARAGRAPH.CENTER
    flush = convert.WD_ALIGN_PARAGRAPH.LEFT

    # How a row sits in the text area (Letter, 1 in margins).
    def row(x0, x1, size=11.0):
        return {"x0": x0, "x1": x1, "columns": [{"size": size}], "text": "x"}

    def kind(x0, x1, right=540.0):
        return convert.row_kind(row(x0, x1), 72.0, right, 612.0)

    passed &= check(kind(116.0, 496.0) == "C", "align: a caption with 44 pt either side is centred", kind(116.0, 496.0))
    passed &= check(kind(88.6, 526.1) == "C",
                    "align: a wrapped centred line whose box includes its trailing space (gaps 16.6 and 13.9) is centred",
                    kind(88.6, 526.1))
    passed &= check(kind(116.1, 495.9, right=543.6) == "C",
                    "align: a justified document's hung right edge does not hide a centred line",
                    kind(116.1, 495.9, right=543.6))
    passed &= check(kind(72.0, 300.0) == "L", "align: a short flush-left line is flush left", kind(72.0, 300.0))
    passed &= check(kind(72.0, 540.0) == "F", "align: a line across the whole column is full", kind(72.0, 540.0))
    passed &= check(kind(150.0, 540.0) == "O", "align: a line indented from the left only is other", kind(150.0, 540.0))

    # Which paragraphs those rows make up.
    passed &= check(convert.centred_paragraph("longfirstlineshorttail", [("longfirstline", "F"), ("shorttail", "C")]),
                    "align: a full first line and a centred tail make a centred paragraph")
    passed &= check(not convert.centred_paragraph("firstlinelastline", [("firstline", "F"), ("lastline", "L")]),
                    "align: a flush-left last line is not a centred paragraph")
    passed &= check(not convert.centred_paragraph("othertext", [("firstline", "F"), ("shorttail", "C")]),
                    "align: a paragraph the rows cannot tile is not decided")
    passed &= check(not convert.centred_paragraph("scopeline", [("scopeline", "C"), ("filler", "L"), ("scopeline", "L")]),
                    "align: the same text tiled where it is flush left is not centred")
    passed &= check(not convert.centred_paragraph("full", [("full", "F")]),
                    "align: a full line alone proves nothing")

    # The pass itself.
    def caption_run(key, kinds=None, indented=True, justified=False):
        return {"lines": [key], "justified": justified, "indented": indented, "page": None, "short": [False],
                "pageRows": kinds}

    document = Document()
    long_caption = document.add_paragraph("Nmap scan of the internal network showing SMB signing disabled on hosts")
    key = convert.compare_squash(long_caption.text)
    declined = collections.Counter()
    convert.align_paragraphs(document, [caption_run(key, [(key, "C")])], declined=declined)
    passed &= check(long_caption.alignment == centre and declined["align-centred"] == 1,
                    "align: a long centred caption is centred, not forced left", (long_caption.alignment, declined))

    document = Document()
    plain_caption = document.add_paragraph("Nmap scan of the internal network showing SMB signing disabled on hosts")
    convert.align_paragraphs(document, [caption_run(key)])
    passed &= check(plain_caption.alignment == flush,
                    "align: without the page's rows the old rule stands", plain_caption.alignment)

    document = Document()
    placed = document.add_paragraph("Nmap scan of the internal network showing SMB signing disabled on hosts")
    placed.paragraph_format.left_indent = Pt(72)
    convert.align_paragraphs(document, [caption_run(key, [(key, "C")])])
    # The rows prove it centred on the text area, so the indent pdf2docx
    # placed it with goes and Word centres it: kept, the caption drifted off
    # centre in another font and read as left-aligned.
    passed &= check(placed.alignment == centre and placed.paragraph_format.left_indent is None,
                    "align: a centred line pdf2docx placed with a left indent is centred without it", placed.alignment)

    document = Document()
    boxed = document.add_paragraph("Nmap scan of the internal network showing SMB signing disabled on hosts")
    boxed.paragraph_format.left_indent = Pt(72)
    boxed.paragraph_format.right_indent = Pt(72)
    convert.align_paragraphs(document, [caption_run(key, [(key, "C")])])
    passed &= check(boxed.alignment != centre and boxed.paragraph_format.left_indent == Pt(72),
                    "align: a centred line placed between two indents is left as it was", boxed.alignment)

    document = Document()
    first = "This engagement covered the external network the internal network and the Active Directory forest"
    last = "and every finding was reproduced"
    wrapped = document.add_paragraph(first + " " + last)
    k1, k2 = convert.compare_squash(first), convert.compare_squash(last)
    rows = [(k1, "F"), (k2, "C")]
    head = caption_run(k1, rows, indented=False, justified=True)
    tail = caption_run(k2, rows)
    convert.align_paragraphs(document, [head, tail])
    passed &= check(wrapped.alignment == centre and not head["justified"],
                    "align: a wrapped centred paragraph keeps its centred tail and is not read as justified",
                    (wrapped.alignment, head["justified"]))

    document = Document()
    ragged = document.add_paragraph(first + " " + last)
    rows = [(k1, "F"), (k2, "L")]
    convert.align_paragraphs(document, [caption_run(k1, rows, indented=False), caption_run(k2, rows, indented=False)])
    passed &= check(ragged.alignment == flush,
                    "align: a paragraph with a flush-left last line is still set flush left", ragged.alignment)

    document = Document()
    grid = document.add_table(rows=1, cols=1)
    cell = grid.rows[0].cells[0].paragraphs[0]
    cell.add_run("Nmap scan of the internal network showing SMB signing disabled on hosts")
    convert.align_paragraphs(document, [caption_run(key, [(key, "C")], indented=False)])
    passed &= check(cell.alignment != centre, "align: a table cell is not judged by the page's centre", cell.alignment)

    # Where a break goes relative to the dashes and spaces the squashed count dropped.
    def shown(paragraph):
        parts = []
        for node in paragraph._p.iter():
            if node.tag == qn("w:t"):
                parts.append(node.text or "")
            elif node.tag == qn("w:br"):
                parts.append("|")
        return "".join(parts)

    def edged(text, last="", first=""):
        line = convert.EdgedLine(convert.compare_squash(text))
        line.last, line.first = last, first
        return line

    def split(text, lines):
        paragraph = Document().add_paragraph(text)
        convert.split_paragraph_at(paragraph._p, lines)
        return shown(paragraph)

    text = "the server-side request-forgery issue"
    got = split(text, [edged("the server-side request-", "-"), edged("forgery issue", "e")])
    passed &= check(got == "the server-side request-|forgery issue", "break: after the hard hyphen the PDF's line ended on", got)
    got = split(text, [convert.compare_squash("the server-side request-"), convert.compare_squash("forgery issue")])
    passed &= check(got == "the server-side request|-forgery issue",
                    "break: plain lines (a table cell's) keep the hyphen where the count put it", got)
    text = "see pages 12\u201314 of the appendix"
    got = split(text, [edged("see pages 12\u2013", "\u2013"), edged("14 of the appendix", "x")])
    passed &= check(got == "see pages 12\u2013|14 of the appendix", "break: after an en dash the PDF's line ended on", got)
    text = "the host \u2014 including the isolated ones"
    got = split(text, [edged("the host \u2014", "\u2014"), edged("including the isolated ones", "s")])
    passed &= check(got == "the host \u2014 |including the isolated ones", "break: after a spaced em dash that ended the line", got)
    got = split(text, [edged("the host", "t"), edged("\u2014 including the isolated ones", "s", "\u2014")])
    passed &= check(got == "the host |\u2014 including the isolated ones", "break: before an em dash that began the next line", got)
    text = "the host\u2014including the isolated ones"
    got = split(text, [edged("the host\u2014", "\u2014"), edged("including the isolated ones", "s")])
    passed &= check(got == "the host\u2014|including the isolated ones", "break: after an unspaced em dash that ended the line", got)
    # A word typst hyphenated: the soft hyphen, a run of its own with pdf2docx's break after it, ends its line,
    # so that break is the one kept (a second one went in before the hyphen and left an empty line), and the
    # hyphen pass puts a visible hyphen at it once the soft one is gone.
    soft = chr(0xAD)
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Exploit: neque eum Torqua")
    paragraph.add_run(soft)
    paragraph.add_run().add_break()
    paragraph.add_run("tum qui hoc primus cognomen")
    convert.split_paragraph_at(paragraph._p, [edged("Exploit: neque eum Torqua" + soft, "a"),
                                              edged("tum qui hoc primus cognomen", "n")])
    got = shown(paragraph)
    passed &= check(got == "Exploit: neque eum Torqua" + soft + "|tum qui hoc primus cognomen",
                    "break: a soft hyphen that ended the PDF's line keeps the converter's one break after it", got)
    for node in paragraph._p.iter(qn("w:t")):
        node.text = (node.text or "").replace(soft, "")
    broken = {(convert.compare_squash("Exploit: neque eum Torqua")[-convert.PAIR_WINDOW:],
               convert.compare_squash("tum qui hoc primus cognomen")[:convert.PAIR_WINDOW])}
    convert.restore_break_hyphens(document, broken)
    got = shown(paragraph)
    passed &= check(got == "Exploit: neque eum Torqua-|tum qui hoc primus cognomen",
                    "hyphens: the hyphen goes on the word before an emptied soft hyphen's run", got)
    # A table cell is cut twice, by the body pass and then with plain lines: both take the soft hyphen, one break.
    paragraph = Document().add_paragraph()
    paragraph.add_run("SMB signing dis" + soft)
    paragraph.add_run("abled")
    convert.split_paragraph_at(paragraph._p, [edged("SMB signing dis" + soft, "s"), edged("abled", "d")])
    convert.split_paragraph_at(paragraph._p, [convert.compare_squash("SMB signing dis"), convert.compare_squash("abled")])
    got = shown(paragraph)
    passed &= check(got == "SMB signing dis" + soft + "|abled",
                    "break: a soft hyphen cut by two passes gets one break", got)
    got = split(text, [edged("the host", "t"), edged("\u2014including the isolated ones", "s", "\u2014")])
    passed &= check(got == "the host|\u2014including the isolated ones", "break: before an unspaced em dash that began the next line", got)
    text = "run nmap -sV against every host"
    got = split(text, [edged("run nmap", "p"), edged("-sV against every host", "t", "-")])
    passed &= check(got == "run nmap |-sV against every host", "break: a dash after the space is the next line's own", got)

    # A space or dash at a run boundary belongs to the end of the line.
    def hyperlink(paragraph, value):
        link = convert.OxmlElement("w:hyperlink")
        run = convert.OxmlElement("w:r")
        node = convert.OxmlElement("w:t")
        node.text = value
        node.set(qn("xml:space"), "preserve")
        run.append(node)
        link.append(run)
        paragraph._p.append(link)

    paragraph = Document().add_paragraph("Read the update guide for")
    hyperlink(paragraph, " CVE-2021-34527 today")
    convert.split_paragraph_at(paragraph._p, [edged("Read the update guide for", "r"), edged("CVE-2021-34527 today", "y")])
    got = shown(paragraph)
    passed &= check(got == "Read the update guide for |CVE-2021-34527 today",
                    "break: the space before a wrapped link stays on the line above", got)

    paragraph = Document().add_paragraph()
    paragraph.add_run("see the host")
    paragraph.add_run(" \u2014 including the rest").bold = True
    convert.split_paragraph_at(paragraph._p, [edged("see the host \u2014", "\u2014"), edged("including the rest", "t")])
    got = shown(paragraph)
    passed &= check(got == "see the host \u2014 |including the rest",
                    "break: an em dash that opens the next run goes back to the line it ended", got)

    paragraph = Document().add_paragraph()
    paragraph.add_run("see the host")
    paragraph.add_run(" \u2014 including the rest").bold = True
    convert.split_paragraph_at(paragraph._p, [edged("see the host", "t"), edged("\u2014 including the rest", "t", "\u2014")])
    got = shown(paragraph)
    passed &= check(got == "see the host |\u2014 including the rest",
                    "break: an em dash that opens the next line stays with it", got)

    paragraph = Document().add_paragraph("alpha beta gamma delta")
    convert.split_paragraph_at(paragraph._p, [edged("alpha beta", "a"), edged("gamma delta", "a")])
    got = shown(paragraph)
    passed &= check(got == "alpha beta |gamma delta", "break: an ordinary break is unchanged", got)

    # The rows a run was made from give the pass what the squashed lines lost.
    pdf = pymupdf.open()
    page = pdf.new_page(width=612, height=792)
    page.insert_text((72, 100), "A paragraph line that begins at the edge and ends on the server-", fontsize=11, fontname="helv")
    page.insert_text((72, 114), "forgery issue the team confirmed again.", fontsize=11, fontname="helv")
    caption = "Nmap scan of the internal network showing SMB signing disabled on hosts"
    shown_width = pymupdf.get_text_length(caption, fontname="helv", fontsize=11)
    page.insert_text(((612 - shown_width) / 2.0, 200), caption, fontsize=11, fontname="helv")
    runs = convert.flush_left_runs(pdf, [], 72.0, 540.0)
    convert.annotate_runs(pdf, [], runs, 72.0, 540.0)
    by_first = {run["lines"][0][:6]: run for run in runs}
    prose = by_first.get("aparag")
    fig = by_first.get("nmapsc")
    passed &= check(prose is not None and prose["lines"][0].last == "-" and prose["lines"][1].last == ".",
                    "annotate: each line keeps the last character its row ended on",
                    prose and [line.last for line in prose["lines"]])
    kinds = dict(fig["pageRows"]) if fig else {}
    passed &= check(fig is not None and kinds.get(convert.compare_squash(caption)) == "C",
                    "annotate: the page's rows are kept with their kind, the centred caption among them", kinds)
    pdf.close()
    return passed


def run_run_checks():
    """Runs of lines: where they end, and when they count as justified."""
    passed = True

    def row(text, x0, x1, y0, size=11.0):
        return {"x0": x0, "x1": x1, "y0": y0, "y1": y0 + size, "base": y0 + 9.0, "text": text,
                "columns": [{"x0": x0, "x1": x1, "y0": y0, "y1": y0 + size, "base": y0 + 9.0, "text": text,
                             "size": size, "font": "Poppins", "psname": "Poppins-Regular", "bold": False, "color": 0}]}

    pitch = 14.0
    rows = []
    y = 100.0
    # A ragged paragraph inside a card, at the card's indent.
    for text, x1 in (("Magna aliqua Ut enim ad minim veniam quis", 500.0), ("ex ea commodo consequat Duis aute irure", 520.0),
                     ("consequat Duis.", 300.0)):
        rows.append(row(text, 64.8, x1, y)); y += pitch
    y += 30.0
    # A justified paragraph at the page's left edge, then its own short last line.
    for text, x1 in (("Nisi ut aliquip ex ea commodo consequat Duis", 558.0), ("esse cillum dolore eu fugiat nulla pariatur", 557.2),
                     ("elit Sed do eiusmod tempor incididunt ut labore", 558.4), ("ut labore et dolore magna.", 250.0)):
        rows.append(row(text, 54.0, x1, y)); y += pitch
    y += 30.0
    # A heading just above its paragraph, within the line gap: one run.
    rows.append(row("12:24 UTC: Unauthenticated access check", 54.0, 287.0, y)); y += pitch
    for text, x1 in (("Quis nostrud exercitation ullamco laboris nisi", 558.0), ("aute irure. Lorem ipsum dolor sit amet", 557.5),
                     ("incididunt. The activity here is recorded", 558.0), ("ullamco laboris nisi ut aliquip.", 320.0)):
        rows.append(row(text, 54.0, x1, y)); y += pitch

    keep = convert.page_rows
    convert.page_rows = lambda page, logical=False, blocks=None: rows
    try:
        doc = pymupdf.open()
        doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        found = convert.flush_left_runs(doc, [], 54.0, 558.0)
    finally:
        convert.page_rows = keep
    shape = [(len(r["lines"]), r["justified"], r.get("indented")) for r in found]
    passed &= check(shape == [(3, False, True), (4, True, False), (5, True, False)],
                    "runs: a ragged card paragraph is one run, a justified one is, a heading joins its paragraph", shape)
    if len(found) == 3:
        passed &= check(found[0]["short"] == [True, False, True],
                        "runs: short is measured against the run's own longest line", found[0]["short"])
        passed &= check(found[2]["short"][0] and not found[2]["short"][1],
                        "runs: the heading is short, the justified line under it is not", found[2]["short"])
    return passed


def run_token_space_checks():
    """A blank span between two coloured tokens is a space the Word file keeps."""
    from pdf2docx.text.Spans import Spans
    from pdf2docx.text.TextSpan import TextSpan
    from pdf2docx.image.ImageSpan import ImageSpan

    passed = True

    def span(text, x0, style=None):
        raw = {"bbox": (x0, 0, x0 + 5 * len(text), 10), "text": text, "font": "Mono", "size": 10, "color": 0}
        if style:
            raw["style"] = style
        return raw

    def picture(x0):
        return {"image": b"", "bbox": (x0, 0, x0 + 10, 10)}

    def restored(raws, declined=None):
        wrapped = convert.spans_keep_inner_blanks(Spans.restore, TextSpan, ImageSpan, declined)
        return [item.text if isinstance(item, TextSpan) else "<image>" for item in wrapped(Spans(), raws)]

    # typst's `def exploit`: keyword, a span of its own holding the space, name.
    passed &= check(restored([span("def", 0), span(" ", 15), span("exploit", 20)]) == ["def", " ", "exploit"],
                    "token spaces: a blank span between two text spans is kept")
    passed &= check(restored([span(" ", 0), span("def", 5), span(" ", 20), span("x", 25), span(" ", 30)])
                    == ["def", " ", "x"],
                    "token spaces: a leading and a trailing blank span are still dropped")
    passed &= check(restored([span("a", 0), span(" ", 5), span(" ", 10)]) == ["a"],
                    "token spaces: blanks after the last text span are dropped")
    style = [{"type": 1, "color": 0}]
    passed &= check(restored([span(" ", 0, style), span("a", 5), span(" ", 10, style)]) == [" ", "a", " "],
                    "token spaces: a blank span with a style is kept wherever it is, as pdf2docx keeps it")
    passed &= check(restored([span(" ", 0), span(" ", 5)]) == [],
                    "token spaces: a line of blanks alone is still nothing")
    passed &= check(restored([span("a", 0), span(" ", 5), picture(10)]) == ["a", " ", "<image>"],
                    "token spaces: a picture counts as the text a blank span stands beside")
    passed &= check(restored([picture(0), span(" ", 10), span("a", 15)]) == ["<image>", " ", "a"],
                    "token spaces: a picture before the blank counts too")
    passed &= check(restored([]) == [], "token spaces: a line with no spans stays empty")

    # A pdf2docx whose spans no longer expose their text is handed to its own
    # method and counted.
    fell = convert.Counter()
    calls = []

    def theirs(self, raws):
        calls.append(raws)
        return self

    class Renamed:
        def __init__(self, raw):
            self.words = raw["text"]

    unreadable = convert.spans_keep_inner_blanks(theirs, Renamed, ImageSpan, fell)
    unreadable(Spans(), [span("x", 0)])
    passed &= check(len(calls) == 1 and fell["spaces-fallthrough"] == 1,
                    "token spaces: spans that cannot be read fall through to pdf2docx's method, counted", dict(fell))

    # End to end: typst draws the space in `def exploit(target):` as a span of
    # its own between two colours. Raw pdf2docx loses it; the converter must not.
    source = """
#set page(paper: "a4", margin: 2cm)
#set text(size: 10pt)
#show raw.where(block: true): it => block(width: 100%, fill: luma(245), inset: 8pt, it)
Exploit used against the CGI endpoint:
```python
import requests
def exploit(target):
    return requests.get(target)
```
```bash
nmap -sV 10.0.0.5 --script vuln && echo done
```
"""
    workdir = tempfile.mkdtemp(prefix="btct-token-space-test-")
    try:
        result, path = run_case(source, workdir, "token-space")
        text = convert.docx_text(path, body_only=True)
        passed &= check("def exploit(target):" in text and "import requests" in text,
                        "token spaces: the Word file keeps the space after def", text[:200])
        passed &= check("echo done" in text and "nmap -sV" in text,
                        "token spaces: the Word file keeps the spaces in a shell line", text[-120:])
        passed &= check("spaces-fallthrough" not in result["declined"],
                        "token spaces: the wrapper applied", result["declined"])
        raw_path = os.path.join(workdir, "token-space-raw.docx")
        convert.convert(os.path.join(workdir, "token-space.pdf"), raw_path, raw=True)
        raw_text = convert.docx_text(raw_path, body_only=True)
        passed &= check("defexploit(target):" in raw_text.replace(" ", "") and "def exploit(target):" not in raw_text,
                        "token spaces: raw pdf2docx still loses it (if not, the wrapper can go)", raw_text[:200])
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return passed


def run_thread_checks():
    """pdf2docx imports under the converter's own address-space cap.

    numpy's OpenBLAS and OpenCV size their thread pools to the CPUs the kernel
    reports, which in a container is the host's. On a 24-core host that import
    mapped 4.4 GB and died of SIGSEGV under the 3 GB cap, so every Word export
    failed before it began. The capped import runs only where `resource`
    exists, which is Linux, where the server runs.
    """
    passed = True
    pinned = {name: os.environ.get(name) for name in ("OPENBLAS_NUM_THREADS", "OMP_NUM_THREADS")}
    passed &= check(all(value == "1" for value in pinned.values()),
                    "threads: the numeric libraries are pinned to one thread", pinned)
    try:
        import resource  # noqa: F401
    except ImportError:
        print("SKIP  threads: no `resource` module here, so the capped import runs on Linux only")
        return passed
    here = os.path.dirname(os.path.abspath(__file__))
    probe = ("import sys; sys.path.insert(0, %r); import convert; "
             "convert.cap_address_space(); import pdf2docx, cv2; print('ok')" % here)
    # Without the variables this process already has, so only convert.py can
    # be what pins the threads in the child.
    env = {k: v for k, v in os.environ.items() if not k.endswith("_NUM_THREADS")}
    child = subprocess.run([sys.executable, "-I", "-B", "-c", probe], capture_output=True, text=True, env=env)
    passed &= check(child.returncode == 0 and child.stdout.strip().endswith("ok"),
                    "threads: pdf2docx and OpenCV import under the address-space cap",
                    (child.returncode, child.stderr[-300:]))
    return passed


def code_cell_paragraph(document, pieces, width_twips=9536, margin_twips=180, font="DejaVu Sans Mono"):
    """A paragraph in a one-cell table of the given width, one run per piece, soft wraps and no breaks."""
    qn = convert.qn
    table = document.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    properties = cell._tc.get_or_add_tcPr()
    size = properties.find(qn("w:tcW"))
    if size is None:
        size = convert.OxmlElement("w:tcW")
        properties.append(size)
    size.set(qn("w:w"), str(width_twips))
    size.set(qn("w:type"), "dxa")
    margins = convert.OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        node = convert.OxmlElement("w:" + side)
        node.set(qn("w:w"), str(margin_twips if side in ("start", "end") else 0))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)
    paragraph = cell.paragraphs[0]
    for piece in pieces:
        run = paragraph.add_run(piece + " ")
        run.font.size = Pt(9)
        run.font.name = font
    return paragraph


def paragraph_lines(element):
    """A paragraph's text split at its line breaks, trailing spaces dropped."""
    qn = convert.qn
    lines = [""]
    for run in element.iter(qn("w:r")):
        for child in run:
            if child.tag == qn("w:t"):
                lines[-1] += child.text or ""
            elif child.tag == qn("w:br"):
                lines.append("")
    return [line.rstrip() for line in lines]


def has_indent(element):
    properties = element.find(convert.qn("w:pPr"))
    return properties is not None and properties.find(convert.qn("w:ind")) is not None


def run_checkbox_checks():
    """A ticked box is a bordered character in its label's line, not a table above it."""
    passed = True
    qn = convert.qn
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    def box_table(width=140, height=140, text="", borders=4):
        sides = "".join('<w:%s w:val="single" w:sz="11" w:color="#1A2A40"/>' % side
                        for side in ("start", "top", "end", "bottom")[:borders])
        return (
            '<w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w="%d"/></w:tblGrid><w:tr>'
            '<w:trPr><w:trHeight w:hRule="exact" w:val="%d"/></w:trPr><w:tc><w:tcPr><w:tcW w:type="dxa" w:w="%d"/>'
            "<w:tcBorders>%s</w:tcBorders></w:tcPr><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:tc></w:tr></w:tbl>"
            % (width, height, width, sides, text)
        )

    def run(text, size=20):
        return '<w:r><w:rPr><w:sz w:val="%d"/></w:rPr><w:t xml:space="preserve">%s</w:t></w:r>' % (size, text)

    def stub():
        return '<w:p><w:pPr><w:spacing w:line="96" w:lineRule="exact"/></w:pPr></w:p>'

    def cell_of(*blocks):
        return parse_xml('<w:tc %s><w:tcPr/>%s</w:tc>' % (nsdecls("w"), "".join(blocks)))

    def label(*runs):
        return "<w:p>%s</w:p>" % "".join(runs)

    # The shape pdf2docx wrote for a ticked item: a spacer, the empty box, then
    # the mark and the label in one paragraph.
    cell = cell_of(stub(), box_table(), label(run("x", 12), run(" "), run("Endpoint protection blocked")))
    document = Document()
    folded = convert.fold_checked_boxes_in(cell)
    blocks = [child.tag.split("}")[1] for child in cell if child.tag != qn("w:tcPr")]
    paragraph = cell.find(qn("w:p"))
    first = paragraph.find(qn("w:r"))
    properties = first.find(qn("w:rPr"))
    fit = properties.find(qn("w:fitText"))
    passed &= check(folded == 1 and blocks == ["p"],
                    "checkbox: the box table and its spacer are gone, the label paragraph stays", (folded, blocks))
    passed &= check(properties.find(qn("w:bdr")) is not None and fit is not None and fit.get(qn("w:val")) == "140",
                    "checkbox: the mark carries a character border fitted to the box's width", (properties.find(qn("w:bdr")), fit))
    passed &= check(convert.paragraph_text(paragraph) == "x Endpoint protection blocked",
                    "checkbox: the paragraph's text is what it was", convert.paragraph_text(paragraph))
    condensed = '<w:r><w:rPr><w:spacing w:val="-2"/><w:sz w:val="12"/></w:rPr><w:t xml:space="preserve">x</w:t></w:r>'
    cell = cell_of(stub(), box_table(), label(condensed, run(" "), run("Endpoint protection blocked")))
    convert.fold_checked_boxes_in(cell)
    mark = cell.find(qn("w:p")).find(qn("w:r")).find(qn("w:rPr"))
    passed &= check(mark.find(qn("w:fitText")) is not None and mark.find(qn("w:spacing")) is None,
                    "checkbox: the box is as wide as fitText says, with no spacing left from its line")

    # No mark in the paragraph: the box is its own run of a no-break space.
    cell = cell_of(box_table(), label(run("Tested quarterly")))
    folded = convert.fold_checked_boxes_in(cell)
    paragraph = cell.find(qn("w:p"))
    runs = paragraph.findall(qn("w:r"))
    passed &= check(folded == 1 and len(runs) == 2 and runs[0].find(qn("w:rPr")).find(qn("w:bdr")) is not None
                    and "".join(t.text for t in runs[0].iter(qn("w:t"))) == " ",
                    "checkbox: with no mark in the paragraph the box is a run of its own", (folded, len(runs)))

    # Things that are not a tick box are left alone.
    for name, blocks in (
        ("a 1x1 table that holds text", [box_table(text="x"), label(run("label"))]),
        ("a box with no label paragraph after it (the unticked shape)", [stub(), box_table(width=140, height=120), label()]),
        ("a bordered cell far larger than a box", [box_table(width=2000, height=400), label(run("label"))]),
        ("a cell with a border on two sides only", [box_table(borders=2), label(run("label"))]),
    ):
        cell = cell_of(*blocks)
        before = len(list(cell))
        folded = convert.fold_checked_boxes_in(cell)
        passed &= check(folded == 0 and len(list(cell)) == before, "checkbox: " + name + " is left alone", folded)
    return passed


def run_overflow_row_checks():
    """A code paragraph pdf2docx wrapped by itself, with a line typst ran off the page."""
    passed = True
    qn = convert.qn
    first = "$ wget -q https://10.10.5.198/downloads/"
    token = "releases/" + ("0123456789abcdef" * 6)[:89]
    third = "[+] india india bravo echo india india foxtrot delta juliet kilo echo charlie kilo"
    fourth = "delta golf foxtrot bravo juliet delta kilo echo charlie alpha echo lima bravo delta"
    rows = [{"text": text, "pieces": [text], "page": 1, "cut": text is token} for text in (first, token, third, fourth)]

    # pdf2docx kept the 98-character token whole (it ends a hair inside the
    # page), the cell is 458.8 pt of text and the token is 537 pt in Word.
    document = Document()
    paragraph = code_cell_paragraph(document, [first, token, third, fourth])
    restored = convert.restore_cut_lines(document, rows)
    breaks = sum(1 for _ in paragraph._p.iter(qn("w:br")))
    passed &= check(restored == 0 and breaks == 3 and paragraph_lines(paragraph._p) == [first, token, third, fourth],
                    "overflow: a whole token's paragraph is broken where the PDF broke it", (restored, breaks))
    indent = paragraph._p.find(qn("w:pPr")).find(qn("w:ind"))
    right = convert.twips_of(indent.get(qn("w:right"))) if indent is not None else None
    passed &= check(right is not None and abs(right + 10.0) < 0.2,
                    "overflow: the right indent reaches the cell's edge and no further, not the whole paragraph's length", right)

    # The token is 12% over the cell's edge, so it is set tighter by what it
    # takes to end there; the lines that fit are left as typst set them.
    def condensed(text):
        out = []
        for run in paragraph._p.iter(qn("w:r")):
            node = run.find(qn("w:t"))
            if node is not None and (node.text or "").strip() == text:
                spacing = run.find(qn("w:rPr")).find(qn("w:spacing"))
                out.append(None if spacing is None else convert.twips_of(spacing.get(qn("w:val"))) * 20.0)
        return out
    twips = condensed(token)
    width = len(token) * (convert.MONO_ADVANCE_EM * 9.0 * (1.0 + convert.WORD_WIDTH_EXCESS) + twips[0] / 20.0) if twips and twips[0] is not None else None
    limit = 458.8 + 9.0 - convert.CELL_EDGE_SLACK_PT
    passed &= check(width is not None and limit - len(token) * 0.05 <= width <= limit,
                    "overflow: a token 12% over the cell's edge is set tighter so that it ends at the edge", (twips, width, limit))
    passed &= check(condensed(third) == [None] and condensed(first) == [None],
                    "overflow: the lines that fit are not touched")

    # A 130-character token is 34% over: tightening it that far would run the
    # glyphs together, so it keeps typst's spacing and Word cuts it at the edge.
    long_token = "releases/" + ("0123456789abcdef" * 8)[:121]
    long_rows = [{"text": text, "pieces": [text], "page": 1, "cut": text is long_token} for text in (first, long_token, third)]
    document = Document()
    paragraph = code_cell_paragraph(document, [first, long_token, third])
    convert.restore_cut_lines(document, long_rows)
    indent = paragraph._p.find(qn("w:pPr")).find(qn("w:ind"))
    right = convert.twips_of(indent.get(qn("w:right"))) if indent is not None else None
    expected = -(len(long_token) * convert.MONO_ADVANCE_EM * 9.0 * (1.0 + convert.WORD_WIDTH_EXCESS) - 458.8 + 2.0)
    passed &= check(condensed(long_token) == [None] and right is not None and abs(right - expected) < 0.2
                    and paragraph_lines(paragraph._p) == [first, long_token, third],
                    "overflow: a token too wide to tighten keeps typst's spacing and gets the indent that keeps its line whole",
                    (condensed(long_token), right, expected))

    # pdf2docx cut the token at the page edge: the end is put back, the
    # paragraph broken and the indent written, in that order.
    document = Document()
    paragraph = code_cell_paragraph(document, [first, token[:84], third])
    restored = convert.restore_cut_lines(document, rows[:3])
    passed &= check(restored == 1 and paragraph_lines(paragraph._p) == [first, token, third],
                    "overflow: a cut token is completed and its paragraph broken at the PDF's lines",
                    (restored, paragraph_lines(paragraph._p)))

    # Rows that do not account for the whole paragraph (a line the PDF has no
    # row for) leave the breaks alone: half a paragraph broken and the rest on
    # one line running off the page is worse than letting Word wrap.
    document = Document()
    paragraph = code_cell_paragraph(document, [first, token, third, "an extra line the rows never held"])
    convert.restore_cut_lines(document, rows[:3])
    passed &= check(sum(1 for _ in paragraph._p.iter(qn("w:br"))) == 0 and not has_indent(paragraph._p),
                    "overflow: rows that are not all of the paragraph leave a whole token's paragraph alone")

    # Prose is never given a negative indent: it would widen every line of it.
    document = Document()
    paragraph = code_cell_paragraph(document, [first, token, third], font="Poppins")
    convert.restore_cut_lines(document, rows[:3])
    passed &= check(sum(1 for _ in paragraph._p.iter(qn("w:br"))) == 0 and not has_indent(paragraph._p),
                    "overflow: a paragraph that is not code is left to wrap")

    # Where the rows are found.
    paragraph = code_cell_paragraph(Document(), [first, token, third, fourth])
    cuts = convert.code_row_cuts(paragraph._p, rows, 1)
    raw = convert.paragraph_text(paragraph._p)
    ends = [raw.index(text) + len(text) for text in (first, token, third)]
    passed &= check(cuts == ends and convert.code_row_cuts(paragraph._p, rows[1:3], 0) is None,
                    "overflow: the cuts are the ends of every row but the last, and None unless the rows are all of it", (cuts, ends))
    return passed


def run_line_shape_checks():
    """Condensed runs, tabs before short lines, and text put back where pdf2docx cut it."""
    passed = True
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("3. Service discovery")
    paragraph.add_run().add_break()
    second = paragraph.add_run("Magna aliqua Ut enim ad minim veniam quis nostrud")
    spacing = convert.OxmlElement("w:spacing")
    spacing.set(convert.qn("w:val"), "40")
    second._element.get_or_add_rPr().append(spacing)
    convert.condense_runs(paragraph._p, 2)
    values = [r.find(convert.qn("w:rPr")).find(convert.qn("w:spacing")).get(convert.qn("w:val"))
              for r in paragraph._p.iter(convert.qn("w:r")) if r.find(convert.qn("w:t")) is not None]
    passed &= check(values == ["-2", "40"], "condense: every text run is tightened, one that had a spacing keeps it", values)
    convert.tab_before_short_breaks(paragraph._p, [convert.compare_squash("3. Service discovery")])
    nodes = [child.tag.split("}")[1] for run in paragraph._p.iter(convert.qn("w:r")) for child in run
             if child.tag in (convert.qn("w:tab"), convert.qn("w:br"), convert.qn("w:t"))]
    passed &= check(nodes == ["t", "tab", "br", "t"], "short line: a tab goes before the break that ends it", nodes)

    paragraph = document.add_paragraph()
    paragraph.add_run("abc def ghi jkl mno")
    convert.split_paragraph_at(paragraph._p, ["abcdef", "ghijkl", "mno"], tabs=[True, False, False])
    nodes = [child.tag.split("}")[1] for run in paragraph._p.iter(convert.qn("w:r")) for child in run
             if child.tag in (convert.qn("w:tab"), convert.qn("w:br"), convert.qn("w:t"))]
    passed &= check(nodes == ["t", "tab", "br", "t", "br", "t"],
                    "split: only the line flagged short gets a tab before its break", nodes)

    # A list item that is only a dash is three runs (the bullet, a space, the
    # dash) with the converter's own break after the dash. The line ends on
    # the dash, two runs past the cut: the space and the dash both go back
    # across the boundary and the converter's break is reused, not doubled.
    class Edged(str):
        pass
    first = Edged("•")
    first.first, first.last = "•", "–"
    second = Edged("•alistitemthatkeeps")
    second.first, second.last = "•", "s"
    paragraph = document.add_paragraph()
    for piece in ("•", " ", "–"):
        paragraph.add_run(piece)
    paragraph.add_run().add_break()
    paragraph.add_run("• A list item that keeps going on")
    convert.split_paragraph_at(paragraph._p, [first, second, "goingon"])
    pieces = [(child.tag.split("}")[1], child.text) for run in paragraph._p.iter(convert.qn("w:r")) for child in run
              if child.tag in (convert.qn("w:br"), convert.qn("w:t")) and (child.tag == convert.qn("w:br") or child.text)]
    before_break = "".join(text for tag, text in pieces[:pieces.index(("br", None))])
    passed &= check(before_break == "• –" and [tag for tag, _ in pieces].count("br") == 2,
                    "split: a line that ends on a dash two runs past the cut takes the dash and reuses the converter's break", pieces)

    # A cut that lands where the converter already broke the line (a heading
    # and its paragraph as one block) reuses that break instead of adding one.
    paragraph = document.add_paragraph()
    paragraph.add_run("4.1.1 Dolore eu fugiat ")
    paragraph.add_run().add_break()
    paragraph.add_run("Veniam quis nostrud exercitation ullamco laboris nisi ut aliquip ex ea")
    added = convert.split_paragraph_at(paragraph._p, ["4.1.1doloreeufugiat", "veniamquisnostrudexercitationullamco", "laborisnisiutaliquipexea"],
                                       tabs=[True, False, False])
    nodes = [child.tag.split("}")[1] for run in paragraph._p.iter(convert.qn("w:r")) for child in run
             if child.tag in (convert.qn("w:tab"), convert.qn("w:br"), convert.qn("w:t"))]
    passed &= check(added == 2 and nodes == ["t", "tab", "br", "t", "br", "t"],
                    "split: the converter's own break is reused, with the tab the short heading needs", (added, nodes))

    # pdf2docx cut both tokens at the page's edge and joined the stumps.
    first = "ff4629fd845fa1cd176399ea65b7059171fff6b3030246c47d407a174ed4463d7e431f0a7bac4fc2769d97834453b40f/"
    second_token = "f04b35443879d9672cf4cab7a0f134e7d3644de471a704d74c6420303b6fff1719507b56ae993671dc1af548df9264ff/"
    paragraph = document.add_paragraph()
    paragraph.add_run(first[:84] + " " + second_token[:84])
    lines = [{"text": first, "pieces": [first], "cut": True}, {"text": second_token, "pieces": [second_token], "cut": True}]
    restored = convert.restore_cut_lines(document, lines)
    # A shorter token that begins like the long one is not a cut of it.
    short_paragraph = document.add_paragraph()
    short_paragraph.add_run("Bearer " + first[:40] + "' done")
    passed &= check(convert.restore_cut_lines(document, [{"text": first, "pieces": [first], "cut": True}]) == 0
                    and convert.paragraph_text(short_paragraph._p) == "Bearer " + first[:40] + "' done",
                    "cut lines: a 40-character token is not completed into the 96-character one")
    passed &= check(convert.restore_cut_lines(document, [{"text": first, "pieces": [first], "cut": False}]) == 0,
                    "cut lines: a line that did not reach the page's edge is never touched")
    text = convert.paragraph_text(paragraph._p)
    passed &= check(restored == 2 and first in text and second_token in text and text.index(first) < text.index(second_token),
                    "cut lines: both tokens are whole again, in order", (restored, text[:60]))
    passed &= check(convert.missing_from_docx(lines, text) == [], "cut lines: and the text check is satisfied")

    # A paragraph's tab STOPS share the element name of the tab character, and
    # a panel line's "leading tabs" were found by walking the paragraph from
    # its top: the stops came first and were removed with the real ones.
    titled = document.add_paragraph()
    titled.paragraph_format.tab_stops.add_tab_stop(Pt(62.3))
    titled.add_run("6.1.2 ")
    titled.add_run().add_tab()
    titled.add_run("Weak configuration setting 2")
    nodes = [n for n in titled._p.iter(convert.qn("w:t"))]
    passed &= check(convert.leading_tabs(nodes) == [], "tabs: a tab stop is not a tab before the first letter")
    indented = document.add_paragraph()
    indented.paragraph_format.tab_stops.add_tab_stop(Pt(62.3))
    indented.add_run().add_tab()
    indented.add_run("listen: 8080")
    nodes = [n for n in indented._p.iter(convert.qn("w:t"))]
    found = convert.leading_tabs(nodes)
    passed &= check(len(found) == 1 and found[0].getparent().tag == convert.qn("w:r"),
                    "tabs: a tab character before the first letter is still one")
    stop = titled._p.find(convert.qn("w:pPr")).find(convert.qn("w:tabs")).find(convert.qn("w:tab"))
    convert.shift_tab_stops(titled._p, -9.6)
    passed &= check(abs(convert.twips_of(stop.get(convert.qn("w:pos"))) - 52.7) < 0.1,
                    "tabs: a stop moves with the text when its container changes")
    convert.shift_tab_stops(titled._p, -60.0)
    passed &= check(titled._p.find(convert.qn("w:pPr")).find(convert.qn("w:tabs")).find(convert.qn("w:tab")) is None,
                    "tabs: a stop that would fall before the text area is dropped")
    return passed & run_run_column_checks() & run_edge_decision_checks()



def run_edge_decision_checks():
    """Justification past hanging punctuation, centred cells, and the text area of each page size."""
    passed = True
    centre = convert.WD_ALIGN_PARAGRAPH.CENTER

    def row(text, x0, x1, y0=0.0, size=11.0):
        return {"x0": x0, "x1": x1, "y0": y0, "y1": y0 + size, "base": y0 + 9.0, "text": text,
                "columns": [{"x0": x0, "x1": x1, "y0": y0, "y1": y0 + size, "base": y0 + 9.0, "text": text,
                             "size": size, "font": "Libertinus Serif", "psname": "Libertinus-Regular", "bold": False, "color": 0}]}

    # A justified paragraph whose widest line ends in a comma (dashes.typ, p1, New Computer Modern 11 pt).
    hung = [row("one two three", 72.0, 543.7), row("four five six", 72.0, 543.7), row("seven eight nine", 72.0, 543.7),
            row("ten eleven share,", 72.0, 546.1), row("the last line.", 72.0, 375.0)]
    reaching = len(convert.lines_reaching_edge(hung))
    passed &= check(reaching == 4, "justify: a line hanging 2.4 pt past the column reaches it, as the ones that do not hang do", reaching)
    # Three Libertinus lines of a ragged paragraph (dashes.typ): one hangs 1.3 pt past, one is 3.6 pt short.
    ragged = [row("first line of servers,", 72.0, 535.2), row("second line of hosts", 72.0, 533.9),
              row("third line of text", 72.0, 530.3), row("same network segment.", 72.0, 257.0)]
    reaching = len(convert.lines_reaching_edge(ragged))
    passed &= check(reaching < 2, "justify: a hanging line alone does not make ragged lines justified", reaching)
    # Arial (old_c.typ): a word, a comma (+2.5), a hyphen (-1.0, no space after it), a word, an em dash (-0.9).
    mixed = [row("out-of-band", 72.0, 543.0), row("portal,", 72.0, 545.5), row("request-", 72.0, 542.0),
             row("the", 72.0, 543.0), row("environment\u2014", 72.0, 542.1), row("segments.", 72.0, 356.9)]
    reaching = len(convert.lines_reaching_edge(mixed))
    passed &= check(reaching == 5, "justify: lines broken after a hyphen or a dash end a space short and still reach", reaching)
    three = [row("guidance rotate the account,", 72.0, 545.4), row("remove the flag from hosts", 72.0, 543.0),
             row("service names for accounts.", 72.0, 431.4)]
    reaching = len(convert.lines_reaching_edge(three))
    passed &= check(reaching == 2, "justify: a three-line paragraph with one hanging line and one firm one is justified", reaching)
    every = [row("a b c,", 72.0, 545.0), row("d e f.", 72.0, 545.2), row("g h i,", 72.0, 545.1), row("j k.", 72.0, 300.0)]
    reaching = len(convert.lines_reaching_edge(every))
    passed &= check(reaching == 3, "justify: when every line hangs the widest stands in for the column", reaching)

    # Through flush_left_runs, with the rows above on a page.
    on_page = [dict(r, y0=100.0 + 14.0 * n, y1=111.0 + 14.0 * n) for n, r in enumerate(hung)]
    keep = convert.page_rows
    convert.page_rows = lambda page, logical=False, blocks=None: on_page
    try:
        doc = pymupdf.open()
        doc.new_page(width=612.0, height=792.0)
        found = convert.flush_left_runs(doc, [], 72.0, 546.1)
    finally:
        convert.page_rows = keep
    passed &= check(len(found) == 1 and found[0]["justified"] is True,
                    "justify: the paragraph that ends a line in a comma is a justified run", [r["justified"] for r in found])

    # A page whose margin differs from the document's (mirrored margins: the
    # even pages stand 36 pt left of the odd ones) forms its runs from its own
    # edge, and a page that has rows on the document's edge keeps it.
    def at(x0, n):
        return {"x0": x0, "x1": x0 + 400.0, "y0": 100.0 + 14.0 * n, "y1": 111.0 + 14.0 * n, "text": "line %d" % n,
                "columns": [{"x0": x0, "x1": x0 + 400.0, "text": "line %d" % n}]}
    passed &= check(convert.page_text_left([at(54.0, n) for n in range(4)], 72.0) == 54.0
                    and convert.page_text_left([at(54.0, n) for n in range(2)], 72.0) == 72.0
                    and convert.page_text_left([at(54.0, n) for n in range(4)] + [at(72.0, 4)], 72.0) == 72.0
                    and convert.page_text_left([at(100.0, n) for n in range(5)], 72.0) == 72.0
                    and convert.page_text_left([at(40.0, 0), at(40.5, 1), at(40.0, 2)] + [at(54.0, n) for n in range(3, 8)], 72.0) == 40.0,
                    "edge: a page with no row on the document's edge takes its own, from three rows, to the left only")
    shifted = [dict(r, x0=r["x0"] - 36.0, x1=r["x1"] - 36.0) for r in on_page]
    convert.page_rows = lambda page, logical=False, blocks=None: shifted
    try:
        doc = pymupdf.open()
        doc.new_page(width=612.0, height=792.0)
        found = convert.flush_left_runs(doc, [], 72.0, 546.1)
    finally:
        convert.page_rows = keep
    passed &= check(len(found) == 1 and found[0]["lines"] == [convert.compare_squash(r["text"]) for r in shifted]
                    and found[0]["justified"] is True and found[0]["indented"] is False,
                    "edge: a page 36 pt left of the document's edge still makes its run", [r["lines"] for r in found])

    # The text area of a page size: the right edge without what hangs, less the space a wrapped line carries.
    edges = convert.text_area_edges([(595.3, [row("a line", 90.7, 547.0), row("hang,", 90.7, 549.5)]),
                                     (841.9, [row("wide line", 56.7, 766.5)])])
    portrait = edges.get(595)
    passed &= check(portrait is not None and portrait[0] == 90.7 and abs(portrait[1] - (547.0 - 0.33 * 11.0)) < 0.01,
                    "edges: a page size's right edge leaves out the hanging line and the space its box carries", portrait)
    passed &= check(edges.get(842) == (56.7, 766.5 - 0.33 * 11.0), "edges: each page size has its own", edges.get(842))
    # Left 3.2 cm, right 1.8 cm (centre.typ): a bold centred line stood 2.8 pt off against the box's end.
    left, right = edges[595]
    centred_line = row("Severity: Critical", 275.7, 359.2)
    got = convert.row_kind(centred_line, left, right, 595.3)
    old = convert.row_kind(centred_line, 90.7, 547.0, 595.3)
    passed &= check(got == "C" and old == "O",
                    "edges: a centred line on a page with unequal margins is found against the page's own text area", (got, old))
    offsets = convert.page_edge_offsets([(595, 547.0), (842, 766.5)], 656.8)
    passed &= check(offsets == {595: 547.0 - 656.8, 842: 766.5 - 656.8},
                    "edges: with as many pages of one size as of the other, neither keeps the median between them", offsets)

    # A cell's text is centred in the cell: stacks of lines on one axis.
    cell = [row("Unconstrained Kerberos", 77.6, 178.9, 100.0), row("delegation on the legacy", 77.2, 179.3, 112.0),
            row("file server cluster", 91.7, 162.3, 124.0)]
    flagged = convert.centred_stacks(cell)
    # The first two lines are 0.4 pt apart at the left, as left-aligned lines are, so they say nothing to each
    # other; the third pairs with the second, and the paragraph is judged by any of its lines.
    passed &= check(flagged == {1, 2}, "stacks: lines of a centred cell that start apart share one axis", flagged)
    left_aligned = [row("SMB signing not required", 200.0, 300.0, 100.0), row("on domain members", 200.0, 270.0, 112.0)]
    passed &= check(convert.centred_stacks(left_aligned) == set(), "stacks: lines that start together are not centred",
                    convert.centred_stacks(left_aligned))
    item = [row("• first item that wraps", 100.0, 200.0, 100.0), row("onto a second line", 112.0, 188.0, 112.0)]
    passed &= check(convert.centred_stacks(item) == set(), "stacks: a list item's wrapped line is not a centred one",
                    convert.centred_stacks(item))
    beside = [row("left cell text", 77.0, 160.0, 100.0), row("right cell text", 300.0, 400.0, 100.0),
              row("more of the left", 80.0, 157.0, 112.0)]
    passed &= check(convert.centred_stacks(beside) == set(),
                    "stacks: a row in another column does not pair with this one", convert.centred_stacks(beside))
    side_rows = [dict(r, columns=r["columns"] * 2) for r in cell]
    passed &= check(convert.centred_stacks(side_rows) == set(), "stacks: a row read across cells is no one cell's line")

    # The pass: a centred cell the converter centred stays centred when one line stands at the table's edge.
    def cell_run(lines, flags):
        return {"lines": lines, "justified": False, "indented": False, "page": None, "short": [False] * len(lines),
                "pageRows": [(line, "O") for line in lines], "centredLines": flags}

    def grid_paragraph(document, text):
        grid = document.add_table(rows=1, cols=1)
        paragraph = grid.rows[0].cells[0].paragraphs[0]
        paragraph.add_run(text)
        paragraph.alignment = centre
        return paragraph

    lines = [convert.compare_squash(t) for t in ("Unconstrained Kerberos", "delegation on the legacy", "file server cluster")]
    document = Document()
    wrapped = grid_paragraph(document, "Unconstrained Kerberos delegation on the legacy file server cluster")
    declined = collections.Counter()
    convert.align_paragraphs(document, [cell_run(lines[:2], [True, True]), cell_run(lines[2:], [True])], declined=declined)
    passed &= check(wrapped.alignment == centre and declined["align-centred"] == 1,
                    "cell: a wrapped, centred cell stays centred when its lines stand on one axis", (wrapped.alignment, declined))
    document = Document()
    wrapped = grid_paragraph(document, "Unconstrained Kerberos delegation on the legacy file server cluster")
    convert.align_paragraphs(document, [cell_run(lines[:2], [False, False]), cell_run(lines[2:], [False])])
    passed &= check(wrapped.alignment != centre, "cell: without the axis the old rule stands", wrapped.alignment)

    # The line breaks of a centred cell keep its two sides equal.
    document = Document()
    wrapped = grid_paragraph(document, "Any account that authenticates to the host leaves a forwardable ticket.")
    wrapped.paragraph_format.left_indent = Pt(14.4)
    wrapped.paragraph_format.right_indent = Pt(14.4)
    cell_lines = [{"page": 1, "text": "Any account that authenticates to the host leaves a forwardable",
                   "pieces": ["Any account that authenticates to the host leaves a forwardable"]},
                  {"page": 1, "text": "ticket.", "pieces": ["ticket."]}]
    convert.break_cell_lines(document, cell_lines, set(), page_count=1)
    left_indent = wrapped.paragraph_format.left_indent
    passed &= check(left_indent is None or abs(left_indent.pt) < 0.1,
                    "cell: zeroing a centred cell's right indent zeroes its left as well", left_indent)
    return passed


def run_run_column_checks():
    """A run's column, from the PDF to the right indent: a landscape page, a padded block, one column of two."""
    passed = True

    def row(text, x0, x1, y0, size=11.0):
        return {"x0": x0, "x1": x1, "y0": y0, "y1": y0 + size, "base": y0 + 9.0, "text": text,
                "columns": [{"x0": x0, "x1": x1, "y0": y0, "y1": y0 + size, "base": y0 + 9.0, "text": text,
                             "size": size, "font": "Poppins", "psname": "Poppins-Regular", "bold": False, "color": 0}]}

    def block(words, x0, x1, y, last=None, lines=4, pitch=14.0):
        """`lines` lines that all end at x1 (justified) and a short last one; returns the rows and the next y."""
        out = []
        for number in range(lines):
            out.append(row("%s line%d consectetur adipiscing elit sed" % (words, number), x0, x1, y))
            y += pitch
        out.append(row("%s ends here." % words, x0, last or x0 + 120.0, y))
        return out, y + pitch + 30.0

    portrait_rows = []
    y = 100.0
    padded, y = block("padded", 90.0, 522.0, y)
    plain, y = block("plain", 54.0, 558.0, y)
    halved, y = block("halved", 54.0, 290.0, y)
    portrait_rows += padded + plain + halved
    landscape_rows = []
    y = 100.0
    long_run, y = block("landscape", 54.0, 738.0, y)
    # Two lines: the first reaches the page's column, or stops short of it.
    reaching = [row("reaching two lines of justified text across a wide page", 54.0, 738.0, y),
                row("reaching ends here.", 54.0, 160.0, y + 14.0)]
    y += 14.0 + 14.0 + 30.0
    ragged = [row("ragged two lines of unjustified text across a wide page", 54.0, 690.0, y),
              row("ragged ends here.", 54.0, 150.0, y + 14.0)]
    landscape_rows += long_run + reaching + ragged
    by_page = {0: portrait_rows, 1: landscape_rows, 2: portrait_rows}

    keep = convert.page_rows
    convert.page_rows = lambda page, logical=False, blocks=None: by_page[page.number]
    try:
        doc = pymupdf.open()
        doc.new_page(width=612.0, height=792.0)
        doc.new_page(width=792.0, height=612.0)
        doc.new_page(width=612.0, height=792.0)
        found = convert.flush_left_runs(doc, [], 54.0, 558.0)
    finally:
        convert.page_rows = keep
    def find(page, word):
        return next((r for r in found if r["page"] == page and r["lines"][0].startswith(word)), None)

    wide = find(1, "reaching")
    narrow = find(1, "ragged")
    passed &= check(wide is not None and wide["justified"] and narrow is not None and not narrow["justified"],
                    "page sizes: two lines that reach the landscape page's own edge are justified, ones that stop 48 pt short are not",
                    (wide and wide["justified"], narrow and narrow["justified"]))
    columns = {name: round(find(page, name)["column"], 1) for page, name in
               ((0, "padded"), (0, "plain"), (0, "halved"), (1, "landscape"))}
    passed &= check(columns == {"padded": 522.0, "plain": 558.0, "halved": 290.0, "landscape": 738.0},
                    "page sizes: each run keeps the right end its own lines share", columns)
    over = round(find(1, "landscape")["overshoot"], 1)
    passed &= check(over == 0.0, "page sizes: overshoot is measured against the page's own size", over)
    offsets = convert.page_edge_offsets([(612, 558.0), (792, 738.0), (612, 558.0)], 558.0)
    passed &= check(offsets == {612: 0.0, 792: 180.0}, "page sizes: the size with most pages keeps the document's edge, the other has its own", offsets)
    passed &= check(convert.page_edge_offsets([(612, 558.0), (612, 560.0)], 559.0) == {},
                    "page sizes: a document of one size has no offsets")
    # Portrait pages that end at 508 or at 542 give a median of 523, which
    # no line ends on; the odd-sized pages are measured against what they share.
    mixed = [(612, 508.1), (612, 542.1), (612, 542.1), (612, 508.1), (792, 785.2), (792, 785.2), (792, 788.7)]
    got = convert.page_edge_offsets(mixed, 523.0)
    passed &= check(got[612] == 0.0 and abs(got[792] - (785.2 - 523.0)) < 0.01,
                    "page sizes: the odd size's edge is the one most of its pages share, not their median", got)
    # An outlier past the others does not move the column, and the widest
    # of the lines that share an edge is the column.
    passed &= check(convert.run_column([521.4, 522.0, 521.8, 560.0], 560.0) == 522.0,
                    "column: a line typst ran past the edge is not the column")
    passed &= check(convert.run_column([], 300.0) == 300.0, "column: no lines to go by is the widest end")
    passed &= check(convert.hangs_into_margin("consectetur adipis-") and convert.hangs_into_margin("elit sed,  ")
                    and not convert.hangs_into_margin("elit sed do") and not convert.hangs_into_margin(""),
                    "column: a line that ends in a hyphen or a comma hangs into the margin, one that ends in a letter does not")
    # Lines that end on a hyphen stand 2.9 pt past the column and say nothing
    # about it: the run's column is read off the lines that end on a letter,
    # and a run with none has no column to give.
    hung = [row("hung line%d consectetur adipiscing elit sed-" % n, 54.0, 560.9, 100.0 + 14.0 * n) for n in range(3)]
    hung += [row("hung line3 consectetur adipiscing elit sed do", 54.0, 558.0, 142.0),
             row("hung ends here.", 54.0, 174.0, 156.0)]
    allhung = [row("all line%d consectetur adipiscing elit sed-" % n, 54.0, 560.9, 300.0 + 14.0 * n) for n in range(3)]
    allhung.append(row("all ends here.", 54.0, 174.0, 342.0))
    convert.page_rows = lambda page, logical=False, blocks=None: hung + allhung
    try:
        single = pymupdf.open()
        single.new_page(width=612.0, height=792.0)
        seen = convert.flush_left_runs(single, [], 54.0, 558.0)
    finally:
        convert.page_rows = keep
    columns = {r["lines"][0][:3]: r["column"] for r in seen}
    passed &= check(columns.get("hun") == 558.0 and columns.get("all") is None,
                    "column: hyphenated lines do not set it, and a run of nothing else has none", columns)
    # The runs of the pages above are the ones the paragraphs below are matched to.
    convert.RUN_COLUMNS.reset(found, 558.0, {612: 0.0, 792: 180.0})

    document = Document()
    for section in document.sections:
        section.page_width, section.page_height = Pt(612.0), Pt(792.0)
        section.left_margin = section.right_margin = Pt(54.0)

    def paragraph_of(rows, align=convert.WD_ALIGN_PARAGRAPH.JUSTIFY):
        paragraph = document.add_paragraph(" ".join(r["text"] for r in rows))
        paragraph.alignment = align
        return paragraph

    pads = paragraph_of(padded)
    full = paragraph_of(plain)
    half = paragraph_of(halved)
    section = document.add_section()
    section.page_width, section.page_height = Pt(792.0), Pt(612.0)
    section.left_margin = section.right_margin = Pt(54.0)
    land = paragraph_of(long_run)
    short = paragraph_of(reaching)
    unclaimed = document.add_paragraph("Nothing in the PDF reads like this paragraph at all")
    got = {name: round(convert.column_right_indent(document, p._p, 558.0), 1)
           for name, p in (("padded", pads), ("plain", full), ("half", half), ("landscape", land), ("two", short),
                           ("unclaimed", unclaimed))}
    passed &= check(got == {"padded": 36.0, "plain": 0.0, "half": 268.0, "landscape": 0.0, "two": 0.0, "unclaimed": 0.0},
                    "indent: a padded block ends 36 pt in, one column of two at its own column, a landscape page at its edge", got)
    # A left-aligned paragraph never stretches, so it keeps the page's edge
    # whatever its run's column was.
    ragged = paragraph_of(halved, convert.WD_ALIGN_PARAGRAPH.LEFT)
    passed &= check(convert.column_right_indent(document, ragged._p, 558.0) == 0.0,
                    "indent: a paragraph that is not justified is measured against the page's edge")
    # The same text on a page of another size answers for that size.
    other = Document()
    other.sections[0].page_width, other.sections[0].page_height = Pt(792.0), Pt(612.0)
    other.sections[0].left_margin = other.sections[0].right_margin = Pt(54.0)
    elsewhere = other.add_paragraph(" ".join(r["text"] for r in halved))
    elsewhere.alignment = convert.WD_ALIGN_PARAGRAPH.JUSTIFY
    passed &= check(round(convert.column_right_indent(other, elsewhere._p, 558.0), 1) == 448.0,
                    "indent: text found only in a run of another page size still answers, with its own section's edge")
    # A paragraph that only overlaps a run (the run goes on into the next
    # paragraph, or starts before this one) is not that run's: another page
    # that repeats a sentence of it must not give it that page's column.
    overlap = document.add_paragraph("Before the run begins. " + " ".join(r["text"] for r in halved[2:]) + " Text that follows.")
    passed &= check(convert.column_right_indent(document, overlap._p, 558.0) == 0.0,
                    "indent: a paragraph holding only part of a run's lines is measured against the page's edge")
    # A line with a dot leader is a table of contents entry: its column is
    # not a text column's.
    leader = document.add_paragraph(" . . . . . . ".join(r["text"] for r in halved))
    leader.alignment = convert.WD_ALIGN_PARAGRAPH.JUSTIFY
    passed &= check(convert.column_right_indent(document, leader._p, 558.0) == 0.0,
                    "indent: a paragraph with a dot leader is measured against the page's edge")
    # The same text set ragged on another page, beside a justified run, is
    # laid out both ways, and no column is claimed.
    both = [dict(r) for r in found if r["page"] == 0 and r["lines"][0].startswith("halved")][:1]
    both.append(dict(both[0], justified=False))
    convert.RUN_COLUMNS.reset(both, 558.0, {792: 180.0})
    passed &= check(convert.column_right_indent(document, half._p, 558.0) == 0.0,
                    "indent: text that is ragged in one run and justified in another has no column")
    convert.RUN_COLUMNS.reset(found, 558.0, {792: 180.0})
    # Two runs of the same text that end on different edges say nothing.
    twins = [dict(r) for r in found if r["page"] == 0 and r["lines"][0].startswith("halved")][:1]
    twins.append(dict(twins[0], column=twins[0]["column"] + 30.0))
    convert.RUN_COLUMNS.reset(twins, 558.0, {})
    passed &= check(convert.column_right_indent(document, half._p, 558.0) == 0.0,
                    "indent: runs of one text that disagree on the column fall back to the page's edge")
    convert.RUN_COLUMNS.reset([], None, {})
    return passed


def toc_entry(title, number, page_index=2, y0=100.0, label=None, gap_before=0.0):
    """One detected contents line, shaped the way detect_toc returns it."""
    look = {"font": None, "psname": None, "size": 11.0, "bold": False, "color": 0}
    segments = [dict(look, offset=0.0, text=title)]
    if label:
        segments = [dict(look, offset=0.0, text=label), dict(look, offset=40.0, text=title)]
    return dict(look, page_index=page_index, y0=y0, x0=72.0, gap_before=gap_before,
                title=(label + " " + title) if label else title, segments=segments, number=str(number),
                leader_extra=0.0)


def toc_paragraph(parent_add, text, number=None, bold=False):
    """A converted line: a run of words, and the dot leader and number when it has them."""
    paragraph = parent_add("")
    run = paragraph.add_run(text if number is None else text + " " + ". " * 12 + str(number))
    run.font.bold = bold
    return paragraph


def run_toc_checks():
    """Rebuilding the contents when the converter merged its lines or left a heading among them."""
    passed = True
    leader = ". " * 12

    # 1. A page of contents joined into one paragraph, a heading glued to its front.
    document = Document()
    merged = document.add_paragraph()
    heading = merged.add_run("Contents ")
    heading.font.bold = True
    heading.font.size = Pt(20)
    for number, (title, page) in enumerate([("Introduction", 1), ("Scope of work", 2), ("Findings", 3)]):
        holder = convert.OxmlElement("w:r")
        holder.append(convert.OxmlElement("w:br"))
        merged._p.append(holder)
        merged.add_run("%s %s%d" % (title, leader, page))
    document.add_paragraph("Body text that follows the contents page.")
    runs = convert.leader_runs(convert.body_paragraphs(document))
    passed &= check(len(runs) == 1, "toc: one paragraph holding three leader lines is a stretch",
                    [len(r) for r in runs])
    entries = [toc_entry("Introduction", 1, y0=100.0), toc_entry("Scope of work", 2, y0=116.0),
               toc_entry("Findings", 3, y0=132.0)]
    declined = collections.Counter()
    written, warnings = convert.rebuild_toc(document, entries, 72.0, 523.0, declined)
    texts = [convert.paragraph_text(p) for p in convert.body_paragraphs(document)]
    passed &= check(written == 3 and not warnings and not declined,
                    "toc: a merged contents page is rebuilt line by line", (written, warnings, dict(declined)))
    passed &= check(texts[0].strip() == "Contents" and texts[1:4] == ["Introduction1", "Scope of work2", "Findings3"],
                    "toc: the heading glued to the first entry stays, its own paragraph, ahead of them", texts)
    kept = convert.body_paragraphs(document)[0]
    first_run = kept.find(convert.qn("w:r"))
    passed &= check(first_run is not None and first_run.find(convert.qn("w:rPr")).find(convert.qn("w:b")) is not None,
                    "toc: the kept heading keeps its own bold")
    passed &= check(texts[-1] == "Body text that follows the contents page.", "toc: what follows is untouched", texts[-1])

    # 2. A heading between two lists, in paragraphs of their own.
    document = Document()
    add = document.add_paragraph
    for text, number in (("Introduction", 1), ("Scope of work", 2), (None, None), ("Figure 1 Network map", 5),
                         ("Figure 2 Open ports", 6)):
        if text is None:
            add("List of figures").runs[0].font.size = Pt(16)
        else:
            toc_paragraph(add, text, number)
    entries = [toc_entry("Introduction", 1, y0=100.0), toc_entry("Scope of work", 2, y0=116.0),
               toc_entry("Network map", 5, y0=170.0, label="Figure 1", gap_before=20.0),
               toc_entry("Open ports", 6, y0=186.0, label="Figure 2")]
    declined = collections.Counter()
    written, warnings = convert.rebuild_toc(document, entries, 72.0, 523.0, declined)
    texts = [convert.paragraph_text(p) for p in convert.body_paragraphs(document)]
    passed &= check(written == 4 and not declined, "toc: four entries around a heading are rebuilt",
                    (written, warnings, dict(declined)))
    passed &= check(texts[:5] == ["Introduction1", "Scope of work2", "List of figures",
                                  "Figure 1Network map5", "Figure 2Open ports6"],
                    "toc: a heading between two lists is kept where it stood", texts)
    heading_run = convert.body_paragraphs(document)[2].find(convert.qn("w:r"))
    passed &= check(heading_run.find(convert.qn("w:rPr")).find(convert.qn("w:sz")).get(convert.qn("w:val")) == "32",
                    "toc: and keeps its size")

    # 3. The same in table rows: the table goes once its last row does, with the heading in a paragraph.
    document = Document()
    table = document.add_table(rows=4, cols=1)
    for row, (text, number) in zip(table.rows, (("Introduction", 1), ("Scope of work", 2), (None, None),
                                                ("Findings", 3))):
        cell = row.cells[0]
        if text is None:
            cell.paragraphs[0].add_run("List of tables")
        else:
            cell.paragraphs[0].add_run("%s %s%d" % (text, leader, number))
    entries = [toc_entry("Introduction", 1, y0=100.0), toc_entry("Scope of work", 2, y0=116.0),
               toc_entry("Findings", 3, y0=170.0)]
    declined = collections.Counter()
    written, warnings = convert.rebuild_toc(document, entries, 72.0, 523.0, declined)
    body = document.element.body
    texts = [convert.paragraph_text(p) for p in convert.body_paragraphs(document)]
    passed &= check(written == 3 and not declined and texts[:4] == ["Introduction1", "Scope of work2",
                                                                    "List of tables", "Findings3"],
                    "toc: a heading row among entry rows is kept", (written, dict(declined), texts))
    passed &= check(not [t for t in body.iter(convert.qn("w:tbl")) if not t.findall(convert.qn("w:tr"))],
                    "toc: no table is left without a row")

    # 4. A heading that cannot be kept apart declines the whole rebuild, and the document is not touched.
    document = Document()
    add = document.add_paragraph
    toc_paragraph(add, "Introduction", 1)
    toc_paragraph(add, "Scope of work", 2)
    picture = add("List of figures")
    picture.add_run()._r.append(convert.OxmlElement("w:drawing"))
    toc_paragraph(add, "Figure 1 Network map", 5)
    toc_paragraph(add, "Figure 2 Open ports", 6)
    before = [convert.paragraph_text(p) for p in convert.body_paragraphs(document)]
    entries = [toc_entry("Introduction", 1), toc_entry("Scope of work", 2),
               toc_entry("Network map", 5, label="Figure 1"), toc_entry("Open ports", 6, label="Figure 2")]
    declined = collections.Counter()
    written, warnings = convert.rebuild_toc(document, entries, 72.0, 523.0, declined)
    after = [convert.paragraph_text(p) for p in convert.body_paragraphs(document)]
    passed &= check(written == 0 and declined["toc-uncovered"] == 1 and warnings and after == before,
                    "toc: a heading holding a picture declines the rebuild and counts it",
                    (written, dict(declined), warnings))

    # 5. The old declines are counted too.
    document = Document()
    for text, number in (("Alpha", 1), ("Beta", 2), ("Gamma", 3)):
        toc_paragraph(document.add_paragraph, text, number)
    declined = collections.Counter()
    convert.rebuild_toc(document, [toc_entry("Alpha", 1), toc_entry("Delta", 4)], 72.0, 523.0, declined)
    passed &= check(declined["toc-unmatched"] == 1, "toc: an entry that cannot be matched is counted", dict(declined))
    document = Document()
    document.add_paragraph("No dot leaders anywhere.")
    declined = collections.Counter()
    convert.rebuild_toc(document, [toc_entry("Alpha", 1)], 72.0, 523.0, declined)
    passed &= check(declined["toc-not-found"] == 1, "toc: a list that is not in the file is counted", dict(declined))

    # 6. A title that appears twice is taken in reading order, and page numbers are not kept as text.
    document = Document()
    for text, number in (("Overview", 1), ("Details", 2), ("Overview", 3)):
        toc_paragraph(document.add_paragraph, text, number)
    entries = [toc_entry("Overview", 1), toc_entry("Details", 2), toc_entry("Overview", 3)]
    declined = collections.Counter()
    written, _ = convert.rebuild_toc(document, entries, 72.0, 523.0, declined)
    texts = [convert.paragraph_text(p) for p in convert.body_paragraphs(document)]
    passed &= check(written == 3 and texts == ["Overview1", "Details2", "Overview3"] and not declined,
                    "toc: a repeated title and the page numbers leave nothing behind", (texts, dict(declined)))
    # 7. A stretch that runs over a page break keeps the break: the paragraph
    # holding the section's end is not an entry, and deleting it with the
    # stretch left a document with one page-starting section fewer than the PDF
    # has pages, which is what the per-page header map reads.
    document = Document()
    for text, number in (("Introduction", 1), ("Scope of work", 2)):
        toc_paragraph(document.add_paragraph, text, number)
    document.add_section(WD_SECTION.NEW_PAGE)
    for text, number in (("Findings", 3), ("Annex", 4)):
        toc_paragraph(document.add_paragraph, text, number)
    sections_before = len(convert.page_sections(document))
    entries = [toc_entry("Introduction", 1, page_index=2, y0=100.0), toc_entry("Scope of work", 2, page_index=2, y0=116.0),
               toc_entry("Findings", 3, page_index=3, y0=100.0), toc_entry("Annex", 4, page_index=3, y0=116.0)]
    declined = collections.Counter()
    written, warnings = convert.rebuild_toc(document, entries, 72.0, 523.0, declined)
    texts = [convert.paragraph_text(p) for p in convert.body_paragraphs(document)]
    holders = [convert.paragraph_text(p) for p in convert.body_paragraphs(document)
               if convert.section_break_of(p) is not None and p.getparent() is document.element.body]
    passed &= check(written == 4 and not declined and len(convert.page_sections(document)) == sections_before == 2,
                    "toc: a stretch over a page break keeps its section break",
                    (written, dict(declined), sections_before, len(convert.page_sections(document))))
    passed &= check(texts.index("Scope of work2") < texts.index("Findings3") and holders == [""],
                    "toc: and the break is an empty paragraph between the two pages' entries", (texts, holders))

    # 8. A contents row that also holds a picture (a typst grid of the outline
    # and a logo is one row of two cells): removing the row for its entries
    # deleted the logo, silently. The rebuild declines, counts it, and leaves
    # the file as it was.
    def row_document(second_cell):
        document = Document()
        table = document.add_table(rows=1, cols=2)
        first = table.rows[0].cells[0]
        first.paragraphs[0].add_run("Introduction " + leader + "1")
        for text, number in (("Scope of work", 2), ("Findings", 3)):
            first.add_paragraph().add_run("%s %s%d" % (text, leader, number))
        second_cell(table.rows[0].cells[1])
        return document

    def logo(cell):
        cell.paragraphs[0].add_run()._r.append(convert.OxmlElement("w:drawing"))

    entries = [toc_entry("Introduction", 1, y0=100.0), toc_entry("Scope of work", 2, y0=116.0),
               toc_entry("Findings", 3, y0=132.0)]
    document = row_document(logo)
    row = document.element.body.find(".//" + convert.qn("w:tr"))
    height = convert.OxmlElement("w:trHeight")
    height.set(convert.qn("w:hRule"), "exact")
    height.set(convert.qn("w:val"), "1758")
    row.get_or_add_trPr().append(height)
    declined = collections.Counter()
    written, warnings = convert.rebuild_toc(document, entries, 72.0, 523.0, declined)
    pictures = len(list(document.element.body.iter(convert.qn("w:drawing"))))
    passed &= check(written == 0 and pictures == 1 and declined["toc-uncovered"] == 1 and warnings
                    and "1 picture" in warnings[0],
                    "toc: a row that holds a picture beside the entries keeps it, declines, and says how many",
                    (written, pictures, dict(declined), warnings))
    passed &= check(height.get(convert.qn("w:hRule")) == "atLeast",
                    "toc: and the declined row may grow, so a wrapped dot leader does not clip its last entry")
    document = row_document(lambda cell: cell.paragraphs[0].add_run("Confidential, do not copy"))
    declined = collections.Counter()
    written, warnings = convert.rebuild_toc(document, entries, 72.0, 523.0, declined)
    passed &= check(written == 0 and declined["toc-uncovered"] == 1 and "text" in warnings[0]
                    and "Confidential, do not copy" in "".join(convert.paragraph_text(p)
                                                               for p in convert.body_paragraphs(document)),
                    "toc: and so does a row that holds other words beside them", (written, dict(declined), warnings))
    # A page number in the neighbouring cell is the entry's own: nothing to lose.
    document = row_document(lambda cell: cell.paragraphs[0].add_run("1 2 3"))
    declined = collections.Counter()
    written, warnings = convert.rebuild_toc(document, entries, 72.0, 523.0, declined)
    passed &= check(written == 3 and not declined and not warnings,
                    "toc: a neighbouring cell with only digits does not stop the rebuild",
                    (written, dict(declined), warnings))

    # 9. Which stretch an entry belongs to is read through the page map. Page 1
    # lays out pictures in columns, so the stretches are in Word sections 3 and
    # 5 while their PDF pages are 1 and 3; compared as they were, the second
    # list's entries (PDF page 3) matched the first list (section 3) as "on
    # their own page", and a title that both lists repeat went to the wrong one.
    document = Document()
    document.add_paragraph("Cover")
    document.add_section(WD_SECTION.CONTINUOUS)
    document.add_paragraph("Left picture")
    document.add_section(WD_SECTION.NEW_COLUMN)
    document.add_paragraph("Right picture")
    document.add_section(WD_SECTION.NEW_PAGE)
    for text, number in (("Alpha", 1), ("Beta", 2), ("Dup", 3)):
        toc_paragraph(document.add_paragraph, text, number)
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("A page between the two lists.")
    document.add_section(WD_SECTION.NEW_PAGE)
    for text, number in (("Dup", 4), ("Gamma", 5), ("Omega", 6)):
        toc_paragraph(document.add_paragraph, text, number)
    groups = convert.sections_by_page(document, [None] * 4)
    entries = [toc_entry("Alpha", 1, page_index=1, y0=100.0), toc_entry("Beta", 2, page_index=1, y0=116.0),
               toc_entry("Dup", 3, page_index=1, y0=132.0), toc_entry("Dup", 4, page_index=3, y0=100.0),
               toc_entry("Gamma", 5, page_index=3, y0=116.0), toc_entry("Omega", 6, page_index=3, y0=132.0)]
    declined = collections.Counter()
    written, warnings = convert.rebuild_toc(document, entries, 72.0, 523.0, declined, page_groups=groups)
    texts = [t for t in (convert.paragraph_text(p) for p in convert.body_paragraphs(document)) if t.strip()]
    passed &= check(groups == [[0, 1, 2], [3], [4], [5]] and written == 6 and not declined
                    and [t for t in texts if t[-1:].isdigit()] == ["Alpha1", "Beta2", "Dup3", "Dup4", "Gamma5", "Omega6"]
                    and texts.index("A page between the two lists.") == texts.index("Dup3") + 1,
                    "toc: a stretch is matched to its entries' page through the page map",
                    (groups, written, dict(declined), texts))

    # 9. A side cell whose text repeats an entry's title is not that entry's
    # (c4: a card reading "Findings / Scope" beside an outline with "1.1 Scope"
    # was deleted with its row, because "scope" was a substring of the
    # stretch's own letters). The stretch holds the entry, so the card's text
    # is a loss: the rebuild declines and nothing is touched.
    entries = [toc_entry("Introduction", 1, y0=100.0), toc_entry("Scope of work", 2, y0=116.0),
               toc_entry("Findings", 3, y0=132.0)]
    document = row_document(lambda cell: cell.paragraphs[0].add_run("Scope of work"))
    before = [convert.paragraph_text(p) for p in convert.body_paragraphs(document)]
    declined = collections.Counter()
    written, warnings = convert.rebuild_toc(document, entries, 72.0, 523.0, declined)
    after = [convert.paragraph_text(p) for p in convert.body_paragraphs(document)]
    passed &= check(written == 0 and declined["toc-uncovered"] == 1 and after == before and warnings
                    and "1 paragraph(s) of text" in warnings[0],
                    "toc: a side cell repeating an entry's title is kept, the rebuild declines and counts it",
                    (written, dict(declined), warnings, after))
    # The same card with words no entry has was always a decline; and a label an
    # entry owns that the stretch does not hold (the "Figure 1" cell of a
    # figure entry whose caption is in the stretch) is still the entry's.
    document = row_document(lambda cell: cell.paragraphs[0].add_run("Executive sponsor"))
    declined = collections.Counter()
    written, _ = convert.rebuild_toc(document, entries, 72.0, 523.0, declined)
    passed &= check(written == 0 and declined["toc-uncovered"] == 1,
                    "toc: a side cell of other words still declines", dict(declined))
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].paragraphs[0].add_run("Network map " + leader + "5")
    table.rows[0].cells[0].add_paragraph().add_run("Open ports " + leader + "6")
    table.rows[0].cells[0].add_paragraph().add_run("Hosts " + leader + "7")
    table.rows[0].cells[1].paragraphs[0].add_run("Figure 1")
    table.rows[0].cells[1].add_paragraph().add_run("Figure 2")
    table.rows[0].cells[1].add_paragraph().add_run("Figure 3")
    entries = [toc_entry("Network map", 5, y0=100.0, label="Figure 1"),
               toc_entry("Open ports", 6, y0=116.0, label="Figure 2"),
               toc_entry("Hosts", 7, y0=132.0, label="Figure 3")]
    declined = collections.Counter()
    written, warnings = convert.rebuild_toc(document, entries, 72.0, 523.0, declined)
    texts = [convert.paragraph_text(p) for p in convert.body_paragraphs(document)]
    passed &= check(written == 3 and not declined and texts[:3] == ["Figure 1Network map5", "Figure 2Open ports6", "Figure 3Hosts7"],
                    "toc: labels in the next cell that the stretch lacks are the entries' own",
                    (written, dict(declined), texts))
    # What the stretch could not find is reported by toc_accounting, in squashed form.
    unfound = []
    convert.toc_accounting("11scope2", [toc_entry("Scope", 2, label="1.1")], unfound)
    passed &= check(unfound == [], "toc: nothing is unfound when the whole title is in the stretch", unfound)
    unfound = []
    convert.toc_accounting("networkmap5", [toc_entry("Network map", 5, label="Figure 1")], unfound)
    passed &= check(unfound == ["figure1"], "toc: a label the stretch lacks is unfound", unfound)
    # A label the stretch lacks is not found inside a longer one it has: "figure1" is not taken from "figure10".
    unfound = []
    convert.toc_accounting("networkmap5figure10routes9",
                           [toc_entry("Network map", 5, label="Figure 1"), toc_entry("Routes", 9, label="Figure 10")],
                           unfound)
    passed &= check(unfound == ["figure1"], "toc: a missing label is not taken from the front of a longer one", unfound)
    # A title that ends in a digit is still found before its own page number, and when no other entry runs on from it.
    unfound = []
    used, starts = convert.toc_accounting("appendix240phase112",
                                          [toc_entry("Appendix 2", 40), toc_entry("Phase 1", 12)], unfound)
    passed &= check(unfound == [] and starts == [0, 11] and all(used),
                    "toc: a title ending in a digit is taken before its page number", (unfound, starts))
    # "Phase 1" on page 11 reads "phase111" in the glued stretch: its own page number, not "Phase 11".
    unfound = []
    used, starts = convert.toc_accounting("phase111phase1112",
                                          [toc_entry("Phase 1", 11), toc_entry("Phase 11", 12)], unfound)
    passed &= check(unfound == [] and starts == [0, 8] and all(used),
                    "toc: a title followed by its own page number is that title, not the front of a longer one",
                    (unfound, starts))
    return passed


def run_page_fill_band_checks():
    """A full-bleed band is drawn behind its text; a panel-wide box or the running footer's strip is not."""
    passed = True
    sheet = pymupdf.open()
    page = sheet.new_page(width=612, height=792)

    def box(rect, colour):
        shape = page.new_shape()
        shape.draw_rect(pymupdf.Rect(*rect))
        shape.finish(fill=colour, color=None)
        shape.commit()

    box((0, 80, 612, 165), (0.07, 0.19, 0.23))       # a chapter band, full bleed
    page.insert_text((60, 130), "Executive summary", fontsize=18, color=(1, 1, 1))
    box((54, 300, 558, 318), (0.35, 0.35, 0.35))     # a table header row, 82% of the page
    page.insert_text((60, 313), "Host Port Service", fontsize=10, color=(1, 1, 1))
    box((0, 770, 612, 792), (0.95, 0.95, 0.95))      # the running footer's strip
    page.insert_text((224, 784), "CONFIDENTIAL", fontsize=9)
    rows = convert.page_rows(page)
    footer = {convert.row_id(0, row) for row in rows if "CONFIDENTIAL" in row["text"]}
    found = convert.page_fills(sheet, footer)[0]
    tops = sorted(round(fill["y0"]) for fill in found)
    passed &= check(tops == [80], "page fills: only the full-bleed band is drawn behind its text", tops)
    sheet.close()
    return passed


def run_margin_note_checks():
    """A note set in the margin leaves the flow for a box where the PDF has it; an empty page-high table goes."""
    passed = True

    def row(text, x0, x1, y0):
        return {"text": text, "key": convert.compare_squash(text), "x0": x0, "x1": x1, "y0": y0, "y1": y0 + 9.5,
                "columns": [{"size": 10.0, "text": text, "x0": x0, "x1": x1}]}

    body_rows = [row("Body line number %d of the finding text" % index, 65.0, 433.0, 100.0 + 14.0 * index)
                 for index in range(8)]
    note_rows = [row("Comment: Please confirm the", 458.5, 562.7, 84.9),
                 row("retest date with the client", 458.5, 546.6, 93.8),
                 row("before we publish.", 458.5, 520.3, 102.8)]
    groups = convert.margin_note_groups(body_rows + note_rows, 65.0, 433.0)
    passed &= check(len(groups) == 1 and len(groups[0]) == 3,
                    "margin notes: rows outside the column, on one edge and a line apart, are one note",
                    [len(group) for group in groups])

    # Three pages, the note on the middle one: the column's edge is the median
    # over a paper size's pages, which a note on one page does not move.
    document = Document()
    document.add_paragraph("Body line number 0 of the finding text")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("Body line number 1 of the finding text")
    holder = document.add_table(rows=1, cols=1).rows[0].cells[0].paragraphs[0]
    holder.add_run("Comment: Please confirm the retest date with the client before we publish.")
    document.add_paragraph("Body line number 2 of the finding text")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("Body line number 3 of the finding text")
    moved = convert.float_margin_notes(document, [body_rows, body_rows + note_rows, body_rows], [595.3] * 3,
                                       convert.Counter())
    boxed = [node for node in document.element.body.iter(convert.qn("w:txbxContent"))]
    passed &= check(moved == 1 and len(boxed) == 1 and "Please confirm" in "".join(t.text or "" for t in boxed[0].iter(convert.qn("w:t")))
                    and document.element.body.find(convert.qn("w:tbl")) is None,
                    "margin notes: the note goes into a box and the table that held it goes", moved)

    # Text the body has too is not a margin note's: nothing outside the column.
    document = Document()
    document.add_paragraph("Comment: Please confirm the retest date with the client before we publish.")
    moved = convert.float_margin_notes(document, [body_rows], [595.3], convert.Counter())
    passed &= check(moved == 0 and not list(document.element.body.iter(convert.qn("w:txbxContent"))),
                    "margin notes: a page with nothing outside its column is left alone", moved)

    def nested(height, words=""):
        document = Document()
        row_cells = document.add_table(rows=1, cols=2).rows[0].cells
        heading = row_cells[0].paragraphs[0].add_run("5.2 A finding heading")
        heading.font.size = Pt(14)
        inner = row_cells[1].add_table(rows=1, cols=1)
        if words:
            inner.rows[0].cells[0].paragraphs[0].add_run(words)
        properties = inner.rows[0]._tr.get_or_add_trPr()
        size = convert.OxmlElement("w:trHeight")
        size.set(convert.qn("w:val"), str(int(height * 20)))
        size.set(convert.qn("w:hRule"), "exact")
        properties.append(size)
        return document

    document = nested(810.8)
    passed &= check(convert.drop_tall_empty_tables(document) == 1,
                    "tall tables: an empty page-high table in a cell is removed")
    # A box with a drawn border is somebody's field (a signature box), not a fill's trace.
    document = nested(810.8)
    inner = next(document.element.body.iter(convert.qn("w:tbl"))).find(".//" + convert.qn("w:tc") + "//" + convert.qn("w:tbl"))
    inner_cell = inner.find(".//" + convert.qn("w:tc"))
    borders = convert.OxmlElement("w:tcBorders")
    edge = convert.OxmlElement("w:top")
    edge.set(convert.qn("w:val"), "single")
    borders.append(edge)
    inner_cell.get_or_add_tcPr().append(borders)
    passed &= check(convert.drop_tall_empty_tables(document) == 0,
                    "tall tables: an empty box with a drawn border is kept")

    # A note with a fill of its own (a review comment's box) takes its cell's fill into the text box.
    document = Document()
    document.add_paragraph("Body line number 0 of the finding text")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("Body line number 1 of the finding text")
    badge_cell = document.add_table(rows=1, cols=1).rows[0].cells[0]
    badge_cell.paragraphs[0].add_run("Comment: Please confirm the retest date with the client before we publish.")
    shading = convert.OxmlElement("w:shd")
    shading.set(convert.qn("w:fill"), "B0001F")
    badge_cell._tc.get_or_add_tcPr().append(shading)
    document.add_paragraph("Body line number 2 of the finding text")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("Body line number 3 of the finding text")
    declined = convert.Counter()
    moved = convert.float_margin_notes(document, [body_rows, body_rows + note_rows, body_rows], [595.3] * 3, declined)
    colours = [node.get("val") for node in document.element.body.iter(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")]
    passed &= check(moved == 1 and colours == ["B0001F"] and not any(True for _ in document.element.body.iter(convert.qn("w:tbl"))),
                    "margin notes: a filled note is moved with its cell's fill as the box's", (moved, colours, dict(declined)))
    document = nested(810.8, "Sidebar text")
    passed &= check(convert.drop_tall_empty_tables(document) == 0,
                    "tall tables: one with text is kept")
    document = nested(40.0)
    passed &= check(convert.drop_tall_empty_tables(document) == 0,
                    "tall tables: a short empty box is kept")
    return passed


def run_cell_condense_checks():
    """A table cell's broken paragraph is condensed line by line from the PDF's widths."""
    passed = True
    first = "Lorem ipsum dolor sit amet consectetur adipiscing elit sed do"
    second = "eiusmod tempor incididunt ut labore et dolore magna aliqua ut"
    last = "enim ad minim veniam."

    def cell_paragraph(lines_text):
        document = Document()
        cell = document.add_table(rows=1, cols=1).rows[0].cells[0]
        # As wide as the PDF's full lines, so they are the ones that would wrap.
        cell.width = Pt(300)
        paragraph = cell.paragraphs[0]
        for index, line in enumerate(lines_text):
            paragraph.add_run(line)
            if index < len(lines_text) - 1:
                paragraph.add_run().add_break()
        return document, paragraph._p

    def spacings(element):
        values = []
        for run in element.findall(convert.qn("w:r")):
            node = run.find(convert.qn("w:rPr") + "/" + convert.qn("w:spacing"))
            if node is not None and run.find(convert.qn("w:t")) is not None:
                values.append(-int(node.get(convert.qn("w:val"))))
        return values

    def pdf_lines(widths):
        return [{"text": line, "pieces": [line], "widths": [width], "page": 1}
                for line, width in zip((first, second, last), widths)]

    document, element = cell_paragraph([first, second, last])
    done = convert.condense_cell_lines(document, pdf_lines([300.0, 300.2, 110.0]), [], page_count=1)
    justified = spacings(element)
    passed &= check(done == 1 and len(justified) == 2 and all(value > convert.CONDENSE_RAGGED_TWIPS for value in justified),
                    "cell lines: a justified cell's full lines get their spaces' shrink back, the short last line nothing",
                    justified)

    document, element = cell_paragraph([first, second, last])
    convert.condense_cell_lines(document, pdf_lines([300.0, 260.0, 110.0]), [], page_count=1)
    ragged = spacings(element)
    passed &= check(len(ragged) == 1 and ragged[0] < justified[0],
                    "cell lines: in a ragged cell only the line that would wrap is condensed, by Word's excess",
                    (ragged, justified))

    document, element = cell_paragraph([first, second, last])
    other = pdf_lines([300.0, 300.2, 110.0])
    other[1]["pieces"] = ["something else entirely"]
    done = convert.condense_cell_lines(document, other, [], page_count=1)
    passed &= check(done == 0 and spacings(element) == [],
                    "cell lines: a paragraph with a line the PDF does not have is left alone", spacings(element))

    # One line is measured too: a checklist label that just fits the PDF's cell wraps in Word without it.
    label = "Credentials shared over an encrypted channel"

    def one_line(width_pt):
        document = Document()
        cell = document.add_table(rows=1, cols=1).rows[0].cells[0]
        cell.width = Pt(300)
        paragraph = cell.paragraphs[0]
        mark = paragraph.add_run("x")
        fit = convert.OxmlElement("w:fitText")
        fit.set(convert.qn("w:val"), "140")
        mark._r.get_or_add_rPr().append(fit)
        paragraph.add_run(" " + label)
        lines = [{"text": "x " + label, "pieces": ["x " + label], "widths": [width_pt], "page": 1}]
        return document, paragraph._p, lines

    document, element, lines = one_line(299.0)
    done = convert.condense_cell_lines(document, lines, [], page_count=1)
    passed &= check(done == 1 and len(spacings(element)) == 2,
                    "cell lines: a one-line cell paragraph that would wrap is condensed", (done, spacings(element)))
    document, element, lines = one_line(250.0)
    passed &= check(convert.condense_cell_lines(document, lines, [], page_count=1) == 0 and spacings(element) == [],
                    "cell lines: a one-line cell paragraph with room to spare is left alone", spacings(element))

    # A checklist label shares its PDF row with the item in the next column: it is two of the row's four pieces.
    document, element, lines = one_line(299.0)
    lines[0]["pieces"] = ["x", "Production data left in place", "x", label]
    lines[0]["widths"] = [5.0, 140.0, 5.0, 287.0]
    lines[0]["ends"] = [60.0, 210.0, 230.0, 529.0]
    lines[0]["text"] = " ".join(lines[0]["pieces"])
    done = convert.condense_cell_lines(document, lines, [], page_count=1)
    passed &= check(done == 1 and len(spacings(element)) == 2,
                    "cell lines: a label matched to two neighbouring pieces of a shared row is condensed",
                    (done, spacings(element)))

    # Pieces a table's gap apart are no run: a risk matrix's "1" and "2", 115 pt apart, are not the cell "12".
    document = Document()
    cell = document.add_table(rows=1, cols=1).rows[0].cells[0]
    cell.width = Pt(14.0)
    cell.paragraphs[0].add_run("12")
    matrix = [{"text": "1 2", "pieces": ["1", "2"], "widths": [5.0, 5.0], "ends": [105.0, 225.0],
               "y0": 100.0, "y1": 110.0, "page": 1}]
    passed &= check(convert.condense_cell_lines(document, matrix, [], page_count=1) == 0
                    and spacings(cell.paragraphs[0]._p) == [],
                    "cell lines: pieces a table's gap apart are not joined into a key")
    # A line far short of its cell is a wrong match or a mismeasured cell: only the usual condensation.
    document, element, lines = one_line(299.0)
    document.tables[0].rows[0].cells[0].width = Pt(150.0)
    convert.condense_cell_lines(document, lines, [], page_count=1)
    plain = convert.condensation_for(299.0, 0, len(convert.compare_squash("x " + label)), False, True)
    passed &= check(spacings(element) and max(spacings(element)) == plain,
                    "cell lines: a large shortfall adds nothing to the usual condensation", (spacings(element), plain))

    # A step's first line, which pdf2docx cuts from the rest of its paragraph (the rest goes after the table), is
    # matched to the PDF's whole row ("1." and its text are two pieces), and is a full justified line when the PDF
    # row under it ends where it does.
    step = "Primum est in quo admirer cur in gravissimis rebus non delectet"
    rest = "eos sermo patrius cum idem fabellas Latinas ad verbum e"

    def step_cell(rest_end):
        document = Document()
        cell = document.add_table(rows=1, cols=1).rows[0].cells[0]
        cell.width = Pt(317.5)
        first = cell.paragraphs[0]
        first.add_run("REPRODUCTION")
        first.add_run().add_break()
        first.add_run("1.")
        first.add_run(" " + step)
        document.add_paragraph(rest)
        lines = [{"text": "REPRODUCTION", "pieces": ["REPRODUCTION"], "widths": [89.1], "ends": [324.4],
                  "y0": 374.9, "y1": 384.6, "page": 1},
                 {"text": "1. " + step, "pieces": ["1.", step], "widths": [7.9, 288.4], "ends": [247.4, 542.1],
                  "y0": 386.7, "y1": 398.7, "page": 1},
                 {"text": rest, "pieces": [rest], "widths": [rest_end - 253.7], "ends": [rest_end],
                  "y0": 400.2, "y1": 412.2, "page": 1}]
        return document, first._p, lines

    document, first, lines = step_cell(542.1)
    convert.condense_cell_lines(document, lines, [], page_count=1)
    full = spacings(first)
    document, first, lines = step_cell(530.0)
    convert.condense_cell_lines(document, lines, [], page_count=1)
    ragged_step = spacings(first)
    passed &= check(len(full) == 2 and len(ragged_step) == 2 and full[0] > ragged_step[0] > 0,
                    "cell lines: a step's first line is matched whole, and justified when the PDF row under it ends there",
                    (full, ragged_step))
    # A right-aligned column ends every row at one edge: the row under starts elsewhere, so the line is not justified
    # and, fitting as ragged text, is left alone.
    document = Document()
    cell = document.add_table(rows=1, cols=1).rows[0].cells[0]
    cell.width = Pt(63.3 + 10.8)
    ports = cell.paragraphs[0]
    ports.add_run("22 80 443 8080")
    lines = [{"text": "22 80 443 8080", "pieces": ["22 80 443 8080"], "widths": [58.65], "ends": [500.0],
              "y0": 100.0, "y1": 110.0, "page": 1},
             {"text": "445 3389", "pieces": ["445 3389"], "widths": [36.0], "ends": [500.0],
              "y0": 112.0, "y1": 122.0, "page": 1}]
    passed &= check(convert.condense_cell_lines(document, lines, [], page_count=1) == 0 and spacings(ports._p) == [],
                    "cell lines: a right-aligned column's line that fits is not taken for a justified one",
                    spacings(ports._p))
    return passed


def run_list_hanging_checks():
    """A list item's text and its wrapped lines stand on the PDF's text edge, not under the marker."""
    passed = True
    qn = convert.qn

    def row(pieces, starts, ends, top, height=13.0):
        return {"text": " ".join(pieces), "pieces": pieces, "ends": ends,
                "widths": [end - start for start, end in zip(starts, ends)],
                "y0": top, "y1": top + height, "page": 1}

    def shape(paragraph):
        parts = []
        for node in (child for run in paragraph._p.iter(qn("w:r")) for child in run):
            if node.tag == qn("w:rPr"):
                continue
            if node.tag == qn("w:t"):
                parts.append(node.text or "")
            elif node.tag == qn("w:tab"):
                parts.append("<tab>")
            else:
                parts.append("|")
        return "".join(parts)

    def stops(paragraph):
        tabs = paragraph._p.find(qn("w:pPr") + "/" + qn("w:tabs"))
        return [round(convert.twips_of(tab.get(qn("w:pos"))), 1) for tab in tabs] if tabs is not None else []

    def item_paragraph(lines_text):
        document = Document()
        paragraph = document.add_paragraph()
        for index, line in enumerate(lines_text):
            paragraph.add_run(line)
            if index < len(lines_text) - 1:
                paragraph.add_run().add_break()
        return document, paragraph

    # Two items pdf2docx joined into one paragraph, the second with a wrapped line.
    document, paragraph = item_paragraph(["2. Voluptate velit esse cillum", "3. Nulla pariatur Lorem ipsum", "tempor."])
    rows = [row(["2.", "Voluptate velit esse cillum"], [54.2, 68.3], [62.7, 300.0], 100.0),
            row(["3.", "Nulla pariatur Lorem ipsum"], [54.0, 68.3], [62.5, 290.0], 114.0),
            row(["tempor."], [68.3], [100.0], 128.0)]
    changed = convert.restore_list_hanging(document, rows, page_count=1)
    passed &= check(changed == 1 and shape(paragraph) == "2.<tab>Voluptate velit esse cillum|3.<tab>Nulla pariatur Lorem ipsum|<tab>tempor.",
                    "hanging: each marker is followed by a tab and the wrapped line starts with one", shape(paragraph))
    passed &= check(stops(paragraph) == [14.1], "hanging: one left stop where the PDF starts the items' text", stops(paragraph))

    # A badge's number is a row of its own, lower than its text; the item is matched all the same.
    document, paragraph = item_paragraph(["1 Ut eius magnitudinem celeritas", "accedit ut neque divinum"])
    rows = [row(["Ut eius magnitudinem celeritas"], [89.8], [435.3], 418.2),
            row(["1"], [70.0], [74.6], 423.4, 10.9),
            row(["accedit ut neque divinum"], [89.8], [431.1], 431.8)]
    changed = convert.restore_list_hanging(document, rows, page_count=1)
    passed &= check(changed == 1 and shape(paragraph) == "1<tab>Ut eius magnitudinem celeritas|<tab>accedit ut neque divinum"
                    and stops(paragraph) == [19.8],
                    "hanging: a badge's number on a row of its own still makes an item", (shape(paragraph), stops(paragraph)))

    # A line that starts on neither edge says the layout is something else: the paragraph is left alone.
    document, paragraph = item_paragraph(["2. Voluptate velit esse cillum", "tempor."])
    rows = [row(["2.", "Voluptate velit esse cillum"], [54.2, 68.3], [62.7, 300.0], 100.0),
            row(["tempor."], [90.0], [120.0], 114.0)]
    declined = convert.Counter()
    changed = convert.restore_list_hanging(document, rows, page_count=1, declined=declined)
    passed &= check(changed == 0 and declined["hanging-declined"] == 1 and "<tab>" not in shape(paragraph),
                    "hanging: a line on another edge declines the paragraph", (changed, dict(declined)))

    # Prose that happens to start with a number is no list: the PDF has no marker piece.
    document, paragraph = item_paragraph(["2 hosts answered on port 445 and", "were signing nothing."])
    rows = [row(["2 hosts answered on port 445 and"], [54.0], [300.0], 100.0),
            row(["were signing nothing."], [54.0], [160.0], 114.0)]
    passed &= check(convert.restore_list_hanging(document, rows, page_count=1) == 0 and "<tab>" not in shape(paragraph),
                    "hanging: a number set as prose is not a marker")

    # A wide gap is the next column of a table, not a list's text edge.
    document, paragraph = item_paragraph(["1 Exposed Jenkins console"])
    rows = [row(["1", "Exposed Jenkins console"], [54.0], [58.0], 100.0)]
    rows[0]["ends"], rows[0]["widths"] = [58.0, 300.0], [4.0, 160.0]
    passed &= check(convert.restore_list_hanging(document, rows, page_count=1) == 0,
                    "hanging: a gap wider than two lines is a table's next column")
    return passed


def run_tab_stop_checks():
    """A line of a block pdf2docx tabs into place reaches its own stop, whatever the float says."""
    import random

    from pdf2docx.common import constants as pdf2docx_constants
    from pdf2docx.text.Lines import Lines

    passed = True
    threshold = 5.0

    class Line:
        def __init__(self, bbox, row):
            self.bbox, self.row, self.tab_stop = bbox, row, 0

        def in_same_row(self, other):
            return self.row == other.row

    class Block:
        def __init__(self, left, horizontal=True):
            self.bbox, self.is_horizontal_text, self.tab_stops = (left, 0.0, 500.0, 100.0), horizontal, None

    class Group:
        def __init__(self, block, lines):
            self.parent, self._instances = block, lines

    wrapped = convert.lines_reach_their_tab_stop(Lines.parse_tab_stop, pdf2docx_constants.MINOR_DIST)

    def tabs_of(step, left, spans):
        block = Block(left)
        group = Group(block, [Line((x0, 0.0, x1, 10.0), row) for x0, x1, row in spans])
        step(group, threshold)
        return [line.tab_stop for line in group._instances], block.tab_stops

    # a2's outline: 81.3 - 65.5 is 15.799999999999997, short of its own stop at 15.8.
    outline = [(65.5, 120.0, 0), (81.3, 140.0, 1), (105.2, 180.0, 2), (137.3, 200.0, 3)]
    raw_tabs, raw_stops = tabs_of(Lines.parse_tab_stop, 65.5, outline)
    passed &= check(raw_tabs[1] == 0 and raw_stops == [15.8, 39.7, 71.8],
                    "tab stops: raw pdf2docx still misses its own stop (if not, the wrapper can go)", (raw_tabs, raw_stops))
    tabs, stops = tabs_of(wrapped, 65.5, outline)
    passed &= check(tabs == [0, 1, 2, 3] and stops == raw_stops,
                    "tab stops: each level reaches its own stop, and the stops are pdf2docx's", (tabs, stops))

    # Two lines in one row: the second counts its tabs from the first's right end, as pdf2docx does.
    # Whole-number edges, where the float cannot differ, so the two methods must agree.
    row_case = [(64.0, 120.0, 0), (150.0, 200.0, 0), (94.0, 140.0, 1)]
    passed &= check(tabs_of(wrapped, 64.0, row_case) == tabs_of(Lines.parse_tab_stop, 64.0, row_case) == ([0, 1, 1], [30.0, 86.0]),
                    "tab stops: a second line in a row counts from the first's right end",
                    (tabs_of(wrapped, 64.0, row_case), tabs_of(Lines.parse_tab_stop, 64.0, row_case)))

    # a1's letter markers: "a." at 88.3 and "b." at 88.1 are two stops 0.2 pt apart in pdf2docx, and Word draws
    # the tab across that step 3.9 pt wide, so they are one stop. Raw pdf2docx gives "b." too few tabs.
    letters = [(65.2, 115.0, 0), (69.4, 130.0, 1), (88.3, 140.0, 2), (88.1, 140.0, 3), (109.8, 150.0, 4)]
    raw_tabs, raw_stops = tabs_of(Lines.parse_tab_stop, 65.2, letters)
    tabs, stops = tabs_of(wrapped, 65.2, letters)
    passed &= check(raw_stops == [4.2, 22.9, 23.1, 44.6] and raw_tabs[3] == 1,
                    "tab stops: raw pdf2docx has two stops 0.2 pt apart and gives the lower line one tab too few",
                    (raw_tabs, raw_stops))
    passed &= check(stops == [4.2, 22.9, 44.6] and tabs == [0, 0, 2, 2, 3],
                    "tab stops: stops closer than a point are one, so both lines of a level take the same tabs", (tabs, stops))

    # Where no two stops are within a point of each other and no offset sits on a float's edge, the wrapper is
    # pdf2docx's step, text direction included (whole multiples of two points, so the float cannot differ).
    chance = random.Random(16)
    differ = 0
    for trial in range(300):
        horizontal = trial % 2 == 0
        spans = []
        for row in range(chance.randint(1, 6)):
            for _ in range(chance.randint(1, 3)):
                a, b = chance.randint(0, 200) * 2.0, chance.randint(0, 200) * 2.0
                spans.append((a, b, row))
        lines = [Line((a, 1.0 * row, b, 20.0 + chance.randint(0, 200) * 2.0), row) for a, b, row in spans]
        left = chance.randint(0, 100) * 2.0
        ours, theirs = [], []
        for step, sink in ((wrapped, ours), (Lines.parse_tab_stop, theirs)):
            for line in lines:
                line.tab_stop = 0
            block = Block(left, horizontal)
            step(Group(block, lines), threshold)
            sink.append(([line.tab_stop for line in lines], block.tab_stops))
        differ += ours != theirs
    passed &= check(differ == 0, "tab stops: on stops a point or more apart the wrapper is pdf2docx's step, either direction",
                    "%d of 300 differ" % differ)

    # A block it cannot read goes to pdf2docx's own method and is counted.
    class Bare:
        pass

    calls, fell = [], convert.Counter()

    def theirs(self, line_threshold):
        calls.append(line_threshold)
        return "theirs"

    answer = convert.lines_reach_their_tab_stop(theirs, 1.0, fell)(Group(Bare(), []), threshold)
    passed &= check(calls == [threshold] and answer == "theirs" and fell["tabs-fallthrough"] == 1,
                    "tab stops: a parent without a box falls through to pdf2docx's method, counted", (calls, dict(fell)))
    return passed


def run_list_break_checks():
    """A list item the PDF starts on its own row gets its own line, even in a paragraph pdf2docx laid out by tabs."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from pdf2docx.text.Lines import Lines

    passed = True
    qn = convert.qn

    def item_row(marker, text, x0, top, height=10.5):
        mark_end = x0 + 5.0 * len(marker)
        text_start = mark_end + 7.3
        return {"text": marker + " " + text, "pieces": [marker, text], "ends": [mark_end, text_start + 5.0 * len(text)],
                "widths": [5.0 * len(marker), 5.0 * len(text)], "y0": top, "y1": top + height, "page": 1}

    def plain_row(text, x0, top, height=10.5):
        return {"text": text, "pieces": [text], "ends": [x0 + 5.0 * len(text)], "widths": [5.0 * len(text)],
                "y0": top, "y1": top + height, "page": 1}

    def build(text, tabs=0, justified=False, font=None, cell=False):
        document = Document()
        host = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0] if cell else document.add_paragraph()
        for _ in range(tabs):
            host.add_run().add_tab()
        run = host.add_run(text)
        if font:
            run.font.name = font
        if justified:
            host.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return document, host

    def shape(paragraph):
        parts = []
        for node in (child for run in paragraph._p.iter(qn("w:r")) for child in run):
            if node.tag == qn("w:t"):
                parts.append(node.text or "")
            elif node.tag == qn("w:tab"):
                parts.append("<tab>")
            elif node.tag == qn("w:br"):
                parts.append("|")
        return "".join(parts)

    deep = "2.2.1.1.1 Targeted UDP checks on discovered hosts 3 Threat modelling and vulnerability analysis"
    deep_rows = [item_row("2.2.1.1.1", "Targeted UDP checks on discovered hosts", 177.5, 100.0),
                 item_row("3", "Threat modelling and vulnerability analysis", 65.5, 114.0)]

    # a2's outline: the deepest item is the widest row, so pdf2docx never broke after it.
    document, paragraph = build(deep, tabs=4)
    before = convert.paragraph_text(paragraph._p)
    added = convert.break_before_list_items(document, deep_rows, page_count=1)
    passed &= check(added == 1 and shape(paragraph) == "<tab>" * 4 + "2.2.1.1.1 Targeted UDP checks on discovered hosts |"
                    "3 Threat modelling and vulnerability analysis" and convert.paragraph_text(paragraph._p) == before,
                    "list breaks: the next outline item starts its own line, the four tabs still first, text unchanged",
                    shape(paragraph))

    # The next row's marker is an outline number.
    document, paragraph = build("2.1 Open-source research on the organisation 2.2 Active enumeration", tabs=1)
    rows = [item_row("2.1", "Open-source research on the organisation", 81.3, 100.0),
            item_row("2.2", "Active enumeration", 81.3, 114.0)]
    passed &= check(convert.break_before_list_items(document, rows, page_count=1) == 1
                    and shape(paragraph) == "<tab>2.1 Open-source research on the organisation |2.2 Active enumeration",
                    "list breaks: a multi-level number opens an item", shape(paragraph))

    # A row with no marker is a wrap, which Word is meant to redo.
    document, paragraph = build("2.2.1.1.1 Targeted UDP checks on discovered hosts amet consectetur adipiscing", tabs=4)
    rows = [deep_rows[0], plain_row("amet consectetur adipiscing", 184.8, 114.0)]
    passed &= check(convert.break_before_list_items(document, rows, page_count=1) == 0 and "|" not in shape(paragraph),
                    "list breaks: a row with no marker is a wrapped line and is left", shape(paragraph))

    # Prose that opens with a number in one piece is no marker.
    document, paragraph = build("2.2.1.1.1 Targeted UDP checks on discovered hosts 12 hosts answered the sweep", tabs=4)
    rows = [deep_rows[0], plain_row("12 hosts answered the sweep", 65.5, 114.0)]
    passed &= check(convert.break_before_list_items(document, rows, page_count=1) == 0 and "|" not in shape(paragraph),
                    "list breaks: a number set as prose is not a marker", shape(paragraph))

    # Two rows side by side share a top: columns, not a list.
    document, paragraph = build(deep, tabs=4)
    rows = [deep_rows[0], dict(deep_rows[1], y0=100.0, y1=110.5)]
    passed &= check(convert.break_before_list_items(document, rows, page_count=1) == 0 and "|" not in shape(paragraph),
                    "list breaks: rows side by side are not consecutive lines", shape(paragraph))

    # The same two rows twice on the page: two stretches could make the line.
    document, paragraph = build(deep, tabs=4)
    again = [dict(row, y0=row["y0"] + 300.0, y1=row["y1"] + 300.0) for row in deep_rows]
    passed &= check(convert.break_before_list_items(document, deep_rows + again, page_count=1) == 0
                    and "|" not in shape(paragraph),
                    "list breaks: a line two stretches of rows could make is left alone", shape(paragraph))

    # Three items on one line.
    document, paragraph = build("2.1 Open research 2.2 Active enumeration 2.3 Reporting phase", tabs=1)
    rows = [item_row("2.1", "Open research", 81.3, 100.0), item_row("2.2", "Active enumeration", 81.3, 114.0),
            item_row("2.3", "Reporting phase", 81.3, 128.0)]
    passed &= check(convert.break_before_list_items(document, rows, page_count=1) == 2
                    and shape(paragraph) == "<tab>2.1 Open research |2.2 Active enumeration |2.3 Reporting phase",
                    "list breaks: every item the PDF starts gets its break", shape(paragraph))

    # A justified paragraph: Word stretches a line that ends in a manual break, so the new break has a tab
    # before it and the paragraph a right tab stop to reach.
    document, paragraph = build(deep, tabs=4, justified=True)
    convert.break_before_list_items(document, deep_rows, page_count=1)
    stops = [(stop.get(qn("w:val")), stop.get(qn("w:pos"))) for stop in paragraph._p.iter(qn("w:tab"))
             if stop.getparent().tag == qn("w:tabs")]
    passed &= check(shape(paragraph).endswith("hosts <tab>|3 Threat modelling and vulnerability analysis")
                    and any(val == "right" for val, _ in stops),
                    "list breaks: a justified paragraph's new break has a tab before it and a right stop", (shape(paragraph), stops))

    # A boxed score opening a row ("9.8" with a 5 pt inset is a piece of its own) after a row typst stretched to
    # the column's edge is a wrap: no break, and pdf2docx's tab before it stays, since Word wraps at it.
    def edge_row(text, top, end=500.0, pieces=None, ends=None, widths=None):
        return {"text": text, "pieces": pieces or [text], "ends": ends or [end], "widths": widths or [end - 62.4],
                "y0": top, "y1": top + 10.5, "page": 1}

    wrap = "The path traversal and the manager both scored high while the relay scored"
    rest = "and the anonymous bind scored low on the same scale"
    rows = [edge_row(wrap, 100.0),
            edge_row("8.1 " + rest, 114.0, pieces=["8.1", rest], ends=[80.0, 500.0], widths=[12.6, 415.0]),
            edge_row("Every other host answered on the usual ports and nothing more", 300.0)]
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run(wrap + " ")
    paragraph.add_run().add_tab()
    paragraph.add_run("8.1 " + rest)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    passed &= check(convert.break_before_list_items(document, rows, page_count=1) == 0
                    and shape(paragraph) == wrap + " <tab>8.1 " + rest,
                    "list breaks: a number opening the row after a stretched one is a wrap", shape(paragraph))

    # The same page's justified rows do not hold back an outline whose row ends short of their edge.
    prose = [edge_row("Lorem ipsum dolor sit amet consectetur adipiscing elit sed do", 300.0 + 14.0 * k, end=520.0)
             for k in range(4)]
    document, paragraph = build(deep, tabs=4)
    passed &= check(convert.break_before_list_items(document, deep_rows + prose, page_count=1) == 1
                    and shape(paragraph).endswith("hosts |3 Threat modelling and vulnerability analysis"),
                    "list breaks: justified prose on the page leaves the outline's break alone", shape(paragraph))

    # Code and a table cell's paragraph are not this pass's.
    document, paragraph = build(deep, tabs=4, font="DejaVu Sans Mono")
    passed &= check(convert.break_before_list_items(document, deep_rows, page_count=1) == 0 and "|" not in shape(paragraph),
                    "list breaks: a code paragraph is left alone", shape(paragraph))
    document, paragraph = build(deep, tabs=4, cell=True)
    passed &= check(convert.break_before_list_items(document, deep_rows, page_count=1) == 0 and "|" not in shape(paragraph),
                    "list breaks: a paragraph in a table cell is left alone", shape(paragraph))

    # End to end: a typst outline five levels deep whose deepest item is the widest row, as in a2.
    source = """
#set page(paper: "a4", margin: (x: 2.2cm, top: 3.2cm, bottom: 3.0cm))
#set text(font: "New Computer Modern", size: 10.5pt, lang: "en")
#set par(justify: false, leading: 0.68em, spacing: 1em)
#set enum(full: true, numbering: "1.1.1", indent: 0.3em, body-indent: 0.7em, spacing: 0.65em)
The engagement followed the phases below, each reported in its own chapter.

+ Pre-engagement interactions
  + Scope, goals and legal agreements
  + Rules of engagement and emergency contacts
+ Intelligence gathering
  + Open-source research on the organisation
    + Public repositories and paste sites
    + Certificate transparency logs
  + Active enumeration
    + Network and service discovery
      + Full TCP sweep of every range
        + Targeted UDP checks on discovered hosts
+ Threat modelling and vulnerability analysis
+ Exploitation and post exploitation
+ Reporting

Each request was repeated three times to rule out a transient response from the cache.
"""
    expected = ["1 Pre-engagement interactions", "1.1 Scope, goals and legal agreements",
                "1.2 Rules of engagement and emergency contacts", "2 Intelligence gathering",
                "2.1 Open-source research on the organisation", "2.1.1 Public repositories and paste sites",
                "2.1.2 Certificate transparency logs", "2.2 Active enumeration",
                "2.2.1 Network and service discovery", "2.2.1.1 Full TCP sweep of every range",
                "2.2.1.1.1 Targeted UDP checks on discovered hosts", "3 Threat modelling and vulnerability analysis",
                "4 Exploitation and post exploitation", "5 Reporting"]

    def outline_lines(path):
        for element in convert.body_paragraphs(Document(path)):
            lines = convert.line_texts(element)
            if lines and any("targetedudpchecks" in convert.compare_squash(line) for line in lines):
                return element, lines
        return None, []

    workdir = tempfile.mkdtemp(prefix="btct-list-break-test-")
    try:
        installed = Lines.parse_tab_stop
        result, path = run_case(source, workdir, "outline")
        passed &= check(Lines.parse_tab_stop is installed, "list breaks: the tab-stop step is put back after a conversion")
        element, lines = outline_lines(path)
        shown = [line.strip("\t ") for line in lines]
        passed &= check(shown == expected, "list breaks: every item of the outline starts a line of its paragraph", shown)
        wanted = [number.split(" ", 1)[0].count(".") for number in expected]
        got = [len(line) - len(line.lstrip("\t")) for line in lines]
        passed &= check(got == wanted, "list breaks: each level starts with one tab per level above it", got)
        passed &= check("tabs-fallthrough" not in result["declined"], "tab stops: the wrapper applied", result["declined"])
        raw_path = os.path.join(workdir, "outline-raw.docx")
        convert.convert(os.path.join(workdir, "outline.pdf"), raw_path, raw=True)
        raw_element, raw_lines = outline_lines(raw_path)
        level_two = [line for line in raw_lines if line.lstrip("\t").startswith("1.1")]
        joined = [line for line in raw_lines if "Targeted" in line and "Threat" in line]
        passed &= check(bool(level_two) and "\t" not in level_two[0] and bool(joined),
                        "list breaks: raw pdf2docx still gives a level-two line no tab and runs item 3 on (if not, the passes can go)",
                        raw_lines[:3] + joined)
        passed &= check(element is not None and raw_element is not None
                        and convert.compare_squash(convert.paragraph_text(element))
                        == convert.compare_squash(convert.paragraph_text(raw_element)),
                        "list breaks: the outline's text is the same as raw pdf2docx's")
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return passed


def run_header_bar_checks():
    """A running header's filled bar is drawn behind it when it is a bar: wide, behind the text, with no rule."""
    passed = True

    def band(span, strip_y=(55.4, 74.4), rule=None, edge="top"):
        return {"edge": edge, "strip_behind": (0.48, 0.12, 0.64), "strip_span": span, "strip_y": strip_y,
                "strip_full_width": False, "rule_under": rule, "text_top": 59.0, "text_bottom": 70.9}

    passed &= check(convert.header_bar(band((62.4, 532.9)), 62.4, 532.9) == (55.4, 74.4),
                    "header bar: a fill the text area's width behind the header's text is its bar")
    passed &= check(convert.header_bar(band((69.3, 82.3)), 64.8, 576.0) is None,
                    "header bar: a logo's box is no bar")
    passed &= check(convert.header_bar(band((62.4, 532.9), rule={"x0": 64.8, "x1": 576.0}), 62.4, 532.9) is None,
                    "header bar: a header with a rule under it keeps the rule")
    passed &= check(convert.header_bar(band((62.4, 532.9), strip_y=(30.0, 45.0)), 62.4, 532.9) is None,
                    "header bar: a fill above the text is not behind it")
    passed &= check(convert.header_bar(band((62.4, 532.9), edge="bottom"), 62.4, 532.9) is None,
                    "header bar: a footer's strip is the footer's business")

    # A footer whose text stands on two baselines is two lines, top first, the page number on the nearer.
    def column(text, base, size=7.5, kind="text"):
        return {"x0": 0.0, "x1": 10.0, "base": base, "size": size,
                "parts": [{"kind": kind, "text": text if kind == "text" else None}]}

    stacked = [column("INTERNAL USE ONLY", 800.7), column("CONFIDENTIAL", 790.3), column(None, 796.1, 9.0, "page")]
    lines = convert.band_lines(stacked)
    shown = [sorted(part.get("text") or part["kind"] for c in members for part in c["parts"]) for _, members in lines]
    passed &= check(shown == [["CONFIDENTIAL"], ["INTERNAL USE ONLY", "page"]] and abs(convert.band_pitch(lines) - 10.4) < 0.05,
                    "band lines: a stacked footer is its lines top to bottom, at their pitch", (shown, convert.band_pitch(lines)))
    flat = [column("Northwind Logistics CONFIDENTIAL", 796.0), column(None, 796.4, 9.0, "page")]
    passed &= check(len(convert.band_lines(flat)) == 1, "band lines: one baseline is one line")
    passed &= check(convert.stacks_apart(lines), "band lines: a footer stacked at its leading stands apart")
    # A title centred beside a two-line block: three baselines 5.5 pt apart in 9 pt type.
    beside = convert.band_lines([column("Client Name", 40.0, 9.0), column("Penetration Test Report", 45.5, 9.0),
                                 column("October 2026", 51.0, 9.0)])
    passed &= check(len(beside) == 3 and not convert.stacks_apart(beside),
                    "band lines: baselines closer than their type are not a stack", beside)
    # "Page 3" between the two labels: the stack's lines come from its words, so the page text joins the nearer.
    slot = lambda c: c.get("slot", "centre")
    page_text = dict(column("Page ", 752.4, 9.0), slot="right",
                     parts=[{"kind": "text", "text": "Page "}, {"kind": "page"}])
    lines = convert.band_lines([column("CONFIDENTIAL", 746.8), column("INTERNAL USE ONLY", 756.9), page_text], slot)
    passed &= check(len(lines) == 2 and abs(convert.band_pitch(lines) - 10.1) < 0.05 and convert.stacks_apart(lines),
                    "band lines: 'Page 3' beside a stack stands on the nearer line", lines)
    top, height = convert.stack_box((lines, convert.band_pitch(lines)))
    passed &= check(abs(top - (746.8 - convert.EXACT_BASELINE_RATIO * 10.1)) < 0.05 and abs(height - 20.2) < 0.05,
                    "band lines: a stack's paragraph starts where its first baseline lands on the PDF's", (top, height))
    two = [column("CONFIDENTIAL", 790.0), dict(column(None, 795.0, 7.5, "page"), slot="left"),
           dict(column(None, 800.0, 7.5, "page"), slot="right"), column("INTERNAL USE ONLY", 805.0)]
    passed &= check(sum(len(members) for _, members in convert.band_lines(two, slot)) == 4,
                    "band lines: every column is on exactly one line")
    uneven = convert.band_lines([column("A", 30.0, 8.0), column("B", 40.8, 8.0), column("C", 56.0, 8.0)])
    passed &= check(abs(convert.band_pitch(uneven) - 10.8) < 0.05, "band lines: the pitch is the smallest gap",
                    convert.band_pitch(uneven))

    def written(text, base, x0, x1, kind="text", size=8.0):
        return dict(column(text, base, size, kind), x0=float(x0), x1=float(x1), font="Arial", psname=None,
                    bold=False, color=0, page_offset=0)

    def footer(columns, strip=None):
        bases = [c["base"] for c in columns]
        return {"edge": "bottom", "columns": columns, "rule_under": None, "strip_behind": strip and (0.9, 0.9, 0.9),
                "strip_y": strip, "strip_full_width": bool(strip), "strip_span": strip and (0.0, 612.0),
                "text_top": min(bases) - 7.0, "text_bottom": max(bases) + 2.0}

    # A strip wider than the text: the indents are the strip's and the first line's indent carries the text
    # back, which Word gives the first line only, so a line after a break tabs to where the first line starts.
    document = Document()
    paragraph = document.add_paragraph()
    convert.fill_band_paragraph(paragraph, footer([written("Acme Security Ltd", 750.0, 72, 150),
                                                   written("Client use only", 761.0, 72, 190),
                                                   written(None, 755.5, 530, 540, "page", 9.0)], (738.0, 770.0)),
                                72.0, 72.0, 72.0, 540.0, 792.0, 612.0)
    nodes = [node.tag.split("}")[1] for run in paragraph._p.iter(convert.qn("w:r")) for node in run
             if node.tag in (convert.qn("w:t"), convert.qn("w:tab"), convert.qn("w:br"))]
    fmt = paragraph.paragraph_format
    start = fmt.left_indent.twips + fmt.first_line_indent.twips
    stops = [(stop.get(convert.qn("w:val")), int(stop.get(convert.qn("w:pos"))))
             for stop in paragraph._p.iter(convert.qn("w:tab")) if stop.getparent().tag == convert.qn("w:tabs")]
    passed &= check("br" in nodes and nodes[nodes.index("br") + 1] == "tab" and ("left", start) in stops,
                    "band lines: a stacked line on a strip starts where the first line does", (nodes, stops, start))
    # A stack tighter than its type stays one line, its labels in reading order with a space between.
    document = Document()
    paragraph = document.add_paragraph()
    convert.fill_band_paragraph(paragraph, footer([written("INTERNAL USE ONLY", 797.6, 240, 350),
                                                   written("CONFIDENTIAL", 790.0, 250, 340),
                                                   written(None, 793.8, 530, 540, "page", 9.0)]),
                                72.0, 72.0, 72.0, 540.0, 792.0, 612.0)
    passed &= check("CONFIDENTIAL INTERNAL USE ONLY" in paragraph.text and "\n" not in paragraph.text,
                    "band lines: one line reads a slot's columns top to bottom, apart", paragraph.text)
    return passed

def run_split_dash_checks():
    """A dash pdf2docx moved to the start of the next paragraph goes back to the line the PDF ends with it."""
    passed = True
    dash = chr(0x2013)
    above = "Torquem illum hosti detraxisse ut aliquam ex eo est consecutus? "
    below = "Laudem et caritatem quae sunt vitae sine metu degendae praesidia "

    def split(rows):
        document = Document()
        first = document.add_paragraph(above)
        second = document.add_paragraph()
        second.add_run(dash + " ")
        second.add_run(below)
        lines = [{"text": row, "page": 1} for row in rows]
        return document, first, second, lines

    document, first, second, lines = split([above + dash + " ", below.strip()])
    moved = convert.return_split_dashes(document, lines, page_count=1)
    passed &= check(moved == 1 and first.text == above + dash + " " and second.text == below,
                    "split dash: a dash the PDF ends a line with goes back to that line",
                    (moved, first.text[-20:], second.text[:20]))
    # A list item or an attribution starts its PDF row with the dash, and keeps it.
    document, first, second, lines = split([above.strip(), dash + " " + below.strip()])
    moved = convert.return_split_dashes(document, lines, page_count=1)
    passed &= check(moved == 0 and second.text.startswith(dash),
                    "split dash: a row that starts with the dash keeps it", (moved, second.text[:20]))
    # The same dash opening a line after a break inside one paragraph (the pitch-cut pass splits it later).
    document = Document()
    joined = document.add_paragraph()
    joined.add_run(above)
    joined.add_run().add_break()
    joined.add_run(dash + " ")
    joined.add_run(below)
    moved = convert.return_split_dashes(document, [{"text": above + dash + " ", "page": 1},
                                                   {"text": below.strip(), "page": 1}], page_count=1)
    shown = "".join((node.text or "") if node.tag == convert.qn("w:t") else "|"
                    for node in joined._p.iter(convert.qn("w:t"), convert.qn("w:br")))
    passed &= check(moved == 1 and shown == above + dash + " |" + below,
                    "split dash: a dash that opens a line after a break goes back before the break", shown[-40:])
    # An em dash the PDF sets without a space before it goes back without one.
    em = chr(0x2014)
    document = Document()
    first = document.add_paragraph("The attacker joined the guest network over the wireless ")
    second = document.add_paragraph()
    second.add_run(em + " ")
    second.add_run(below)
    lines = [{"text": "The attacker joined the guest network over the wireless" + em, "page": 1},
             {"text": below.strip(), "page": 1}]
    moved = convert.return_split_dashes(document, lines, page_count=1)
    passed &= check(moved == 1 and first.text == "The attacker joined the guest network over the wireless" + em + " ",
                    "split dash: an unspaced dash goes back unspaced", (moved, first.text[-12:]))
    # A page that also starts a row with the dash and this text cannot say which one this paragraph is.
    document, first, second, lines = split([above + dash + " ", below.strip(), dash + " " + below.strip()])
    passed &= check(convert.return_split_dashes(document, lines, page_count=1) == 0,
                    "split dash: a row that opens with the dash and the same text declines the move")
    # Nothing in the PDF to say where the dash was: left alone.
    document, first, second, lines = split(["Something else entirely on this page."])
    passed &= check(convert.return_split_dashes(document, lines, page_count=1) == 0,
                    "split dash: with no row ending on the dash nothing moves")
    return passed


def run_side_picture_checks():
    """Two screenshots side by side keep one row, and a row holding a picture is measured."""
    passed = True
    workdir = tempfile.mkdtemp(prefix="side-pictures-")
    try:
        png = os.path.join(workdir, "shot.png")
        pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 20), False).save(png)

        def figure(top_b):
            document = Document()
            first = document.add_paragraph()
            first.add_run().add_picture(png, width=Pt(100), height=Pt(50))
            second = document.add_paragraph()
            second.add_run("(a) Left caption")
            second.add_run().add_tab()
            second.add_run().add_picture(png, width=Pt(100), height=Pt(60))
            pictures = [{"page": 0, "rect": pymupdf.Rect(60, 100, 160, 150), "visible": pymupdf.Rect(60, 100, 160, 150)},
                        {"page": 0, "rect": pymupdf.Rect(170, top_b, 270, top_b + 60),
                         "visible": pymupdf.Rect(170, top_b, 270, top_b + 60)}]
            return document, first._p, second._p, pictures

        def drawings(element):
            return len(list(element.iter(convert.qn("w:drawing"))))

        document, first, second, pictures = figure(100.0)
        moved = convert.rejoin_side_pictures(document, pictures, convert.Counter())
        stops = [convert.twips_of(tab.get(convert.qn("w:pos"))) for tab in first.iter(convert.qn("w:tab"))
                 if tab.get(convert.qn("w:pos"))]
        passed &= check(moved == 1 and drawings(first) == 2 and drawings(second) == 0
                        and convert.paragraph_text(second).strip() == "(a) Left caption",
                        "side pictures: the taller picture leaves the caption's line for its partner's",
                        (moved, drawings(first), drawings(second), convert.paragraph_text(second)))
        passed &= check(stops and abs(stops[0] - 110.0) < 0.1,
                        "side pictures: a tab stop where the second picture starts, from the first one's start", stops)

        document, first, second, pictures = figure(130.0)
        moved = convert.rejoin_side_pictures(document, pictures, convert.Counter())
        passed &= check(moved == 0 and drawings(first) == 1 and drawings(second) == 1,
                        "side pictures: pictures whose tops do not line up are left as written", moved)

        # A paragraph whose line height is automatic is measured by its
        # tallest thing, so an exact row that cannot hold a picture grows.
        document = Document()
        cell = document.add_table(rows=1, cols=1).rows[0].cells[0]
        cell.paragraphs[0].add_run().add_picture(png, width=Pt(100), height=Pt(60))
        metrics = convert.automatic_paragraph_metrics(cell.paragraphs[0]._p)
        passed &= check(metrics is not None and metrics["line"] >= 60.0,
                        "row height: a picture's paragraph is as tall as the picture", metrics)
        row = cell._tc.getparent()
        properties = row.get_or_add_trPr()
        height = convert.OxmlElement("w:trHeight")
        height.set(convert.qn("w:val"), "1000")
        height.set(convert.qn("w:hRule"), "exact")
        properties.append(height)
        convert.relax_overfull_rows(document)
        passed &= check(height.get(convert.qn("w:hRule")) == "atLeast",
                        "row height: an exact row too short for its picture is let grow", height.get(convert.qn("w:hRule")))
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return passed


def run_orphan_char_checks():
    """A table's cells take every character of a line that straddles their row edges."""
    from pdf2docx import Converter
    from pdf2docx.page.RawPageFactory import RawPageFactory
    from pdf2docx.table.Cell import Cell
    from pdf2docx.table.TableBlock import TableBlock

    passed = True
    workdir = tempfile.mkdtemp(prefix="orphans-")
    try:
        sheet = pymupdf.open()
        page = sheet.new_page(width=300, height=300)
        page.insert_text((60, 150), "infinitum ligature", fontsize=10)
        path = os.path.join(workdir, "line.pdf")
        sheet.save(path)
        sheet.close()
        converter = Converter(path)
        try:
            settings = converter.default_settings
            raw = RawPageFactory.create(page_engine=converter.fitz_doc[0], backend="PyMuPDF")
            raw.restore(**settings)
            raw.clean_up(**settings)
            source = list(raw.blocks)[0]
        finally:
            converter.close()

        class Table:
            """What `TableBlock.assign_blocks` reads: rows of cells."""
            def __init__(self, edges):
                self._rows = [[Cell({"bbox": (50, top, 200, bottom)})] for top, bottom in zip(edges, edges[1:])]

            def cell_text(self):
                return ["".join(block.text for block in row[0].blocks) for row in self._rows]

        def handed_over(edges, line, declined=None, wrapped=True):
            table = Table(edges)
            step = TableBlock.assign_blocks
            if wrapped:
                step = convert.straddling_lines_keep_their_chars(step, declined)
            step(table, [line])
            return table.cell_text()

        # The line is 13.8 pt high and the edges cut it into 6.3, 4.5 and 3.0:
        # no cell holds half of any character, so pdf2docx's own hand-over
        # drops the line without a trace.
        edges = [100, 145.5, 150, 160]
        text = source.text
        lost = handed_over(edges, source.copy(), wrapped=False)
        passed &= check("".join(lost) == "", "orphan chars: pdf2docx drops a line no cell holds half of (if not, the wrapper can go)", lost)
        kept = handed_over(edges, source.copy())
        passed &= check(kept[0] == text and "".join(kept[1:]) == "",
                        "orphan chars: the whole line goes to the cell it overlaps most", kept)

        # A zero-width character (the `i` of a `fi` ligature) is lost wherever
        # a line is cut at all, whatever the cell holds.
        cut = [100, 144, 160]
        def with_ligature():
            line = source.copy()
            chars = [char for span in line.spans for char in span.chars]
            x0, y0, _, y1 = chars[4].bbox
            chars[4].update_bbox((x0, y0, x0, y1))
            return line
        lost = handed_over(cut, with_ligature(), wrapped=False)
        passed &= check(lost[1] == text[:4] + text[5:] and lost[0] == "",
                        "orphan chars: pdf2docx drops a zero-width character of a cut line (if not, the wrapper can go)", lost)
        kept = handed_over(cut, with_ligature())
        passed &= check(kept[1] == text and kept[0] == "",
                        "orphan chars: a zero-width character comes back after the one before it", kept)

        # `ffi`: two zero-width characters in a row keep their order.
        def with_triple():
            line = source.copy()
            chars = [char for span in line.spans for char in span.chars]
            for char in chars[4:6]:
                x0, y0, _, y1 = char.bbox
                char.update_bbox((x0, y0, x0, y1))
            return line
        kept = handed_over(cut, with_triple())
        passed &= check(kept[1] == text and kept[0] == "",
                        "orphan chars: two zero-width characters in a row keep their order", kept)

        # A zero-width character at the start of a line goes before the next one.
        def with_leading():
            line = source.copy()
            chars = [char for span in line.spans for char in span.chars]
            x0, y0, _, y1 = chars[0].bbox
            chars[0].update_bbox((x0, y0, x0, y1))
            return line
        kept = handed_over(cut, with_leading())
        passed &= check(kept[1] == text, "orphan chars: a leading zero-width character stays first", kept)

        # Nothing is added where nothing was lost: a line one cell holds is
        # handed over as before, once.
        whole = handed_over([100, 140, 160], source.copy())
        passed &= check(whole[0] == "" and whole[1] == text, "orphan chars: a line one cell holds is not duplicated", whole)
        both = handed_over([100, 146, 160], source.copy())
        passed &= check("".join(both).count("infinitum") == 1,
                        "orphan chars: a line cut across two cells keeps each character once", both)

        # A cell that holds a span's whole box gets a deep copy of it, so a
        # neighbour is found by its character and box, not by identity.
        chars = [char for span in source.copy().spans for char in span.chars]
        twin = chars[3].copy()
        passed &= check(convert.same_char(twin, chars[3]) and not convert.same_char(chars[4], chars[3]),
                        "orphan chars: a copied character is the same character, a neighbour is not")

        # One line read across two cells, the second cell's characters cut by
        # sliver rows: they used to follow the first cell's last character into
        # its column. They are left out now, for the text check to name.
        def across_columns():
            line = source.copy()
            chars = [char for span in line.spans for char in span.chars]
            for char in chars[10:]:
                x0, y0, x1, y1 = char.bbox
                char.update_bbox((x0 + 80.0, y0, x1 + 80.0, y1))
            # The span's and the line's boxes follow their characters, or the
            # left cell holds most of the stale line box and pdf2docx hands it
            # the whole line before the wrapper sees anything.
            for span in line.spans:
                box = pymupdf.Rect(span.chars[0].bbox)
                for char in span.chars[1:]:
                    box |= char.bbox
                span.update_bbox(box)
            line.update_bbox(pymupdf.Rect(line.spans[0].bbox))
            return line

        class Columns:
            """A left cell over the whole height beside a column cut into sliver rows."""
            def __init__(self):
                self._rows = [[Cell({"bbox": (50, 100, 125, 160)}), Cell({"bbox": (130, 100, 260, 145.5)})],
                              [Cell(), Cell({"bbox": (130, 145.5, 260, 150)})],
                              [Cell(), Cell({"bbox": (130, 150, 260, 160)})]]

            def cell_text(self):
                return ["".join(block.text for block in cell.blocks) for row in self._rows for cell in row if cell]

        declined = convert.Counter()
        table = Columns()
        convert.straddling_lines_keep_their_chars(TableBlock.assign_blocks, declined)(table, [across_columns()])
        cells = table.cell_text()
        passed &= check("ligature" not in cells[0] and cells[0].strip() == text[:10].strip()
                        and declined["orphan-chars-other-column"] > 0,
                        "orphan chars: a character in another cell's column does not join this one", (cells, dict(declined)))

        # A table that no longer looks like one is left to pdf2docx and counted.
        declined = convert.Counter()

        class Bare:
            pass

        convert.straddling_lines_keep_their_chars(lambda self, blocks: None, declined)(Bare(), [source.copy()])
        passed &= check(declined["orphan-chars-fallthrough"] == 1,
                        "orphan chars: a table with no rows is counted as a fall-through, not an error", dict(declined))
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return passed


def run_split_cell_checks():
    """A line that a table edge cuts through ends up in one cell, merged with the empty one beside it."""
    from pdf2docx import Converter
    from pdf2docx.page.RawPageFactory import RawPageFactory
    from pdf2docx.table.TableBlock import TableBlock

    passed = True
    workdir = tempfile.mkdtemp(prefix="splitcells-")
    try:
        sheet = pymupdf.open()
        page = sheet.new_page(width=300, height=300)
        page.insert_text((60, 150), "T1190", fontsize=10)
        path = os.path.join(workdir, "line.pdf")
        sheet.save(path)
        sheet.close()
        converter = Converter(path)
        try:
            settings = converter.default_settings
            raw = RawPageFactory.create(page_engine=converter.fitz_doc[0], backend="PyMuPDF")
            raw.restore(**settings)
            raw.clean_up(**settings)
            source = list(raw.blocks)[0]
        finally:
            converter.close()

        def line(top, height=10.6):
            """The line, moved so that it spans `top` to `top + height`."""
            moved = source.copy()
            moved.update_bbox((50, top, 120, top + height))
            return moved

        def table(spec, width=100.0):
            """One column, cells top to bottom: (top, bottom, fill, border widths, line top or None)."""
            rows = []
            for top, bottom, fill, border, at in spec:
                rows.append({"bbox": (50, top, 50 + width, bottom), "height": bottom - top, "cells": [{
                    "bbox": (50, top, 50 + width, bottom), "bg_color": fill, "border_width": border,
                    "border_color": (0, 0, 0, 0), "merged_cells": (1, 1)}]})
            grid = TableBlock({"bbox": (50, spec[0][0], 50 + width, spec[-1][1]), "rows": rows})
            for row, (_, _, _, _, at) in zip(grid._rows, spec):
                if at is not None:
                    row[0].blocks.append(line(at))
            return grid

        plain = (0, 0, 0, 0)
        # A card's value line, 10.6 pt high, in a 5.3 pt cell with a 4.7 pt empty one under it.
        down = table([(0, 40, 7, plain, 2), (40, 45.3, 7, (0, 0, 0.5, 0), 41), (45.3, 50, 7, (0.5, 0, 0, 0), None),
                      (50, 70, 7, plain, 52)])
        convert.lines_keep_one_cell(down)
        top, spare = down._rows[1][0], down._rows[2][0]
        passed &= check(top.merged_cells == (2, 1) and abs(top.bbox.y0 - 40) < 0.1 and abs(top.bbox.y1 - 50) < 0.1,
                        "split cells: the cell and the empty one under it become one", (top.merged_cells, tuple(top.bbox)))
        passed &= check(not spare and len(top.blocks) == 1, "split cells: the spare cell is a placeholder, the line stays", len(top.blocks))
        passed &= check(tuple(top.border_width) == (0, 0, 0, 0),
                        "split cells: the rule between the two (the stripes' rule) is gone", top.border_width)
        passed &= check(down._rows[0][0].merged_cells == (1, 1) and down._rows[3][0].merged_cells == (1, 1),
                        "split cells: cells whose line fits are left as they are")

        # A tight ruled cell: the line box overflows by 1.5 pt (its descender) across a drawn rule into the
        # empty cell under it. That is no cut, and the rule between the two rows stays.
        tight = table([(0, 40, 7, plain, 2), (40, 50.6, 7, (0, 0, 0.5, 0), 41.5), (50.6, 70, 7, (0.5, 0, 0, 0), None)])
        convert.lines_keep_one_cell(tight)
        passed &= check(tight._rows[1][0].merged_cells == (1, 1) and bool(tight._rows[2][0]),
                        "split cells: a line box that only overflows a tight ruled cell is not merged across the rule",
                        tight._rows[1][0].merged_cells)

        # The same line reaching up into an empty cell: the upper cell takes it over.
        up = table([(0, 10, 7, plain, None), (10, 15.3, 7, plain, 3), (15.3, 40, 7, plain, 20)])
        convert.lines_keep_one_cell(up)
        head = up._rows[0][0]
        passed &= check(head.merged_cells == (2, 1) and len(head.blocks) == 1 and not up._rows[1][0],
                        "split cells: a line that reaches up is moved into the empty cell above", (head.merged_cells, len(head.blocks)))

        # The line may reach a little (0.8 pt) without anything being merged.
        near = table([(0, 20, 7, plain, 0), (20, 29, 7, plain, 19.2), (29, 40, 7, plain, None)])
        convert.lines_keep_one_cell(near)
        passed &= check(all(row[0].merged_cells == (1, 1) for row in near._rows) and all(row[0] for row in near._rows),
                        "split cells: a reach under the tolerance merges nothing")

        # A neighbour with a line of its own, another fill or another width is never merged.
        full = table([(0, 20, 7, plain, 1), (20, 25.3, 7, plain, 20), (25.3, 40, 7, plain, 26)])
        convert.lines_keep_one_cell(full)
        passed &= check(all(row[0].merged_cells == (1, 1) and row[0] for row in full._rows),
                        "split cells: a cell with a line of its own is not merged")
        other = table([(0, 10, 3, plain, None), (10, 15.3, 7, plain, 3)])
        convert.lines_keep_one_cell(other)
        passed &= check(all(row[0].merged_cells == (1, 1) for row in other._rows) and other._rows[0][0],
                        "split cells: a neighbour with another fill is not merged")

        # A cell of two rows already (the stripes' rowspan) is not eaten: the row under it is a placeholder.
        from pdf2docx.table.Cell import Cell
        spanned = table([(0, 10, 7, plain, None), (10, 15.3, 7, plain, None), (15.3, 25, 7, plain, 9)])
        spanned._rows[0][0].merged_cells = (2, 1)
        spanned._rows[1]._cells._instances[0] = Cell()
        convert.lines_keep_one_cell(spanned)
        passed &= check(spanned._rows[0][0].merged_cells == (2, 1) and spanned._rows[2][0].merged_cells == (1, 1),
                        "split cells: a placeholder row is never taken for an empty cell")

        # It runs with the hand-over of lines to cells, and a table that does not look like one is counted.
        declined = convert.Counter()

        class Bare:
            pass

        convert.straddling_lines_keep_their_chars(lambda self, blocks: None, declined)(Bare(), [])
        passed &= check(declined["split-cells-fallthrough"] == 1,
                        "split cells: a table with no rows is counted as a fall-through, not an error", dict(declined))
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return passed


def run_overprint_checks():
    """Text drawn on top of different text, and the shapes that only look like it."""
    passed = True
    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    # A pile: three different rows at one baseline.
    for text in ("Stale user accounts 27", "Missing patch baseline 28", "Legacy protocol enabled 30"):
        page.insert_text((LEFT, 700), text, fontsize=11)
    # Tight leading: lines 11 pt apart in 11 pt type.
    for k in range(4):
        page.insert_text((LEFT, 200 + 11 * k), "An ordinary line of prose number %d" % k, fontsize=11)
    # The same word twice in one place, which is how some writers fake bold.
    page.insert_text((LEFT, 400), "Critical", fontsize=11)
    page.insert_text((LEFT + 0.3, 400), "Critical", fontsize=11)
    lines = convert.body_lines(doc, [])
    stacked = [line["text"] for line in lines if line["overprinted"]]
    passed &= check(len(stacked) == 1 and "Stale user accounts 27" in stacked[0],
                    "overprint: rows set on one another are flagged", stacked)
    passed &= check(all(line["page"] == 1 for line in lines), "overprint: every line says its page")
    passed &= check(not any(line["overprinted"] for line in lines if "ordinary line" in line["text"]),
                    "overprint: tightly leaded prose is not")
    passed &= check(not any(line["overprinted"] for line in lines if line["text"].startswith("Critical")),
                    "overprint: a word drawn twice for bold is not")
    doc.close()
    return passed


WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
DML = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def tilted_pdf(path):
    """Two pages: a stamp at 4 degrees, a ribbon at 45, a quarter turn each way, and a faint watermark.

    Page 1 draws a 9% watermark first (the way a page background is laid
    under everything) and its stamp last. Page 2 has the ribbon and the two
    quarter turns, of which only the one that does not read bottom to top is
    dropped by pdf2docx.
    """
    doc = pymupdf.open()
    first = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    mark = pymupdf.Point(150, 560)
    first.insert_text(mark, "CONFIDENTIAL", fontsize=60, fontname="hebo", color=(0.48, 0.12, 0.24),
                      fill_opacity=0.09, morph=(mark, pymupdf.Matrix(35)))
    for k in range(6):
        first.insert_text((LEFT, 120 + 14 * k), "An ordinary line of the report, number %d, set upright." % k,
                          fontsize=11)
    stamp = pymupdf.Point(120, 764)
    first.insert_text(stamp, "F O R A U T H O R I S E D R E A D E R S O N L Y", fontsize=14, fontname="hebo",
                      color=(0.48, 0.12, 0.24), morph=(stamp, pymupdf.Matrix(4)))
    second = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    second.insert_text((LEFT, 120), "A second page of ordinary text, set upright.", fontsize=11)
    ribbon = pymupdf.Point(500, 60)
    second.insert_text(ribbon, "CLIENT COPY", fontsize=9, fontname="hebo", color=(0.1, 0.1, 0.1),
                       morph=(ribbon, pymupdf.Matrix(-45)))
    second.insert_text((300, 500), "Reads bottom to top", fontsize=11, rotate=90)
    second.insert_text((330, 300), "Reads top to bottom", fontsize=11, rotate=270)
    doc.save(path)
    doc.close()


def anchored_boxes(docx_path):
    """Every anchored text box of a Word file: its text, turn in degrees, centre and z-order flag."""
    with zipfile.ZipFile(docx_path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    boxes = []
    for anchor in root.iter(WP + "anchor"):
        content = next(anchor.iter(W + "txbxContent"), None)
        if content is None:
            continue
        offset_x = int(anchor.find(WP + "positionH").find(WP + "posOffset").text)
        offset_y = int(anchor.find(WP + "positionV").find(WP + "posOffset").text)
        extent = anchor.find(WP + "extent")
        width, height = int(extent.get("cx")), int(extent.get("cy"))
        transform = next(anchor.iter(DML + "xfrm"))
        boxes.append({"text": "".join(node.text or "" for node in content.iter(W + "t")),
                      "turn": int(transform.get("rot")) / 60000,
                      "cx": (offset_x + width / 2) / 12700, "cy": (offset_y + height / 2) / 12700,
                      "behind": anchor.get("behindDoc") == "1",
                      "color": next((c.get(W + "val") for c in content.iter(W + "color")), None)})
    return boxes


def run_cover_checks():
    """A cover: the first line's offset kept, a cell's labels broken where the PDF broke them, plain fills drawn."""
    from docx.enum.text import WD_LINE_SPACING

    passed = True
    qn = convert.qn

    def exact_paragraph(document, text, before):
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(11.4)
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.space_after = Pt(0)
        return paragraph

    # The offset of a first line that pdf2docx wrote as spacing before (margin 0 under a full-bleed band).
    document = Document()
    section = document.sections[0]
    section.top_margin = Pt(0)
    kicker = exact_paragraph(document, "PENETRATION TEST REPORT", 82.7)
    exact_paragraph(document, "Halcyon Freight Group", 40.0)
    moved = convert.fold_leading_space_into_margin(document)
    passed &= check(moved == 1 and abs(section.top_margin.pt - 82.7) < 0.1
                    and kicker.paragraph_format.space_before.pt == 0,
                    "cover: a first line's offset becomes the page's top margin",
                    (moved, section.top_margin.pt, kicker.paragraph_format.space_before))

    # An ordinary page: the first line is within a line of the margin and is left as the walk models it.
    document = Document()
    document.sections[0].top_margin = Pt(54)
    near = exact_paragraph(document, "A first line under an ordinary margin", 6.0)
    moved = convert.fold_leading_space_into_margin(document)
    passed &= check(moved == 0 and document.sections[0].top_margin.pt == 54 and near.paragraph_format.space_before.pt == 6.0,
                    "cover: a first line within one line of the margin is not moved", moved)

    # Word starts the body below a header whatever the margin says, so the two would no longer agree.
    document = Document()
    document.sections[0].top_margin = Pt(0)
    document.sections[0].header.is_linked_to_previous = False
    under_header = exact_paragraph(document, "A first line under a running header", 60.0)
    moved = convert.fold_leading_space_into_margin(document)
    passed &= check(moved == 0 and under_header.paragraph_format.space_before.pt == 60.0,
                    "cover: a page with a header keeps its first line's spacing", moved)

    # Labels pdf2docx broke after two of four lines: the PDF's other line ends are added, the right indent dropped.
    pages_doc = Document()
    grid = pages_doc.add_table(rows=1, cols=2)
    labels = grid.rows[0].cells[0].paragraphs[0]
    labels.add_run("Prepared by ")
    labels.add_run().add_break()
    labels.add_run("Engagement window ")
    labels.add_run("Report version ")
    labels.add_run().add_break()
    labels.add_run("Classification")
    labels.paragraph_format.right_indent = Pt(36)
    grid.rows[0].cells[1].paragraphs[0].add_run("Corvid Security Group")
    rows = [("Prepared by", "Corvid Security Group"), ("Engagement window", "2026-03-02 to 2026-03-13"),
            ("Report version", "1.0 (final)"), ("Classification", "Confidential, TLP:AMBER")]
    cell_lines = [{"text": "%s %s" % pair, "page": 1, "pieces": list(pair)} for pair in rows]
    done = convert.break_cell_lines(pages_doc, cell_lines, set(), page_count=1)
    breaks = sum(1 for _ in labels._p.iter(qn("w:br")))
    passed &= check(done == 1 and breaks == 3 and labels.paragraph_format.right_indent == 0,
                    "cover: a cell paragraph with some of its breaks gets the rest at the PDF's line ends",
                    (done, breaks, labels.paragraph_format.right_indent))

    # A break the PDF's lines do not have (mid line) leaves the paragraph alone.
    odd_doc = Document()
    odd_grid = odd_doc.add_table(rows=1, cols=2)
    odd = odd_grid.rows[0].cells[0].paragraphs[0]
    odd.add_run("Prepared by Engagement ")
    odd.add_run().add_break()
    odd.add_run("window Report version Classification")
    odd.paragraph_format.right_indent = Pt(36)
    done = convert.break_cell_lines(odd_doc, cell_lines, set(), page_count=1)
    passed &= check(done == 0 and sum(1 for _ in odd._p.iter(qn("w:br"))) == 1 and odd.paragraph_format.right_indent == Pt(36),
                    "cover: a break the PDF's lines do not have is not finished around", done)

    # The plain fills of a page: a band and a stripe come back, a rule, a swatch, a fill under text and white do not.
    workdir = tempfile.mkdtemp(prefix="btct-cover-test-")
    try:
        path = os.path.join(workdir, "cover.pdf")
        doc = pymupdf.open()
        page = doc.new_page(width=595.0, height=842.0)
        page.draw_rect(pymupdf.Rect(0, 0, 56.7, 842), color=None, fill=(0.478, 0.121, 0.239))
        page.draw_rect(pymupdf.Rect(56.7, 0, 65.2, 842), color=None, fill=(0.788, 0.545, 0.169))
        page.draw_rect(pymupdf.Rect(100, 300, 500, 301), color=None, fill=(0.2, 0.2, 0.2))
        page.draw_rect(pymupdf.Rect(100, 400, 130, 430), color=None, fill=(0.2, 0.6, 0.2))
        page.draw_rect(pymupdf.Rect(0, 700, 595, 842), color=None, fill=(1, 1, 1))
        # A fill under text that stops short of the page's edges (a panel); a
        # full-bleed band is drawn behind its text (`run_page_fill_band_checks`).
        page.draw_rect(pymupdf.Rect(40, 500, 555, 640), color=None, fill=(0.9, 0.9, 0.9))
        page.insert_text((200, 570), "Text over a fill", fontsize=12)
        page.draw_rect(pymupdf.Rect(150, 100, 400, 200), color=(0, 0, 0), fill=(0.9, 0.2, 0.2))
        fills = convert.page_fills(doc)
        doc.close()
        found = [(round(f["x0"]), round(f["x1"]), f["rgb"]) for f in fills[0]]
        passed &= check(found == [(0, 57, "7A1F3D"), (57, 65, "C98B2B")],
                        "cover: only the band and the stripe are plain fills", found)
        document = Document()
        document.add_paragraph("Kicker")
        declined = collections.Counter()
        written = convert.restore_page_fills(document, fills, declined)
        anchors = list(document.element.body.iter(WP + "anchor"))
        offsets = [int(a.find(WP + "positionH").find(WP + "posOffset").text) for a in anchors]
        looks = [a.get("behindDoc") for a in anchors]
        colours = [next(a.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")).get("val")
                   for a in anchors]
        passed &= check(written == 2 and not declined and looks == ["1", "1"] and offsets == [0, round(56.7 * 12700)]
                        and colours == ["7A1F3D", "C98B2B"],
                        "cover: each is written behind the text at the PDF's place, in drawing order",
                        (written, dict(declined), looks, offsets, colours))
        declined = collections.Counter()
        passed &= check(convert.restore_page_fills(Document(), [[], fills[0]], declined) == 0
                        and declined["page-fills-no-page"] == 1,
                        "cover: a page the file does not have is declined by name", dict(declined))
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return passed


def run_rotated_checks():
    """Lines pdf2docx drops for their direction are written back as text boxes, turned and placed."""
    from pdf2docx.text.Line import Line

    passed = True
    # The rule is pdf2docx's own, so a release that changes it fails here.
    for direction in [(1.0, 0.0), (0.0, -1.0), (0.0, 1.0), (-1.0, 0.0), (0.9975640177726746, -0.06975647062063217),
                      (0.7071067690849304, 0.7071067690849304), (0.9999999, 0.0)]:
        theirs = Line({"dir": list(direction), "bbox": (0, 0, 10, 10), "spans": []}).dir
        kept = Line({"dir": list(direction), "bbox": (0, 0, 10, 10), "spans": []}).text_direction.value != -1
        passed &= check(convert.carried_by_pdf2docx(theirs) == kept,
                        "rotated: the carried test agrees with pdf2docx on %s" % (direction,), (theirs, kept))
    passed &= check(convert.blend_with_white(0x7A1F3D, 1.0) == 0x7A1F3D and convert.blend_with_white(0, 0.0) == 0xFFFFFF
                    and convert.blend_with_white(0x000000, 0.5) == 0x808080,
                    "rotated: opacity is blended onto white")

    workdir = tempfile.mkdtemp(prefix="btct-rotated-test-")
    try:
        pdf_path = os.path.join(workdir, "tilted.pdf")
        tilted_pdf(pdf_path)
        with pymupdf.open(pdf_path) as doc:
            lines = convert.rotated_lines(doc)
        texts = sorted(line["spans"][0]["text"] for line in lines)
        passed &= check(texts == ["CLIENT COPY", "CONFIDENTIAL", "F O R A U T H O R I S E D R E A D E R S O N L Y",
                                  "Reads top to bottom"],
                        "rotated: the lines pdf2docx drops are found, the upright and bottom-to-top ones are not", texts)
        by_text = {line["spans"][0]["text"]: line for line in lines}
        for text, expected in (("F O R A U T H O R I S E D R E A D E R S O N L Y", -4.0), ("CLIENT COPY", 45.0),
                               ("CONFIDENTIAL", -35.0), ("Reads top to bottom", 90.0)):
            passed &= check(abs(by_text[text]["angle"] - expected) < 0.1,
                            "rotated: %s is turned %.0f degrees" % (text[:12], expected), by_text[text]["angle"])
        for text, size, face in (("F O R A U T H O R I S E D R E A D E R S O N L Y", 14, "hebo"), ("CLIENT COPY", 9, "hebo")):
            wanted = pymupdf.get_text_length(text, fontname=face, fontsize=size)
            passed &= check(abs(by_text[text]["length"] - wanted) < 1.5,
                            "rotated: %s measures %.1f pt along its direction" % (text[:12], wanted),
                            by_text[text]["length"])
        passed &= check(by_text["CONFIDENTIAL"]["behind"] and not by_text["CLIENT COPY"]["behind"]
                        and not by_text["F O R A U T H O R I S E D R E A D E R S O N L Y"]["behind"],
                        "rotated: only what the PDF draws before its page's upright text goes behind it",
                        {t: line["behind"] for t, line in by_text.items()})
        faint = by_text["CONFIDENTIAL"]["spans"][0]["color"]
        passed &= check(faint > 0xE0E0E0, "rotated: the 9% watermark reads as a pale colour", hex(faint))

        raw = convert.convert(pdf_path, os.path.join(workdir, "raw.docx"), raw=True)
        passed &= check(raw["textCheck"]["missing"] == 4 and not anchored_boxes(os.path.join(workdir, "raw.docx")),
                        "rotated: pdf2docx's own file lacks all four lines", raw["textCheck"])
        full_path = os.path.join(workdir, "full.docx")
        full = convert.convert(pdf_path, full_path)
        passed &= check(full["textCheck"]["missing"] == 0 and not full["textCheck"]["fellBack"],
                        "rotated: the text check is clean on the first rung", (full["textCheck"], full["warnings"]))
        passed &= check(not [name for name in full["declined"] if name.startswith("rotated-")],
                        "rotated: nothing was declined", full["declined"])
        boxes = {box["text"]: box for box in anchored_boxes(full_path)}
        passed &= check(sorted(boxes) == sorted(["CLIENT COPY", "CONFIDENTIAL", "Reads top to bottom",
                                                 "F O R A U T H O R I S E D R E A D E R S O N L Y"]),
                        "rotated: one text box per line, with the PDF's text", sorted(boxes))
        for text, line in by_text.items():
            box = boxes[text.replace("\xad", "-")]
            passed &= check(abs(box["turn"] % 360 - line["angle"] % 360) < 0.01,
                            "rotated: %s is turned by the line's angle" % text[:12], (box["turn"], line["angle"]))
            passed &= check(abs(box["cx"] - line["cx"]) < 0.05 and abs(box["cy"] - line["cy"]) < 0.05,
                            "rotated: %s is turned about the line's own centre" % text[:12],
                            ((box["cx"], box["cy"]), (line["cx"], line["cy"])))
            passed &= check(box["behind"] == line["behind"], "rotated: %s is placed in the PDF's order" % text[:12])
        # The watermark is on page 1 and the ribbon on page 2: each box is
        # anchored in a paragraph of its own page, which is what puts it there.
        document = Document(full_path)
        pages = convert.page_blocks(document.element.body)
        homes = {}
        for number, blocks in enumerate(pages):
            for block_ in blocks:
                for anchor in block_.iter(WP + "anchor"):
                    homes[convert.paragraph_text(anchor).strip()] = number
        passed &= check(homes["CONFIDENTIAL"] == 0 and homes["CLIENT COPY"] == 1,
                        "rotated: each box is anchored on its own page", homes)
        # A pass that was switched off leaves the lines to the check, which names them.
        skipped = convert.convert(pdf_path, os.path.join(workdir, "skipped.docx"), skip=frozenset({"rotated"}))
        passed &= check(skipped["textCheck"]["missing"] >= 1 and skipped["textCheck"]["fellBack"],
                        "rotated: without the pass the text check falls back, as it did", skipped["textCheck"])
        passed &= check("rotated" in convert.SKIPPABLE, "rotated: the pass can be switched off by name")
        passed &= run_tail_page_checks(workdir)

        if shutil.which("typst") is None:
            print("SKIP  the header watermark case needs the typst CLI on PATH")
            return passed
        # The common way to put a diagonal watermark on every page: inside the
        # header. Without the pass the six lost lines sent the whole report to
        # the bottom rung, which also gave up lifting the header and footer.
        header_source = HEADER_WATERMARK
        result, docx_path = run_case(header_source, workdir, "header-mark")
        passed &= check(result["textCheck"]["missing"] == 0 and not result["textCheck"]["fellBack"]
                        and result["headerFooter"]["header"] and result["headerFooter"]["footer"],
                        "rotated: a watermark in the header keeps the first rung, with its header and footer lifted",
                        (result["textCheck"], result["headerFooter"], result["warnings"]))
        marks = [box for box in anchored_boxes(docx_path) if box["text"] == "DRAFT"]
        passed &= check(len(marks) == result["pages"] and all(box["behind"] for box in marks),
                        "rotated: one watermark box per page, behind the text", (len(marks), result["pages"]))
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return passed


def tail_page_pdf(path, upright_on_last):
    """Two pages: ordinary text, then one tilted line, with an upright line beside it when asked."""
    doc = pymupdf.open()
    first = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    first.insert_text((LEFT, 120), "An ordinary line of the report, set upright.", fontsize=11)
    last = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    if upright_on_last:
        last.insert_text((LEFT, 120), "An upright line on the last page.", fontsize=11)
    mark = pymupdf.Point(150, 560)
    last.insert_text(mark, "END OF REPORT", fontsize=40, fontname="hebo", color=(0.5, 0.5, 0.5),
                     morph=(mark, pymupdf.Matrix(45)))
    doc.save(path)
    doc.close()


def two_tail_page_pdf(path):
    """Three pages: ordinary text, then two pages that hold one tilted line each."""
    doc = pymupdf.open()
    first = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    first.insert_text((LEFT, 120), "An ordinary line of the report, set upright.", fontsize=11)
    for text in ("END OF REPORT", "APPENDIX BLANK"):
        page = doc.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        mark = pymupdf.Point(150, 560)
        page.insert_text(mark, text, fontsize=40, fontname="hebo", color=(0.5, 0.5, 0.5),
                         morph=(mark, pymupdf.Matrix(45)))
    doc.save(path)
    doc.close()


def run_tail_page_checks(workdir):
    """A last page that holds only tilted text gets a page of its own in the Word file."""
    passed = True

    def converted_like_pdf2docx():
        # What pdf2docx writes for it: the first page and a section break, then
        # nothing, so the body's own section properties open an empty page.
        document = Document()
        document.add_paragraph("An ordinary line of the report, set upright.")
        document.add_section(WD_SECTION.NEW_PAGE)
        return document

    def stub_of(document):
        return next(b for b in document.element.body if b.tag == convert.qn("w:p")
                    and next(b.iter(WP + "anchor"), None) is not None)

    for upright in (False, True):
        path = os.path.join(workdir, "tail-%s.pdf" % upright)
        tail_page_pdf(path, upright)
        with pymupdf.open(path) as doc:
            lines = convert.rotated_lines(doc)
            rows = convert.page_anchor_rows(doc, [])
        document = converted_like_pdf2docx()
        declined = collections.Counter()
        written = convert.restore_rotated_lines(document, lines, declined, page_rows_of=rows)
        pages = convert.page_blocks(document.element.body)
        homes = [number for number, blocks in enumerate(pages) for block_ in blocks
                 for _ in block_.iter(WP + "anchor")]
        if not upright:
            passed &= check(written == 1 and not declined and homes == [1] and len(pages) == 2,
                            "rotated: a last page of only tilted text gets its own page in the file",
                            (written, dict(declined), homes, len(pages)))
            passed &= check(stub_of(document).find(convert.qn("w:pPr")).find(convert.qn("w:pageBreakBefore")) is None,
                            "rotated: and needs no page break when the body's section already starts a page")
            # A body section that does not begin a page needs the break from the paragraph.
            document = converted_like_pdf2docx()
            document.sections[-1].start_type = WD_SECTION.CONTINUOUS
            convert.restore_rotated_lines(document, lines, collections.Counter(), page_rows_of=rows)
            passed &= check(stub_of(document).find(convert.qn("w:pPr")).find(convert.qn("w:pageBreakBefore")) is not None,
                            "rotated: a continuous body section gets the page break from the paragraph itself")
            document = converted_like_pdf2docx()
            declined = collections.Counter()
            written = convert.restore_rotated_lines(document, lines, declined)
            passed &= check(written == 0 and declined["rotated-no-page"] == 1,
                            "rotated: without the PDF's rows no page is made", (written, dict(declined)))
        else:
            passed &= check(written == 0 and declined["rotated-no-page"] == 1 and len(pages) == 1,
                            "rotated: a last page with upright text the file lost is still a decline",
                            (written, dict(declined), len(pages)))

    # Two trailing tilted pages: `pages` grows with each stub, so the test for
    # "not the first new page" compared against a length that moved with it and
    # the second stub never got its break (r1: 3 PDF pages, 2 in Word).
    path = os.path.join(workdir, "tail-two.pdf")
    two_tail_page_pdf(path)
    with pymupdf.open(path) as doc:
        lines = convert.rotated_lines(doc)
        rows = convert.page_anchor_rows(doc, [])
    document = converted_like_pdf2docx()
    declined = collections.Counter()
    written = convert.restore_rotated_lines(document, lines, declined, page_rows_of=rows)
    stubs = [b for b in document.element.body if b.tag == convert.qn("w:p") and next(b.iter(WP + "anchor"), None) is not None]
    breaks = [b.find(convert.qn("w:pPr")).find(convert.qn("w:pageBreakBefore")) is not None for b in stubs]
    passed &= check(written == 2 and not declined and len(stubs) == 2 and breaks == [False, True],
                    "rotated: two trailing tilted pages each get a page, the second with its own break",
                    (written, dict(declined), len(stubs), breaks))
    document = converted_like_pdf2docx()
    document.sections[-1].start_type = WD_SECTION.CONTINUOUS
    convert.restore_rotated_lines(document, lines, collections.Counter(), page_rows_of=rows)
    stubs = [b for b in document.element.body if b.tag == convert.qn("w:p") and next(b.iter(WP + "anchor"), None) is not None]
    breaks = [b.find(convert.qn("w:pPr")).find(convert.qn("w:pageBreakBefore")) is not None for b in stubs]
    passed &= check(breaks == [True, True], "rotated: and both break when the body's section does not start a page", breaks)
    return passed


HEADER_WATERMARK = """
#set page(paper: "a4", margin: (top: 3cm, bottom: 2.5cm, x: 2.5cm),
  header: [
    #place(top + left, dx: 4cm, dy: 6cm, rotate(-35deg, text(size: 60pt, fill: luma(200), weight: "bold")[DRAFT]))
    #text(size: 9pt)[ACME Security Assessment] #h(1fr) #text(size: 9pt)[Contoso Ltd]
    #line(length: 100%)
  ],
  footer: context [#align(center)[#text(size: 9pt)[Page #counter(page).display() of #counter(page).final().first()]]])
#set text(size: 10pt)
= Findings
#for i in range(1, 260) [
  This is an ordinary paragraph of the report, number #i. It explains an observation at some length so that the page fills with text.
]
"""


# ── the document cases ────────────────────────────────────────────────────
# These compile a Typst source and run the whole of convert() on it, which
# needs the typst CLI on PATH. They are the ones that pin behaviour rather
# than a function: every one of them is a document that lost content.


def compile_case(source, pdf_path):
    with open(pdf_path + ".typ", "w", encoding="utf-8") as handle:
        handle.write(source)
    fonts = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "dist", "fonts")
    args = ["typst", "compile", "--ignore-system-fonts"]
    if os.path.isdir(fonts):
        args += ["--font-path", fonts]
    args += [pdf_path + ".typ", pdf_path]
    subprocess.run(args, check=True, capture_output=True)


def run_case(source, workdir, name):
    """Compile, convert, and hand back what `result.json` says."""
    pdf_path = os.path.join(workdir, name + ".pdf")
    docx_path = os.path.join(workdir, name + ".docx")
    compile_case(source, pdf_path)
    result = convert.convert(pdf_path, docx_path)
    return result, docx_path


def compile_untagged(source, pdf_path):
    """Compile, then remove the MarkInfo so the PDF reads as untagged."""
    tagged_path = pdf_path + ".tagged.pdf"
    compile_case(source, tagged_path)
    with pymupdf.open(tagged_path) as doc:
        doc.xref_set_key(doc.pdf_catalog(), "MarkInfo", "null")
        doc.save(pdf_path)


def run_untagged_case(source, workdir, name):
    """`run_case` on the same PDF with its MarkInfo removed, so it reads as untagged."""
    pdf_path = os.path.join(workdir, name + ".pdf")
    docx_path = os.path.join(workdir, name + ".docx")
    compile_untagged(source, pdf_path)
    result = convert.convert(pdf_path, docx_path)
    return result, docx_path


def squashed_docx(path, body_only):
    return convert.compare_squash(convert.docx_text(path, body_only=body_only))


def section_band_parts(path, part):
    """Per Word section, the header (or footer) Word shows on its first page."""
    document = Document(path)
    shown, current, current_first = [], None, None
    for section in document.sections:
        holder = getattr(section, part)
        if not holder.is_linked_to_previous:
            current = holder
        first = getattr(section, "first_page_" + part)
        if not first.is_linked_to_previous:
            current_first = first
        shown.append(current_first if section.different_first_page_header_footer else current)
    return len(document.sections), shown


def section_band_texts(path, part):
    count, shown = section_band_parts(path, part)
    return count, [" ".join(p.text for p in holder.paragraphs).strip() if holder is not None else ""
                   for holder in shown]


def section_band_xml(path, part):
    count, shown = section_band_parts(path, part)
    return count, [holder._element.xml if holder is not None else "" for holder in shown]


def face_entry(path, family, **extra):
    """A staged face as `fonts` describes one, for the refusal checks."""
    entry = {"path": path, "family": family, "subfamily": "Regular", "postscript": "",
             "full": "", "typoFamily": "", "typoSubfamily": "", "bold": False,
             "italic": False, "fsType": 0, "damaged": False, "postscriptOutlines": False,
             "licenceText": "This Font Software is licensed under the SIL Open Font License"}
    entry.update(extra)
    return entry


def one_family_plan(family, entry):
    """A plan holding one family of one face, which is all a refusal needs."""
    return {"faces": {family: entry}, "missing": [], "families": {family: {(False, False): entry}},
            "byFamily": {family.lower(): family}, "byFace": {}}


def run_font_conversion_checks(font_tools, font_dir, workdir):
    """A PostScript-outline face is redrawn, and only where it may be.

    The redrawn face is what a reader will be shown, so the checks are on the
    bytes that go in the file rather than on the decision to make them: it
    has to parse, it has to have quadratic outlines and no PostScript ones,
    and every advance width has to be the number it was, because the line
    endings in this Word file are the PDF's and a width that moved is a line
    that wraps twice.
    """
    import io

    from fontTools.ttLib import TTFont

    passed = True
    source = next((os.path.join(font_dir, n) for n in sorted(os.listdir(font_dir))
                   if n.lower().endswith(".otf")), None)
    if source is None:
        print("SKIP  no OpenType font was staged, so the conversion is untested")
        return passed
    entry = font_tools.read_font(source)
    passed &= check(entry["postscriptOutlines"], "fonts: the staged .otf reads as PostScript outlines")
    passed &= check(font_tools.licence_of(entry) is not None,
                    "fonts: and its licence is one this recognises", font_tools.licence_of(entry))

    # A deadline already past has to stop the loop at its first check, not
    # somewhere after a face's worth of glyphs have already been drawn.
    started = time.perf_counter()
    past = font_tools.to_truetype(source, deadline=time.monotonic() - 1.0)
    spent = time.perf_counter() - started
    passed &= check(past is None, "fonts: a deadline already past yields no face", past)
    passed &= check(spent < 1.0, "fonts: and it stops well under a second", "%.3f s" % spent)

    data = font_tools.to_truetype(source)
    passed &= check(data is not None, "fonts: it converts")
    if data is None:
        return passed
    before, after = TTFont(source, fontNumber=0), TTFont(io.BytesIO(data), fontNumber=0)
    try:
        passed &= check("glyf" in after and "CFF " not in after,
                        "fonts: the converted face has quadratic outlines and no PostScript ones")
        passed &= check(after.sfntVersion == "\000\001\000\000",
                        "fonts: and says so in its header", repr(after.sfntVersion))
        moved = [g for g in before.getGlyphOrder() if before["hmtx"][g][0] != after["hmtx"][g][0]]
        passed &= check(not moved, "fonts: every advance width is the number it was",
                        (len(moved), len(before.getGlyphOrder())))
        vertical = [(t, f) for t, f in (("hhea", "ascent"), ("hhea", "descent"), ("hhea", "lineGap"),
                                        ("OS/2", "sTypoAscender"), ("OS/2", "sTypoDescender"),
                                        ("OS/2", "usWinAscent"), ("OS/2", "usWinDescent"),
                                        ("head", "unitsPerEm"))
                    if getattr(before[t], f) != getattr(after[t], f)]
        passed &= check(not vertical, "fonts: and so is every vertical metric", vertical)
        passed &= check(all(t in after for t in ("GPOS", "GSUB", "kern") if t in before),
                        "fonts: the kerning and shaping tables came across")
        passed &= check(all(before["name"].getDebugName(i) == after["name"].getDebugName(i)
                            for i in (1, 2, 6)),
                        "fonts: and it still calls itself what it did",
                        (after["name"].getDebugName(1), after["name"].getDebugName(2)))
    finally:
        before.close()
        after.close()

    # It goes in the file, in the slot its own subfamily names.
    target = os.path.join(workdir, "converted.docx")
    blank = Document()
    blank.add_paragraph("x")
    blank.save(target)
    family = entry["family"]
    embedded, refused, redrawn = font_tools.embed(target, one_family_plan(family, entry))
    passed &= check(embedded == [family] and redrawn == [family],
                    "fonts: a PostScript-outline face is carried, redrawn", (embedded, redrawn))
    with zipfile.ZipFile(target) as archive:
        table = archive.read("word/fontTable.xml").decode("utf-8")
    slot = font_tools.STYLE_ELEMENTS[(entry["bold"], entry["italic"])]
    passed &= check(('<w:font w:name="%s">' % family) in table and ("<" + slot) in table,
                    "fonts: under its own name and in its own slot", slot)

    # And only where it may be: the option off, an unreadable licence, and a
    # margin too thin to redraw anything in.
    embedded, refused, redrawn = font_tools.embed(target, one_family_plan(family, entry),
                                                  convert_outlines=False)
    passed &= check(embedded == [] and "turned off" in refused[0][1],
                    "fonts: with conversion off it is refused", refused)
    unreadable = dict(entry, licenceText="All rights reserved, Example Type Foundry")
    embedded, refused, redrawn = font_tools.embed(target, one_family_plan(family, unreadable))
    passed &= check(embedded == [] and "does not clearly allow" in refused[0][1],
                    "fonts: a licence this cannot read is refused", refused)
    embedded, refused, redrawn = font_tools.embed(target, one_family_plan(family, entry),
                                                 convert_seconds=0.0)
    passed &= check(embedded == [] and "not enough time" in refused[0][1],
                    "fonts: and so is a margin too thin to redraw it in", refused)
    return passed


def run_font_control_char_checks(font_tools, font_dir, workdir):
    """A family carrying raw control characters must not break the table.

    `escape` only quotes `&`, `<` and `"`, and nothing stops a font's own name
    table from saying its family is `Ev\\x01il\\x07Fam`: written straight into
    `w:font w:name="..."` that is not well-formed XML. `read_font` is where
    every later use of a family or subfamily reads it, so it is where the
    characters come out.
    """
    from fontTools.ttLib import TTFont

    passed = True
    source = next((os.path.join(font_dir, n) for n in sorted(os.listdir(font_dir))
                   if n.lower().endswith(".ttf")), None)
    if source is None:
        print("SKIP  no TrueType font was staged, so the control-character check is untested")
        return passed

    renamed = os.path.join(workdir, "control-char-font.ttf")
    font = TTFont(source)
    for record in font["name"].names:
        if record.nameID == 1:
            record.string = "Ev\x01il\x07Fam"
        elif record.nameID == 2:
            record.string = "Reg\x0cular"
    font.save(renamed)
    font.close()

    entry = font_tools.read_font(renamed)
    passed &= check(entry is not None, "fonts: a font renamed with control characters still reads")
    if entry is None:
        return passed
    passed &= check(entry["family"] == "EvilFam",
                    "fonts: control characters are stripped from the family", entry["family"])
    passed &= check(entry["subfamily"] == "Regular",
                    "fonts: and from the subfamily", entry["subfamily"])

    target = os.path.join(workdir, "control-char.docx")
    blank = Document()
    blank.add_paragraph("x")
    blank.save(target)
    embedded, refused, _ = font_tools.embed(target, one_family_plan(entry["family"], entry))
    passed &= check(embedded == [entry["family"]], "fonts: it is still carried", (embedded, refused))
    with zipfile.ZipFile(target) as archive:
        table = archive.read("word/fontTable.xml").decode("utf-8")
    try:
        ElementTree.fromstring(table)
        well_formed = True
    except ElementTree.ParseError as err:
        well_formed = False
        print("     %s" % err)
    passed &= check(well_formed, "fonts: the fontTable.xml it wrote parses as XML")

    # A name table is as long as the font wants it to be, and the family
    # reaches `w:name` in every run of the document.
    long_name = os.path.join(workdir, "long-name-font.ttf")
    font = TTFont(source)
    for record in font["name"].names:
        if record.nameID in (1, 16):
            record.string = "L" * 4000
    font.save(long_name)
    font.close()
    entry = font_tools.read_font(long_name)
    passed &= check(entry is not None and len(entry["family"]) == font_tools.MAX_FAMILY_CHARS,
                    "fonts: a 4,000 character family is cut to the cap",
                    entry and len(entry["family"]))
    return passed


def run_font_failure_checks(font_tools, font_dir, workdir):
    """A face, or the whole embed step, that raises must not cost the document.

    `to_truetype` and `fonts.embed` both run against a font this server did
    not choose, so a bug or a fontTools edge case in either has to degrade to
    a warning: the caller already has a Word file that passed the text check
    and the link check, and losing it over a font is the worse failure.
    """
    import cases

    passed = True
    source = next((os.path.join(font_dir, n) for n in sorted(os.listdir(font_dir))
                   if n.lower().endswith(".otf")), None)
    if source is None:
        print("SKIP  no OpenType font was staged, so the failure paths are untested")
        return passed
    entry = font_tools.read_font(source)
    if not entry or not entry.get("postscriptOutlines"):
        print("SKIP  the staged .otf did not read as PostScript outlines")
        return passed

    target = os.path.join(workdir, "raiser.docx")
    blank = Document()
    blank.add_paragraph("x")
    blank.save(target)

    def raise_to_truetype(*_args, **_kwargs):
        raise RuntimeError("synthetic to_truetype failure")

    original_to_truetype = font_tools.to_truetype
    font_tools.to_truetype = raise_to_truetype
    try:
        embedded, refused, redrawn = font_tools.embed(target, one_family_plan(entry["family"], entry))
    finally:
        font_tools.to_truetype = original_to_truetype
    passed &= check(embedded == [] and redrawn == [],
                    "fonts: a face whose conversion raises is skipped, not fatal", (embedded, refused))
    passed &= check(bool(refused) and "could not be redrawn" in refused[0][1],
                    "fonts: and gets the same warning a declined conversion gives", refused)

    def raise_embed(*_args, **_kwargs):
        raise RuntimeError("synthetic embed failure")

    original_embed = font_tools.embed
    font_tools.embed = raise_embed
    try:
        pdf_path = os.path.join(workdir, "embedfail.pdf")
        docx_path = os.path.join(workdir, "embedfail.docx")
        compile_case(cases.HYPHENS, pdf_path)
        result = convert.convert(pdf_path, docx_path, font_dirs=[font_dir])
    finally:
        font_tools.embed = original_embed
    passed &= check(os.path.exists(docx_path) and os.path.getsize(docx_path) > 0,
                    "fonts: a raising embed step still leaves the finished document")
    passed &= check(result["fontsEmbedded"] == [], "fonts: nothing is reported embedded",
                    result["fontsEmbedded"])
    passed &= check(any("internal error" in w for w in result["warnings"]),
                    "fonts: and one plain warning says fonts were not embedded", result["warnings"])
    return passed


RUN_FONTS = W + "rFonts"
RUN_STYLE = {W + "b": "bold", W + "i": "italic"}


def run_flag(properties, tag):
    node = properties.find(tag)
    if node is None:
        return False
    return (node.get(W + "val") or "true").lower() not in ("0", "false", "off")


def run_font_naming_checks(font_tools, docx_path, font_dir, result):
    """Every run names a family Word can resolve, and the table agrees.

    This is what the whole change is for, so it is asserted on the file
    rather than on the functions: the name in a run has to be a family name
    out of some staged font file's own name table, the bold and italic bits
    have to be the ones that file states in name ID 2, and there has to be
    one `w:font` for each of those families carrying the faces it declares.
    """
    passed = True
    staged = font_tools.collect([font_dir])
    by_family = {}
    for entry in staged:
        by_family.setdefault(entry["family"], {})[(entry["bold"], entry["italic"])] = entry
    with zipfile.ZipFile(docx_path) as archive:
        parts = [n for n in archive.namelist()
                 if re.match(r"word/(document|header\d*|footer\d*)\.xml$", n)]
        table = ElementTree.fromstring(archive.read("word/fontTable.xml"))
        runs = []
        for part in sorted(parts):
            for paragraph in ElementTree.fromstring(archive.read(part)).iter(W + "p"):
                for run in paragraph.iter(W + "r"):
                    drawn = [child.tag for child in run if child.tag in VISIBLE_RUN_CHILDREN]
                    if drawn:
                        runs.append((part, run))

    bare = [part for part, run in runs
            if (run.find(W + "rPr") is None
                or run.find(W + "rPr").find(RUN_FONTS) is None
                or not run.find(W + "rPr").find(RUN_FONTS).get(W + "ascii"))]
    passed &= check(not bare, "fonts: every run that draws something names a font", bare[:3])

    wrong = []
    for _, run in runs:
        properties = run.find(W + "rPr")
        fonts_node = properties.find(RUN_FONTS) if properties is not None else None
        if fonts_node is None:
            continue
        name = fonts_node.get(W + "ascii")
        slots = by_family.get(name)
        if slots is None:
            wrong.append((name, "no staged file calls itself that"))
            continue
        want = (run_flag(properties, W + "b"), run_flag(properties, W + "i"))
        if want not in slots:
            wrong.append((name, "no face of it is %s" % (want,)))
        for attribute in (W + "hAnsi", W + "cs"):
            if fonts_node.get(attribute) != name:
                wrong.append((name, "%s disagrees" % attribute.split("}")[1]))
    passed &= check(not wrong, "fonts: every run names a family out of a font file, with its own flags",
                    wrong[:3])

    declared = {}
    for node in table.iter(W + "font"):
        slots = [child.tag.split("}")[1] for child in node if child.tag.startswith(W + "embed")]
        if slots:
            declared[node.get(W + "name")] = sorted(slots)
    named = sorted({fonts_node.get(W + "ascii") for _, run in runs
                    for properties in [run.find(W + "rPr")] if properties is not None
                    for fonts_node in [properties.find(RUN_FONTS)] if fonts_node is not None})
    passed &= check(sorted(declared) == sorted(result["fontsEmbedded"]),
                    "fonts: one w:font per carried family", (sorted(declared), result["fontsEmbedded"]))
    for family, slots in declared.items():
        want = sorted(font_tools.STYLE_ELEMENTS[style].split(":")[1]
                      for style in by_family.get(family, {})
                      if font_tools.carriable(by_family[family][style]))
        passed &= check(slots == want, "fonts: %s carries the faces its files declare" % family,
                        (slots, want))
    passed &= check(all(name in declared or name not in by_family for name in named),
                    "fonts: no run names a carried family the table left out", named)
    return passed


def run_documents():
    import cases

    if shutil.which("typst") is None:
        print("SKIP  the document cases need the typst CLI on PATH")
        return True
    passed = True
    workdir = tempfile.mkdtemp(prefix="btct-convert-test-")
    try:
        # 1. A long table whose rows differ only by numbers.
        result, path = run_case(cases.TABLE, workdir, "table")
        passed &= check(result["textCheck"]["missing"] == 0, "table: no line of the report was lost",
                        result["textCheck"])
        passed &= check(not result["headerFooter"]["header"] and not result["headerFooter"]["footer"],
                        "table: no table row was mistaken for a band", result["headerFooter"])
        body = squashed_docx(path, body_only=False)
        passed &= check(body.count("host111") == 1 and body.count("host073") == 1,
                        "table: Host-111 and Host-073 each appear once",
                        (body.count("host111"), body.count("host073")))

        # 1b. The rung below the fallback. A pass that rearranges what the
        #     converter wrote is vandalised so it loses a line of a table, and
        #     the only thing that gets the line back is turning every
        #     restructuring and placement pass off. This document has no band
        #     and no table of contents, so the two repairs are already off and
        #     that rung is the one left.
        keep_cells = convert.align_table_cells
        calls, damaged = [], set()

        def vandalise(table, top, left, lines, anchors=(), depth=convert.RHYTHM_NESTING, seams=None):
            calls.append(table)
            root = id(table.getroottree().getroot())
            if root not in damaged:
                # The last long line of the table, because every row of this
                # document reads almost the same and a shorter one is a
                # substring of a row that is still there, which the check
                # counts as present.
                victim = None
                for paragraph in table.iter(convert.qn("w:p")):
                    nodes = list(paragraph.iter(convert.qn("w:t")))
                    if len(convert.compare_squash("".join(n.text or "" for n in nodes))) >= 20:
                        victim = nodes
                if victim:
                    for node in victim:
                        node.text = ""
                    damaged.add(root)
            return keep_cells(table, top, left, lines, anchors, depth, seams)

        convert.align_table_cells = vandalise
        try:
            ladder = convert.convert(os.path.join(workdir, "table.pdf"),
                                     os.path.join(workdir, "ladder.docx"))
        finally:
            convert.align_table_cells = keep_cells
        passed &= check(bool(calls) and bool(damaged),
                        "ladder: the vandalised table pass really ran", (len(calls), len(damaged)))
        passed &= check(ladder["textCheck"]["missing"] == 0 and ladder["textCheck"]["fellBack"],
                        "ladder: the line came back once the layout passes went off",
                        ladder["textCheck"])
        passed &= check(convert.PLAIN_LAYOUT in ladder["warnings"],
                        "ladder: and the file says it carries the converter's own layout",
                        ladder["warnings"])

        # 2. A repeated first body line under a real running header.
        result, path = run_case(cases.FIRST_BODY_LINE, workdir, "firstline")
        passed &= check(result["textCheck"]["missing"] == 0, "first line: nothing was lost", result["textCheck"])
        passed &= check(result["headerFooter"]["header"] and result["headerFooter"]["footer"],
                        "first line: the real header and footer were lifted", result["headerFooter"])
        opener = convert.compare_squash("Lorem ipsum dolor sit amet, consectetur adipiscing elit")
        passed &= check(opener in squashed_docx(path, body_only=True),
                        "first line: the body line is still in the body")
        with zipfile.ZipFile(path) as archive:
            bands = "".join(
                archive.read(n).decode("utf-8", "replace")
                for n in archive.namelist()
                if n.startswith("word/header") or n.startswith("word/footer")
            )
        passed &= check("Lorem ipsum dolor" not in bands, "first line: it was not pasted into a header")
        passed &= check("Acme Security Assessment" in bands, "first line: the real header was")

        # 3. Odd and even headers. The tagged path gives each Word section the
        #    header of its own page, so the two texts alternate as the PDF has them.
        result, path = run_case(cases.ODD_EVEN, workdir, "oddeven")
        passed &= check(result["textCheck"]["missing"] == 0, "odd/even: nothing was lost", result["textCheck"])
        passed &= check(result["headerFooter"]["header"] and not result["textCheck"]["fellBack"],
                        "odd/even: the header was lifted", result["headerFooter"])
        count, headers = section_band_texts(path, "header")
        passed &= check(count == result["pages"] and all(
            ("Acme Security Assessment" in text) == (index % 2 == 0)
            and ("Contoso Consulting Group" in text) == (index % 2 == 1) for index, text in enumerate(headers)),
            "odd/even: each page keeps the header of its own parity, never both", headers)

        # 4. The shape the repairs exist for: both must run, and say nothing.
        result, _ = run_case(cases.REPORT, workdir, "report")
        passed &= check(result["textCheck"]["missing"] == 0, "report: nothing was lost", result["textCheck"])
        passed &= check(not result["textCheck"]["fellBack"], "report: no fallback was needed")
        passed &= check(result["headerFooter"]["header"] and result["headerFooter"]["footer"],
                        "report: header and footer were lifted", result["headerFooter"])
        # The rebuild only engages when the converter keeps the entries as
        # separate paragraphs. On a page this sparse it merges the whole list
        # into one, and then the repair declines and says so, which is the
        # right answer: the list is ugly and complete. The reference report is
        # dense enough that its 27 entries are rebuilt.
        expected = ("rebuilt", "could not be found in the converted file")
        noise = [w for w in result["warnings"] if not any(e in w for e in expected)]
        passed &= check(noise == [], "report: no warning beyond the table of contents", noise)

        # 5. A table of contents in three scripts.
        result, path = run_case(cases.MIXED_SCRIPT_TOC, workdir, "scripts")
        passed &= check(result["textCheck"]["missing"] == 0, "scripts: nothing was lost", result["textCheck"])
        passed &= check(not result["textCheck"]["fellBack"], "scripts: no fallback was needed")
        body = squashed_docx(path, body_only=False)
        for title in ("Глава первая",
                      "Κεφάλαιο δύο"):
            passed &= check(convert.compare_squash(title) in body, "scripts: %r survived" % title)

        # 6. Links.
        result, path = run_case(cases.LINKS, workdir, "links")
        targets = []
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name.endswith(".rels"):
                    for rel in ElementTree.fromstring(archive.read(name)):
                        if rel.get("TargetMode") == "External":
                            targets.append(rel.get("Target"))
        passed &= check(sorted(targets) == ["https://example.com/ok", "mailto:team@example.com"],
                        "links: only http(s) and mailto survive", targets)
        passed &= check(any("unsupported addresses" in w for w in result["warnings"]),
                        "links: the removal is reported", result["warnings"])
        passed &= check(convert.unsafe_docx_targets(path) == [], "links: the output check agrees")

        # 7. A repeating rule under the header with a distinct title above it.
        #    The rectangle that gets erased is the whole page width, so growing
        #    it to reach the rule used to delete every title on every page,
        #    with the text check blind to it.
        result, path = run_case(cases.RULE_OVER_TITLES, workdir, "rule")
        passed &= check(result["textCheck"]["missing"] == 0, "rule: nothing was lost", result["textCheck"])
        body = squashed_docx(path, body_only=False)
        kept = [i + 1 for i in range(8) if convert.compare_squash("unique title %d" % (i + 1)) in body]
        passed &= check(kept == list(range(1, 9)), "rule: every finding title survived", kept)
        passed &= check(result["headerFooter"]["header"], "rule: the header was still lifted")

        # 8. A document whose every page is nothing but the candidate.
        result, _ = run_case(cases.NOTHING_BUT_A_BAND, workdir, "bandonly")
        passed &= check(result["textCheck"]["missing"] == 0, "band only: nothing was lost", result["textCheck"])
        passed &= check(result["textCheck"]["pdfLines"] > 0, "band only: there was something to check",
                        result["textCheck"])
        passed &= check(not result["headerFooter"]["header"], "band only: the header was declined")
        passed &= check(any("told apart from the body" in w for w in result["warnings"]),
                        "band only: the warning says why", result["warnings"])

        # 9. Soft hyphens are the typesetter's line breaks and have to go;
        #    a real hyphen at a line end is part of the word and must stay.
        result, path = run_case(cases.HYPHENS, workdir, "hyphens")
        passed &= check(result["textCheck"]["missing"] == 0, "hyphens: nothing was lost", result["textCheck"])
        text = convert.docx_text(path, body_only=False)
        passed &= check("­" not in text, "hyphens: no soft hyphen reached the Word file",
                        text.count("­"))
        passed &= check("non-critical" in text, "hyphens: a real hyphen survived")
        passed &= check("employed" in text or "employ" in text, "hyphens: the broken word was joined")

        # 10. The marker repair must not reach into code. A space inserted
        #     into a command is invisible to the text check, which squashes
        #     whitespace, so this is the only thing that catches it.
        result, path = run_case(cases.CODE_AND_MARKERS, workdir, "code")
        passed &= check(result["textCheck"]["missing"] == 0, "code: nothing was lost", result["textCheck"])
        text = convert.docx_text(path, body_only=False)
        passed &= check("./deploy 1.{{X}}" in text, "code: the command kept its marker unspaced")
        passed &= check("check version1.2.3" in text, "code: the command kept its version unspaced")
        # The marker is followed by a space, or by the tab the hanging-indent pass puts there (docx_text reads a tab as a line end).
        passed &= check(re.search(r"1\.\s\{\{X\}\} must be rebuilt", text) is not None,
                        "prose: the list marker got its space back")
        passed &= check(re.search(r"2\.\sversion1\.2\.3 is the baseline", text) is not None,
                        "prose: a version number inside a word was left alone")

        # 11. Line breaks. The default reproduces the PDF's own line endings
        #     with manual breaks; `word` lets Word re-flow. Neither may put a
        #     break inside code or inside a table cell.
        pdf_path = os.path.join(workdir, "breaks.pdf")
        compile_case(cases.HYPHENS, pdf_path)
        forced = os.path.join(workdir, "breaks-pdf.docx")
        flowed = os.path.join(workdir, "breaks-word.docx")
        result = convert.convert(pdf_path, forced)
        passed &= check(result["lineBreaks"] == "pdf", "breaks: pdf is the default", result["lineBreaks"])
        passed &= check(result["textCheck"]["missing"] == 0, "breaks: nothing was lost", result["textCheck"])
        other = convert.convert(pdf_path, flowed, line_breaks="word")
        passed &= check(other["lineBreaks"] == "word", "breaks: word mode is available")
        with zipfile.ZipFile(forced) as archive:
            forced_xml = archive.read("word/document.xml").decode("utf-8")
        with zipfile.ZipFile(flowed) as archive:
            flowed_xml = archive.read("word/document.xml").decode("utf-8")
        passed &= check(forced_xml.count("<w:br/>") > flowed_xml.count("<w:br/>"),
                        "breaks: pdf mode adds manual breaks",
                        (forced_xml.count("<w:br/>"), flowed_xml.count("<w:br/>")))

        result, path = run_case(cases.CODE_AND_MARKERS, workdir, "codebreaks")
        with zipfile.ZipFile(path) as archive:
            body = archive.read("word/document.xml").decode("utf-8")
        panel = body[body.find("deploy") - 900 : body.find("deploy") + 400] if "deploy" in body else ""
        passed &= check("<w:br/>" not in panel, "breaks: none inside a code panel")

        # 12. Fonts. The parts have to be there, obfuscated with the key the
        #     table names, and de-obfuscate to the file they came from.
        import fonts as font_tools

        font_dir = os.path.join(workdir, "fonts")
        os.makedirs(font_dir, exist_ok=True)
        source = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "dist", "fonts")
        if os.path.isdir(source):
            for name in sorted(os.listdir(source)):
                if name.lower().endswith((".ttf", ".otf")):
                    shutil.copyfile(os.path.join(source, name), os.path.join(font_dir, name))
        pdf_path = os.path.join(workdir, "fontcase.pdf")
        docx_path = os.path.join(workdir, "fontcase.docx")
        # Set in a family the Word file can actually carry, whichever one of
        # the staged files that turns out to be. typst's own default is an
        # OpenType face with PostScript outlines, which Word declines, and a
        # test that read the default would be testing the refusal twice.
        staged = font_tools.collect([font_dir])
        carried = next((e for e in staged if font_tools.carriable(e)), None)
        prelude = '#set text(font: "%s")\n' % carried["family"] if carried else ""
        compile_case(prelude + cases.HYPHENS, pdf_path)
        result = convert.convert(pdf_path, docx_path, font_dirs=[font_dir])
        if not result["fontsEmbedded"]:
            print("SKIP  no font file matched the families this PDF names")
        else:
            with zipfile.ZipFile(docx_path) as archive:
                names = archive.namelist()
                parts = [n for n in names if n.startswith("word/fonts/")]
                table = archive.read("word/fontTable.xml").decode("utf-8")
                rels = archive.read("word/_rels/fontTable.xml.rels").decode("utf-8")
                settings = archive.read("word/settings.xml").decode("utf-8")
                types = archive.read("[Content_Types].xml").decode("utf-8")
                passed &= check(bool(parts), "fonts: a font part was written", parts)
                passed &= check("<w:embedTrueTypeFonts/>" in settings, "fonts: settings say they are embedded")
                passed &= check("odttf" in types, "fonts: the content type is declared")
                keys = re.findall(r'<w:embed\w+ r:id="(\w+)" w:fontKey="([^"]+)"', table)
                passed &= check(bool(keys), "fonts: the table names a key per face", len(keys))
                same = 0
                for rid, key in keys:
                    target = re.search(r'Id="%s"[^>]*Target="([^"]+)"' % rid, rels)
                    if not target:
                        continue
                    raw = font_tools.obfuscate(archive.read("word/" + target.group(1)), key)
                    if raw[:4] in (bytes([0, 1, 0, 0]), b"OTTO", b"true"):
                        for name in os.listdir(font_dir):
                            with open(os.path.join(font_dir, name), "rb") as handle:
                                if handle.read() == raw:
                                    same += 1
                                    break
                passed &= check(same == len(keys), "fonts: every part is its source font, obfuscated",
                                (same, len(keys)))
            passed &= check(result["textCheck"]["missing"] == 0, "fonts: nothing was lost")
            passed &= run_font_naming_checks(font_tools, docx_path, font_dir, result)

        # A font the foundry forbids must be skipped rather than embedded.
        passed &= check(not font_tools.embeddable({"fsType": 0x0002}), "fonts: a restricted font is refused")
        passed &= check(not font_tools.embeddable({"fsType": 0x0200}), "fonts: a bitmap-only font is refused")
        passed &= check(font_tools.embeddable({"fsType": 0}), "fonts: an installable font is allowed")
        passed &= check(font_tools.embeddable({"fsType": 8}), "fonts: an editable font is allowed")
        # Preview-and-print is refused for Word's reason, not the foundry's:
        # a document carrying one opens read-only and cannot be exported.
        passed &= check(not font_tools.embeddable({"fsType": 0x0004}),
                        "fonts: a preview-and-print font is refused")
        passed &= check("read-only" in font_tools.refusal([face_entry(__file__, "Locked", fsType=0x0004)]),
                        "fonts: and the reason says why")

        # Name ID 2 decides which of Word's four slots a face fills, and both
        # of the ways of misreading it lose a face.
        for subfamily, want in (("Regular", (False, False)), ("Bold", (True, False)),
                                ("Italic", (False, True)), ("Bold Italic", (True, True)),
                                ("BoldItalic", (True, True)), ("BoldOblique", (True, True)),
                                ("Book", (False, False)), ("Oblique", (False, True)),
                                ("SemiBold", (False, False)), ("SemiBoldItalic", (False, True))):
            passed &= check(font_tools.style_of(subfamily) == want,
                            "fonts: %r is %s" % (subfamily, want), font_tools.style_of(subfamily))

        # A font file cut short parses lazily and would be embedded whole.
        whole = [name for name in sorted(os.listdir(font_dir)) if name.lower().endswith((".ttf", ".otf"))]
        if whole:
            with open(os.path.join(font_dir, whole[0]), "rb") as handle:
                raw = handle.read()
            hurt_dir = os.path.join(workdir, "hurtfonts")
            os.makedirs(hurt_dir, exist_ok=True)
            hurt = os.path.join(hurt_dir, whole[0])
            with open(hurt, "wb") as handle:
                handle.write(raw[: len(raw) // 3])
            entry = font_tools.read_font(hurt)
            passed &= check(entry is None or entry.get("damaged"),
                            "fonts: a truncated font is read as damaged", entry and entry["family"])
            passed &= check(font_tools.read_font(os.path.join(font_dir, whole[0])).get("damaged") is False,
                            "fonts: a whole font is not")
            hurt_entry = dict(face_entry(hurt, "Hurt"), damaged=True)
            _, refused, _ = font_tools.embed(None, one_family_plan("Hurt", hurt_entry))
            passed &= check([why for _, why in refused] == ["the font file is damaged"],
                            "fonts: and it says the file is damaged", refused)

        passed &= run_font_conversion_checks(font_tools, font_dir, workdir)
        passed &= run_font_control_char_checks(font_tools, font_dir, workdir)
        passed &= run_font_failure_checks(font_tools, font_dir, workdir)

        # And there is a bound on the bytes a Word file may carry.
        huge = face_entry(os.path.join(font_dir, whole[0]) if whole else __file__, "Huge")
        keep = font_tools.MAX_FONT_BYTES
        font_tools.MAX_FONT_BYTES = 1
        try:
            embedded, refused, _ = font_tools.embed(None, one_family_plan("Huge", huge))
        finally:
            font_tools.MAX_FONT_BYTES = keep
        passed &= check(embedded == [] and refused and "may take" in refused[0][1],
                        "fonts: a face over the byte cap is refused", refused)

        # 12b. A budget too small for the fidelity passes skips them, says so,
        #      and still checks that the Word file kept the PDF's text.
        pdf_path = os.path.join(workdir, "budget.pdf")
        compile_case(cases.REPORT, pdf_path)
        tight = convert.convert(pdf_path, os.path.join(workdir, "budget.docx"), budget_seconds=0.001)
        passed &= check(convert.CROWDED_OUT in tight["warnings"],
                        "budget: a thin margin skips the layout passes", tight["warnings"][:1])
        passed &= check(tight["textCheck"]["missing"] == 0,
                        "budget: the text check still runs", tight["textCheck"])
        roomy = convert.convert(pdf_path, os.path.join(workdir, "roomy.docx"), budget_seconds=600)
        passed &= check(convert.CROWDED_OUT not in roomy["warnings"],
                        "budget: a real margin keeps them", roomy["warnings"])
        passed &= check(roomy["pages"] == tight["pages"],
                        "budget: skipping them does not move the page count",
                        (roomy["pages"], tight["pages"]))
        passed &= check(convert.affordable(time.monotonic(), None, 999.0),
                        "budget: no budget means no limit")

        # 13. Too many shapes for the converter.
        pdf_path = os.path.join(workdir, "shapes.pdf")
        compile_case(cases.SHAPES, pdf_path)
        started = time.perf_counter()
        try:
            convert.convert(pdf_path, os.path.join(workdir, "shapes.docx"))
            passed &= check(False, "shapes: the document was refused")
        except convert.ConvertError as err:
            spent = time.perf_counter() - started
            passed &= check("too complex" in str(err), "shapes: the document was refused", err)
            passed &= check(spent < 1.0, "shapes: refused in under a second", "%.3f s" % spent)

        # 14. A paragraph the converter dropped, simulated by taking one out of
        #    the PDF's side of the comparison's counterpart.
        result, path = run_case(cases.REPORT, workdir, "report2")
        text = convert.docx_text(path, body_only=True)
        cut = text.replace("Scope of the engagement", "", 1)
        doc = pymupdf.open(os.path.join(workdir, "report2.pdf"))
        lines = convert.body_lines(doc, [])
        doc.close()
        passed &= check(convert.missing_from_docx(lines, cut) != [],
                        "a heading removed from the Word file is caught")

        # 15. A two-piece footer that pdf2docx turns into a borderless table
        #     reaching up the page. Raw still loses the prose, which is what
        #     shows the case still has the shape; the converter must not.
        #     This PDF is tagged, so its top rung lifts the footer through
        #     the tags and converts a PDF without it, and a clean top rung
        #     says nothing about the page-order wrapper. Only the raw rung
        #     still meets this footer, and the raw check is the one that says
        #     whether the wrapper can go. Do not retire it on a tagged run.
        result, path = run_case(cases.FOOTER_TABLE_REACH, workdir, "footer-reach")
        passed &= check(result["textCheck"]["missing"] == 0 and not result["textCheck"]["fellBack"],
                        "footer-reach: no prose line lost, at the top rung", result["textCheck"])
        passed &= check("page-order-fallthrough" not in result["declined"],
                        "footer-reach: the page-order wrapper sorted every page", result["declined"])
        raw = convert.convert(os.path.join(workdir, "footer-reach.pdf"),
                              os.path.join(workdir, "footer-reach-raw.docx"), raw=True)
        passed &= check(raw["textCheck"]["missing"] > 0,
                        "footer-reach: raw pdf2docx still loses it (if not, the fix can go)",
                        raw["textCheck"])

        # A pdf2docx whose blocks no longer sort (here, blocks with no bbox)
        # gets its own borderless-table step, unpatched, rather than an error.
        class Unsortable(list):
            def reset(self, blocks):
                raise AssertionError("reset must not be reached")

        class Constructor:
            _blocks = Unsortable([object(), object()])

        calls = []
        fell_through = convert.Counter()
        wrapped = convert.stream_tables_in_page_order(lambda self, *a, **k: calls.append(self) or "done",
                                                      fell_through)
        answer = wrapped(Constructor())
        passed &= check(answer == "done" and len(calls) == 1,
                        "footer-reach: a block list that cannot be sorted falls through to pdf2docx's own step")
        # The fall-through is counted, so a pdf2docx that stops sorting shows
        # in result.json rather than silently losing the footer fix.
        passed &= check(fell_through["page-order-fallthrough"] == 1,
                        "footer-reach: an unsortable block list is counted as a fall-through", dict(fell_through))

        class NoBlocks:
            pass

        wrapped(NoBlocks())
        passed &= check(fell_through["page-order-fallthrough"] == 2 and len(calls) == 2,
                        "footer-reach: a constructor with no block list is counted as a fall-through",
                        dict(fell_through))

        # 16. A letter whose rhythm walk would anchor on an address block that
        #     Word sets at another pitch than the PDF. The walk keeps this page
        #     top as its anchor now (case 19), so the relative walk, which is
        #     where the anchor is chosen, is run on the converter's own spacing
        #     directly.
        result, path = run_case(cases.UNEVEN_ANCHOR, workdir, "uneven-anchor")
        passed &= check(result["textCheck"]["missing"] == 0, "uneven-anchor: no line lost", result["textCheck"])
        plain = os.path.join(workdir, "uneven-anchor-plain.docx")
        convert.convert(os.path.join(workdir, "uneven-anchor.pdf"), plain, skip=frozenset({"rhythm"}))
        document = Document(plain)
        doc = pymupdf.open(os.path.join(workdir, "uneven-anchor.pdf"))
        rows = convert.page_anchor_rows(doc, [])[0]
        height = doc[0].rect.height
        doc.close()
        tried = convert.Counter()
        undo, changed, _, _ = convert.place_page(
            document.sections[0], convert.section_blocks(document.element.body)[0], rows,
            convert.text_columns(rows), height, False, declined=tried)
        passed &= check(tried.get("rhythm-uneven-anchor", 0) > 0 and not undo and not changed,
                        "uneven-anchor: a relative walk declines an anchor off the PDF's pitch", dict(tried))

        # 17. A heading the converter joined to the paragraph under it, in a
        #     block already set at the PDF's pitch.
        result, path = run_case(cases.HEADING_IN_PARAGRAPH, workdir, "heading-in-paragraph")
        # The line pass gives the joined block's run to the heading and the
        # paragraph each, so the heading ends on a line of its own and the
        # vertical pass has no tall row left to decline.
        with zipfile.ZipFile(path) as archive:
            root = ElementTree.fromstring(archive.read("word/document.xml"))
        starts = [convert.compare_squash("".join(t.text or "" for t in p.iter(W + "t")))[:12] for p in root.iter(W + "p")]
        passed &= check("2.conclusion" in starts and result["declined"].get("rhythm-tall-row", 0) == 0,
                        "heading-in-paragraph: the joined heading ends on a line of its own, nothing declined for it",
                        (sorted(set(starts)), dict(result["declined"])))
        passed &= check(result["textCheck"]["missing"] == 0, "heading-in-paragraph: no line lost",
                        result["textCheck"])

        # 18. A code panel the converter coloured in part.
        result, path = run_case(cases.PANEL_PARTLY_SHADED, workdir, "panel-partly-shaded")
        # pdf2docx colours three of the four lines as cells of its own in the
        # panel's grey; those cells are the panel, so it is made whole.
        with pymupdf.open(os.path.join(workdir, "panel-partly-shaded.pdf")) as doc:
            rects = [d["rect"] for d in doc[0].get_drawings() if d.get("fill") and d["rect"].width > 40]
        whole = []
        with zipfile.ZipFile(path) as archive:
            root = ElementTree.fromstring(archive.read("word/document.xml"))
        for table in root.iter(W + "tbl"):
            rows = table.findall(W + "tr")
            inner = "".join(t.text or "" for t in table.iter(W + "t"))
            if len(rows) == 1 and "POST /ar/products" in inner and "comment=hello" in inner:
                height = rows[0].find(W + "trPr").find(W + "trHeight")
                whole.append(convert.twips_of(height.get(W + "val")) if height is not None else None)
        passed &= check(result["declined"].get("panels-partial", 0) == 0 and rects and whole
                        and any(h is not None and abs(h - rects[0].height) < 2.0 for h in whole),
                        "panel-partly-shaded: the part-coloured panel is one table as tall as the PDF's rectangle",
                        (result["declined"], [round(r.height, 1) for r in rects], whole))
        passed &= check(result["textCheck"]["missing"] == 0, "panel-partly-shaded: no line lost",
                        result["textCheck"])

        # 19. An empty paragraph with spacing after it above the page's first
        #     heading. The walk keeps the page top as its anchor, so the
        #     heading block's first baseline, added up the way Word lays the
        #     page out (the larger of two touching spacings), lands where the
        #     walk straddles the PDF's rows for that block.
        result, path = run_case(cases.SPACER_ABOVE_HEADING, workdir, "spacer-above-heading")
        passed &= check("rhythm-uneven-anchor" not in result["declined"],
                        "spacer-above-heading: the walk anchors on the page top", result["declined"])
        document = Document(path)
        cursor, pending, baseline = document.sections[0].top_margin.pt, 0.0, None
        for order, paragraph in enumerate(document.paragraphs):
            metrics = convert.paragraph_metrics(paragraph._element)
            if metrics is None:
                break
            cursor += max(pending, metrics["before"] if order else 0.0)
            if paragraph.text.startswith("Stored"):
                baseline = cursor + convert.EXACT_BASELINE_RATIO * metrics["line"]
                held = dict(metrics, text=paragraph.text)
                break
            cursor += metrics["lines"] * metrics["line"]
            pending = metrics["after"]
        doc = pymupdf.open(os.path.join(workdir, "spacer-above-heading.pdf"))
        rows = convert.page_anchor_rows(doc, [])[0]
        doc.close()
        wanted = None
        if baseline is not None:
            text = convert.compare_squash(held["text"])
            mine = [row for row in rows if row["key"] and row["key"] in text]
            wanted = convert.first_baseline_target(mine, held["line"])
        passed &= check(wanted is not None and abs(baseline - wanted) < 0.5,
                        "spacer-above-heading: the heading block's first baseline is where the PDF puts it",
                        (baseline, wanted))
        passed &= check(result["textCheck"]["missing"] == 0, "spacer-above-heading: no line lost",
                        result["textCheck"])

        # 20. Paragraphs with an indented first line: the alignment pass must
        #     keep the indent the converter read off the PDF.
        result, path = run_case(cases.INDENTED_FIRST_LINE, workdir, "indented-first-line")
        passed &= check(result["declined"].get("align-earlier-line", 0) > 0,
                        "indented-first-line: align leaves a paragraph whose first line is not in the run",
                        result["declined"])
        indents = [p.paragraph_format.first_line_indent for p in Document(path).paragraphs
                   if p.text.startswith("Lorem")]
        passed &= check(indents and all(indent is not None and indent.pt > 5 for indent in indents),
                        "indented-first-line: every paragraph keeps its first-line indent",
                        [indent and indent.pt for indent in indents])
        passed &= check(result["textCheck"]["missing"] == 0, "indented-first-line: no line lost",
                        result["textCheck"])

        # 20b. A table too tall for its figure: the overflow rows are piled up
        #      in the PDF, every rung loses them, and the top rung is kept.
        climbs = []
        original_redo = convert.redo_plainly

        def counted(*args, **kwargs):
            climbs.append(kwargs.get("place"))
            return original_redo(*args, **kwargs)

        convert.redo_plainly = counted
        try:
            result, path = run_case(cases.STACKED_TABLE_FIGURE, workdir, "stacked")
        finally:
            convert.redo_plainly = original_redo
        text = result["textCheck"]
        passed &= check(text["missing"] >= 1 and text["overprinted"] == text["missing"],
                        "stacked: every line lost is one the PDF overprints", text)
        passed &= check(climbs == [True, False] and not text["fellBack"],
                        "stacked: the ladder climbed down, recovered nothing and kept the top rung", (climbs, text))
        passed &= check(result["headerFooter"]["header"] and result["headerFooter"]["footer"]
                        and convert.PLAIN_LAYOUT not in result["warnings"],
                        "stacked: so the header and footer stay lifted", result["headerFooter"])
        passed &= check(any("printed on top of other text on PDF page 2" in w for w in result["warnings"]),
                        "stacked: the warning names the page the PDF cannot be read on", result["warnings"])

        # 20c. A screenshot clipped by its figure's frame: cropped to what the
        #      PDF shows and put inside the frame, inline, with the caption
        #      following it in the flow.
        shot = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 900, 506), 0)
        shot.set_rect(shot.irect, (40, 44, 52))
        shot.set_rect(pymupdf.IRect(100, 300, 800, 480), (120, 160, 220))
        shot.save(os.path.join(workdir, "shot.png"))
        result, path = run_case(cases.FRAMED_FIGURE, workdir, "framed-figure")
        with pymupdf.open(os.path.join(workdir, "framed-figure.pdf")) as doc:
            pictures = convert.pdf_pictures(doc)
        passed &= check(len(pictures) == 2 and all(p["rect"].height > p["visible"].height + 50
                                                    and abs(p["visible"].height - 157) < 2.0 for p in pictures),
                        "framed: the visible part of each picture is the frame's clip, not the whole drawing",
                        [(round(p["rect"].height, 1), round(p["visible"].height, 1)) for p in pictures])
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
            root = ElementTree.fromstring(archive.read("word/document.xml"))
        inlines_in_cells = sum(1 for cell in root.iter(W + "tc") for _ in cell.iter(W + "drawing"))
        passed &= check("<wp:anchor" not in xml and inlines_in_cells == 2 and "<a:srcRect" in xml,
                        "framed: both pictures are inline in their frame cells and cropped",
                        (xml.count("<wp:anchor"), inlines_in_cells, xml.count("<a:srcRect")))
        passed &= check(result["textCheck"]["missing"] == 0 and not result["textCheck"]["fellBack"],
                        "framed: nothing lost", result["textCheck"])

        # 20d. A token typst could not break runs into the margin; pdf2docx
        #      cuts it at its page box. The end is put back from the PDF, so
        #      the report stays on the top rung with nothing missing.
        result, path = run_case(cases.CUT_TOKEN, workdir, "cut-token")
        text = convert.docx_text(path, body_only=False)
        passed &= check("2769d97834453b40f" in convert.compare_squash(text),
                        "cut-token: the whole token is in the Word file")
        passed &= check(result["textCheck"]["missing"] == 0 and not result["textCheck"]["fellBack"]
                        and result["headerFooter"]["header"],
                        "cut-token: nothing missing, no fallback, header lifted", result["textCheck"])

        # 21. A tagged PDF: the header and footer are what typst marked as
        #     Header and Footer artifacts, not a guess.
        result, path = run_case(cases.TAGGED_SKIP_PAGES, workdir, "tagged-skip")
        passed &= check(result["declined"].get("bands-tagged", 0) == 1, "tagged-skip: the band came from the tags", result["declined"])
        passed &= check(result["headerFooter"]["header"] and result["headerFooter"]["footer"], "tagged-skip: header and footer lifted", result["headerFooter"])
        passed &= check(result["textCheck"]["missing"] == 0 and not result["textCheck"]["fellBack"], "tagged-skip: nothing lost", result["textCheck"])
        body = convert.docx_text(path, body_only=True)
        passed &= check("Blue Harbor" not in body, "tagged-skip: header text is not in the body")
        with zipfile.ZipFile(path) as archive:
            page_breaks = archive.read("word/document.xml").count(b'w:type="page"')
        passed &= check(page_breaks == 0 and result["pageBreaksInserted"] == 0, "tagged-skip: pdf2docx's own sections already start each page, so no break is added", page_breaks)

        # The rule and the strip are drawn inside the artifacts, so they are
        # gone from the stripped PDF and only the band can bring them back.
        result, path = run_case(cases.TAGGED_DECORATED, workdir, "tagged-decorated")
        with pymupdf.open(os.path.join(workdir, "tagged-decorated.pdf")) as doc:
            found = convert.detect_tagged_bands(doc, collections.Counter())
        edges = {band["edge"]: band for band in (found[0] if found else [])}
        passed &= check("top" in edges and edges["top"]["rule_under"] is not None,
                        "tagged-decorated: the rule under the header joined the tagged band",
                        edges.get("top", {}).get("rule_under"))
        passed &= check("bottom" in edges and edges["bottom"]["strip_behind"] is not None
                        and edges["bottom"]["strip_full_width"],
                        "tagged-decorated: the full-width strip joined the tagged footer",
                        edges.get("bottom", {}).get("strip_behind"))
        passed &= check(result["declined"].get("bands-tagged", 0) == 1 and not result["textCheck"]["fellBack"]
                        and result["textCheck"]["missing"] == 0,
                        "tagged-decorated: lifted by the tags with nothing lost", (result["declined"], result["textCheck"]))
        with zipfile.ZipFile(path) as archive:
            headers = "".join(archive.read(n).decode("utf-8") for n in archive.namelist() if n.startswith("word/header"))
            footers = "".join(archive.read(n).decode("utf-8") for n in archive.namelist() if n.startswith("word/footer"))
        passed &= check("<w:pBdr><w:bottom" in headers, "tagged-decorated: the Word header carries the rule")
        passed &= check('w:fill="F3F3F3"' in footers, "tagged-decorated: the Word footer carries the strip")
        # The strip runs to the page's edge, so the footer has to as well:
        # Word shades a paragraph's own line only, and a footer the height of
        # its text left white above and below the grey.
        strip = edges.get("bottom", {}).get("strip_y")
        passed &= check(strip is not None and strip[1] >= PAGE_HEIGHT - 1.0,
                        "tagged-decorated: the strip is known to reach the page's edge", strip)
        line = round((PAGE_HEIGHT - strip[0]) * 20) if strip else None
        sections = Document(path).sections
        passed &= check(line is not None and 'w:lineRule="exact" w:line="%d"' % line in footers
                        and all(s.footer_distance.pt == 0 for s in sections[1:]),
                        "tagged-decorated: the footer starts at the edge with a line as tall as the strip",
                        (line, [s.footer_distance.pt for s in sections]))
        passed &= check("<w:position w:val=" in footers,
                        "tagged-decorated: and its text is lifted back onto the PDF's baseline")

        result, path = run_case(cases.TAGGED_FOOTER_TEXT_IN_BODY, workdir, "tagged-footer-text")
        passed &= check(result["headerFooter"]["footer"] and not result["headerFooter"]["header"],
                        "tagged-footer-text: the footer is lifted and no header is made up", result["headerFooter"])
        body = squashed_docx(path, body_only=True)
        passed &= check(body.count(convert.compare_squash("CONFIDENTIAL")) == 2,
                        "tagged-footer-text: each page's body line stays in the body",
                        body.count(convert.compare_squash("CONFIDENTIAL")))
        passed &= check(result["textCheck"]["missing"] == 0 and not result["textCheck"]["fellBack"],
                        "tagged-footer-text: nothing lost", result["textCheck"])

        # Anything that goes wrong reading the tags sends the document to the
        # heuristic, never out of convert(). Reading a row of a stripped copy
        # raises; reading the document convert() opened does not.
        keep_rows, keep_copy = convert.page_rows, convert.tagged.stripped_copy
        copies = []

        def recorded_copy(doc):
            made = keep_copy(doc)
            if made is not None:
                copies.append(made[0])
            return made

        def rows_but_not_of_the_copy(page, logical=False, blocks=None):
            if any(page.parent is copy for copy in copies):
                raise RuntimeError("forced: the stripped copy cannot be read")
            return keep_rows(page, logical, blocks)

        convert.page_rows, convert.tagged.stripped_copy = rows_but_not_of_the_copy, recorded_copy
        try:
            result, path = run_case(cases.FIRST_BODY_LINE, workdir, "tagged-raises")
        finally:
            convert.page_rows, convert.tagged.stripped_copy = keep_rows, keep_copy
        passed &= check(bool(copies), "tagged-raises: the tagged path was tried")
        passed &= check(result["ok"] and result["declined"].get("bands-tagged-inconsistent", 0) == 1
                        and result["declined"].get("bands-heuristic", 0) == 1
                        and not result["declined"].get("bands-tagged"),
                        "tagged-raises: an error reading the tags falls to the heuristic", result["declined"])
        passed &= check(result["headerFooter"] == {"header": True, "footer": True, "pagesWithBands": 4}
                        and result["textCheck"]["missing"] == 0,
                        "tagged-raises: the heuristic's bands, nothing lost",
                        (result["headerFooter"], result["textCheck"]))

        result, path = run_case(cases.TAGGED_TITLE_IN_BODY, workdir, "tagged-title")
        body = convert.compare_squash(convert.docx_text(path, body_only=True))
        passed &= check(body.count(convert.compare_squash("Quarterly Status Report")) == 1, "tagged-title: the heading copy stays in the body")
        passed &= check(result["declined"].get("bands-tagged", 0) == 1, "tagged-title: tagged path used", result["declined"])

        # A header artifact on one page of five is not a running band, so that
        # edge is declined and no other page is given it. The page number
        # beside it is judged on its own: it is still lifted from the tags
        # (not handed to the heuristic), and the second copy removes only the
        # footer artifacts, so the notice's text stays in the body.
        result, path = run_case(cases.TAGGED_ONE_PAGE_HEADER, workdir, "tagged-one-page")
        passed &= check(result["declined"].get("bands-tagged-partial", 0) == 1
                        and result["declined"].get("bands-tagged", 0) == 1
                        and not result["declined"].get("bands-heuristic"),
                        "tagged-one-page: a one-page header artifact declines its edge and not the footer",
                        result["declined"])
        count, footers = section_band_xml(path, "footer")
        passed &= check(count == result["pages"] and all("PAGE" in xml for xml in footers),
                        "tagged-one-page: the page number is still a footer field on every page",
                        ["PAGE" in xml for xml in footers])
        passed &= check(convert.compare_squash("Draft notice for page three") in squashed_docx(path, body_only=True),
                        "tagged-one-page: the declined header's words stay in the body")
        count, headers = section_band_texts(path, "header")
        passed &= check(all("Draft notice" not in text for i, text in enumerate(headers) if i != 2),
                        "tagged-one-page: no Word section carries a header its PDF page lacks", headers)
        passed &= check(result["textCheck"]["missing"] == 0, "tagged-one-page: nothing lost", result["textCheck"])

        # A header artifact inside page three's body, on one page of the
        # band's four: it stays in the body and in no header.
        result, path = run_case(cases.TAGGED_ARTIFACT_IN_BODY, workdir, "tagged-artifact-body")
        passed &= check(result["declined"].get("bands-tagged-partial", 0) == 1
                        and not result["declined"].get("bands-tagged"),
                        "tagged-artifact-body: a group on fewer pages than its band declines the tags",
                        result["declined"])
        body = squashed_docx(path, body_only=True)
        passed &= check(convert.compare_squash("Placeholder note set apart on page three") in body,
                        "tagged-artifact-body: the body text stays in the body")
        count, headers = section_band_texts(path, "header")
        passed &= check(all("Placeholder" not in text for text in headers),
                        "tagged-artifact-body: no Word section carries it as a header", headers)
        passed &= check(result["textCheck"]["missing"] == 0, "tagged-artifact-body: nothing lost",
                        result["textCheck"])

        # The thesis shape: the footer is on every page and the header starts
        # on page two, so Word's first section gets the footer and no header.
        result, path = run_case(cases.TAGGED_LATE_HEADER, workdir, "tagged-late-header")
        passed &= check(result["declined"].get("bands-tagged", 0) == 1, "tagged-late-header: tagged path used",
                        result["declined"])
        count, headers = section_band_texts(path, "header")
        passed &= check(count == result["pages"] and headers and headers[0] == ""
                        and all("Appendix material" in text for text in headers[1:]),
                        "tagged-late-header: no Word section carries a header its PDF page lacks",
                        (count, headers))
        count, footers = section_band_xml(path, "footer")
        passed &= check(count == result["pages"] and all("PAGE" in xml for xml in footers),
                        "tagged-late-header: every section keeps its page-number footer",
                        ["PAGE" in xml for xml in footers])
        passed &= check(result["textCheck"]["missing"] == 0 and not result["textCheck"]["fellBack"],
                        "tagged-late-header: nothing lost", result["textCheck"])

        # The same shape when the converter writes one section for the whole
        # document: Word cannot give the first page a footer and no header,
        # so the bands stay in the body and the declines say why.
        keep_run = convert.run_pdf2docx
        runs = []

        def one_section(source, part, *args, **kwargs):
            keep_run(source, part, *args, **kwargs)
            runs.append(source)
            if len(runs) == 1:
                merged = Document(part)
                body = merged.element.body
                for ppr in list(body.iter(convert.qn("w:pPr"))):
                    for inner in ppr.findall(convert.qn("w:sectPr")):
                        ppr.remove(inner)
                merged.save(part)

        convert.run_pdf2docx = one_section
        try:
            result, path = run_case(cases.TAGGED_LATE_HEADER, workdir, "tagged-late-one-section")
        except Exception as exc:  # noqa: BLE001 - the check reports it
            result = {"ok": False, "error": repr(exc)}
        finally:
            convert.run_pdf2docx = keep_run
        passed &= check(result.get("ok") and result["declined"].get("bands-pages-apart", 0) == 1
                        and result["textCheck"]["fellBack"] and result["textCheck"]["missing"] == 0
                        and not result["headerFooter"]["header"],
                        "tagged-late-one-section: header and footer on different pages are not lifted "
                        "without a section per page",
                        result.get("error") or (result["declined"], result["textCheck"], result["headerFooter"]))

        # A tagged band carries the paper it was read on, so a report with a
        # sideways page keeps its header; the heuristic's one set of numbers
        # still cannot, and its bands are dropped for a mixed page size and
        # not counted as lifted.
        result, path = run_case(cases.TAGGED_MIXED_SIZES, workdir, "tagged-mixed-sizes")
        passed &= check(result["declined"].get("bands-tagged", 0) == 1
                        and not result["declined"].get("bands-mixed-sizes")
                        and not result["declined"].get("bands-pages-apart")
                        and not result["declined"].get("bands-tagged-partial"),
                        "tagged-mixed-sizes: a tagged header survives a sideways page", result["declined"])
        passed &= check(result["textCheck"]["missing"] == 0 and not result["textCheck"]["fellBack"],
                        "tagged-mixed-sizes: nothing lost", result["textCheck"])
        result, path = run_untagged_case(cases.TAGGED_MIXED_SIZES, workdir, "untagged-mixed-sizes")
        passed &= check(result["declined"].get("bands-mixed-sizes", 0) == 1
                        and not result["declined"].get("bands-tagged")
                        and not result["declined"].get("bands-heuristic"),
                        "untagged-mixed-sizes: heuristic bands dropped for page size are counted as such",
                        result["declined"])
        passed &= check(result["textCheck"]["missing"] == 0, "untagged-mixed-sizes: nothing lost",
                        result["textCheck"])

        # The second stripped copy, the one the converter reads, can fail or
        # disagree with the first. Either way the export goes on with the
        # bands in the body, and says so.
        for name, second in (("none", lambda doc, made: None),
                             ("counts", lambda doc, made: (made[0], {**made[1], 0: made[1].get(0, 0) + 1}))):
            calls = []

            def flaky_copy(doc, second=second):
                made = keep_copy(doc)
                calls.append(made)
                if len(calls) == 1 or made is None:
                    return made
                answer = second(doc, made)
                if answer is None:
                    made[0].close()
                return answer

            convert.tagged.stripped_copy = flaky_copy
            try:
                result, path = run_case(cases.FIRST_BODY_LINE, workdir, "tagged-restrip-" + name)
            except Exception as exc:  # noqa: BLE001 - the check reports it
                result = {"ok": False, "error": repr(exc)}
            finally:
                convert.tagged.stripped_copy = keep_copy
            passed &= check(result.get("ok") and result["declined"].get("bands-tagged-restrip", 0) == 1
                            and not result["declined"].get("bands-tagged")
                            and not result["headerFooter"]["header"] and not result["headerFooter"]["footer"]
                            and result["textCheck"]["missing"] == 0
                            and any("second time" in w for w in result["warnings"]),
                            "tagged-restrip-%s: a second copy that fails keeps the bands in the body" % name,
                            result.get("error") or (result["declined"], result["headerFooter"]))

        # A tag tree that cannot be read at the contents gate goes to detect_toc.
        keep_types = convert.tagged.struct_types

        def broken_types(doc):
            raise RuntimeError("forced: the tag tree cannot be read")

        convert.tagged.struct_types = broken_types
        try:
            result, path = run_case(cases.DOT_LEADERS_NOT_TOC, workdir, "toc-gate-raises")
        except Exception as exc:  # noqa: BLE001 - the check reports it
            result = {"ok": False, "error": repr(exc)}
        finally:
            convert.tagged.struct_types = keep_types
        passed &= check(result.get("ok") and not result["declined"].get("toc-untagged"),
                        "toc-gate-raises: an unreadable tag tree falls back to detect_toc",
                        result.get("error") or result["declined"])

        # Every existing case: untagged-equivalent behaviour. cases.TABLE has
        # no header or footer at all, so the heuristic must still find none.
        result, path = run_case(cases.TABLE, workdir, "table-tagged")
        passed &= check(not result["headerFooter"]["header"] and not result["headerFooter"]["footer"], "table: tagged PDF with no band artifacts lifts nothing", result["headerFooter"])

        # 22. The same documents with their tags removed go through the
        #     heuristic, with the result they had before the tagged path
        #     existed. TAGGED_SKIP_PAGES has too few band pages for the
        #     heuristic to lift anything, so it counts neither path;
        #     FIRST_BODY_LINE is lifted by the heuristic and counts it.
        for source, name, expected, heuristic in (
            (cases.TAGGED_SKIP_PAGES, "untagged-skip", {"header": False, "footer": False, "pagesWithBands": 0}, 0),
            (cases.FIRST_BODY_LINE, "untagged-firstline", {"header": True, "footer": True, "pagesWithBands": 4}, 1),
        ):
            result, path = run_untagged_case(source, workdir, name)
            passed &= check(result["declined"].get("bands-tagged", 0) == 0
                            and result["declined"].get("bands-heuristic", 0) == heuristic,
                            "%s: the heuristic path ran" % name, result["declined"])
            passed &= check(result["headerFooter"] == expected, "%s: the same bands as before" % name,
                            result["headerFooter"])
            passed &= check(result["textCheck"]["missing"] == 0, "%s: nothing lost" % name, result["textCheck"])

        # 23. A contents page is rebuilt only where the tag tree has a TOC.
        result, path = run_case(cases.DOT_LEADERS_NOT_TOC, workdir, "price-list")
        passed &= check(result["tocEntries"] == 0 and result["declined"].get("toc-untagged", 0) == 1,
                        "toc: a tagged dot-leader list with no TOC element is left alone",
                        (result["tocEntries"], result["declined"]))
        passed &= check(result["textCheck"]["missing"] == 0, "toc: the price list lost nothing", result["textCheck"])
        result, path = run_case(cases.OUTLINE_TOC, workdir, "outline")
        # The entries are found and rebuilt. On a page this sparse the converter
        # merges them into one paragraph, which used to leave the rebuild
        # declining with a warning that it could not place them.
        passed &= check(result["declined"].get("toc-untagged", 0) == 0 and result["tocEntries"] > 0
                        and any("table-of-contents entries were rebuilt" in w for w in result["warnings"]),
                        "toc: a tagged #outline still has its contents page found and rebuilt",
                        (result["tocEntries"], result["declined"], result["warnings"]))
        passed &= check(result["textCheck"]["missing"] == 0, "toc: the rebuilt outline lost nothing", result["textCheck"])
        result, path = run_untagged_case(cases.DOT_LEADERS_NOT_TOC, workdir, "price-list-untagged")
        passed &= check(result["declined"].get("toc-untagged", 0) == 0,
                        "toc: an untagged PDF still goes through detect_toc", result["declined"])
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return passed


def run_switch_checks():
    """The two dev-only switches, and that the server never passes them."""
    import cases

    passed = True
    here = os.path.dirname(os.path.abspath(__file__))
    with open(os.path.join(here, "..", "export.mjs"), encoding="utf-8") as handle:
        export_source = handle.read()
    passed &= check("--raw" not in export_source and "--skip" not in export_source,
                    "switches: export.mjs never passes --raw or --skip")
    passed &= check(convert.main(["convert.py", "--skip=nope", "a.pdf", "b.docx", "c.json"]) == 2,
                    "switches: an unknown pass name is a usage error")

    if shutil.which("typst") is None:
        print("SKIP  the switch checks need the typst CLI on PATH")
        return passed
    workdir = tempfile.mkdtemp(prefix="btct-switch-test-")
    try:
        pdf = os.path.join(workdir, "report.pdf")
        # Untagged: REPORT's contents page is hand-made, and a tagged PDF
        # without a TOC element is (rightly) not searched for one.
        compile_untagged(cases.REPORT, pdf)
        full = convert.convert(pdf, os.path.join(workdir, "full.docx"))
        # A decline is a pass refusing one shape on its own, which a plain run
        # may do (this case's rhythm walk declines an uneven anchor); each one
        # is named after the pass that declined.
        passed &= check(full["mode"] == "full" and full["skipped"] == []
                        and all(any(name.startswith(p + "-") for p in convert.SKIPPABLE)
                                for name in full["declined"]),
                        "switches: a plain run says full, skipped nothing, and names each decline by its pass",
                        (full["mode"], full["skipped"], full["declined"]))
        # The case is too sparse for the rebuild to engage (the converter merges
        # its entries into one paragraph), so a contents page that was found is
        # told apart by the warning that says the rebuild could not place it.
        toc_words = [w for w in full["warnings"] if "table of contents" in w or "table-of-contents" in w]
        passed &= check(full["headerFooter"]["header"] and toc_words,
                        "switches: the case has a header and a contents page to switch off",
                        (full["headerFooter"], full["warnings"]))

        skipped = convert.convert(pdf, os.path.join(workdir, "skip.docx"), skip=frozenset({"bands", "toc"}))
        passed &= check(not skipped["headerFooter"]["header"] and not skipped["headerFooter"]["footer"],
                        "switches: --skip=bands keeps the bands in the body", skipped["headerFooter"])
        passed &= check(skipped["tocEntries"] == 0
                        and not [w for w in skipped["warnings"] if "table of contents" in w or "table-of-contents" in w],
                        "switches: --skip=toc looks for no contents page", skipped["warnings"])
        passed &= check(skipped["skipped"] == ["bands", "toc"], "switches: the result names what was skipped",
                        skipped["skipped"])
        passed &= check(skipped["textCheck"]["missing"] == 0, "switches: skipping loses no text",
                        skipped["textCheck"])

        raw = convert.convert(pdf, os.path.join(workdir, "raw.docx"), raw=True)
        passed &= check(raw["mode"] == "raw" and raw["tocEntries"] == 0
                        and not raw["headerFooter"]["header"],
                        "switches: --raw rebuilds nothing", (raw["mode"], raw["tocEntries"], raw["headerFooter"]))
        passed &= check(raw["textCheck"]["missing"] == 0 and not raw["textCheck"]["fellBack"],
                        "switches: --raw still runs the text check", raw["textCheck"])
        passed &= check(sorted(raw["skipped"]) == sorted(convert.SKIPPABLE),
                        "switches: --raw is every pass skipped", raw["skipped"])
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return passed


def check_corpus_document(name, source, workdir):
    """(passed, why) for one corpus document, without Word."""
    import corpus

    pdf = os.path.join(workdir, name + ".pdf")
    out = os.path.join(workdir, name + ".docx")
    try:
        corpus.compile_document(source, pdf)
    except subprocess.CalledProcessError as err:
        return False, "typst failed: " + (err.stderr or "")[:200]
    try:
        result = convert.convert(pdf, out, font_dirs=corpus.font_dirs())
    except convert.ConvertError as err:
        return False, "refused: " + str(err)[:200]
    text = result["textCheck"]
    if text["fellBack"] or text["missing"]:
        return False, "text check %s" % text
    try:
        document = Document(out)
    except Exception as err:  # noqa: BLE001 - any failure to open is the finding
        return False, "python-docx cannot open it: %s" % err
    unsafe = convert.unsafe_docx_targets(out)
    if unsafe:
        return False, "unsafe link targets %s" % unsafe[:2]
    sections = len(document.sections)
    if not 1 <= sections <= result["pages"] + 1:
        return False, "%d sections for %d PDF pages" % (sections, result["pages"])
    return True, "%d pages, %d sections, %d warnings" % (result["pages"], sections, len(result["warnings"]))


def run_corpus_checks():
    import corpus

    if shutil.which("typst") is None:
        print("SKIP  the corpus needs the typst CLI on PATH")
        return True
    passed = True
    for name, _category, source in corpus.documents():
        passed &= check(os.path.isfile(os.path.join(source, "main.typ")), "corpus: %s has a main.typ" % name)
    workdir = tempfile.mkdtemp(prefix="btct-corpus-test-")
    started = time.monotonic()
    try:
        for name, _category, source in corpus.documents():
            ok, why = check_corpus_document(name, source, workdir)
            expected = corpus.EXPECTED_FAILURES.get(name)
            if expected is None:
                passed &= check(ok, "corpus: %s converts cleanly" % name, why)
            else:
                passed &= check(not ok, "corpus: %s still fails as expected (%s)" % (name, expected), why)
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    print("      corpus took %.1f s" % (time.monotonic() - started))
    return passed


def strip_rtl_spaces(src, dst):
    """Copy a Word file with the spaces removed from every run that holds RTL text."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                tree = ElementTree.fromstring(data)
                for node in tree.iter(W + "t"):
                    if node.text and convert.is_rtl(node.text):
                        node.text = node.text.replace(" ", "")
                data = ElementTree.tostring(tree, xml_declaration=True, encoding="UTF-8")
            zout.writestr(item, data)


def faithful_docx(lines, path):
    """A Word file holding exactly the PDF's reading, one paragraph per piece.

    What a perfect conversion would contain. The text check has to pass it,
    and it is the file the negative check strips, so that check does not
    depend on what pdf2docx does with right-to-left text.
    """
    document = Document()
    for line in lines:
        for piece in line["pieces"] or [line["text"]]:
            document.add_paragraph(piece)
    document.save(path)


def run_rtl_check_checks():
    import cases
    import corpus

    passed = True
    passed &= check(convert.is_rtl("أجرى فريقنا تقييمًا"), "rtl: an Arabic line is right-to-left")
    passed &= check(convert.is_rtl("במהלך הבדיקה נמצאה חולשת SQL injection"),
                    "rtl: a Hebrew line with two English words is right-to-left")
    passed &= check(not convert.is_rtl('The policy, titled "מדיניות אבטחת מידע", was revised in 2024.'),
                    "rtl: an English line quoting Hebrew stays left-to-right")
    passed &= check(not convert.is_rtl("1. 2. 3."), "rtl: digits and punctuation are not right-to-left")
    passed &= check(convert.compare_words("  أجرى   فريقنا تقييمًا ") == "أجرى فريقنا تقييمًا",
                    "rtl: words are squashed and joined by one space", convert.compare_words("أجرى   فريقنا"))
    # The PDF's spans arrive left to right on the page, each one already in
    # reading order inside. A right-to-left line reads them the other way
    # round, except for a run of English inside it, which keeps its order.
    visual = ["חולשת", " ", "SQL", " ", "injection", " ", "נמצאה", " ", "במהלך"]
    passed &= check(convert.bidi_order(visual, True)
                    == ["במהלך", " ", "נמצאה", " ", "SQL", " ", "injection", " ", "חולשת"],
                    "rtl: a right-to-left line reads its spans right to left, English kept in order",
                    convert.bidi_order(visual, True))
    quoted = ["titled “", "מידע", " ", "אבטחת", " ", "מדיניות", "”, was"]
    passed &= check(convert.bidi_order(quoted, False)
                    == ["titled “", "מדיניות", " ", "אבטחת", " ", "מידע", "”, was"],
                    "rtl: an English line reverses only the Hebrew it quotes",
                    convert.bidi_order(quoted, False))
    passed &= check(convert.bidi_order(["one", " ", "two"], False) == ["one", " ", "two"],
                    "rtl: a left-to-right line keeps its spans as they are")
    passed &= check(convert.bidi_order(["וכן", " ,", "ההזמנות"], True) == ["ההזמנות", ", ", "וכן"],
                    "rtl: punctuation read right to left is turned round with its line",
                    convert.bidi_order(["וכן", " ,", "ההזמנות"], True))
    for visual, logical, label in (
            (["א", " ", "2024", " ", "ב"], ["ב", " ", "2024", " ", "א"], "a year"),
            (["פגיע", " ", "10.0.0.5", " ", "בכתובת"], ["בכתובת", " ", "10.0.0.5", " ", "פגיע"], "an address"),
            (["חשופה", " ", "2.4.1", " ", "גרסה"], ["גרסה", " ", "2.4.1", " ", "חשופה"], "a version"),
            (["א", " ", "Apache", " ", "2.4", " ", "ב"], ["ב", " ", "Apache", " ", "2.4", " ", "א"],
             "a version after an English word")):
        passed &= check(convert.bidi_order(visual, True) == logical,
                        "rtl: %s in a Hebrew line keeps its digits in order" % label,
                        convert.bidi_order(visual, True))
    passed &= check(not convert.reorder_is_safe(["1.4.2", "גרסה"], "גרסה 2.4.1"),
                    "rtl: a reorder that would turn a number round is refused")
    passed &= check(not convert.reorder_is_safe(["שלבי", "דו"], "דו-שלבי"),
                    "rtl: a reorder that would add or drop a dash is refused")
    passed &= check(convert.reorder_is_safe(["חשופה", " ", "2.4.1", " ", "גרסה"], "גרסה 2.4.1 חשופה"),
                    "rtl: a reorder that moves whole words and numbers is allowed")
    # The last rung restructures nothing: runs pdf2docx wrote in page order
    # stay in page order there, and only the rungs above may reorder them.
    for place in (False, True):
        document = Document()
        paragraph = document.add_paragraph()
        for word in ("מידע", "אבטחת", "מדיניות"):
            paragraph.add_run(word)
        look = {"panels": [], "swatches": [], "face": None, "aligned": [], "pictures": [], "pairs": set(),
                "broken": [], "rules": [], "boxes": [], "height": PAGE_HEIGHT, "lineBreaks": "pdf",
                "plan": None, "skip": frozenset(), "declined": convert.Counter(),
                "rtl": ["מדיניות אבטחת מידע"]}
        convert.decorate_document(document, look, {}, time.monotonic(), None, 0.0, place=place)
        text = convert.paragraph_text(paragraph._p)
        if place:
            passed &= check(text == "מדיניות אבטחת מידע", "rtl: the first rung writes page-order Hebrew in reading order",
                            text)
        else:
            passed &= check(text == "מידעאבטחתמדיניות" and look["declined"]["rtl-reorder-plain-rung"] == 1,
                            "rtl: the last rung does not reorder runs", (text, dict(look["declined"])))
    # The rungs above the last mark a paragraph they rewrote as right-to-left
    # and turn its indent and alignment round with the mark. The last rung
    # moves nothing: the paragraph keeps the properties the converter wrote.
    # (Measured through Word: the mark without the mirroring put rtl-arabic
    # 0.0022 below raw, so the last rung does not mark at all.)
    for place in (False, True):
        document = Document()
        paragraph = document.add_paragraph()
        for word in ("מדיניות", "אבטחת", "מידע"):
            paragraph.add_run(word)
        paragraph.paragraph_format.left_indent = Pt(20)
        paragraph.alignment = convert.WD_ALIGN_PARAGRAPH.LEFT
        look = {"panels": [], "swatches": [], "face": None, "aligned": [], "pictures": [], "pairs": set(),
                "broken": [], "rules": [], "boxes": [], "height": PAGE_HEIGHT, "lineBreaks": "pdf",
                "plan": None, "skip": frozenset(), "declined": convert.Counter(),
                "rtl": ["מדיניות אבטחת מידע"]}
        convert.decorate_document(document, look, {}, time.monotonic(), None, 0.0, place=place)
        ppr = paragraph._p.find(W + "pPr")
        indent, align = ppr.find(W + "ind"), ppr.find(W + "jc")
        state = (ppr.find(W + "bidi") is not None, dict(indent.attrib), align.get(W + "val"))
        if place:
            passed &= check(state[0] and indent.get(W + "right") == "400" and state[2] == "right",
                            "rtl: the first rung turns a marked paragraph's indent and alignment round", state)
        else:
            runs_marked = all(run._r.find(W + "rPr/" + W + "rtl") is not None for run in paragraph.runs)
            passed &= check(not state[0] and indent.get(W + "left") == "400" and indent.get(W + "right") is None
                            and state[2] == "left" and runs_marked,
                            "rtl: the last rung marks the runs and leaves the paragraph as written", state)
    # A number that follows right-to-left words in a left-to-right line is
    # drawn to their left, so the page cannot say which side of them it was
    # typed on. Such a line is ambiguous; a number in a right-to-left line,
    # or away from the Hebrew, is not.
    for text, want in (("Server שרת ראשי 12 and", True), ('The policy "מדיניות אבטחה 2024" was reviewed', True),
                       ("the host (שרת 5) and", True), ("the host (5 שרת) and", True),
                       ("titled \"מדיניות\", was revised in 2024", False),
                       ("גרסה 2.4.1 חשופה", False), ("version 2.4.1 only", False),
                       ("كلمة ١٢ فقط", False), ("the word كلمة ١٢ here", True)):
        passed &= check(convert.ambiguous_number(text) == want,
                        "rtl: %r is %sambiguous" % (text, "" if want else "not "))
    english_inside = ["א", " ", "SQL", "; ", "TLS", " ", "ב"]
    passed &= check(convert.bidi_order(english_inside, True) == ["ב", " ", "SQL", "; ", "TLS", " ", "א"],
                    "rtl: punctuation inside a run of English keeps its order",
                    convert.bidi_order(english_inside, True))
    # The squashed test alone passes a right-to-left line whose spaces are
    # gone; the word test is what catches it.
    lines = [{"text": "أجرى فريقنا تقييمًا", "pieces": ["أجرى فريقنا تقييمًا"]}]
    passed &= check(convert.missing_from_docx(lines, "أجرى فريقنا تقييمًا") == [],
                    "rtl: an Arabic line with its spaces is found")
    passed &= check(convert.missing_from_docx(lines, "أجرىفريقناتقييمًا") == ["أجرى فريقنا تقييمًا"],
                    "rtl: an Arabic line run together is reported missing")
    english = [{"text": "run together here", "pieces": ["run together here"]}]
    passed &= check(convert.missing_from_docx(english, "runtogetherhere") == [],
                    "rtl: a left-to-right line is still compared squashed")

    passed &= check(convert.respace(["أجرىفريقنا", "تقييمًا"], "أجرى فريقنا تقييمًا") == ["أجرى فريقنا ", "تقييمًا"],
                    "respace: a space goes back inside a run and at a run boundary",
                    convert.respace(["أجرىفريقنا", "تقييمًا"], "أجرى فريقنا تقييمًا"))
    passed &= check(convert.respace(["أجرى فريقنا"], "أجرى فريقنا") == ["أجرى فريقنا"],
                    "respace: a run that already has its spaces is unchanged")
    passed &= check(convert.respace(["شيء آخر"], "أجرى فريقنا") is None,
                    "respace: text that does not match the line is left alone")
    passed &= check(convert.respace(["SQLinjection", "חולשת"], "חולשת SQL injection") is None,
                    "respace: a different word order is not guessed at")
    passed &= check(convert.respace(["أجج رى"], "أججرى") == ["أججرى"],
                    "respace: a space the PDF does not have inside a word is taken out",
                    convert.respace(["أجج رى"], "أججرى"))
    passed &= check(convert.respace([" מידע", "אבטחת "], "מידע אבטחת") == [" מידע ", "אבטחת "],
                    "respace: whitespace before the first word and after the last is kept",
                    convert.respace([" מידע", "אבטחת "], "מידע אבטחת"))
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("מידע")
    run.bold = False
    run.font.size = Pt(11)
    paragraph.paragraph_format.left_indent = Pt(20)
    paragraph.alignment = convert.WD_ALIGN_PARAGRAPH.LEFT
    convert.mark_run_rtl(run._r)
    convert.mark_paragraph_rtl(paragraph._p)
    run_tags = [child.tag.replace(W, "") for child in run._r.find(W + "rPr")]
    passed &= check(run_tags == ["b", "bCs", "sz", "szCs", "rtl"],
                    "rtl: a marked run carries its complex-script size and weight, in schema order", run_tags)
    ppr = paragraph._p.find(W + "pPr")
    para_tags = [child.tag.replace(W, "") for child in ppr]
    indent, align = ppr.find(W + "ind"), ppr.find(W + "jc")
    passed &= check(para_tags.index("bidi") < para_tags.index("ind") < para_tags.index("jc")
                    and indent.get(W + "right") == "400" and indent.get(W + "left") is None
                    and align.get(W + "val") == "right",
                    "rtl: a marked paragraph has its indent and alignment turned round with it",
                    (para_tags, dict(indent.attrib), dict(align.attrib)))
    passed &= check(convert.direction_chunks("במהלך נמצאה SQL injection חולשת", True)
                    == [("R", "במהלך נמצאה "), ("L", "SQL injection"), ("R", " חולשת")],
                    "rtl: a Hebrew line splits into its Hebrew and English stretches",
                    convert.direction_chunks("במהלך נמצאה SQL injection חולשת", True))
    passed &= check(convert.direction_chunks("titled “מדיניות אבטחת מידע”, was", False)
                    == [("L", "titled “"), ("R", "מדיניות אבטחת מידע"), ("L", "”, was")],
                    "rtl: the quotes around Hebrew in an English line stay English",
                    convert.direction_chunks("titled “מדיניות אבטחת מידע”, was", False))

    if shutil.which("typst") is None:
        print("SKIP  the RTL document checks need the typst CLI on PATH")
        return passed
    workdir = tempfile.mkdtemp(prefix="btct-rtl-test-")
    try:
        for name, sample in (("rtl-arabic", "أججرى ففريققننا"), ("rtl-mixed", "מדיניות אבטחת מידע")):
            pdf = os.path.join(workdir, name + ".pdf")
            corpus.compile_document(os.path.join(corpus.CORPUS, name), pdf)
            doc = pymupdf.open(pdf)
            lines = convert.body_lines(doc, [])
            doc.close()
            text = "\n".join(line["text"] for line in lines)
            passed &= check(sample in text, "rtl: %s is read in reading order with its own spaces" % name,
                            [line["text"] for line in lines if convert.has_rtl(line["text"])][:2])
            # A Word file that holds the PDF's text exactly passes, and the
            # same file with the spaces taken out of its RTL runs does not.
            faithful = os.path.join(workdir, name + "-faithful.docx")
            faithful_docx(lines, faithful)
            clean = convert.missing_from_docx(lines, convert.docx_text(faithful, body_only=False))
            passed &= check(clean == [], "rtl: %s read back from a faithful Word file is complete" % name, clean[:2])
            stripped = os.path.join(workdir, name + "-stripped.docx")
            strip_rtl_spaces(faithful, stripped)
            missing = convert.missing_from_docx(lines, convert.docx_text(stripped, body_only=False))
            passed &= check(len(missing) > 0, "rtl: %s with its RTL spaces stripped is reported missing" % name,
                            len(missing))

        # Numbers in Hebrew lines survive the repair, and a Word file with one
        # of them turned round is caught.
        source = os.path.join(workdir, "rtl-numbers")
        os.makedirs(source)
        with open(os.path.join(source, "main.typ"), "w", encoding="utf-8") as handle:
            handle.write(cases.RTL_NUMBERS)
        pdf = os.path.join(workdir, "rtl-numbers.pdf")
        corpus.compile_document(source, pdf)
        converted = os.path.join(workdir, "numbers.docx")
        result = convert.convert(pdf, converted, font_dirs=corpus.font_dirs())
        text = convert.docx_text(converted, body_only=False)
        passed &= check(all(number in text for number in ("2.4.1", "10.0.0.5", "2024"))
                        and not any(number in text for number in ("1.4.2", "5.0.0.01", "4202")),
                        "rtl: numbers in Hebrew lines keep their digits in order", text)
        passed &= check(result["textCheck"]["missing"] == 0 and not result["textCheck"]["fellBack"],
                        "rtl: the numbers document is complete at the top rung", result["textCheck"])
        passed &= check("גרסה 2.4.1 חשופה" in text and "השרת בכתובת 10.0.0.5 פגיע" in text,
                        "rtl: Hebrew reads in its own word order, words apart, numbers in place")
        with zipfile.ZipFile(converted) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
        passed &= check("<w:bidi" in xml and "<w:rtl" in xml,
                        "rtl: at the top rung Hebrew paragraphs and runs are marked right-to-left")
        doc = pymupdf.open(pdf)
        lines = convert.body_lines(doc, [])
        doc.close()
        turned = os.path.join(workdir, "numbers-turned.docx")
        with zipfile.ZipFile(converted) as zin, zipfile.ZipFile(turned, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "word/document.xml":
                    data = data.replace(b"10.0.0.5", b"5.0.0.01")
                zout.writestr(item, data)
        missing = convert.missing_from_docx(lines, convert.docx_text(turned, body_only=False))
        passed &= check(any("10.0.0.5" in line for line in missing),
                        "rtl: an address turned round in the Word file is reported missing", missing)

        # Numbers after Hebrew in English lines: the page cannot say which
        # side of the Hebrew they were typed on, so the repair must not write
        # its guess into Word and certify it.
        source = os.path.join(workdir, "rtl-ambiguous")
        os.makedirs(source)
        with open(os.path.join(source, "main.typ"), "w", encoding="utf-8") as handle:
            handle.write(cases.RTL_AMBIGUOUS_NUMBERS)
        pdf = os.path.join(workdir, "rtl-ambiguous.pdf")
        corpus.compile_document(source, pdf)
        converted = os.path.join(workdir, "ambiguous.docx")
        result = convert.convert(pdf, converted, font_dirs=corpus.font_dirs())
        text = convert.docx_text(converted, body_only=False)
        guesses = ("Server 12 שרת ראשי", "“2024 מדיניות אבטחה”", "( 5שרת)")
        passed &= check(not any(guess in text for guess in guesses),
                        "rtl-ambiguous: the Word file does not carry the reordered guess", text)
        passed &= check(result["declined"].get("rtl-ambiguous-number", 0) == 3,
                        "rtl-ambiguous: each ambiguous line is declined by name", result["declined"])
        named = [w for w in result["warnings"] if "could not be found in the Word file" in w]
        passed &= check(result["textCheck"]["missing"] == 0 or named,
                        "rtl-ambiguous: the result is complete or names what it lost",
                        (result["textCheck"], named))
        doc = pymupdf.open(pdf)
        default = [row["text"] for row in convert.page_rows(doc[0])]
        lines = [line["text"] for line in convert.body_lines(doc, [])]
        doc.close()
        passed &= check(lines == default, "rtl-ambiguous: the check reads those rows as the page shows them",
                        (lines, default))
        # The same lines with the guess written in are not certified either.
        guessed = os.path.join(workdir, "ambiguous-guessed.docx")
        document = Document()
        for guess in ("Server 12 שרת ראשי and 7שרת משני were tested.",
                      "The policy “2024 מדיניות אבטחה” was reviewed.",
                      "We reached the host ( 5שרת) and stopped there."):
            document.add_paragraph(guess)
        document.save(guessed)
        doc = pymupdf.open(pdf)
        missing = convert.missing_from_docx(convert.body_lines(doc, []), convert.docx_text(guessed, body_only=False))
        doc.close()
        # (The third line's one Hebrew word reads the same either way.)
        passed &= check(sorted(line.split()[1] for line in missing) == ["12", "policy"],
                        "rtl-ambiguous: a Word file holding the guess is reported missing", missing)

        # Arabic-Indic numbers read the way they are drawn, and a Word file
        # that says 21 where the PDF says 12 is never certified.
        source = os.path.join(workdir, "rtl-arabic-digits")
        os.makedirs(source)
        with open(os.path.join(source, "main.typ"), "w", encoding="utf-8") as handle:
            handle.write(cases.RTL_ARABIC_DIGITS)
        pdf = os.path.join(workdir, "rtl-arabic-digits.pdf")
        corpus.compile_document(source, pdf)
        doc = pymupdf.open(pdf)
        reading = "\n".join(line["text"] for line in convert.body_lines(doc, []))
        doc.close()
        passed &= check("١٢" in reading and "٥٠" in reading and "٢١" not in reading and "٠٥" not in reading
                        and "34" in reading,
                        "rtl-digits: Arabic-Indic numbers are read in their own order", reading)
        converted = os.path.join(workdir, "arabic-digits.docx")
        result = convert.convert(pdf, converted, font_dirs=corpus.font_dirs())
        text = convert.docx_text(converted, body_only=False)
        certified_wrong = [n for n in ("٢١", "٠٥") if n in text] and result["textCheck"]["missing"] == 0
        passed &= check(not certified_wrong, "rtl-digits: a turned number is never certified",
                        (result["textCheck"], text))
        named = [w for w in result["warnings"] if "could not be found in the Word file" in w]
        passed &= check(result["textCheck"]["missing"] == 0 or named,
                        "rtl-digits: the result is complete or names what it lost", (result["textCheck"], named))

        # declined describes the file that was kept: a document that falls
        # back reports its last rung's declines only, and --raw none at all.
        mixed_pdf = os.path.join(workdir, "rtl-mixed.pdf")
        rungs = []
        original_decorate = convert.decorate_document

        def tagged(document, look, *args, **kwargs):
            rungs.append(kwargs.get("place", args[4] if len(args) > 4 else True))
            look["declined"]["rung-%d" % len(rungs)] += 1
            return original_decorate(document, look, *args, **kwargs)

        convert.decorate_document = tagged
        try:
            result = convert.convert(mixed_pdf, os.path.join(workdir, "mixed-rungs.docx"),
                                     font_dirs=corpus.font_dirs())
        finally:
            convert.decorate_document = original_decorate
        tags = sorted(k for k in result["declined"] if k.startswith("rung-"))
        passed &= check(result["textCheck"]["fellBack"] and len(rungs) >= 2 and tags == ["rung-%d" % len(rungs)],
                        "declined: a document that falls back reports the kept rung only",
                        (rungs, result["declined"]))
        passed &= check(result["declined"].get("align-rtl", 0) >= 1,
                        "declined: the kept rung keeps the align-rtl it was built with", result["declined"])
        raw = convert.convert(mixed_pdf, os.path.join(workdir, "mixed-raw.docx"), font_dirs=corpus.font_dirs(),
                              raw=True)
        passed &= check(raw["declined"] == {}, "declined: --raw reports no pass declines", raw["declined"])

        # A run that holds English as well as Arabic is never marked w:rtl,
        # which would lay the English out right to left too.
        source = os.path.join(workdir, "rtl-in-code")
        os.makedirs(source)
        with open(os.path.join(source, "main.typ"), "w", encoding="utf-8") as handle:
            handle.write(cases.RTL_IN_CODE)
        pdf = os.path.join(workdir, "rtl-in-code.pdf")
        corpus.compile_document(source, pdf)
        converted = os.path.join(workdir, "in-code.docx")
        result = convert.convert(pdf, converted, font_dirs=corpus.font_dirs())
        with zipfile.ZipFile(converted) as zin:
            tree = ElementTree.fromstring(zin.read("word/document.xml"))
        mixed_marked = ["".join(t.text or "" for t in run.iter(W + "t")) for run in tree.iter(W + "r")
                        if run.find(W + "rPr/" + W + "rtl") is not None
                        and convert.strong_direction("".join(t.text or "" for t in run.iter(W + "t")))[1]]
        passed &= check(mixed_marked == [], "rtl-in-code: no run with English in it is marked w:rtl", mixed_marked)
        # The caption's Arabic words are one run of their own. Matched by its
        # letters rather than its words: the PDF's text layer repeats some of
        # this font's letters, and the Word file carries them as found.
        caption_marked = [text for text in ("".join(t.text or "" for t in run.iter(W + "t"))
                                            for run in tree.iter(W + "r")
                                            if run.find(W + "rPr/" + W + "rtl") is not None)
                          if convert.strong_direction(text)[0] >= 10 and not convert.strong_direction(text)[1]]
        passed &= check(caption_marked != [], "rtl-in-code: the caption's pure-Arabic run is still marked w:rtl",
                        caption_marked)
        passed &= check(result["declined"].get("rtl-mixed-run", 0) > 0,
                        "rtl-in-code: the repair declines to mark a mixed run", result["declined"])
        passed &= check(result["textCheck"]["missing"] == 0, "rtl-in-code: no line lost", result["textCheck"])

        # The repair, on real conversions.
        pdf = os.path.join(workdir, "rtl-in-shapes.pdf")
        corpus.compile_document(os.path.join(corpus.CORPUS, "rtl-in-shapes"), pdf)
        converted = os.path.join(workdir, "shapes.docx")
        result = convert.convert(pdf, converted, font_dirs=corpus.font_dirs())
        passed &= check(result["textCheck"]["missing"] == 0 and not result["textCheck"]["fellBack"],
                        "rtl: the Arabic caption keeps every word apart at the top rung", result["textCheck"])
        with zipfile.ZipFile(converted) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
        passed &= check("<w:rtl" in xml, "rtl: the caption's Arabic runs are marked right-to-left")


        pdf = os.path.join(workdir, "rtl-arabic.pdf")
        doc = pymupdf.open(pdf)
        lines = convert.body_lines(doc, [])
        doc.close()
        converted = os.path.join(workdir, "ar.docx")
        result = convert.convert(pdf, converted, font_dirs=corpus.font_dirs())
        with zipfile.ZipFile(converted) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
        # rtl-arabic used to end on the last rung, which marks runs but no
        # paragraph, because pdf2docx dropped letters from its table's cells.
        # With those letters given back (`straddling_lines_keep_their_chars`)
        # it holds the top rung, where the repair sets paragraphs too.
        passed &= check(not result["textCheck"]["fellBack"] and result["textCheck"]["missing"] == 0,
                        "rtl: rtl-arabic holds the top rung with no line lost", result["textCheck"])
        passed &= check("<w:rtl" in xml and "<w:bidi" in xml,
                        "rtl: on the top rung Arabic runs are marked and their paragraphs set right-to-left")
        missing = convert.missing_from_docx(lines, convert.docx_text(converted, body_only=False))
        body = [line["text"] for line in lines if line["text"].startswith("أججرى")]
        passed &= check(body and body[0] not in missing, "rtl: an Arabic body line comes back word for word",
                        len(missing))
        # Stripping the spaces the repair put back loses lines that were found.
        stripped = os.path.join(workdir, "ar-stripped.docx")
        strip_rtl_spaces(converted, stripped)
        lost = set(convert.missing_from_docx(lines, convert.docx_text(stripped, body_only=False))) - set(missing)
        passed &= check(lost, "rtl: the converted Arabic with its spaces stripped is reported missing", len(lost))
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return passed


def pdf_band_lines(pdf_path, needle):
    """Per PDF page, the text of the line that holds `needle`, or None."""
    with pymupdf.open(pdf_path) as doc:
        return [next((" ".join(line.split()) for line in page.get_text().splitlines() if needle in line), None)
                for page in doc]


def run_band_variant_documents():
    """A header and footer whose text changes per page, a sideways page, a mirrored number and
    a header on one page, on real typst output."""
    import cases

    if shutil.which("typst") is None:
        print("SKIP  the band variant cases need the typst CLI on PATH")
        return True
    passed = True
    workdir = tempfile.mkdtemp(prefix="btct-band-variants-")
    try:
        # 1. The chapter's title in the header: one band per chapter, each
        #    given to the sections of its own pages.
        result, path = run_case(cases.TAGGED_CHAPTER_HEADER, workdir, "chapter-header")
        declined = result["declined"]
        passed &= check(declined.get("bands-tagged", 0) == 1 and not declined.get("bands-tagged-varies")
                        and not declined.get("bands-tagged-partial") and not declined.get("bands-pages-apart"),
                        "chapter header: a header that changes per chapter is lifted", declined)
        passed &= check(result["textCheck"]["missing"] == 0 and not result["textCheck"]["fellBack"],
                        "chapter header: nothing lost, no fallback", result["textCheck"])
        count, headers = section_band_texts(path, "header")
        wanted = pdf_band_lines(os.path.join(workdir, "chapter-header.pdf"), "Status Report |")
        passed &= check(count == result["pages"] and len(set(w for w in wanted if w)) >= 3,
                        "chapter header: one section per page and several distinct chapter titles",
                        (count, result["pages"], sorted(set(w for w in wanted if w))))
        passed &= check(all(convert.compare_squash(want or "") == convert.compare_squash(text)
                            for want, text in zip(wanted, headers)),
                        "chapter header: every section carries the title its own PDF page has",
                        [(w, t) for w, t in zip(wanted, headers) if convert.compare_squash(w or "") != convert.compare_squash(t)][:3])
        count, footers = section_band_xml(path, "footer")
        passed &= check(count == result["pages"] and all("PAGE" in xml for xml in footers),
                        "chapter header: the page number stays a field on every page",
                        ["PAGE" in xml for xml in footers])
        passed &= check(not any("could not be lifted" in w or "No repeating" in w for w in result["warnings"]),
                        "chapter header: no warning that the band was not lifted", result["warnings"])

        # 2. Odd and even footers: the same words, in the other order.
        result, path = run_case(cases.TAGGED_ODD_EVEN_FOOTER, workdir, "odd-even-footer")
        declined = result["declined"]
        passed &= check(declined.get("bands-tagged", 0) == 1 and not declined.get("bands-tagged-varies")
                        and not declined.get("bands-pages-apart"),
                        "odd/even footer: lifted, not declined as varying", declined)
        passed &= check(result["textCheck"]["missing"] == 0 and not result["textCheck"]["fellBack"],
                        "odd/even footer: nothing lost, no fallback", result["textCheck"])
        count, footers = section_band_texts(path, "footer")
        passed &= check(count == result["pages"], "odd/even footer: one section per page", (count, result["pages"]))
        order_ok = all(("Page" in text and "CONFIDENTIAL" in text
                        and (text.index("Page") < text.index("CONFIDENTIAL")) == (index % 2 == 0))
                       for index, text in enumerate(footers))
        passed &= check(order_ok, "odd/even footer: each page has its own order of the two halves", footers)
        count, xml = section_band_xml(path, "footer")
        passed &= check(all("PAGE" in item for item in xml), "odd/even footer: the number is a field on both kinds",
                        ["PAGE" in item for item in xml])

        # 3. A sideways page in the middle: its band is read on its own paper.
        result, path = run_case(cases.TAGGED_LANDSCAPE_BAND, workdir, "landscape-band")
        declined = result["declined"]
        passed &= check(declined.get("bands-tagged", 0) == 1 and not declined.get("bands-mixed-sizes")
                        and not declined.get("bands-pages-apart"),
                        "landscape: a sideways page does not decline the bands", declined)
        passed &= check(result["textCheck"]["missing"] == 0 and not result["textCheck"]["fellBack"],
                        "landscape: nothing lost, no fallback", result["textCheck"])
        passed &= check(not any("mixes page sizes" in w for w in result["warnings"]),
                        "landscape: no warning that the sizes are mixed", result["warnings"])
        document = Document(path)
        sideways = [i for i, s in enumerate(document.sections) if s.page_width.pt > s.page_height.pt]
        count, headers = section_band_texts(path, "header")
        passed &= check(sideways and all("Blue Harbor" in headers[i] for i in sideways),
                        "landscape: the sideways section has its own header", (sideways, headers))
        stops = {}
        for index in (0, sideways[0] if sideways else 0):
            paragraph = document.sections[index].header.paragraphs[0]
            stops[index] = max((stop.position.pt for stop in paragraph.paragraph_format.tab_stops
                                if stop.position is not None), default=0.0)
        passed &= check(sideways and stops[sideways[0]] > stops[0] + 100.0,
                        "landscape: the sideways header's right edge is that paper's, not the portrait's", stops)

        # 4. Without a section per page Word cannot give a page its own
        #    header, so the chapter bands are declined and the text is kept.
        keep_run = convert.run_pdf2docx
        runs = []

        def one_section(source, part, *args, **kwargs):
            keep_run(source, part, *args, **kwargs)
            runs.append(source)
            if len(runs) == 1:
                merged = Document(part)
                body = merged.element.body
                for ppr in list(body.iter(convert.qn("w:pPr"))):
                    for inner in ppr.findall(convert.qn("w:sectPr")):
                        ppr.remove(inner)
                merged.save(part)

        convert.run_pdf2docx = one_section
        try:
            result, path = run_case(cases.TAGGED_CHAPTER_HEADER, workdir, "chapter-one-section")
        except Exception as exc:  # noqa: BLE001 - the check reports it
            result = {"ok": False, "error": repr(exc)}
        finally:
            convert.run_pdf2docx = keep_run
        passed &= check(result.get("ok") and result["declined"].get("bands-pages-apart", 0) == 1
                        and result["textCheck"]["missing"] == 0 and not result["headerFooter"]["header"],
                        "chapter header: one Word section for the whole report keeps the headers in the body",
                        result.get("error") or (result["declined"], result["textCheck"], result["headerFooter"]))

        # 5. A number on the right of odd pages and the left of even ones is
        #    two footers, not one: the field comes after the tabs on one kind
        #    and before them on the other.
        result, path = run_case(cases.TAGGED_MIRRORED_NUMBER, workdir, "mirrored-number")
        declined = result["declined"]
        passed &= check(declined.get("bands-tagged", 0) == 1 and not declined.get("bands-tagged-varies")
                        and not declined.get("bands-tagged-partial") and not declined.get("bands-pages-apart"),
                        "mirrored number: lifted, as one band per side", declined)
        passed &= check(result["textCheck"]["missing"] == 0 and not result["textCheck"]["fellBack"],
                        "mirrored number: nothing lost, no fallback", result["textCheck"])
        count, footers = section_band_xml(path, "footer")
        sides = []
        for xml in footers:
            field, tab = xml.find("PAGE"), xml.find("<w:tab/>")
            sides.append("right" if 0 <= tab < field else "left" if 0 <= field < tab else "?")
        passed &= check(count == result["pages"] and sides == ["right", "left"] * (count // 2) + ["right"] * (count % 2),
                        "mirrored number: odd pages carry it after the tabs and even pages before them", sides)
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return passed


def run_band_paragraph_checks():
    """Where a band's text starts and stops in the Word paragraph, from the columns the PDF drew."""
    from docx.shared import Pt

    def column(text, x0, x1):
        return {"x0": x0, "x1": x1, "text": text, "parts": [{"kind": "text", "text": text}], "font": "Arial",
                "psname": "Arial", "size": 8.0, "bold": False, "color": None, "base": None, "page_offset": None}

    def paragraph_of(columns, strip, left_margin, right_margin=72.0):
        document = Document()
        band = {"edge": "bottom", "columns": columns, "text_top": 745.0, "text_bottom": 756.0, "rule_under": None,
                "strip_behind": strip, "strip_full_width": False, "strip_span": (64.8, 576.0) if strip else None,
                "strip_y": None}
        paragraph = document.add_paragraph()
        convert.fill_band_paragraph(paragraph, band, left_margin, right_margin, 64.8, 576.0, 792.0, page_width=612.0)
        return paragraph

    passed = True
    # b2's footer on a page whose margin is 60.7 pt: the strip starts 4.1 pt in, the text is at the strip's
    # start, and a left stop at the indent was passed over by the first tab, which centred the text.
    footer = paragraph_of([column("Meridian CONFIDENTIAL", 64.8, 267.0), column("9 / 48", 550.0, 576.0)],
                          (0.88, 0.88, 0.88), 60.7)
    runs = [t for t in footer._p.iter(convert.qn("w:tab")) if t.getparent().tag == convert.qn("w:r")]
    passed &= check(len(runs) == 2, "band paragraph: two tabs between the three slots and no leading one", runs)
    passed &= check(footer.paragraph_format.first_line_indent in (None, 0),
                    "band paragraph: text at the strip's start needs no first-line indent",
                    footer.paragraph_format.first_line_indent)
    # The left text starts 31 pt inside the rule (a logo stands before it): a first-line indent, not a stop.
    header = paragraph_of([column("Meridian", 95.8, 192.0), column("Halcyon", 450.0, 576.0)],
                          (0.88, 0.88, 0.88), 64.8)
    passed &= check(header.paragraph_format.first_line_indent is not None
                    and abs(header.paragraph_format.first_line_indent.pt - 31.0) < 0.2,
                    "band paragraph: the left text starts where the PDF's left column does",
                    header.paragraph_format.first_line_indent)
    stops = sorted(stop.position.pt for stop in header.paragraph_format.tab_stops
                   if stop.alignment != convert.WD_TAB_ALIGNMENT.CLEAR)
    passed &= check(len(stops) == 2 and abs(stops[1] - (576.0 - 64.8)) < 0.2,
                    "band paragraph: the right stop is the right column's own edge", stops)
    # A margin wider than the text area (a page whose text sits left of the margin): the indent goes negative
    # rather than being held at the margin.
    wide = paragraph_of([column("Meridian", 64.8, 192.0), column("9", 570.0, 576.0)], (0.88, 0.88, 0.88), 70.0)
    passed &= check(wide.paragraph_format.left_indent is not None and wide.paragraph_format.left_indent.pt < -5.0,
                    "band paragraph: text left of the margin is placed there", wide.paragraph_format.left_indent)
    # The widest fill is the strip: a progress bar's fill over its track does not replace it.
    class Rect:
        def __init__(self, x0, y0, x1, y1):
            self.x0, self.y0, self.x1, self.y1 = x0, y0, x1, y1
            self.width, self.height = x1 - x0, y1 - y0

    class Page:
        class rect:
            width = 792.0
    track = {"rect": Rect(57.6, 563.4, 734.4, 565.9), "fill": (0.88, 0.88, 0.88), "type": "f"}
    fill = {"rect": Rect(57.6, 563.4, 649.8, 565.9), "fill": (0.05, 0.49, 0.52), "type": "f"}
    for drawings in ([track, fill], [fill, track]):
        found = convert.band_decorations([Page], drawings, [0], "bottom", 567.0, 578.0)
        passed &= check(found[3] == (0.88, 0.88, 0.88) and found[5] == (57.6, 734.4),
                        "band decorations: the widest fill is the strip, whatever the drawing order", found[3:6])
    return passed


def run_page_section_checks():
    """Which Word sections belong to which PDF page, and the check on a band's text."""
    from docx.enum.section import WD_SECTION
    from docx.shared import Pt

    passed = True
    document = Document()
    first = document.sections[0]
    first.page_width, first.page_height = Pt(612), Pt(792)
    document.add_paragraph("page one")
    # A page laid out in columns: a continuous section and a next-column one.
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("page two, left")
    document.add_section(WD_SECTION.NEW_COLUMN)
    document.add_paragraph("page two, right")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("page three")
    portrait = (612.0, 792.0)
    groups = convert.sections_by_page(document, [portrait, portrait, portrait])
    passed &= check(groups == [[0], [1, 2], [3]],
                    "page sections: a column section belongs to the page it opens on", groups)
    passed &= check(convert.sections_by_page(document, [portrait] * 4) is None,
                    "page sections: four pages from three starts is not a map")
    passed &= check(convert.sections_by_page(document, [portrait] * 2) is None,
                    "page sections: two pages from three starts is not a map")
    # With a sideways page the sizes must agree, not only the count.
    sideways = (792.0, 612.0)
    passed &= check(convert.sections_by_page(document, [portrait, sideways, portrait]) is None,
                    "page sections: a section of the wrong paper breaks the map")
    document.sections[1].page_width, document.sections[1].page_height = Pt(792), Pt(612)
    document.sections[2].page_width, document.sections[2].page_height = Pt(792), Pt(612)
    passed &= check(convert.sections_by_page(document, [portrait, sideways, portrait]) == [[0], [1, 2], [3]],
                    "page sections: sections of the page's own paper hold")

    # bands_need_pages: several texts on one edge need a section per page.
    one = {"edge": "top", "pages": [0, 1], "source": "tagged", "columns": []}
    other = dict(one, pages=[2, 3])
    footer = {"edge": "bottom", "pages": [0, 1, 2, 3], "source": "tagged", "columns": []}
    passed &= check(not convert.bands_need_pages([dict(one, pages=[0, 1, 2, 3]), footer]),
                    "bands need pages: one text per edge on the same pages does not")
    passed &= check(convert.bands_need_pages([one, other, footer]),
                    "bands need pages: two header texts do")
    passed &= check(convert.bands_need_pages([one, footer]), "bands need pages: edges on different pages do")

    # band_text_gaps: the words of a band must be in the section it begins in.
    sections = Document()
    sections.add_paragraph("a")
    sections.add_section(WD_SECTION.NEW_PAGE)
    sections.add_paragraph("b")
    for index, text in enumerate(("Alpha chapter", "Beta chapter")):
        holder = sections.sections[index].header
        holder.is_linked_to_previous = False
        holder.paragraphs[0].text = text
    band = lambda text, pages: {"edge": "top", "source": "tagged", "pages": pages,
                                "columns": [{"parts": [{"kind": "text", "text": text}]}]}
    passed &= check(convert.band_text_gaps(sections, [band("Alpha chapter", [0]), band("Beta chapter", [1])],
                                           [[0], [1]]) == [],
                    "band text: each page's words are in its own section")
    gaps = convert.band_text_gaps(sections, [band("Alpha chapter", [0, 1])], [[0], [1]])
    passed &= check(len(gaps) == 1 and "page 2" in gaps[0],
                    "band text: a page whose words are in another section's header is named", gaps)

    # A page number seen on one page only takes the offset its text showed elsewhere.
    def row(text, columns):
        return {"text": text, "y0": 780.0, "y1": 790.0, "x0": 72.0, "x1": 300.0, "base": 788.0,
                "columns": [{"text": c, "x0": 72.0 + 100.0 * i, "x1": 150.0 + 100.0 * i, "font": "F", "psname": "F",
                             "size": 9.0, "bold": False, "color": 0, "base": 788.0} for i, c in enumerate(columns)]}
    many = [(i, row("Page %d" % (i + 1), ["Page %d" % (i + 1)])) for i in (2, 4, 6)]
    alone = [(9, row("Page 10", ["Page 10"]))]
    groups = {((595.0, 842.0), "Page #", 0): many, ((842.0, 595.0), "Page #", 1): alone}
    described = convert.describe_groups(groups)
    parts = [c["parts"] for c in described[((842.0, 595.0), "Page #", 1)]]
    passed &= check(described is not None and any(p["kind"] == "page" for ps in parts for p in ps),
                    "describe groups: a page number on a single page is a field when other pages show the offset", parts)
    lone = convert.describe_groups({((842.0, 595.0), "Page #", 1): alone})
    parts = [c["parts"] for c in lone[((842.0, 595.0), "Page #", 1)]]
    passed &= check(not any(p["kind"] == "page" for ps in parts for p in ps),
                    "describe groups: with nothing to agree with, the digits stay as printed", parts)
    return passed


def run_page_number_checks():
    """Page numbers that restart, and a total that follows the document (structure review findings 1 and 7)."""
    from docx.enum.section import WD_SECTION
    from docx.oxml.ns import qn
    from docx.shared import Pt

    passed = True

    def band(pages, offset, parts, edge="bottom"):
        column = {"parts": parts, "page_offset": offset, "x0": 72.0, "x1": 300.0, "font": "F", "psname": "F",
                  "size": 9.0, "bold": False, "color": 0, "base": 780.0}
        return {"edge": edge, "source": "tagged", "pages": pages, "columns": [column], "rule_under": None,
                "strip_behind": None, "text_top": 770.0, "text_bottom": 782.0}

    page = {"kind": "page"}
    text = lambda value: {"kind": "text", "text": value}

    # printed_page_numbers and page_number_starts: the p10 shape prints 1, 2, 1, 2.
    p10 = [band([0, 1], 1, [text("Page "), page]), band([2, 3], -1, [text("Page "), page])]
    printed = convert.printed_page_numbers(p10)
    passed &= check(printed == {0: 1, 1: 2, 2: 1, 3: 2}, "page numbers: what the PDF printed on each page", printed)
    passed &= check(convert.page_number_starts(printed) == {2: 1},
                    "page numbers: a body that restarts at 1 gets a start on its first page",
                    convert.page_number_starts(printed))
    passed &= check(convert.page_number_starts({0: 1, 1: 2, 2: 3}) == {},
                    "page numbers: a count that never restarts needs no start")
    # p1: a cover without a number, roman static pages, then arabic from 1 on page index 3.
    passed &= check(convert.page_number_starts({3: 1, 4: 2}) == {3: 1},
                    "page numbers: a body that opens after unnumbered pages restarts at 1, never below 0")
    passed &= check(convert.page_number_starts({0: 5, 1: 6}) == {0: 5},
                    "page numbers: a first page that prints 5 starts at 5")
    passed &= check(convert.page_number_starts({1: -1}) == {},
                    "page numbers: a start below 0 is not written")
    passed &= check(convert.page_number_starts({0: 1, 1: 2, 2: 7, 3: 8, 4: 4}) == {2: 7, 4: 4},
                    "page numbers: each change of the running count is one start")

    # The "of 42" is a PAGEREF to a bookmark in the body: a pass that drops that paragraph leaves a footer
    # that prints "Error! Bookmark not defined.", and band_text_gaps names it.
    from docx.enum.section import WD_SECTION
    ref_document = Document()
    ref_document.add_paragraph("one")
    ref_document.add_section(WD_SECTION.NEW_PAGE)
    ref_document.add_paragraph("two")
    convert.add_last_page_bookmark(ref_document)
    footer_part = ref_document.sections[0].footer
    footer_part.is_linked_to_previous = False
    convert.add_total_field(footer_part.paragraphs[0], " PAGEREF %s " % convert.LAST_PAGE_BOOKMARK, "2")
    ref_band = {"edge": "bottom", "source": "tagged", "pages": [0, 1], "columns": []}
    gaps = convert.band_text_gaps(ref_document, [ref_band], [[0], [1]])
    passed &= check(not any("bookmark" in gap for gap in gaps),
                    "page ref: a footer whose bookmark is in the body is not a gap", gaps)
    for node in list(ref_document.element.body.iter(convert.qn("w:bookmarkStart"))):
        node.getparent().remove(node)
    gaps = convert.band_text_gaps(ref_document, [ref_band], [[0], [1]])
    passed &= check(any("bookmark" in gap for gap in gaps),
                    "page ref: a footer whose bookmark was dropped is named as a gap", gaps)

    # split_total_parts: only the number right after the page field, only when it is the last number.
    parts = [text("Page "), page, text(" of 42")]
    out, changed = convert.split_total_parts(parts, 42)
    passed &= check(changed and [x["kind"] for x in out] == ["text", "page", "text", "total"]
                    and out[2]["text"] == " of ", "total: ' of 42' after the page field becomes a field", out)
    # Connectors: "out of" and a capital "Of" join a total; a label ("Team", "Rev") does not.
    passed &= check(all(convert.is_total_connector(c) for c in (" of ", " / ", " out of ", " Of ", " von "))
                    and not any(convert.is_total_connector(c) for c in (" | Team ", " Rev ", " Draft ")),
                    "total: connectors are separators or short function words, labels are not")
    out, changed = convert.split_total_parts(parts, 41)
    passed &= check(not changed and out == parts, "total: a constant that is not the last number stays text")
    out, changed = convert.split_total_parts([text("3 / 42 pages"), page], 42)
    passed &= check(not changed, "total: digits before the page field are not a total")
    out, changed = convert.split_total_parts([page, text(" / 42 Rev A")], 42)
    passed &= check(changed and out[-1] == {"kind": "text", "text": " Rev A"},
                    "total: what follows the total stays text", out)
    out, changed = convert.split_total_parts([page, text(" of 420")], 42)
    passed &= check(not changed, "total: 420 is not 42")
    out, changed = convert.split_total_parts([page, text(" is a very long run of words 42")], 42)
    passed &= check(not changed, "total: a number far from the page field is not its total")
    # The connector has to read as one: a separator, or a word of up to three letters.
    for connector in (" of ", " OF ", " / ", " von ", " | ", " \u00b7 ", " \u2013 ", " - ", "/"):
        out, changed = convert.split_total_parts([page, text(connector + "42")], 42)
        passed &= check(changed and out[-1] == {"kind": "total"} and out[-2]["text"] == connector,
                        "total: %r before the last number makes it the total" % connector, out)
    for connector in (" | Team ", " Rev ", " \u00b7 Team ", " Appendix "):
        out, changed = convert.split_total_parts([page, text(connector + "42")], 42)
        passed &= check(not changed and out == [page, text(connector + "42")],
                        "total: a static number after %r stays text" % connector, out)

    # total_page_field: NUMPAGES only when the file's page count is the last printed number.
    field = convert.total_page_field({0: 1, 1: 2, 2: 3}, 3, {})
    passed &= check(field and "NUMPAGES" in field["instruction"] and field["last"] == 3 and not field["bookmark"],
                    "total field: no restart and numbering from 1 is NUMPAGES", field)
    field = convert.total_page_field({0: 1, 1: 2, 2: 1, 3: 2}, 4, {2: 1})
    passed &= check(field and "PAGEREF" in field["instruction"] and field["last"] == 2 and field["bookmark"],
                    "total field: a restart makes it a PAGEREF to the last page", field)
    field = convert.total_page_field({1: 1, 2: 2}, 3, {1: 1})
    passed &= check(field and "PAGEREF" in field["instruction"] and field["last"] == 2,
                    "total field: an unnumbered cover makes the total two for three pages", field)
    field = convert.total_page_field({0: 1, 1: 2}, 5, {})
    passed &= check(field and field["last"] == 5 and "NUMPAGES" in field["instruction"],
                    "total field: pages after the last band carry the count forward", field)
    passed &= check(convert.total_page_field({}, 4, {}) is None, "total field: no page number, no total")

    # word_page_numbers and band_text_gaps: the starts in the file must give the printed numbers.
    document = Document()
    document.sections[0].page_width, document.sections[0].page_height = Pt(612), Pt(792)
    document.add_paragraph("one")
    for label in ("two", "three", "four"):
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_paragraph(label)
    groups = [[0], [1], [2], [3]]
    passed &= check(convert.word_page_numbers(document, groups) == [1, 2, 3, 4],
                    "word numbers: no start counts on from 1", convert.word_page_numbers(document, groups))
    numbering = lambda gaps: [g for g in gaps if "page number" in g]
    gaps = numbering(convert.band_text_gaps(document, p10, groups))
    passed &= check(len(gaps) == 1 and "PDF page 3 prints page number 1" in gaps[0] and "print 3" in gaps[0],
                    "band text: a restart the file does not carry is named", gaps)
    convert.set_page_numbering_start(document.sections[2], 1)
    passed &= check(convert.word_page_numbers(document, groups) == [1, 2, 1, 2],
                    "word numbers: a start on a lead section restarts the count",
                    convert.word_page_numbers(document, groups))
    passed &= check(numbering(convert.band_text_gaps(document, p10, groups)) == [],
                    "band text: with the start written the numbers agree")

    # apply_bands writes the starts, and the total as a field with its result kept.
    document = Document()
    document.sections[0].page_width, document.sections[0].page_height = Pt(612), Pt(792)
    document.add_paragraph("one")
    for label in ("two", "three", "four"):
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_paragraph(label)
    for section in document.sections:
        section.page_width, section.page_height = Pt(612), Pt(792)
    sizes = [(612.0, 792.0)] * 4
    bands = [band([0, 1], 1, [text("Front "), page, text(" of 2")]),
             band([2, 3], -1, [text("Body "), page, text(" of 2")])]
    convert.apply_bands(document, bands, 792.0, 72.0, 540.0, 4, sizes)
    starts = [(i, s._sectPr.find(qn("w:pgNumType")).get(qn("w:start")))
              for i, s in enumerate(document.sections) if s._sectPr.find(qn("w:pgNumType")) is not None]
    passed &= check(starts == [(0, "1"), (2, "1")], "apply bands: the body section restarts at 1", starts)
    footer = document.sections[2].footer.paragraphs[0]._p
    code = [n.text for n in footer.iter(qn("w:instrText"))]
    passed &= check(code == [" PAGE ", " PAGEREF _btct_last_page "],
                    "apply bands: a total after a restart is a PAGEREF", code)
    shown = "".join(n.text for n in document.sections[2].footer.paragraphs[0]._p.iter(qn("w:t")))
    passed &= check(shown == "Body  of 2",
                    "apply bands: the field keeps the last printed number as its result", shown)
    marks = [n.get(qn("w:name")) for n in document.element.body.iter(qn("w:bookmarkStart"))]
    passed &= check(marks == ["_btct_last_page"], "apply bands: one bookmark names the last page", marks)
    passed &= check(convert.band_text_gaps(document, bands, [[0], [1], [2], [3]]) == [],
                    "apply bands: the file gives every page the number the PDF printed")

    # Without a restart the total is NUMPAGES and no bookmark is added.
    document = Document()
    document.sections[0].page_width, document.sections[0].page_height = Pt(612), Pt(792)
    document.add_paragraph("one")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("two")
    for section in document.sections:
        section.page_width, section.page_height = Pt(612), Pt(792)
    bands = [band([0, 1], 1, [text("Page "), page, text(" of 2")])]
    convert.apply_bands(document, bands, 792.0, 72.0, 540.0, 2, [(612.0, 792.0)] * 2)
    code = [n.text for n in document.sections[0].footer.paragraphs[0]._p.iter(qn("w:instrText"))]
    passed &= check(code == [" PAGE ", " NUMPAGES "], "apply bands: a plain count totals with NUMPAGES", code)
    passed &= check(not list(document.element.body.iter(qn("w:bookmarkStart"))),
                    "apply bands: NUMPAGES needs no bookmark")
    return passed


def describe(path):
    doc = pymupdf.open(path)
    print("pages: %d" % doc.page_count)
    print("content bounds: %.1f .. %.1f" % convert.content_bounds(doc))
    bands = []
    for edge in ("top", "bottom"):
        band, reason = convert.detect_band(doc, edge)
        if band is None:
            print("%s: no repeating band%s" % (edge, (" (declined: %s)" % reason) if reason else ""))
            continue
        bands.append(band)
        print(
            "%s: y %.1f..%.1f on %d page(s), rule=%s strip=%s"
            % (edge, band["top"], band["bottom"], len(band["pages"]), band["rule_under"], band["strip_behind"])
        )
        for column in band["columns"]:
            print(
                "   x %.1f..%.1f %s %.1fpt offset=%s %s"
                % (
                    column["x0"],
                    column["x1"],
                    column["font"],
                    column["size"],
                    column["page_offset"],
                    [p.get("text", "<PAGE>") for p in column["parts"]],
                )
            )
    entries = convert.detect_toc(doc, bands)
    print("table-of-contents entries: %d" % len(entries))
    for entry in entries[:40]:
        print("   p%-3d x%.1f  %-52s %s" % (entry["page_index"], entry["x0"], entry["title"][:52], entry["number"]))
    doc.close()


def run_alignment_fix_checks():
    """A centred figure and caption stay centred in Word, and a list item in a cell is never centred."""
    import io

    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    passed = True
    png = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 20), False)
    picture = io.BytesIO(png.tobytes("png"))

    document = Document()
    section = document.sections[0]
    section.page_width, section.left_margin, section.right_margin = Inches(8.5), Inches(1), Inches(1)
    centred = document.add_paragraph()
    centred.add_run().add_picture(picture, width=Inches(4))
    centred.paragraph_format.left_indent = Inches(1.25)  # (6.5 - 4) / 2: centred by its indent
    picture.seek(0)
    flush = document.add_paragraph()
    flush.add_run().add_picture(picture, width=Inches(4))
    table = document.add_table(rows=3, cols=1)
    item, address, label = (row.cells[0].paragraphs[0] for row in table.rows)
    item.add_run("2. Apply input and output guardrails")
    address.add_run("129.21.249.5")
    label.add_run("Critical")
    for cell in (item, address, label):
        cell.alignment = WD_ALIGN_PARAGRAPH.CENTER

    declined = collections.Counter()
    passed &= check(convert.centre_lone_pictures(document, declined) == 1,
                    "align: one picture centred by its indent is centred", dict(declined))
    passed &= check(centred.alignment == WD_ALIGN_PARAGRAPH.CENTER and centred.paragraph_format.left_indent is None,
                    "align: the centred picture loses its indent and is centred")
    passed &= check(flush.alignment != WD_ALIGN_PARAGRAPH.CENTER,
                    "align: a picture at the margin is left where it is")
    passed &= check(convert.left_align_cell_list_items(document, declined) == 1,
                    "align: one list item in a cell set left", dict(declined))
    passed &= check(item.alignment == WD_ALIGN_PARAGRAPH.LEFT, "align: the remediation step is left-aligned")
    passed &= check(address.alignment == WD_ALIGN_PARAGRAPH.CENTER and label.alignment == WD_ALIGN_PARAGRAPH.CENTER,
                    "align: an address and a label stay centred")
    return passed


def run_label_box_checks():
    """An empty screenshot slot and a rounded tag come out of Word the shape the PDF draws them."""
    passed = True
    pdf = pymupdf.open()
    page = pdf.new_page(width=612, height=792)
    slot = pymupdf.Rect(81, 100, 531, 258.4)
    page.draw_rect(slot, color=None, fill=(0.96, 0.96, 0.96))
    page.draw_rect(slot, color=(0.706, 0.706, 0.706), fill=None, width=1)
    label = "[ Screenshot placeholder ]"
    size = 10
    width = pymupdf.get_text_length(label, fontsize=size)
    page.insert_text(((slot.x0 + slot.x1 - width) / 2, (slot.y0 + slot.y1) / 2 + size * 0.35), label, fontsize=size)
    tag = pymupdf.Rect(460, 300, 540, 312)
    page.draw_rect(tag, color=None, fill=(0.878, 0.4, 0.4), radius=2 / 12)
    page.draw_rect(tag, color=(1, 1, 1), fill=None, width=1.5, radius=2 / 12)
    width = pymupdf.get_text_length("REDISCOVERED", fontsize=7)
    page.insert_text(((tag.x0 + tag.x1 - width) / 2, (tag.y0 + tag.y1) / 2 + 7 * 0.35), "REDISCOVERED", fontsize=7)
    found = {box["key"]: box for box in convert.label_boxes(pdf)}
    placeholder = found.get(convert.compare_squash(label))
    rediscovered = found.get(convert.compare_squash("REDISCOVERED"))
    passed &= check(placeholder is not None and rediscovered is not None,
                    "labels: the empty slot and the tag are both found", sorted(found))
    if placeholder is None or rediscovered is None:
        return False
    passed &= check(placeholder["radius"] < convert.LABEL_ROUND_MIN_PT and placeholder["stroke"] == "B4B4B4"
                    and placeholder["centred"],
                    "labels: the slot is square, outlined in its grey and centred on the page", placeholder)
    passed &= check(rediscovered["radius"] >= convert.LABEL_ROUND_MIN_PT and rediscovered["stroke"] == "FFFFFF",
                    "labels: the tag is rounded with a white outline", rediscovered)

    document = Document()
    slot_table = document.add_table(rows=1, cols=1)
    cell = slot_table.rows[0].cells[0].paragraphs[0]
    cell.paragraph_format.space_before = convert.Pt(70)
    cell.add_run("            " + label)
    document.add_paragraph("between")
    tag_table = document.add_table(rows=1, cols=1)
    tag_table.rows[0].cells[0].paragraphs[0].add_run("REDISCOVERED")
    declined = collections.Counter()
    passed &= check(convert.rebuild_label_boxes(document, [placeholder, rediscovered], 72.0, declined) == 2,
                    "labels: both boxes are rebuilt", dict(declined))
    body = document.element.body
    tables = body.findall(convert.qn("w:tbl"))
    slot_xml = tables[0] if tables else None
    passed &= check(slot_xml is not None
                    and slot_xml.find(".//" + convert.qn("w:vAlign")).get(convert.qn("w:val")) == "center"
                    and slot_xml.find(".//" + convert.qn("w:trHeight")).get(convert.qn("w:val")) == "3168",
                    "labels: the slot is a box its own height whose cell centres the label")
    text = "".join(t.text for t in slot_xml.iter(convert.qn("w:t"))) if slot_xml is not None else ""
    spacing = slot_xml.find(".//" + convert.qn("w:spacing")) if slot_xml is not None else None
    passed &= check(text == label and spacing is not None and spacing.get(convert.qn("w:before")) == "0",
                    "labels: the label loses the spaces and spacing that pushed it into place", text)
    shape = body.find(".//" + convert.qn("a:prstGeom"))
    passed &= check(len(tables) == 1 and shape is not None and shape.get("prst") == "roundRect",
                    "labels: the tag is a rounded rectangle, not a table")
    outline = body.find(".//" + convert.qn("a:ln"))
    passed &= check(outline is not None and outline.get("w") == "19050",
                    "labels: the tag keeps its 1.5pt outline")
    return passed


if __name__ == "__main__":
    if len(sys.argv) > 1:
        describe(sys.argv[1])
        sys.exit(0)
    ok = run_built_in()
    ok = run_documents() and ok
    ok = run_band_paragraph_checks() and ok
    ok = run_page_section_checks() and ok
    ok = run_page_number_checks() and ok
    ok = run_band_variant_documents() and ok
    ok = run_switch_checks() and ok
    ok = run_rtl_check_checks() and ok
    ok = run_alignment_fix_checks() and ok
    ok = run_label_box_checks() and ok
    ok = run_corpus_checks() and ok
    sys.exit(0 if ok else 1)
